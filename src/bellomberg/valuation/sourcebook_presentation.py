"""Read-only workbook view of already acquired financial, consensus and guidance data.

No provider is called here. Raw units and provider period labels are deliberately
preserved: a displayed history is not a certified comparable time series.
"""
from math import ceil, isfinite
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from bellomberg.core.language import text as tr

from .documented_presentation import _finish, _header, _line, _put, _sheet

# General retains small ratios and lets Excel use scientific notation for raw magnitudes.
AMOUNT = 'General'


MISSING = 'n.d.'
SOURCE_NAMES = ('financials', 'consensus', 'guidance')
_METADATA = frozenset(('source_id', 'as_of', 'status', 'unit', 'valid_until'))


def _literal(ws, row, col, value, fmt=None):
    """Provider text is never a formula, including leading =, +, - or @."""
    cell = _put(ws, row, col, value, fmt)
    if isinstance(value, str):
        cell.data_type = 's'
    return cell


def _unit(source, value=None):
    explicit = (value.get('unit') if isinstance(value, dict) else None) or source.get('unit')
    return explicit if isinstance(explicit, str) and explicit.strip() else MISSING


def _metadata(ws, source, as_of):
    for row, (name, value) in enumerate((
        (tr('Fonte', 'Source'), source.get('source_id') or MISSING),
        (tr('Osservata il', 'Observed on'), source.get('as_of') or MISSING),
        (tr('Cutoff', 'Cutoff'), as_of or MISSING),
        (tr('Stato', 'Status'), source.get('status') or 'data_missing'),
        (tr('Scadenza', 'Valid until'), source.get('valid_until') or MISSING),
        (tr('Nota provider', 'Provider note'), source.get('message') or MISSING),
    ), 5):
        _literal(ws, row, 2, name)
        _literal(ws, row, 4, value)


def _walk(value, path=(), inherited=None):
    """Yield raw leaves; period/metric labels are identifiers, never inferred dates."""
    inherited = inherited or {}
    if isinstance(value, dict):
        context = {**inherited, **{key: value[key] for key in _METADATA if key in value}}
        if value.get('table') is True and {'index', 'columns', 'data'} <= set(value):
            index, columns, cells = (value[key] for key in ('index', 'columns', 'data'))
            if (isinstance(index, list) and isinstance(columns, list) and isinstance(cells, list)
                    and len(index) == len(cells) and all(isinstance(line, list) and len(line) == len(columns) for line in cells)):
                for label, line in zip(index, cells):
                    for column, cell in zip(columns, line):
                        yield (*path, label, column), cell, context
                yield (*path, 'missing_cells'), value.get('missing_cells', MISSING), context
            else:
                yield (*path, 'malformed_table'), MISSING, context
        else:
            for key, child in value.items():
                if key not in _METADATA:
                    yield from _walk(child, (*path, key), context)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            # List order remains visible. A period label is displayed in its own row.
            yield from _walk(child, (*path, index), inherited)
    else:
        yield path, value, inherited


def _display(value):
    if value is None:
        return MISSING
    if isinstance(value, float) and not isfinite(value):
        return MISSING
    return value


def _history(ws, source):
    data = source.get('data')
    row = 14
    if not isinstance(data, dict):
        _literal(ws, row, 2, tr('Bilanci assenti', 'Financial statements unavailable'))
        return row + 1
    errors = data.get('errors') if isinstance(data.get('errors'), dict) else {}
    for name in ('income_stmt', 'balance_sheet', 'cashflow'):
        _line(ws, row, name, 8, section=True)
        row += 1
        table = data.get(name)
        if not isinstance(table, dict) or table.get('table') is not True:
            _literal(ws, row, 2, MISSING)
            _literal(ws, row, 4, errors.get(name) or tr('Tabella non acquisita', 'Table not acquired'))
            row += 2
            continue
        index, columns, cells = (table.get(key) for key in ('index', 'columns', 'data'))
        if (not isinstance(index, list) or not isinstance(columns, list) or not isinstance(cells, list)
                or len(index) != len(cells) or any(not isinstance(line, list) or len(line) != len(columns) for line in cells)):
            _literal(ws, row, 2, tr('Tabella malformata', 'Malformed table'))
            row += 2
            continue
        _header(ws, row, [tr('Voce originale', 'Original item'), tr('Unita dichiarata', 'Declared unit'), *columns])
        row += 1
        unit = _unit(source, table)
        for label, line in zip(index, cells):
            _literal(ws, row, 2, label)
            _literal(ws, row, 3, unit)
            for column, value in enumerate(line, 4):
                _literal(ws, row, column, _display(value), AMOUNT if isinstance(value, (int, float)) and not isinstance(value, bool) else None)
            row += 1
        _literal(ws, row, 2, tr('Celle assenti dal provider', 'Provider missing cells'))
        _literal(ws, row, 4, table.get('missing_cells', MISSING))
        row += 2
    return row


def _structured(ws, source):
    _header(ws, 13, [tr('Campo / percorso provider', 'Provider field / path'),
                     tr('Valore grezzo', 'Raw value'), tr('Unita', 'Unit'),
                     tr('Fonte', 'Source'), tr('Data', 'Date'), tr('Stato', 'Status'),
                     tr('Scadenza', 'Valid until')])
    data = source.get('data')
    row = 14
    if data is None or data == {} or data == []:
        _literal(ws, row, 2, tr('Dati assenti', 'Data unavailable'))
        _literal(ws, row, 3, MISSING)
        return row + 1
    count = 0
    for path, value, context in _walk(data):
        _literal(ws, row, 2, '/'.join(map(str, path)) or 'value')
        _literal(ws, row, 3, _display(value), AMOUNT if isinstance(value, (int, float)) and not isinstance(value, bool) else None)
        _literal(ws, row, 4, context.get('unit') or _unit(source))
        _literal(ws, row, 5, context.get('source_id') or source.get('source_id') or MISSING)
        _literal(ws, row, 6, context.get('as_of') or source.get('as_of') or MISSING)
        _literal(ws, row, 7, context.get('status') or source.get('status') or 'data_missing')
        _literal(ws, row, 8, context.get('valid_until') or source.get('valid_until') or MISSING)
        row += 1
        count += 1
    if not count:
        _literal(ws, row, 2, tr('Nessun valore acquisito', 'No acquired values'))
        row += 1
    return row


def present_sourcebook(wb, payload):
    """Append three literal source views from acquisition_snapshot; return False if absent."""
    case = (payload.get('acquisition_snapshot') or {}).get('case')
    if not isinstance(case, dict):
        return False
    sources = case.get('sources') if isinstance(case.get('sources'), dict) else {}
    for name, title in (('financials', tr('Storico contabile', 'Financial history')),
                        ('consensus', 'Consensus'), ('guidance', 'Guidance')):
        source = sources.get(name) if isinstance(sources.get(name), dict) else {'status': 'data_missing'}
        ws = _sheet(wb, {'financials': 'Source History', 'consensus': 'Source Consensus',
                         'guidance': 'Source Guidance'}[name], title, 9)
        _line(ws, 4, tr('Dati grezzi dello snapshot: periodi e unita non normalizzati; nessuna comparabilita certificata.',
                        'Raw snapshot data: periods and units unnormalized; comparability is not certified.'), 9, height=40)
        _metadata(ws, source, case.get('as_of'))
        last = _history(ws, source) if name == 'financials' else _structured(ws, source)
        ws.column_dimensions['B'].width = 58
        ws.column_dimensions['C'].width = 18 if name == 'financials' else 32
        if name == 'financials':
            for column in range(4, ws.max_column+1):
                ws.column_dimensions[get_column_letter(column)].width=24
        else:
            ws.column_dimensions['E'].width = 38
        for cells in ws.iter_rows(min_row=5):
            lines = 1
            for cell in cells:
                if cell.value is None:
                    continue
                numeric=isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool)
                cell.alignment = Alignment(vertical='top', wrap_text=not numeric,
                    horizontal='right' if numeric else 'left')
                width = ws.column_dimensions[get_column_letter(cell.column)].width
                lines = max(lines, sum(max(1, ceil(len(part) / max(1, width-2))) for part in str(cell.value).split('\n')))
            ws.row_dimensions[cells[0].row].height = min(409, max(22, 16*lines+6))
        ws.freeze_panes = 'D14'
        ws.print_title_rows = '13:13' if name != 'financials' else '2:4'
        _finish(ws, last, max(9, ws.max_column))
    return True
