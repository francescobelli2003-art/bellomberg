"""Dated primary market references for analyst judgments, never valuation inputs.

Reuse the trusted immutable downloader. Keep Treasury yields and each published
ERP variant separate; no issuer beta, WACC, terminal rate or default is inferred.
"""
from copy import deepcopy
import csv
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO, StringIO
import json
from math import isfinite
from pathlib import Path
import re
from urllib.parse import urlsplit

from .document_evidence import source_dates


_ERP_COLUMNS = {
    'T.Bond Rate': 'treasury_rate_paired_with_erp',
    '$ Riskfree Rate': 'default_adjusted_usd_rate_paired_with_erp',
    'ERP (T12m)': 'erp_trailing_twelve_months',
    'ERP (T12m) with adj riskfree rate': 'erp_trailing_twelve_months_adjusted_rf',
    'ERP (T12 m with sustainable payout)': 'erp_sustainable_payout',
}
_LIMITATION = ('Primary market references only. ERP variants and their paired rates use different '
    'conventions and must not be mixed silently. Observation dates are not publication dates. '
    'No issuer exposure, beta, debt pricing, target weights, WACC, terminal rate or method record is inferred.')


def reference_urls(as_of):
    year = date.fromisoformat(as_of).year
    return {'treasury': 'https://home.treasury.gov/resource-center/data-chart-center/interest-rates/'
        f'daily-treasury-rates.csv/{year}/all?type=daily_treasury_yield_curve&field_tdr_date_value={year}&page&_format=csv',
        'erp': 'https://pages.stern.nyu.edu/~adamodar/pc/implprem/ERPbymonth.xlsx'}


def _latest(rows, cutoff):
    if len({row[0] for row in rows}) != len(rows):
        raise ValueError('ambiguous duplicate market observation date')
    eligible = [row for row in rows if row[0] <= cutoff]
    if not eligible:
        raise ValueError('no market observation on or before cutoff')
    return max(eligible, key=lambda row: row[0])


def _treasury(raw, cutoff):
    from bellomberg.market_data.market_inputs import _bdays_between, _STALE_BDAYS
    reader = csv.DictReader(StringIO(raw.decode('utf-8-sig')))
    if not reader.fieldnames or any(reader.fieldnames.count(key) != 1 for key in ('Date', '10 Yr')):
        raise ValueError('unique Treasury Date/10 Yr headers required')
    rows = [(datetime.strptime(row['Date'], '%m/%d/%Y').date(), number, row)
            for number, row in enumerate(reader, 2)]
    day, number, row = _latest(rows, cutoff)
    if _bdays_between(day, cutoff) > _STALE_BDAYS:
        raise ValueError('STALE Treasury observation')
    try:
        value = Decimal(row['10 Yr']) / 100
    except (InvalidOperation, TypeError) as exc:
        raise ValueError('Treasury yield missing or nonnumeric') from exc
    if not value.is_finite() or not isfinite(float(value)):
        raise ValueError('Treasury yield is not finite')
    return [{'concept': 'treasury_10_year_par_yield', 'value': float(value), 'unit': 'ratio',
        'end': day.isoformat(), 'source_value': row['10 Yr'], 'source_unit': 'percent',
        'locator': {'row': number, 'date_column': 'Date', 'value_column': '10 Yr'}}]


def _erp(raw, cutoff):
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    try:
        if 'Historical ERP' not in wb.sheetnames:
            raise ValueError('Historical ERP worksheet missing')
        ws = wb['Historical ERP']
        header = list(next(ws.iter_rows(min_row=1, max_row=1)))
        names = [cell.value for cell in header]
        if any(names.count(key) != 1 for key in ('Start of month', *_ERP_COLUMNS)):
            raise ValueError('unique dated ERP variant headers required')
        date_index = names.index('Start of month')
        rows = []
        for row in ws.iter_rows(min_row=2):
            observed = row[date_index].value
            if observed is None:
                continue
            if isinstance(observed, str):
                # The primary workbook mixes Excel dates with explicit d-Mon-yy
                # text. Parse English month labels without the host's locale.
                match = re.fullmatch(r'(\d{1,2})-([A-Za-z]{3})-(\d{2})', observed)
                months = 'jan feb mar apr may jun jul aug sep oct nov dec'.split()
                if not match or match[2].lower() not in months:
                    raise ValueError('unrecognized ERP text date')
                month = months.index(match[2].lower()) + 1
                observed = datetime.strptime(f'{match[1]}-{month}-{match[3]}', '%d-%m-%y')
            if not isinstance(observed, datetime) or observed.day != 1:
                raise ValueError('ERP start-of-month date required')
            rows.append((observed.date(), row))
        day, row = _latest(rows, cutoff)
        # Same freshness convention as damodaran_data.get_erp; no stale snapshot fallback.
        if (cutoff - day).days > 60:
            raise ValueError('STALE monthly ERP observation')
        facts = []
        for name, concept in _ERP_COLUMNS.items():
            index = names.index(name); cell = row[index]; value = cell.value
            if type(value) not in (int, float) or not isfinite(value):
                raise ValueError('ERP variant missing/nonfinite: ' + name)
            facts.append({'concept': concept, 'value': value, 'unit': 'ratio', 'end': day.isoformat(),
                'source_label': name, 'locator': {'sheet': ws.title, 'value_cell': cell.coordinate,
                    'date_cell': row[date_index].coordinate, 'header_cell': header[index].coordinate}})
        return facts
    finally:
        wb.close()


def normalize_reference(raw, *, kind, as_of, retrieval):
    cutoff = date.fromisoformat(as_of)
    urls = reference_urls(as_of)
    if kind not in urls or not isinstance(retrieval, dict) or retrieval.get('url') != urls[kind]:
        raise ValueError('exact primary market dataset URL required')
    digest = sha256(raw).hexdigest()
    if retrieval.get('document_sha256') != digest:
        raise ValueError('market reference byte SHA256 mismatch')
    document = {'url': urls[kind], 'document_sha256': digest, 'published_at': None,
        'availability_basis': 'observed_download', 'retrieval': deepcopy(retrieval)}
    document.update(source_dates(document, cutoff))
    facts = _treasury(raw, cutoff) if kind == 'treasury' else _erp(raw, cutoff)
    text = json.dumps({'dataset': kind, 'currency': 'USD', 'facts': facts, 'limitation': _LIMITATION},
                      ensure_ascii=False, separators=(',', ':'), allow_nan=False)
    text_sha = sha256(text.encode('utf-8')).hexdigest()
    document.update(id='market-' + kind + '-' + text_sha, text=text, sha256=text_sha,
        origin='primary_market_reference', metadata={'normalizer': 'market_reference_v1',
        'currency': 'USD', 'dataset': kind, 'reference_cutoff': as_of},
        extraction_coverage={'status': 'dated_reference_observations', 'raw_sha256_checked': True,
                             'source_locators': True, 'limitation': _LIMITATION})
    return document


def collect_market_references(*, currency, as_of, archive_root, download=None, now=None, ticker=None):
    from bellomberg.market_data.lettore_trimestrali import scarica_documento
    result = {'status': 'incomplete', 'documents': [], 'issues': [], 'currency': currency,
              'limitation': _LIMITATION}
    if currency != 'USD':
        result['issues'].append({'source': 'market references',
            'reason': 'Primary market-reference acquisition currently supports explicit USD only; currency=' + str(currency)})
        return result
    root = Path(archive_root).resolve()
    download = download or scarica_documento
    for kind, url in reference_urls(as_of).items():
        try:
            fetched = download(url, str(root), host_consentiti=[urlsplit(url).hostname], public_only=True)
            if fetched.get('stato') != 'ok' or fetched.get('url_finale') != url:
                raise ValueError('primary market download failed or redirected: ' + str(fetched.get('motivo')))
            path = Path(fetched['path']).resolve()
            if not path.is_relative_to(root):
                raise ValueError('market source outside verified archive')
            raw = path.read_bytes()
            stamp = now if now is not None else datetime.now(timezone.utc)
            document = normalize_reference(raw, kind=kind, as_of=as_of, retrieval={
                'url': url, 'document_sha256': fetched['sha256'], 'retrieved_at': stamp.isoformat()})
            result['documents'].append(document)
        except Exception as exc:
            result['issues'].append({'source': 'market references ' + kind, 'reason': type(exc).__name__ + ': ' + str(exc)})
    if ticker is not None:
        from .beta_reference_evidence import collect_beta_reference
        beta = collect_beta_reference(ticker, currency=currency, as_of=as_of, archive_root=root, now=now)
        result['documents'].extend(beta['documents'])
        result['issues'].extend(beta['issues'])
        result['beta_reference'] = {k: deepcopy(v) for k, v in beta.items() if k != 'documents'}
    else:
        result['beta_reference'] = {'status': 'not_requested', 'reason': 'No verified issuer ticker supplied.'}
    result['status'] = 'ready' if not result['issues'] else 'partial' if result['documents'] else 'incomplete'
    return result
