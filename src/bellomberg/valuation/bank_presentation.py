"""Bank Summary in the reviewed FCFF layout, retaining the bank's live formulas."""
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.views import Selection

from bellomberg.core.language import text as tr
from .documented_formulas import _formula
from .documented_presentation import _sheet, _line, _header, _put, _finish, _number, col, PRICE, PERCENT, BLUE
from .market_quote import market_quote_view


def apply_bank_summary(model):
    wb, payload = model.wb, model.payload
    old = wb['Summary']
    values = {c: old.cell(8, c).value for c in (4, 5, 6)}
    guards = {c: old.cell(13, c).value for c in (4, 5, 6)}
    if any(not isinstance(v, str) or not v.startswith('=') for v in [*values.values(), *guards.values()]):
        raise ValueError('bank Summary requires the original guarded formulas')
    del wb['Summary']
    ws = _sheet(wb, 'Summary', payload['ticker'] + ' — ' + tr('Valutazione bancaria', 'Bank valuation'), 6, repeat_header=False)
    ws.freeze_panes = None
    ws.sheet_view.selection = [Selection(activeCell='B2', sqref='B2')]
    _line(ws, 4, model.value('model', 'perimeter')['entity'], 6, height=30)
    _line(ws, 5, tr('Formule collegate. Controlli del modello e approvazione del PM sono distinti.',
                    'Linked formulas. Model checks and PM approval are distinct.'), 6, section=True)
    _header(ws, 7, [tr('Scenari di valutazione', 'Valuation scenarios'), tr('Unita', 'Units'),
                    tr('Ribassista', 'Bear'), 'Base', tr('Rialzista', 'Bull')])
    cutoff = payload.get('acquisition_snapshot', {}).get('case', {}).get('as_of')
    raw = payload.get('market_quote') or {}
    quote = market_quote_view(raw, as_of=raw.get('acquired_as_of') or cutoff, usable=True)
    for row, label, unit in (
            (8, 'Fair value', payload.get('currency')),
            (9, tr('Upside sul prezzo osservato', 'Upside to observed price'), '%'),
            (11, tr('Costo del capitale proprio (Ke)', 'Cost of equity (Ke)'), '%'),
            (12, tr('Crescita perpetua', 'Perpetual growth'), '%'),
            (13, tr('Validita degli input', 'Input validity'), None)):
        _put(ws, row, 2, label, bold=row == 8)
        if unit is not None:
            _put(ws, row, 3, unit)
    observed_price = quote.get('price')
    for c, scenario in enumerate(('bear', 'base', 'bull'), 4):
        _formula(ws, 8, c, values[c][1:], PRICE)
        if quote['status_at_read'] == 'ok' and _number(observed_price) and observed_price > 0:
            _formula(ws, 9, c, f'IFERROR(IF(AND(ISNUMBER($E$19),$E$19>0),{col(c)}8/$E$19-1,"n.d."),"n.d.")', PERCENT)
        else:
            _put(ws, 9, c, None, PERCENT)
        _formula(ws, 11, c, model.ref(scenario, 'capital.ke'), PERCENT)
        _formula(ws, 12, c, model.ref(scenario, 'terminal_growth'), PERCENT)
        _formula(ws, 13, c, guards[c][1:], 'General')
    _line(ws, 15, tr('Date e riferimenti di mercato', 'Dates and market references'), 6, section=True)
    missing = tr('n.d.', 'n.a.')
    fields = [(tr('Data della valutazione', 'Model valuation date'), payload.get('valuation_date')),
              (tr('Prezzo del modello', 'Model price'), None),
              (tr('Prezzo osservato', 'Observed price'), observed_price),
              (tr('Ora del prezzo osservato (UTC)', 'Observed quote time (UTC)'), quote.get('observed_at')),
              (tr('Fonte / mercato', 'Source / exchange'), ' / '.join(str(quote.get(k) or missing) for k in ('source_id', 'exchange'))),
              (tr('Stato quotazione alla data acquisita', 'Quote status at acquisition date'), quote['status_at_read'])]
    for row, (label, value) in enumerate(fields, 17):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        _put(ws, row, 2, label)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        _put(ws, row, 5, value, PRICE if _number(value) else None)
    _formula(ws, 18, 5, model.ref('model', 'quotation', 'price'), PRICE)
    _line(ws, 23, tr('Valore alla data del modello, non rivalutato al prezzo osservato. La copia resta datata; Excel ricalcola le simulazioni.',
                     'Value at the model date, not rolled forward to the observed price. This is a dated copy; Excel recalculates simulations.'), 6, height=46)
    evidence = [r.get('evidence', {}) for r in payload['analytical_quality']['rows']]
    dates = sorted({r['source_date'] for r in evidence if isinstance(r.get('source_date'), str) and r['source_date']})
    expiry = sorted({r['valid_until'] for r in evidence if isinstance(r.get('valid_until'), str) and r['valid_until']})
    _line(ws, 26, tr('Date delle fonti citate: ', 'Cited source dates: ') + (' / '.join(dict.fromkeys((dates[0], dates[-1]))) if dates else missing)
          + '\n' + tr('Input verificati al: ', 'Inputs checked as of: ') + str(cutoff or missing)
          + tr(' · validi fino al: ', ' · valid through: ') + (expiry[0] if expiry else missing)
          + '\n' + tr('Generazione: ', 'Generation: ') + str(payload.get('generation_id') or missing), 6, height=55)
    chart = BarChart(); chart.type = 'bar'; chart.style = 13; chart.varyColors = False
    chart.add_data(Reference(ws, min_col=4, max_col=6, min_row=8, max_row=8), from_rows=True)
    chart.set_categories(Reference(ws, min_col=4, max_col=6, min_row=7, max_row=7))
    chart.title = tr('Valore per azione', 'Value per share') + ' (' + str(payload.get('currency') or missing) + ')'
    chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[0].graphicalProperties.line.solidFill = BLUE
    chart.legend = None; chart.height = 6.5; chart.width = 19
    chart.dLbls = DataLabelList(showVal=True, showCatName=True, showSerName=False,
                               showLegendKey=False, separator=': ', dLblPos='outEnd')
    ws.add_chart(chart, 'B29')
    _finish(ws, 40, 6)
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToHeight = 1
    wb.move_sheet(ws, -wb.index(ws)); wb.active = 0
