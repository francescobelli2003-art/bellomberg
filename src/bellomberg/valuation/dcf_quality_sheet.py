"""DCF-Q1 evidence beside resolved inputs; generated snapshot, not a new DCF."""
from math import ceil, isfinite

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bellomberg.valuation.dcf_quality import MATERIAL_DRIVERS, _finite


def append_sector_quality_sheet(wb, payload):
    """Expose the same method/acquisition decision before any draft figures."""
    from bellomberg.valuation.dcf_quality import assess_valuation_usability
    bundle = payload.get("acquisition_snapshot") or {}
    decision = payload.get("valuation_decision") or {}
    usability = assess_valuation_usability(payload, as_of=bundle.get("case", {}).get("as_of"))
    ws = wb.create_sheet("Metodo e dati", 0)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    for column, width in (("A", 26), ("B", 27), ("C", 90)):
        ws.column_dimensions[column].width = width
    rows = [
        ["VALUTAZIONE SETTORIALE", payload.get("ticker"), payload.get("company")],
        ["Fair value", "UTILIZZABILE" if usability["usable"] else "NON UTILIZZABILE / BOZZA",
         "; ".join(usability["reasons"]) or "Controlli del contratto superati"],
        ["Metodo", decision.get("method_id"), decision.get("rationale")],
        ["Decisione / supporto", decision.get("decision_status"), decision.get("support_status")],
        ["Requisiti", decision.get("requirements_status"), ", ".join(decision.get("missing_fields") or [])],
        ["Snapshot / generazione", payload.get("snapshot_id"), payload.get("generation_id")],
        ["Nota", "Dati e ipotesi distinti", "I calcoli nei fogli successivi non sono FV validati se questa scheda dichiara BOZZA."],
        ["FONTE / CAMPO", "STATO / DATA", "DETTAGLIO"],
    ]
    for source, value in bundle.get("case", {}).get("sources", {}).items():
        rows.append([source, str(value.get("status")) + " / " + str(value.get("as_of") or "n.d."),
                     value.get("message") or value.get("source_id")])
    for task in payload.get("acquisition_tasks") or []:
        rows.append([task.get("field"), task.get("status"), task.get("reason")])
    for row_number, values in enumerate(rows, 1):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_number, col, str(value) if value is not None else "n.d.")
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Calibri", size=11, bold=row_number in (1, 2, 8))
        ws.row_dimensions[row_number].height = max(32, 16 * ceil(max(len(str(v)) for v in values) / 90))
    wb.active = 0
    return ws


def append_quality_sheet(wb, report):
    """Append the documentary report; current scenario edits visibly invalidate it.

    Evidence is always written as text, including strings starting with '='.
    Numerical bridge values are explicitly a generation-time snapshot.
    """
    if report.get('record_adapter') is True:
        return _append_documented_quality(wb, report)
    if report.get("method_id") == "managed_care_distributable_equity":
        return _append_managed_care_quality(wb, report)
    ws = wb.create_sheet("Qualita e revisioni")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C8"
    widths = {"A": 13, "B": 23, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 23}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    def text(row, col, value, bold=False, color="1A1A1A"):
        cell = ws.cell(row, col)
        cell.value = str(value) if value is not None else "n.d."
        cell.data_type = "s"
        cell.font = Font(name="Calibri", size=11, bold=bold, color=color)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        return cell

    def line(row, value, heading=False):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        text(row, 1, value, bold=heading, color="FFFFFF" if heading else "1A1A1A")
        if heading:
            ws.cell(row, 1).fill = PatternFill("solid", fgColor="0B2545")
        # Width-based wrapping with extra space for explicit newlines.
        lines = sum(max(1, ceil(len(s) / 112)) for s in str(value).split("\n"))
        ws.row_dimensions[row].height = max(26, 17 * lines + 8)

    snap = report.get("snapshot") or {}
    bridge = report.get("revision_bridge") or {}
    line(1, f"{snap.get('ticker') or 'DCF'} | ASSUNZIONI E REVISIONI", True)
    line(2, f"{report.get('status', 'n.d.')} | Data analisi: {snap.get('as_of') or 'n.d.'}")
    line(3, report.get("note") or "La completezza documentale non certifica il fair value.")
    line(4, "Snapshot alla generazione: modifiche ai driver richiedono rivalutazione delle prove e del ponte.")
    line(5, "PONTE DELLE REVISIONI", True)
    line(6, bridge.get("note") or "Ponte n.d.: nessuna attribuzione disponibile.")
    row = 7
    if bridge.get("status") == "CALCOLATO":
        headers = ["Driver", "FV prima", "FV dopo", "Effetto FV"]
        for col, label in enumerate(headers, 1):
            text(row, col, label, True)
        row += 1
        for step in bridge.get("steps", []):
            text(row, 1, step["driver"])
            for col, key in enumerate(("previous", "current", "delta"), 2):
                ws.cell(row, col, step[key]).number_format = "#,##0.00"
            row += 1
        line(row, f"Valuta {bridge.get('currency')} | delta {bridge.get('delta')} | "
             f"residuo {bridge.get('residual')} | valori fotografati alla generazione")
    else:
        line(row, "EFFETTO SUL FAIR VALUE: n.d. — " + str(bridge.get("reason") or "dati insufficienti"))
    row += 1
    if report.get("previous_snapshot_note"):
        line(row, report["previous_snapshot_note"])
        row += 1
    for rev in report.get("revision_rows", []):
        operator = {"minimum": ">=", "maximum": "<=", "point": "="}.get(rev.get("value_type"), "n.d.")
        line(row, f"{rev.get('metric', 'n.d.')} | {rev.get('period', 'n.d.')} | {rev.get('basis', 'n.d.')} | "
             f"{rev.get('unit', 'n.d.')}\n"
             f"{rev.get('previous_date', 'n.d.')} {operator} {rev.get('previous_value', 'n.d.')} -> "
             f"{rev.get('source_date', 'n.d.')} {operator} {rev.get('current_value', 'n.d.')}\n"
             f"{rev.get('rationale', 'Motivazione n.d.')}\n"
             f"Prima: {rev.get('previous_source', 'n.d.')}\nDopo: {rev.get('source', 'n.d.')}\n"
             f"Buchi: {'; '.join(rev.get('issues') or []) or 'nessuno nel contratto'}")
        row += 1
    # Global errors (periods, unexpected records, unsupported controls) must be
    # visible too; driver-prefixed duplicates already appear next to each input.
    prefixes = tuple(f"{r.get('scenario')}/{r.get('driver')}: " for r in report.get("rows", []))
    for issue in report.get("issues", []):
        if not str(issue).startswith(prefixes):
            line(row, "BUCO: " + str(issue))
            row += 1
    for name, rationale in (report.get("scenario_rationale") or {}).items():
        line(row, f"SCENARIO {name}: {rationale}")
        row += 1
    controls = report.get("model_controls") or {}
    if controls:
        line(row, "CONTROLLI DEL MODELLO (snapshot)", True)
        row += 1
        for key, value in controls.items():
            line(row, f"{key}: {value}")
            row += 1
    line(row, "DRIVER RISOLTI E PROVE DELLE ASSUNZIONI", True)
    row += 1
    headers = ["Scenario", "Driver"] + (snap.get("forecast_years") or ["Y1", "Y2", "Y3", "Y4", "Y5"]) + ["Confronto col foglio"]
    for col, label in enumerate(headers, 1):
        text(row, col, label, True)
    row += 1
    first_driver_row = row
    for item in report.get("rows", []):
        scenario, driver = item.get("scenario"), item.get("driver")
        text(row, 1, scenario, True)
        text(row, 2, driver)
        values = item.get("values")
        values = values if isinstance(values, list) else [values]
        for col, value in enumerate(values[:5], 3):
            if type(value) in (int, float) and isfinite(value):
                ws.cell(row, col, value).number_format = "0.00%" if scenario != "model" or driver in ("wacc", "terminal_g") else "#,##0.00"
            else:
                text(row, col, "n.d.")
        sheet = "Scenario " + str(scenario).capitalize()
        if driver in MATERIAL_DRIVERS and sheet in wb.sheetnames and len(values) == 5:
            source_row = 4 + MATERIAL_DRIVERS.index(driver)
            conditions = [f"'{sheet}'!{get_column_letter(6+i)}{source_row}={get_column_letter(3+i)}{row}" for i in range(5)]
            ws.cell(row, 8, '=IF(AND(' + ','.join(conditions) + '),"INVARIATO","MODIFICATO")')
        elif scenario == "model" and type(values[0]) in (int, float) and isfinite(values[0]):
            reference = {"wacc": "'DCF'!B4", "terminal_g": "'DCF'!B5", "ronic": "'DCF'!E6",
                         "net_debt": "'DCF'!B6", "shares_m": "'DCF'!B7",
                         "last_revenue": "'Scenario Base'!E18", "nwc0": "'Scenario Base'!E38"}.get(driver)
            if driver == "shares_m" and "DCF" in wb.sheetnames:
                diluted = wb["DCF"]["B8"].value
                if type(diluted) in (int, float) and isfinite(diluted) and diluted:
                    reference = "'DCF'!B8"
            if reference and reference.split("!")[0].strip("'") in wb.sheetnames:
                ws.cell(row, 8, f'=IF({reference}=C{row},"INVARIATO","MODIFICATO")')
            else:
                text(row, 8, "snapshot / n.d.")
        else:
            text(row, 8, "snapshot / n.d.")
        ws.row_dimensions[row].height = 32
        row += 1
        evidence = item.get("evidence") or {}
        description = (f"{evidence.get('kind', 'Fonte n.d.')} | metrica: {evidence.get('metric', 'n.d.')} | "
            f"base: {evidence.get('basis', 'n.d.')}\n"
            f"Fonte: {evidence.get('source', item.get('provenance') or 'n.d.')}\n"
            f"Data: {evidence.get('source_date', 'n.d.')} | valida fino: {evidence.get('valid_until', 'n.d.')} | "
            f"riferimento: {evidence.get('source_locator', 'n.d.')}\n"
            f"Motivazione: {evidence.get('rationale', 'n.d.')}\n"
            f"Buchi: {'; '.join(item.get('issues') or []) or 'nessuno nel contratto'}")
        line(row, description)
        row += 1
    ws["A4"] = (f'=IF(COUNTIF(H{first_driver_row}:H{row},"MODIFICATO")>0,'
                '"MODIFICATO: rivalidare prove e ponte prima di usare il fair value",'
                '"Snapshot alla generazione; fonti, ancore e ponte richiedono nuova verifica a ogni revisione")')
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 38
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:H{row - 1}"
    return ws


def _append_documented_quality(wb, report):
    """Flatten structured source contracts without Excel's silent cell truncation."""
    rows=[]
    def leaves(value,path):
        if isinstance(value,dict) and value:
            for key,child in value.items(): yield from leaves(child,path+'.'+str(key))
        elif isinstance(value,list) and value:
            for index,child in enumerate(value): yield from leaves(child,path+'['+str(index)+']')
        else: yield path,value
    for item in report.get('rows',[]):
        value=item.get('values')
        structured=isinstance(value,dict) or isinstance(value,list) and any(isinstance(v,(dict,list)) for v in value)
        for path,leaf in leaves(value,item.get('driver','')) if structured else [(item.get('driver'),value)]:
            rows.append({**item,'driver':path,'values':leaf})
    prepared={**report,'rows':rows}
    return _append_managed_care_quality(wb, prepared)


def _append_managed_care_quality(wb, report):
    """Dated engine snapshot with any explicit horizon; never truncate to five years."""
    import json
    ws = wb.create_sheet("Qualita e revisioni")
    next_row=1
    def append(values):
        nonlocal next_row
        # 30k leaves room below Excel's 32,767 UTF-16-unit limit, including
        # supplementary characters (split by UTF-16 units, never mid-character).
        chunks=[]
        for value in values:
            parts=[];current=[];units=0
            if isinstance(value,str):
                for char in value:
                    width=2 if ord(char)>0xffff else 1
                    if units+width>30000:parts.append(''.join(current));current=[];units=0
                    current.append(char);units+=width
                parts.append(''.join(current))
            else:parts=[value]
            chunks.append(parts)
        base_row=next_row
        ws.append([parts[0] for parts in chunks])
        next_row+=1
        for col,parts in enumerate(chunks,1):
            for part,content in enumerate(parts[1:],2):
                ws.append(['TEXT CONTINUATION',f'row={base_row}; column={col}; part={part}/{len(parts)}',content])
                next_row+=1
        return base_row
    years = (report.get('snapshot') or {}).get('forecast_years') or []
    append(['ASSUNZIONI E REVISIONI', report.get('status'), report.get('note')])
    append(['Snapshot alla generazione; variazioni agli input richiedono un nuovo calcolo e nuovi controlli.'])
    for issue in report.get('issues', []):
        append(['BUCO', str(issue)])
    for name, rationale in (report.get('scenario_rationale') or {}).items():
        append(['SCENARIO', name, rationale])
    append(['Scenario', 'Driver'] + years + ['Unita / periodo / entita'])
    for item in report.get('rows', []):
        values = item.get('values')
        values = values if isinstance(values, list) else [values]
        evidence = item.get('evidence') or {}
        def cell_value(v):
            if type(v) in (int,float):
                if not _finite(v) or isinstance(v,int) and abs(v)>999999999999999:
                    return 'SOURCE AS TEXT (outside Excel numeric precision): '+str(v)
                return v
            return v if isinstance(v,str) else json.dumps(v,ensure_ascii=False)
        data = [cell_value(v) for v in values]
        data_row=append([item.get('scenario'), item.get('driver')] + data + [None] * max(0, len(years) - len(data)) + [
            ' | '.join(str(evidence.get(k) or 'n.d.') for k in ('unit', 'period', 'entity'))])
        for col in range(3, 3 + len(data)):
            ws.cell(data_row, col).number_format = '0.00%' if evidence.get('unit') == 'ratio' else '#,##0.0000'
        append(['Fonte', evidence.get('source'), evidence.get('source_date'), evidence.get('valid_until'),
                   evidence.get('basis'), evidence.get('rationale')])
    for revision in report.get('revision_rows', []):
        append(['Revisione', json.dumps(revision, ensure_ascii=False)])
    for column in range(1, max(len(years) + 4, ws.max_column + 1)):
        ws.column_dimensions[get_column_letter(column)].width = 25 if column > 2 else 35
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[row[0].row].height = 42
    ws.sheet_view.showGridLines = False
    return ws
