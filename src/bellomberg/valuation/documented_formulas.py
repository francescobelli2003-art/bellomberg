"""Linked FCFF schedules from documented inputs, with an immutable engine baseline."""
from datetime import date
from math import isfinite

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.workbook.properties import CalcProperties

from bellomberg.core.language import text as tr
from .documented_presentation import _sheet, _put, _line, _header, _finish, _label, col, AMOUNT, PRICE, PERCENT


def _formula(ws, row, column, expression, fmt=AMOUNT):
    cell = _put(ws, row, column, '=' + expression, fmt)
    cell.data_type = 'f'
    cell.font = Font(name='Arial', size=10, color='17324D')
    cell.alignment = Alignment(horizontal='right', vertical='center')


def operating_formula_ready(payload):
    """Only expose a live model when its independently computed baseline exists."""
    def finite(value):
        return isinstance(value, (int, float)) and not isinstance(value, bool) and isfinite(value)
    if not payload.get('valuation_usability', {}).get('usable'):
        return False
    evidence = payload.get('analytical_quality', {}).get('rows', [])
    values = {(r['scenario'], r['driver']): r['values'] for r in evidence}
    calendar = values.get(('model', 'calendar'), {})
    n = len(calendar.get('periods', []))
    if not n or calendar.get('discount_convention') not in ('annual_end', 'ACT/365F'):
        return False
    for scenario in ('bear', 'base', 'bull'):
        data = payload.get('calculation_details', {}).get('scenarios', {}).get(scenario, {})
        if not finite(data.get('fair_value_per_share')):
            return False
        for key in ('revenue', 'ebitda', 'ebit', 'ufcf', 'research_amortization'):
            series = data.get('rows', {}).get(key, [])
            if len(series) != n or not all(finite(v) for v in series):
                return False
        if (scenario, 'revenue_build') not in values:
            return False
    return True


def _input_rows(value, path=()):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from _input_rows(child, path + (key,))
    elif isinstance(value, list):
        if all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in value):
            yield path, value
        else:
            for index, child in enumerate(value):
                yield from _input_rows(child, path + (index,))
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        yield path, value


def _inputs(wb, evidence, years):
    ws = _sheet(wb, 'Model Inputs', tr('Input del modello e provenienza', 'Model inputs and provenance'), len(years) + 3)
    _line(ws, 4, tr('Blu: input modificabili. Le modifiche sono simulazioni locali, non nuove ipotesi approvate.',
                    'Blue: editable inputs. Changes are local simulations, not newly approved assumptions.'), len(years) + 3, height=40)
    refs, row = {}, 7
    for scenario in ('model', 'bear', 'base', 'bull'):
        _header(ws, row, [scenario, tr('Unita', 'Units'), *years])
        row += 1
        for record in (r for r in evidence if r['scenario'] == scenario):
            for path, values in _input_rows(record['values']):
                label = '.'.join(map(str, (record['driver'], *path)))
                _put(ws, row, 2, _label(record['driver']) if not path else label)
                unit = record.get('evidence', {}).get('unit')
                ratio = unit == 'ratio' or path[-1:] in (('growth',), ('utilization',))
                _put(ws, row, 3, '%' if ratio else unit)
                vector = isinstance(values, list)
                fmt = PERCENT if ratio else PRICE
                for index, value in enumerate(values if vector else [values]):
                    c = _put(ws, row, index + 4, value, fmt)
                    c.font = Font(name='Arial', size=10, color='0000FF')
                    c.comment = Comment(str(record.get('evidence', {})), 'Source')
                    refs[(scenario, record['driver'], *path, *((index,) if vector else ()))] = f"'Model Inputs'!{c.coordinate}"
                row += 1
        row += 2
    _finish(ws, row, len(years) + 3)
    # The readable assumptions table owns its input cells; the detailed table links to it.
    drivers = ('revenue_growth', 'gross_margin', 'rnd_pct', 'sga_pct', 'capdev_pct', 'da_tan_pct', 'tax_rate', 'capex_pct', 'nwc_pct')
    assumptions = wb['Assumptions']
    for scenario, first in (('base', 8), ('bear', 20), ('bull', 32)):
        for row, driver in enumerate(drivers, first):
            for i in range(len(years)):
                key = (scenario, driver, i)
                original = refs[key].split('!')[1]
                refs[key] = f"'Assumptions'!{col(i + 4)}{row}"
                _formula(ws, ws[original].row, ws[original].column, refs[key], PERCENT)
                assumptions.cell(row, i + 4).font = Font(name='Arial', size=10, color='0000FF')
    for row, driver in enumerate(('wacc', 'terminal_growth', 'terminal_ronic', 'net_debt', 'equity_adjustments', 'shares', 'opening_nwc'), 44):
        for c, scenario in enumerate(('bear', 'base', 'bull'), 4):
            key = ('model' if driver in ('shares', 'opening_nwc') else scenario, driver)
            if key[0] == 'model' and c != 4:
                _formula(assumptions, row, c, f'D{row}', PRICE)
                continue
            original = refs[key].split('!')[1]
            refs[key] = f"'Assumptions'!{col(c)}{row}"
            _formula(ws, ws[original].row, ws[original].column, refs[key], PERCENT if row < 47 else PRICE if driver == 'shares' else AMOUNT)
            assumptions.cell(row, c).font = Font(name='Arial', size=10, color='0000FF')
    _line(assumptions, 4, tr('Blu: ipotesi modificabili. Le modifiche non sono approvate dal modello nell’app.',
                            'Blue: editable assumptions. Changes are not approved by the application model.'), len(years) + 3, height=36)
    return refs


def _revenue_build(wb, values, refs, years):
    ws = _sheet(wb, 'Revenue Build', tr('Costruzione dei ricavi', 'Revenue build'), len(years) + 3)
    _line(ws, 4, tr('Ricavi da driver operativi. La crescita del gruppo e una formula, non un CAGR aggiunto.',
                    'Revenue from operating drivers. Group growth is derived, not an added CAGR.'), len(years) + 3, height=38)
    row, segment_row, bases = 7, 7, {}
    for scenario, assumption_row in (('base', 8), ('bear', 20), ('bull', 32)):
        build = values[(scenario, 'revenue_build')]
        _header(ws, row, [scenario, tr('Base', 'Opening') if build['basis'] == 'segment_guidance' else tr('Unita', 'Units'), *years])
        first = row + 1
        if build['basis'] == 'segment_guidance':
            for j, segment in enumerate(build['segments']):
                row += 1
                display = segment_row + 1 + j
                _put(ws, row, 2, segment['label'])
                base = f"'Segments'!C{display}"
                if scenario == 'base':
                    bases[segment['segment_id']] = base
                else:
                    _formula(wb['Segments'], display, 3, bases[segment['segment_id']])
                _formula(ws, row, 3, base)
                address = refs[(scenario, 'revenue_build', 'segments', j, 'base')].split('!')[1]
                _formula(wb['Model Inputs'], wb['Model Inputs'][address].row, wb['Model Inputs'][address].column, base)
                refs[(scenario, 'revenue_build', 'segments', j, 'base')] = bases[segment['segment_id']]
                for i in range(len(years)):
                    growth = f"'Segments'!{col(i + 4)}{display}"
                    wb['Segments'].cell(display, i + 4).font = Font(name='Arial', size=10, color='0000FF')
                    key = (scenario, 'revenue_build', 'segments', j, 'growth', i)
                    address = refs[key].split('!')[1]
                    _formula(wb['Model Inputs'], wb['Model Inputs'][address].row, wb['Model Inputs'][address].column, growth, PERCENT)
                    refs[key] = growth
                    opening = base if i == 0 else f'{col(i + 3)}{row}'
                    _formula(ws, row, i + 4, f'{opening}*(1+{growth})')
            segment_row += len(build['segments']) + 3
            row += 1
            for i in range(len(years)):
                c = col(i + 4)
                _formula(ws, row, i + 4, f'SUM({c}{first}:{c}{row - 1})')
        else:
            for driver in ('volume', 'unit_price', 'utilization', 'other_revenue'):
                row += 1
                _put(ws, row, 2, driver)
                for i in range(len(years)):
                    _formula(ws, row, i + 4, refs[(scenario, 'revenue_build', driver, i)], PERCENT if driver == 'utilization' else PRICE)
            row += 1
            for i in range(len(years)):
                c = col(i + 4)
                _formula(ws, row, i + 4, f'{c}{first}*{c}{first + 1}*{c}{first + 2}+{c}{first + 3}')
        _put(ws, row, 2, tr('Ricavi totali', 'Total revenue'), bold=True)
        for i in range(len(years)):
            previous = refs[('model', 'historical_revenue')] if i == 0 else f"'Revenue Build'!{col(i + 3)}{row}"
            _formula(wb['Assumptions'], assumption_row, i + 4, f"'Revenue Build'!{col(i + 4)}{row}/{previous}-1", PERCENT)
        row += 3
    _finish(ws, row, len(years) + 3)


def apply_operating_formulas(wb, payload):
    evidence = payload['analytical_quality']['rows']
    values = {(r['scenario'], r['driver']): r['values'] for r in evidence}
    calendar = values[('model', 'calendar')]
    years = [p['end'] for p in calendar['periods']]
    n, end = len(years), len(years) + 3
    refs = _inputs(wb, evidence, years)
    _revenue_build(wb, values, refs, years)
    checks_sheet = _sheet(wb, 'Input Checks', tr('Vincoli numerici degli input', 'Numeric input constraints'), 5)
    _header(checks_sheet, 7, [tr('Scenario', 'Scenario'), tr('Gruppo', 'Group'), tr('Esito', 'Result'), ''])
    check_row = 8

    def ref(scenario, driver, *path):
        return refs[(scenario, driver, *path)]

    def input_at(ws, row, key, fmt=AMOUNT):
        _formula(ws, row, 4, ref(*key), fmt)

    labels = {
        8: ('Ricavi', 'Revenue'), 9: ('Costo del venduto', 'Cost of sales'),
        10: ('Utile lordo', 'Gross profit'), 11: ('Ricerca e sviluppo', 'Research and development'),
        12: ('Costi commerciali e generali', 'Selling and general costs'),
        13: ('Sviluppo capitalizzato', 'Capitalized development'), 14: ('EBITDA', 'EBITDA'),
        15: ('Ammortamenti materiali', 'Tangible depreciation'), 16: ('Ammortamenti sviluppo', 'Development amortization'),
        17: ('EBIT', 'EBIT'), 18: ('Imposte operative', 'Operating taxes'), 19: ('NOPAT', 'NOPAT'),
        20: ('Ammortamenti totali', 'Total depreciation/amortization'), 21: ('Investimenti materiali', 'Capital expenditure'),
        22: ('Investimenti in sviluppo', 'Development investment'), 23: ('Capitale circolante', 'Working capital'),
        24: ('Variazione capitale circolante', 'Change in working capital'),
        25: ('Reinvestimento netto', 'Net reinvestment'), 26: ('FCFF', 'FCFF'),
        28: ('Margine EBITDA', 'EBITDA margin'), 29: ('Margine EBIT', 'EBIT margin'),
        30: ('FCFF / ricavi', 'FCFF / revenue'), 33: ('WACC', 'WACC'),
        34: ('Periodo di sconto', 'Discount period'), 35: ('Valore attuale FCFF', 'Present value of FCFF'),
        39: ('EBIT terminale normalizzato', 'Normalized terminal EBIT'), 40: ('Aliquota terminale', 'Terminal tax rate'),
        41: ('Crescita perpetua', 'Perpetual growth'), 42: ('RONIC terminale', 'Terminal RONIC'),
        43: ('Reinvestimento terminale / NOPAT', 'Terminal reinvestment / NOPAT'),
        44: ('FCFF primo anno stabile', 'First stable-year FCFF'), 45: ('Valore terminale', 'Terminal value'),
        46: ('Valore attuale terminale', 'Present value of terminal value'),
        48: ('Valore attuale flussi espliciti', 'PV of explicit cash flows'), 49: ('Enterprise value', 'Enterprise value'),
        50: ('Debito netto', 'Net debt'), 51: ('Rettifiche equity', 'Equity adjustments'),
        52: ('Equity value', 'Equity value'), 53: ('Azioni diluite (mln)', 'Diluted shares (million)'),
        54: ('Valore / azione - valuta bilancio', 'Value / share - financial currency'),
        55: ('Valore / azione - unita quotata', 'Value / share - quoted unit'),
        56: ('Peso terminale / enterprise value', 'Terminal share / enterprise value'),
    }
    money = tr('mln ', 'million ') + payload['financial_currency']
    for scenario in ('bear', 'base', 'bull'):
        del wb[scenario]
        ws = _sheet(wb, scenario, payload['ticker'] + ' — ' + scenario, end)
        _line(ws, 4, tr('Formule collegate. Ipotesi: Assumptions e Model Inputs. Simulazioni Excel non approvate.',
                        'Linked formulas. Assumptions: Assumptions and Model Inputs. Excel simulations are unapproved.'), end, height=38)
        _put(ws, 5, 2, tr('Validita numerica degli input', 'Numeric input validity'))
        first_check = check_row
        for group, expression in enumerate(_validity(refs, values, scenario, n), 1):
            _put(checks_sheet, check_row, 2, scenario); _put(checks_sheet, check_row, 3, group)
            _formula(checks_sheet, check_row, 4, f'IFERROR(IF({expression},"OK","KO"),"KO")')
            check_row += 1
        _formula(ws, 5, 4, f'IF(COUNTIF(\'Input Checks\'!D{first_check}:D{check_row-1},"KO")=0,"OK","KO")')
        _header(ws, 7, [tr('Conto economico e cassa', 'Operating statement and cash flow'), tr('Unita', 'Units'), *years])
        for row, caption in labels.items():
            _put(ws, row, 2, tr(*caption), bold=row in (8, 14, 17, 19, 26, 49, 52, 55))
            unit = '%' if row in (28, 29, 30, 33, 40, 41, 42, 43, 56) else money
            if row in (34, 53):
                unit = tr('anni', 'years') if row == 34 else tr('mln azioni', 'million shares')
            if row in (54, 55):
                unit = payload['financial_currency'] if row == 54 else payload['currency']
            _put(ws, row, 3, unit)
        for i in range(n):
            c, previous = col(i + 4), col(i + 3)
            initial = ref('model', 'historical_revenue') if i == 0 else previous + '8'
            nwc = ref('model', 'opening_nwc') if i == 0 else previous + '23'
            amortization = ref('model', 'capdev_amortization_years')
            amortization_terms = '+'.join(f'IF({i - j}<{amortization},{col(j + 4)}13/{amortization},0)' for j in range(i + 1))
            expressions = {
                8: f'{initial}*(1+{ref(scenario,"revenue_growth",i)})',
                9: f'{c}8*(1-{ref(scenario,"gross_margin",i)})', 10: f'{c}8-{c}9',
                11: f'{c}8*{ref(scenario,"rnd_pct",i)}', 12: f'{c}8*{ref(scenario,"sga_pct",i)}',
                13: f'{c}8*{ref(scenario,"capdev_pct",i)}', 14: f'{c}10-{c}11-{c}12+{c}13',
                15: f'{c}8*{ref(scenario,"da_tan_pct",i)}',
                16: f'{amortization_terms}+{ref(scenario,"opening_intangible_amortization",i)}',
                17: f'{c}14-{c}15-{c}16', 18: f'MAX({c}17,0)*{ref(scenario,"tax_rate",i)}',
                19: f'{c}17-{c}18', 20: f'{c}15+{c}16', 21: f'{c}8*{ref(scenario,"capex_pct",i)}',
                22: f'{c}13', 23: f'{c}8*{ref(scenario,"nwc_pct",i)}', 24: f'{c}23-{nwc}',
                25: f'{c}21+{c}22+{c}24-{c}20', 26: f'{c}19-{c}25',
                33: ref(scenario, 'wacc'), 35: f'{c}26/(1+{c}33)^{c}34',
            }
            for row, numerator in ((28, 14), (29, 17), (30, 26)):
                expressions[row] = f'IF({c}8=0,"n.d.",{c}{numerator}/{c}8)'
            for row, expression in expressions.items():
                _formula(ws, row, i + 4, f'IF($D$5="OK",{expression},"n.d.")', PERCENT if row in (28, 29, 30, 33) else AMOUNT)
            time = i + 1 if calendar['discount_convention'] == 'annual_end' else (date.fromisoformat(years[i]) - date.fromisoformat(calendar['valuation_date'])).days / 365
            _put(ws, 34, i + 4, time, '0.000')
        last = col(end)
        terminal = lambda key: ref(scenario, 'terminal_bridge', key)
        _line(ws, 37, tr('Regime stabile e ponte al valore azionario', 'Stable state and equity value bridge'), end, True)
        _formula(ws, 39, 4, f'IF($D$5="OK",{last}17+{last}16-{last}13+{terminal("cycle_adjustment")}-{terminal("expiring_product_loss")}+{terminal("replacement_product_income")}+{terminal("other_adjustment")},"n.d.")')
        for row, key in ((40, (scenario, 'tax_rate', n - 1)), (41, (scenario, 'terminal_growth')), (42, (scenario, 'terminal_ronic'))):
            input_at(ws, row, key, PERCENT)
        gate = '$D$5="OK"'
        for row, expression in {
            43: 'D41/D42', 44: 'D39*(1+D41)*(1-D40)*(1-D43)',
            45: f'D44/({ref(scenario,"wacc")}-D41)', 46: f'D45/(1+{ref(scenario,"wacc")})^{last}34',
            48: f'SUM(D35:{last}35)', 49: 'D48+D46', 52: 'D49-D50+D51', 54: 'D52/D53',
            55: f'D54*{ref("model","quotation","financial_to_quote_rate")}*{ref("model","quotation","quote_units_per_currency")}*{ref("model","quotation","shares_per_quote")}',
            56: 'IF(D49=0,"n.d.",D46/D49)',
        }.items():
            _formula(ws, row, 4, f'IF({gate},{expression},"n.d.")', PERCENT if row in (43, 56) else PRICE if row in (54, 55) else AMOUNT)
        for row, key in ((50, (scenario, 'net_debt')), (51, (scenario, 'equity_adjustments')), (53, ('model', 'shares'))):
            input_at(ws, row, key, PRICE if row == 53 else AMOUNT)
        _line(ws, 58, tr('I controlli Excel non sostituiscono la revisione economica e i controlli dell’app.',
                         'Excel checks do not replace economic review and application validation.'), end, height=32)
        _finish(ws, 59, end)
        for row in (8, 14, 17, 19, 26, 49, 52, 55):
            for column in range(4, end + 1):
                ws.cell(row, column).font = Font(name='Arial', size=10, color='17324D', bold=True)
        for key, expression in (('normalized_ebit', f"'{scenario}'!D39"),
                                ('capitalized_research_adjustment', f"'{scenario}'!{last}16-'{scenario}'!{last}13")):
            address = refs[(scenario, 'terminal_bridge', key)].split('!')[1]
            _formula(wb['Model Inputs'], wb['Model Inputs'][address].row, wb['Model Inputs'][address].column,
                     f'IF(\'{scenario}\'!D5="OK",{expression},"n.d.")')
    _finish(checks_sheet, check_row, 5)
    _formula(wb['Summary'], 18, 5, ref('model', 'quotation', 'price'), PRICE)
    if payload['valuation_usability']['usable']:
        for c, scenario in enumerate(('bear', 'base', 'bull'), 4):
            _formula(wb['Summary'], 8, c, f'IF(ISNUMBER(\'{scenario}\'!D55),ROUND(\'{scenario}\'!D55,2),"n.d.")', PRICE)
            _formula(wb['Summary'], 11, c, ref(scenario, 'wacc'), PERCENT)
            _formula(wb['Summary'], 12, c, ref(scenario, 'terminal_growth'), PERCENT)
            if isinstance(wb['Summary']['E19'].value, (int, float)):
                _formula(wb['Summary'], 9, c, f'IF(AND(ISNUMBER({col(c)}8),$E$19>0),{col(c)}8/$E$19-1,"n.d.")', PERCENT)
    if payload['valuation_usability']['usable']:
        _line(wb['Summary'], 5, tr('Formule interattive. Le modifiche richiedono una nuova validazione.',
                                 'Interactive formulas. Changes require fresh validation.'), 6, True, height=36)
    _line(wb['Summary'], 23, tr('Valore alla data del modello. Excel ricalcola le formule all apertura. Gli edit sono simulazioni locali; la quotazione osservata resta datata.',
                              'Value at the model date. Excel recalculates formulas on opening. Edits are local simulations; the observed quote remains dated.'), 6, height=52)
    if 'Segments' in wb:
        _line(wb['Segments'], 4, tr('Basi storiche in milioni della valuta di bilancio. Blu: crescita modificabile; gli edit richiedono nuova revisione.',
                                  'Historical bases in millions of financial currency. Blue: editable growth; edits require fresh review.'), end, height=38)
        build = values[('base', 'revenue_build')]
        for row in range(8, 8 + len(build['segments'])):
            wb['Segments'].cell(row, 3).font = Font(name='Arial', size=10, color='0000FF')
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    _terminal_bridge(wb, refs, n, payload['financial_currency'])
    _sensitivity(wb, refs, n)
    _checks(wb, payload, n)
    wb._model_link = {'refs': refs, 'raw': {s: f"'{s}'!D54" for s in ('bear','base','bull')},
        'shares': {s: refs[('model','shares')] for s in ('bear','base','bull')},
        'guards': {s: f"'{s}'!D5" for s in ('bear','base','bull')},
        'calls': {s: [f"MAX(0,-'{s}'!{col(i+4)}26)" for i in range(n)] + [f"MAX(0,{refs[(s,'net_debt')]})"] for s in ('bear','base','bull')}}
    ordered = ['Summary', 'Analysis', 'Revenue Build', 'base', 'bear', 'bull', 'Assumptions',
               'Segments', 'Model Inputs', 'Terminal Bridge', 'Sensitivity', 'Input Checks', 'Model Checks',
               'Valuation', 'Metodo e dati', 'Qualita e revisioni']
    for index, name in enumerate(name for name in ordered if name in wb):
        wb.move_sheet(wb[name], index - wb.index(wb[name]))


def _validity(refs, values, scenario, n):
    def ref(driver, *path, model=False):
        return refs[('model' if model else scenario, driver, *path)]
    checks = []
    for key, address in refs.items():
        if key[0] not in ('model', scenario) or key[1:] in (
                ('terminal_bridge', 'normalized_ebit'), ('terminal_bridge', 'capitalized_research_adjustment')):
            continue
        checks.append(f'ISNUMBER({address})')
    life = ref('capdev_amortization_years', model=True)
    growth, wacc, ronic = (ref(k) for k in ('terminal_growth', 'wacc', 'terminal_ronic'))
    checks += [f'{life}>0', f'{life}=INT({life})', f'{growth}>-1',
               f'{wacc}>MAX(0,{growth})', f'{ronic}>MAX(0,{growth})',
               f'{ref("shares",model=True)}>0', f'{ref("historical_revenue",model=True)}>0']
    for i in range(n):
        checks += [f'{ref("revenue_growth",i)}>-1', f'{ref("tax_rate",i)}>=0',
                   f'{ref("tax_rate",i)}<=1', f'{ref("capdev_pct",i)}<={ref("rnd_pct",i)}']
        checks += [f'{ref(k,i)}>=0' for k in ('rnd_pct','sga_pct','capdev_pct','da_tan_pct','capex_pct','opening_intangible_amortization')]
    for key in ('financial_to_quote_rate', 'quote_units_per_currency', 'shares_per_quote', 'price'):
        checks.append(f'{ref("quotation",key,model=True)}>0')
    checks += [f'{ref("terminal_bridge",key)}>=0' for key in ('expiring_product_loss', 'replacement_product_income')]
    build = values[(scenario, 'revenue_build')]
    if build['basis'] == 'segment_guidance':
        bases = []
        for j, segment in enumerate(build['segments']):
            base = ref('revenue_build', 'segments', j, 'base'); bases.append(base)
            checks.append(f'{base}{">0" if segment["role"] == "segment" else "<>0"}')
            bound = '>=-1' if segment['role'] == 'consolidation' else '>-1'
            checks += [f'{ref("revenue_build","segments",j,"growth",i)}{bound}' for i in range(n)]
        historical = ref('historical_revenue', model=True)
        checks.append(f'ABS(({"+".join(bases)})-{historical})<=MAX(1E-8,ABS({historical})*1E-9)')
    else:
        for i in range(n):
            checks += [f'{ref("revenue_build",k,i)}>=0' for k in ('volume','unit_price','utilization','other_revenue')]
            checks.append(f'{ref("revenue_build","utilization",i)}<=1')
    # Separate helper cells respect both Excel's formula-length and argument limits.
    return ['AND('+','.join(checks[i:i+40])+')' for i in range(0,len(checks),40)]


def _terminal_bridge(wb, refs, n, currency):
    ws = _sheet(wb, 'Terminal Bridge', tr('Transizione al regime stabile', 'Transition to stable operations'), 6)
    _line(ws, 4, tr('Spiega il salto di cassa: risultato operativo normalizzato, crescita e reinvestimento. La sostenibilita economica richiede una tesi.',
                    'Explains the cash-flow change: normalized operating earnings, growth and reinvestment. Economic sustainability requires a thesis.'), 6, height=42)
    _header(ws, 7, [tr('Ponte economico', 'Economic bridge'), tr('Unita', 'Units'), 'bear', 'base', 'bull'])
    captions = {
        8: ('Crescita ultimo anno esplicito', 'Final explicit-year growth'), 9: ('Crescita stabile', 'Stable growth'),
        10: ('EBIT ultimo anno esplicito', 'Final explicit-year EBIT'), 11: ('Ricerca interamente spesata', 'Fully expensed research adjustment'),
        12: ('Rettifica ciclo', 'Cycle adjustment'), 13: ('Perdita prodotti in scadenza', 'Expiring product loss'),
        14: ('Utili prodotti sostitutivi', 'Replacement product income'), 15: ('Altre rettifiche', 'Other adjustments'),
        16: ('EBIT normalizzato', 'Normalized EBIT'), 18: ('NOPAT ultimo anno esplicito', 'Final explicit-year NOPAT'),
        19: ('NOPAT primo anno stabile', 'First stable-year NOPAT'), 20: ('Variazione NOPAT', 'NOPAT change'),
        22: ('Reinvestimento netto finale', 'Final net reinvestment'), 23: ('Reinvestimento netto stabile', 'Stable net reinvestment'),
        24: ('Minore / (maggiore) reinvestimento', 'Lower / (higher) reinvestment'), 26: ('FCFF ultimo anno esplicito', 'Final explicit-year FCFF'),
        27: ('Contributo del NOPAT', 'Contribution from NOPAT'), 28: ('Contributo del reinvestimento', 'Contribution from reinvestment'),
        29: ('FCFF primo anno stabile', 'First stable-year FCFF'), 31: ('Differenza dal DCF (attesa zero)', 'Difference from DCF (expected zero)'),
        33: ('Peso terminale / EV', 'Terminal share / EV'),
    }
    for row, label in captions.items():
        _put(ws, row, 2, tr(*label), bold=row in (16, 26, 29, 33))
        _put(ws, row, 3, '%' if row in (8, 9, 33) else tr('mln ', 'million ') + currency)
    last = col(n + 3)
    for column, scenario in enumerate(('bear', 'base', 'bull'), 4):
        c, s = col(column), f"'{scenario}'!"
        expressions = {8: refs[(scenario, 'revenue_growth', n-1)], 9: s+'D41', 10: s+last+'17',
            11: s+last+'16-'+s+last+'13', 16: f'SUM({c}10:{c}15)', 18: s+last+'19',
            19: f'{c}16*(1+{c}9)*(1-{s}D40)', 20: f'{c}19-{c}18', 22: s+last+'25',
            23: f'{c}19*{c}9/{s}D42', 24: f'{c}22-{c}23', 26: f'{c}18-{c}22',
            27: c+'20', 28: c+'24', 29: f'SUM({c}26:{c}28)', 31: f'{c}29-{s}D44', 33: s+'D56'}
        for row, key in ((12, 'cycle_adjustment'), (13, 'expiring_product_loss'),
                         (14, 'replacement_product_income'), (15, 'other_adjustment')):
            expressions[row] = ('-' if row == 13 else '') + refs[(scenario, 'terminal_bridge', key)]
        for row, expression in expressions.items():
            _formula(ws, row, column, f'IF({s}D5="OK",{expression},"n.d.")', PERCENT if row in (8, 9, 33) else AMOUNT)
    _finish(ws, 35, 6)
    ws.page_setup.orientation='portrait'
    ws.page_setup.fitToHeight=1


def _sensitivity(wb, refs, n):
    ws = _sheet(wb, 'Sensitivity', tr('Sensibilita e attese implicite', 'Sensitivity and implied expectations'), 7)
    _line(ws, 4, tr('Scenario base: shock esplorativi, non probabilita. Tutte le altre ipotesi restano ferme.',
                    'Base case: exploratory shocks, not probabilities. All other assumptions held fixed.'), 7, height=36)
    _header(ws, 7, ['WACC / g', '', '', '', '', ''])
    for c, shock in enumerate((-.005, -.0025, 0., .0025, .005), 3):
        _formula(ws, 7, c, f"'base'!D41+({shock})", '0.00%')
        ws.cell(7, c).font = Font(name='Arial', size=10, color='FFFFFF', bold=True)
    last = col(n + 3)
    for row, shock in enumerate((-.01, -.005, 0., .005, .01), 8):
        _formula(ws, row, 2, f"'base'!D33+({shock})", '0.00%')
        for c in range(3, 8):
            rate, growth = f'$B{row}', f'{col(c)}$7'
            pv = '+'.join(f"'base'!{col(i+4)}26/(1+{rate})^'base'!{col(i+4)}34" for i in range(n))
            tv = f"'base'!D39*(1+{growth})*(1-'base'!D40)*(1-{growth}/'base'!D42)/({rate}-{growth})/(1+{rate})^'base'!{last}34"
            factor = '*'.join(refs[('model', 'quotation', k)] for k in ('financial_to_quote_rate', 'quote_units_per_currency', 'shares_per_quote'))
            fv = f"(({pv})+({tv})-'base'!D50+'base'!D51)/'base'!D53*({factor})"
            valid = f"AND('base'!D5=\"OK\",{rate}>MAX({growth},0),'base'!D42>MAX({growth},0),{growth}>-1,'base'!D53>0)"
            _formula(ws, row, c, f'IF({valid},{fv},"n.d.")', PRICE)
    _line(ws, 15, tr('EBIT stabile implicito nel prezzo osservato', 'Stable EBIT implied by the observed price'), 7, True)
    _line(ws, 17, tr('Inferenza condizionata: stessi flussi espliciti, WACC, g, RONIC, imposte e bridge del base. Non e consensus.',
                     'Conditional inference: same base explicit cash flows, WACC, g, RONIC, taxes and bridge. This is not consensus.'), 7, height=44)
    _put(ws, 19, 2, tr('EBIT normalizzato del modello', 'Model normalized EBIT'))
    _formula(ws, 19, 4, "'base'!D39")
    _put(ws, 20, 2, tr('EBIT stabile implicito', 'Implied stable EBIT'))
    numerator = f"('Summary'!E19/({factor})*'base'!D53+'base'!D50-'base'!D51-'base'!D48)*(1+'base'!D33)^'base'!{last}34*('base'!D33-'base'!D41)"
    denominator = "(1+'base'!D41)*(1-'base'!D40)*(1-'base'!D43)"
    condition = f"AND(ISNUMBER('Summary'!E19),'Summary'!E19>0,'base'!D33>'base'!D41,'base'!D42>'base'!D41,({denominator})>0)"
    _formula(ws, 20, 4, f'IF(\'base\'!D5="OK",IF({condition},({numerator})/({denominator}),"n.d."),"n.d.")')
    _finish(ws, 22, 7)


def _checks(wb, payload, n):
    ws = _sheet(wb, 'Model Checks', tr('Controlli e confronto con il motore', 'Checks and engine comparison'), 6)
    _line(ws, 4, tr('Scostamenti dal motore dopo un edit indicano una simulazione da validare. Non una nuova approvazione.',
                    'Differences from the engine after edits indicate a simulation to validate, not a new approval.'), 6, height=38)
    _header(ws, 7, [tr('Controllo', 'Check'), tr('Scenario', 'Scenario'), tr('Motore originale', 'Original engine'), 'Excel', tr('Esito', 'Result')])
    row = 8
    for scenario in ('bear', 'base', 'bull'):
        data = payload['calculation_details']['scenarios'][scenario]
        baseline = [('Fair value', data['fair_value_per_share'], f"'{scenario}'!D54")]
        baseline += [(f'FCFF {i + 1}', value, f"'{scenario}'!{col(i + 4)}26") for i, value in enumerate(data['rows']['ufcf'])]
        for label, expected, cell in baseline:
            _put(ws, row, 2, label); _put(ws, row, 3, scenario)
            _put(ws, row, 4, expected, PRICE); _formula(ws, row, 5, cell, PRICE)
            _formula(ws, row, 6, f'IF(ISNUMBER(E{row}),IF(ABS(E{row}-D{row})<=MAX(1E-8,ABS(D{row})*1E-10),"OK","SIMULAZIONE"),"KO")')
            row += 1
    _line(ws, row + 2, tr('Il terminale elevato richiede una tesi sulla stabilizzazione; non dimostra da solo un errore.',
                          'High terminal dependence requires a stabilization thesis; it does not by itself prove an error.'), 6, height=36)
    _finish(ws, row + 4, 6)
