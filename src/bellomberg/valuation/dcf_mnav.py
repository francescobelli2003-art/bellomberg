"""
dcf_mnav.py - Motore canonico mNAV/NAV per i VEICOLI (DAT + CEF) — V5 Lotto 1
(design audit/18 approvato dal PM voce per voce 23/07/2026, decisioni D1-D4).

PRINCIPIO: su un veicolo il DCF resta VIETATO (dcf_engine rifiuta): questo
motore NON e' un DCF — e' la formalizzazione canonica della valutazione a
mNAV/NAV che il sistema gia' predica. Numeri SOLO dai tool ufficiali
(dat_metrics per le DAT, cef_nav per i fondi chiusi): il motore li mette in un
workbook a formule VIVE + mirror Python (parita' come banca/rab), con fonte
e vintage per OGNI input. Dato mancante o stale oltre soglia = RIFIUTO
DICHIARATO del canonico (regola 14/07), mai un modello con numeri vecchi.
Il tool live resta la fonte del giorno-per-giorno: il canonico e' la FOTO
della tesi (V5.4).

SEMANTICA FV (decisione PM D2, 23/07): un fair value su un veicolo NON e' mai
un target price. FV SOLO se l'analista dichiara `nav_target` (target di
premio/sconto sul NAV): senza target il canonico e' una scheda informativa
completa con FV n.d. DICHIARATO col motivo. Il target e' inteso:
- cef_nav (fondo chiuso): rapporto prezzo/NAV (es. 0,80 = sconto strutturale 20%);
- dat_hype: mNAV target su Adjusted NAV per azione FD;
- dat_bitcoin: target su mNAV **EV** (doctrine del tool: davanti alle
  ordinarie ci sono debt+preferred) -> FV = (target x BTC-NAV - debt - pref
  + cash) / azioni basic; a target 1,0 coincide col NAV equity per azione.
  NB: NON e' "NAV/azione x target" — con decine di mld $ di debt+pref davanti la
  moltiplicazione nuda sposterebbe anche la parte dei creditori (dichiarato
  nel foglio; e' il perfezionamento EV della D2 per il solo kind dat_bitcoin).
Banda di buon senso sul target DICHIARATA (WARN fuori TARGET_WARN_BAND,
hard bound TARGET_HARD_BOUNDS oltre cui il FV e' n.d. dichiarato):
mai BLOCK automatico su una view (D2).

STALENESS (soglie dichiarate): NAV settimanale del fondo chiuso -> STALE dal tool
(>14g, cef_nav.NAV_STALE_DAYS) = rifiuto; record della fonte bitcoin piu' vecchio di
RECORD_STALE_DAYS_BTC = rifiuto; per dat_hype gli input sono di bilancio con cadenza
dichiarata dal SITO (effective_date + next_update + lag note nel foglio
Fonti & Vintage): nessun rifiuto sull'eta' — e' la cadenza ufficiale della
fonte, il lag e' dichiarato riga per riga.

VALUTA: tutto nella valuta del NAV (USD per i tre kind). Per un fondo chiuso
quotato in GBp il prezzo usato da modello e sanity e' quello
GIA' convertito dal tool cef_nav (stesso tasso del payload del tool — mai
due misure dello stesso cambio, lezione cross-valuta 23/07), conversione
dichiarata in cella. payload_currency = valuta del NAV.
"""
from typing import Any, Dict, Optional

import bellomberg.storage.classificazione as cl
from bellomberg.valuation.dcf_bank import _safe  # riuso, mai riscrivere

try:
    from openpyxl.styles import Font
except ImportError:
    Font = None  # il builder workbook e' comunque guarded sull'import openpyxl

# perimetro D1: i KIND che il motore sa calcolare, per SOTTOSTANTE (estensione = codice del
# motore + fonte ufficiale nuova, decisione PM). La copertura per SIMBOLO viene dal NEGOZIO
# dei veicoli (05/09, lotto 2b): voci `tipo=dat` col `sottostante`; i fondi chiusi da
# cef_nav.CEF_SOURCES, a sua volta derivato dal negozio. Ieri era un dict cablato col book.
# Vista del contratto comune per compatibilita'; il routing legge il KIND dichiarato.
KIND_DA_SOTTOSTANTE = {s: k for k, s in cl.SOTTOSTANTI_KIND_DAT.items()}


def _dat_del_negozio(negozio=None):
    n = negozio if negozio is not None else cl.carica_veicoli()
    return {t: v for t, v in n["veicoli"].items() if v["tipo"] == "dat"}


def mnav_kinds_del_negozio(negozio=None) -> dict:
    """{TICKER: kind} dichiarato e coerente nel negozio; non dedotto dal sottostante."""
    return {t: v["kind"] for t, v in _dat_del_negozio(negozio).items()
            if v.get("kind") in cl.SOTTOSTANTI_KIND_DAT
            and v["sottostante"] == cl.SOTTOSTANTI_KIND_DAT[v["kind"]]}


def dat_senza_kind_del_negozio(negozio=None) -> dict:
    """{TICKER: motivo} delle DAT dichiarate che il motore NON sa calcolare: sottostante
    assente o non supportato. Dichiarate, non taciute."""
    supportati = ", ".join(sorted(KIND_DA_SOTTOSTANTE))
    out = {}
    for t, v in _dat_del_negozio(negozio).items():
        if not v["sottostante"]:
            out[t] = "sottostante assente nella voce del negozio (supportati: %s)" % supportati
        elif v["sottostante"] not in KIND_DA_SOTTOSTANTE:
            out[t] = ("sottostante %r non supportato dal motore mNAV (supportati: %s)"
                      % (v["sottostante"], supportati))
        elif v.get("kind") is None:
            out[t] = "kind assente nella voce del negozio: dichiara il motore DAT, non viene inferito"
        elif cl.SOTTOSTANTI_KIND_DAT.get(v["kind"]) != v["sottostante"]:
            out[t] = "kind incoerente con il sottostante dichiarato"
    return out


def __getattr__(name):
    # Compatibility view is acquired only by explicit legacy callers. Importing
    # the pure NAV arithmetic must not read a personal vehicle registry.
    if name=='MNAV_KINDS':
        return mnav_kinds_del_negozio()
    raise AttributeError(name)

RECORD_STALE_DAYS_BTC = 45      # ~5 settimane senza record nuovo = fonte da verificare
TARGET_HARD_BOUNDS = (0.10, 3.00)   # oltre: input rotto -> FV n.d. dichiarato
TARGET_WARN_BAND = (0.50, 1.50)     # fuori: WARN dichiarato, mai BLOCK (D2)
SENS_SHOCKS = (-0.40, -0.20, 0.0, 0.20, 0.40)          # sensitivity DAT sul sottostante
CEF_DISCOUNT_GRID = (0.60, 0.70, 0.80, 0.90, 1.00)     # griglia prezzo/NAV per i CEF

_SUBSECTOR = {"dat_bitcoin": "DAT Bitcoin (mNAV)",
              "dat_hype": "DAT HYPE (mNAV)",
              "cef_nav": "Fondo chiuso (NAV/sconto)"}
_METHOD = {"dat_bitcoin": "mNAV canonico DAT Bitcoin (dati strategy.com + prezzi live; "
                          "FV solo con nav_target analista su mNAV EV — D2)",
           "dat_hype": "mNAV canonico DAT HYPE (mirror computeDashboard hypestrat.xyz; "
                       "FV solo con nav_target analista su Adjusted NAV/FD — D2)",
           "cef_nav": "NAV/sconto canonico fondo chiuso (NAV ufficiale settimanale; "
                      "FV solo con nav_target analista prezzo/NAV — D2)"}


def _validate_target(nav_target, src: Dict[str, str], warnings: list):
    """Target dell'analista validato: None se assente o fuori hard bound (FV n.d.
    dichiarato col motivo); fuori banda di buon senso = WARN dichiarato, mai BLOCK."""
    if nav_target is None:
        src["nav_target"] = ("ASSENTE: FV n.d. DICHIARATO (D2: senza un target di "
                             "premio/sconto dell'analista il fair value non esce — "
                             "dire FV=NAV implicherebbe 'il premio chiude a 1,0')")
        return None, "nessun nav_target dell'analista (D2): FV n.d. dichiarato"
    t = _safe(nav_target)
    lo_h, hi_h = TARGET_HARD_BOUNDS
    if t is None or not lo_h <= t <= hi_h:
        motivo = (f"nav_target '{nav_target}' non numerico" if t is None else
                  f"nav_target {t:g} fuori hard bound {lo_h:g}-{hi_h:g}")
        src["nav_target"] = f"{motivo}: IGNORATO, FV n.d. DICHIARATO (input rotto, non una view)"
        warnings.append(f"{motivo}: FV n.d. dichiarato — correggere l'input.")
        return None, motivo + ": FV n.d. dichiarato"
    lo, hi = TARGET_WARN_BAND
    src["nav_target"] = "ANALISTA: target premio/sconto %.2f (fonte nella variant view)" % t
    if not lo <= t <= hi:
        warnings.append("nav_target %.2f FUORI dalla banda di buon senso %.1f-%.1f "
                        "(dichiarata): view forte da argomentare nella variant view — "
                        "WARN, mai BLOCK su una view (D2)." % (t, lo, hi))
    return t, None


def _parse_iso_date(s):
    from datetime import date
    try:
        return date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return None


def build_mnav_spec(ticker, info, nav_target=None, variant_view=None, today=None,
                    tool_payload=None, negozio=None) -> Dict[str, Any]:
    """Spec canonico mNAV. `tool_payload` iniettabile per i test OFFLINE; in
    produzione si chiama il tool ufficiale del ticker (get_dat_metrics /
    get_cef_nav — numeri SOLO dai tool). Errore/stale = {'error': ...}
    RIFIUTO dichiarato, mai un canonico con numeri vecchi o proxy.
    Il perimetro viene dal negozio dei veicoli (`negozio` iniettabile; default riletto):
    ogni rifiuto dice PERCHE' il titolo e' fuori, con un conteggio e mai un elenco."""
    from datetime import date
    from bellomberg.valuation.cef_nav import cef_sources_del_negozio, cef_senza_fonte_del_negozio
    t = str(ticker or "").upper().strip()
    _today = today or date.today()
    n = negozio if negozio is not None else cl.carica_veicoli()
    if n["origine"] in ("assente", "illeggibile"):
        return {"error": (f"{t}: negozio dei veicoli {n['origine'].upper()} ({n['motivo']}): il "
                          "perimetro mnav non e' determinabile finche' il file non c'e'/non e' "
                          "corretto — canonico RIFIUTATO dichiarato, non un proxy.")}
    kinds = mnav_kinds_del_negozio(n)
    cef = cef_sources_del_negozio(n)
    kind = kinds.get(t)
    if kind is None and t in cef:
        kind = "cef_nav"
    if kind is None:
        senza_kind = dat_senza_kind_del_negozio(n)
        if t in senza_kind:
            return {"error": (f"{t}: DAT dichiarata nel negozio dei veicoli ma fuori dal perimetro "
                              f"mnav: {senza_kind[t]} — estensione = codice del motore con fonte "
                              "ufficiale nuova (D1), non un proxy.")}
        senza_fonte = cef_senza_fonte_del_negozio(n)
        if t in senza_fonte:
            return {"error": (f"{t}: fondo chiuso dichiarato nel negozio dei veicoli ma fuori dal "
                              f"perimetro mnav: senza fonte NAV riconosciuta ({senza_fonte[t]}) — "
                              "serve nav_fonte = una chiave di cef_nav.FETCH_NAV e nav_valuta, "
                              "non un proxy.")}
        _nc = len(cef)
        return {"error": (f"{t}: fuori dal perimetro mnav: nel negozio dei veicoli non e' una DAT "
                          f"con sottostante supportato ne' un fondo chiuso con fonte NAV "
                          f"({len(kinds)} DAT e {_nc} fond{'o' if _nc == 1 else 'i'} "
                          f"chius{'o' if _nc == 1 else 'i'} copert{'o' if _nc == 1 else 'i'} oggi). "
                          "Estensione = voce nel negozio con fonte ufficiale (D1), non un proxy.")}

    if tool_payload is None:
        try:
            if kind in ("dat_bitcoin", "dat_hype"):
                from bellomberg.valuation.dat_metrics import get_dat_metrics
                tool_payload = get_dat_metrics(t, negozio=n)
            else:
                from bellomberg.valuation.cef_nav import get_cef_nav
                tool_payload = get_cef_nav(t, negozio=n)
        except Exception as e:
            return {"error": (f"{t}: fetch del tool ufficiale fallito "
                              f"({type(e).__name__}: {e}) — canonico mnav RIFIUTATO, "
                              "riprovare (mai un modello senza la fonte)")}
    if not isinstance(tool_payload, dict) or tool_payload.get("error"):
        return {"error": (f"{t}: la fonte ufficiale ha risposto con errore "
                          f"({(tool_payload or {}).get('error') or 'payload non valido'}) "
                          "— canonico mnav RIFIUTATO dichiarato, riprovare")}

    src: Dict[str, str] = {}
    warnings: list = []
    tgt, fv_note = _validate_target(nav_target, src, warnings)

    base = {"ticker": t, "engine": "mnav", "kind": kind,
            "profile_key": kind, "subsector": _SUBSECTOR[kind],
            "company_name": (info or {}).get("longName") or (info or {}).get("shortName")
                            or tool_payload.get("name") or t,
            "nav_target": tgt, "fv_note": fv_note,
            "variant_view": variant_view, "_sources": src, "warnings_spec": warnings}

    if kind == "dat_bitcoin":
        return _spec_mstr(base, tool_payload, _today)
    if kind == "dat_hype":
        return _spec_purr(base, tool_payload)
    return _spec_cef(base, tool_payload, info)


def _spec_mstr(base, p, today) -> Dict[str, Any]:
    """dat_bitcoin: record ufficiale della fonte + prezzi live gia' nel derived_mnav del
    tool (mai rifetchati qui: stessa misura del tool, mai due misure)."""
    src, warnings = base["_sources"], base["warnings_spec"]
    latest = p.get("latest") or {}
    dm = p.get("derived_mnav") or {}
    if p.get("mnav_error") or not dm:
        return {"error": (f"{base['ticker']}: derived_mnav non disponibile dal tool "
                          f"({p.get('mnav_error') or 'campo assente'}) — canonico "
                          "RIFIUTATO dichiarato (senza prezzi live niente mNAV, no proxy)")}
    inp = dm.get("inputs") or {}
    btc = _safe(latest.get("btc_holdings"))
    sh = _safe(latest.get("basic_shares_outstanding"))
    px_btc = _safe(inp.get("price_btc"))
    px_mstr = _safe(inp.get("price_mstr"))
    if not btc or not sh or not px_btc or not px_mstr or min(btc, sh, px_btc, px_mstr) <= 0:
        return {"error": (f"{base['ticker']}: input essenziali mancanti o non positivi nel payload "
                          "del tool (btc_holdings/azioni/prezzi) — canonico RIFIUTATO")}
    as_of = inp.get("record_as_of") or latest.get("as_of_date")
    d_asof = _parse_iso_date(as_of)
    if d_asof is not None:
        age = (today - d_asof).days
        if age > RECORD_STALE_DAYS_BTC:
            return {"error": ("%s: record della fonte del %s (%d giorni fa, soglia "
                              "%dg): STALE — canonico RIFIUTATO dichiarato, verificare "
                              "la fonte (mai un modello su un tesoro vecchio)"
                              % (base["ticker"], as_of, age, RECORD_STALE_DAYS_BTC))}
        src["record"] = ("%s record btcTrackerData[0], as-of %s (%d giorni)"
                         % (p.get("fonte") or "fonte ufficiale", as_of, age))
    else:
        src["record"] = "%s record btcTrackerData[0]" % (p.get("fonte") or "fonte ufficiale")
        warnings.append("%s: data as-of del record non parsabile (%r): staleness NON "
                        "verificabile — dichiarato." % (base["ticker"], as_of))
    src["prezzi"] = ("yfinance fast_info (dal tool, computed_at %s, cache %s min) — "
                     "stessi prezzi del payload live, mai rifetchati"
                     % (inp.get("computed_at_utc") or "n.d.", inp.get("cache_ttl_min") or "n.d."))
    # campo assente != zero (review 21/07 del tool): mnav_ev/NAV equity n.d. dichiarati
    ev_missing = [k for k in ("debt", "pref", "cash") if latest.get(k) is None]
    if ev_missing:
        warnings.append("%s: campi %s assenti nel record (assente != zero): mNAV EV, "
                        "NAV equity per azione e FV n.d. DICHIARATI." % (base["ticker"], ev_missing))
    base.update({
        "currency": "USD", "price": px_mstr,
        "price_note": "prezzo azione USD dal tool (yfinance fast_info)",
        "btc_holdings": btc, "shares_basic": sh, "px_btc": px_btc,
        "debt": _safe(latest.get("debt")), "pref": _safe(latest.get("pref")),
        "cash": _safe(latest.get("cash")), "ev_missing": ev_missing,
        "record_as_of": as_of,
        "tool_mnav_equity": _safe(dm.get("mnav_equity_basic")),
        "tool_mnav_ev": _safe(dm.get("mnav_ev")),
        "leggimi": dm.get("leggimi"),
        "vintage": {"record_as_of": as_of,
                    "computed_at_utc": inp.get("computed_at_utc"),
                    "cache_ttl_min": inp.get("cache_ttl_min")},
    })
    if ev_missing:
        src["ev_fields"] = ("debt/pref/cash: campi %s ASSENTI nel record %s "
                            "(assente != zero, layout cambiato?)" % (ev_missing, p.get("fonte") or "della fonte"))
    else:
        src["ev_fields"] = "debt/pref/cash dal record %s (stesso as-of del record)" % (p.get("fonte") or "della fonte")
    return base


def _spec_purr(base, p) -> Dict[str, Any]:
    """dat_hype: input di bilancio dal JSON IR ufficiale + prezzi live del tool; il
    foglio rende in celle la computeDashboard del sito (stessa formula, C7..C39)."""
    src, warnings = base["_sources"], base["warnings_spec"]
    inp = dict(p.get("dat_fundamentals_musd") or {})
    if not _safe(inp.get("basicShares")) or not _safe(inp.get("hypeHeld")):
        return {"error": (f"{base['ticker']}: input incompleti dal JSON IR (basicShares/hypeHeld) "
                          "— canonico RIFIUTATO dichiarato")}
    px_purr = _safe((p.get("purr_quote") or {}).get("last"))
    px_hype = _safe((p.get("hype_quote") or {}).get("last"))
    if not px_purr or not px_hype or px_purr <= 0 or px_hype <= 0:
        return {"error": ("%s: prezzo live mancante (azione: %s | HYPE: %s) — canonico "
                          "RIFIUTATO dichiarato, riprovare (mai mNAV su prezzi vecchi)"
                          % (base["ticker"], (p.get("purr_quote") or {}).get("fonte"),
                             (p.get("hype_quote") or {}).get("fonte")))}
    # la formula del sito tratta i campi assenti come 0 (mirror ESATTO): si DICHIARA
    # quali erano assenti, mai in silenzio (regola 14/07 sul canonico)
    zero_keys = [k for k in ("bookNAV", "cashFromOps", "cashFromFin", "treasuryDeploy",
                             "reportedDigital", "taxRate", "taxBasis", "reportedDTL")
                 if _safe(inp.get(k)) is None]
    if zero_keys:
        warnings.append("%s: campi %s assenti nel JSON IR -> 0 nella formula del sito "
                        "(mirror esatto della computeDashboard, dichiarato)." % (base["ticker"], zero_keys))
    vin = p.get("vintage") or {}
    src["input_bilancio"] = ("%s, effective %s, next update %s — valori in $M dichiarati"
                             % (p.get("fonte") or "JSON IR ufficiale della fonte",
                                vin.get("effective_date") or "n.d.",
                                vin.get("next_update") or "n.d."))
    src["prezzi"] = ("azione: %s | HYPE: %s (dal tool, stessi prezzi del payload live)"
                     % ((p.get("purr_quote") or {}).get("fonte"),
                        (p.get("hype_quote") or {}).get("fonte")))
    warrants = []
    for w in (p.get("warrants") or []):
        s_, a_ = _safe(w.get("strike")), _safe(w.get("amount"))
        if s_ and a_ and s_ > 0 and a_ > 0:
            warrants.append({"strike": s_, "amount": a_})
    src["warrants"] = ("%d tranche dal JSON IR, diluizione a TREASURY METHOD solo se "
                       "ITM (formula del sito)" % len(warrants)) if warrants else \
                      "nessuna tranche warrant nel JSON IR (FD = basic, dichiarato)"
    base.update({
        "currency": "USD", "price": px_purr,
        "price_note": "prezzo azione USD dal tool (fonte dichiarata in _sources)",
        "inputs_musd": {k: (_safe(inp.get(k)) or 0.0)
                        for k in ("bookNAV", "cashFromOps", "cashFromFin", "treasuryDeploy",
                                  "reportedDigital", "hypeHeld", "taxRate", "taxBasis",
                                  "reportedDTL", "basicShares")},
        "inputs_zero_keys": zero_keys,
        "warrants": warrants, "px_hype": px_hype,
        "tool_mnav": _safe((p.get("derived") or {}).get("mnav")),
        "lag_note": vin.get("lag_note"),
        "vintage": {"effective_date": vin.get("effective_date"),
                    "next_update": vin.get("next_update"),
                    "lag_note": vin.get("lag_note")},
    })
    return base


def _spec_cef(base, p, info) -> Dict[str, Any]:
    """CEF: NAV ufficiale + sconto del tool cef_nav. Il prezzo usato e'
    quello GIA' convertito nella valuta del NAV dal tool (stesso tasso del payload,
    mai una seconda misura del cambio)."""
    src = base["_sources"]
    if p.get("nav_staleness"):
        return {"error": ("%s: NAV STALE dalla fonte ufficiale (%s) — canonico "
                          "RIFIUTATO dichiarato (pubblicazione settimanale ferma: "
                          "verificare la fonte, mai un modello su un NAV vecchio)"
                          % (base["ticker"], p["nav_staleness"]))}
    nav = _safe(p.get("nav_per_share_usd"))
    if not nav or nav <= 0:
        return {"error": (f"{base['ticker']}: NAV per azione mancante/non positivo dal "
                          "tool cef_nav — canonico RIFIUTATO dichiarato")}
    px_nav_ccy = _safe(p.get("market_price_in_nav_ccy"))
    px_quote = _safe(p.get("market_price"))
    quote_ccy = p.get("market_price_currency")
    if px_nav_ccy is None or px_nav_ccy <= 0:
        return {"error": ("%s: prezzo nella valuta del NAV non disponibile dal tool "
                          "(%s) — sconto e sanity non calcolabili: canonico RIFIUTATO "
                          "dichiarato, riprovare"
                          % (base["ticker"], p.get("discount") or "conversione n.d."))}
    if p.get("nav_as_of") is None:
        base["warnings_spec"].append("%s: data as-of del NAV non trovata dalla fonte: "
                                     "trattare il NAV con cautela (dichiarato dal tool)."
                                     % base["ticker"])
    src["nav"] = ("%s, as-of %s (%s giorni)"
                  % (p.get("nav_source") or "fonte NAV ufficiale",
                     p.get("nav_as_of") or "n.d.",
                     p.get("nav_age_days") if p.get("nav_age_days") is not None else "n.d."))
    fx_rate = None
    if quote_ccy and quote_ccy != "USD":
        src["prezzo"] = ("quotazione %.2f %s -> %.2f USD col tasso del tool cef_nav "
                         "(%s) — stessa misura del payload live, mai un secondo fetch"
                         % (px_quote or 0.0, quote_ccy, px_nav_ccy,
                            p.get("fx_conversion") or "conversione dichiarata dal tool"))
        if quote_ccy in ("GBp", "GBX") and px_quote:
            # tasso IMPLICITO nei numeri del tool (px_usd / (GBp/100)): il foglio
            # riproduce ESATTAMENTE il prezzo convertito del payload
            fx_rate = round(px_nav_ccy / (px_quote / 100.0), 6)
    else:
        src["prezzo"] = "quotazione gia' in USD (yfinance via tool cef_nav)"
    base.update({
        "currency": "USD", "price": px_nav_ccy,
        "price_note": ("prezzo convertito in USD dal tool cef_nav (quotazione %s %s)"
                       % (px_quote if px_quote is not None else "n.d.", quote_ccy or "?")),
        "price_quote": px_quote, "price_quote_ccy": quote_ccy, "fx_rate": fx_rate,
        "nav_per_share": nav, "nav_as_of": p.get("nav_as_of"),
        "nav_age_days": p.get("nav_age_days"),
        "tool_discount_pct": _safe(p.get("discount_to_nav_pct")),
        "returns": {"MTD": p.get("mtd_return_pct"), "QTD": p.get("qtd_return_pct"),
                    "YTD": p.get("ytd_return_pct")},
        "discount_note": p.get("discount_note"),
        "vintage": {"nav_as_of": p.get("nav_as_of"), "nav_age_days": p.get("nav_age_days"),
                    "fx_conversion": p.get("fx_conversion")},
    })
    return base


# ---------- MIRROR PYTHON (payload = fonte di verita'; i fogli lo riproducono) ----------
def _common_equity_nav(gross_assets,cash,debt,preferred,other_claims,equity_adjustments):
    return gross_assets+cash-debt-preferred-other_claims+equity_adjustments


def compute_mnav_values(spec) -> Dict[str, Any]:
    """Misure mNAV/NAV + FV (solo con target, D2) + warnings dichiarati. Il mirror
    ricalcola dagli INPUT del tool e si confronta con i derivati del tool stesso:
    uno scarto oltre tolleranza = warning dichiarato (cintura di parita')."""
    if spec.get('documented_inputs'):
        c=spec['components']; cash=c['cash']; debt=c['debt']; shares=spec['shares']
        other=sum(c[k] for k in ('other_liabilities','accrued_fees','distributions_payable','tax'))
        common=_common_equity_nav(c['gross_assets'],cash,debt,c['preferred'],other,c['equity_adjustments'])
        fv=common*spec['nav_target'] if spec['target_basis']=='equity_nav' else _common_equity_nav(
            c['gross_assets']*spec['nav_target'],cash,debt,c['preferred'],other,c['equity_adjustments'])
        return {'common_equity_nav':common,'nav_per_share':common/shares,
                'fair_value_nav':fv/shares,'shares':shares,'components':c,'target_basis':spec['target_basis']}
    warnings = list(spec.get("warnings_spec") or [])
    tgt = spec.get("nav_target")
    out: Dict[str, Any] = {"warnings": warnings}
    fv = None

    if spec["kind"] == "dat_bitcoin":
        btc_nav = spec["btc_holdings"] * spec["px_btc"]
        mktcap = spec["shares_basic"] * spec["price"]
        out["btc_nav_usd"] = round(btc_nav)
        out["mnav_equity"] = round(mktcap / btc_nav, 3)
        if spec.get("ev_missing"):
            out["mnav_ev"] = None
            out["nav_per_share"] = None
        else:
            d_, p_, c_ = spec["debt"], spec["pref"], spec["cash"]
            out["mnav_ev"] = round((mktcap + d_ + p_ - c_) / btc_nav, 3)
            out["nav_per_share"] = round(_common_equity_nav(btc_nav,c_,d_,p_,0.,0.) / spec["shares_basic"], 2)
            if tgt is not None:
                fv = (tgt * btc_nav - d_ - p_ + c_) / spec["shares_basic"]
        if tgt is not None and spec.get("ev_missing"):
            out["fv_note"] = ("FV n.d.: target su mNAV EV ma campi %s assenti nel "
                              "record (dichiarato)" % spec["ev_missing"])
        # parita' mirror<->tool (stessi input, stessa formula: scarto = bug)
        for mine, tool, lbl in ((out["mnav_equity"], spec.get("tool_mnav_equity"), "mNAV equity"),
                                (out.get("mnav_ev"), spec.get("tool_mnav_ev"), "mNAV EV")):
            if mine is not None and tool is not None and abs(mine - tool) > 0.005:
                warnings.append("PARITA' ROTTA %s: mirror %.3f vs tool %.3f — bug da "
                                "investigare, numeri NON affidabili." % (lbl, mine, tool))

    elif spec["kind"] == "dat_hype":
        i = spec["inputs_musd"]
        px, hy = spec["price"], spec["px_hype"]
        hype_val = i["hypeHeld"] * hy                                    # C7
        dil = 0.0
        for w in spec["warrants"]:
            if w["strike"] < px:                                          # treasury method
                dil += (px - w["strike"]) * w["amount"] / px
        fd = i["basicShares"] + dil                                       # C11
        tax = i["taxRate"] / 100.0
        dtl = -tax * (hype_val - i["taxBasis"])                           # C38
        dtl_change = i["reportedDTL"] - tax * (hype_val - i["taxBasis"])  # C21
        adj_nav = (i["bookNAV"] - i["cashFromOps"] + i["cashFromFin"]
                   - i["treasuryDeploy"] - i["reportedDigital"] + hype_val + dtl_change)  # C22
        out.update({"hype_value_musd": round(hype_val, 1),
                    "dtl_musd_signed": round(dtl, 1),
                    "adjusted_nav_musd": round(adj_nav, 1),
                    "fd_shares_m": round(fd, 3),
                    "warrants_itm": sum(1 for w in spec["warrants"] if w["strike"] < px)})
        if fd <= 0 or adj_nav <= 0:
            out["nav_per_share"] = None
            out["mnav"] = None
            out["mnav_dtl_addback"] = None
            warnings.append("%s: Adjusted NAV o azioni FD non positivi (%.1f $M / %.3f M): "
                            "mNAV e FV n.d. DICHIARATI — leggere gli input, non mediare."
                            % (spec["ticker"], adj_nav, fd))
            if tgt is not None:
                out["fv_note"] = "FV n.d.: Adjusted NAV/FD non positivo (dichiarato)"
        else:
            anav_ps = adj_nav / fd                                        # C23
            anav_dtl = (adj_nav - dtl) / fd                               # C24
            out["nav_per_share"] = round(anav_ps, 4)
            out["nav_per_share_dtl_addback"] = round(anav_dtl, 4)
            out["mnav"] = round(px / anav_ps, 3)
            out["mnav_dtl_addback"] = round(px / anav_dtl, 3) if anav_dtl > 0 else None
            if tgt is not None:
                fv = anav_ps * tgt
            tool_m = spec.get("tool_mnav")
            if tool_m is not None and abs(out["mnav"] - tool_m) > 0.005:
                warnings.append("PARITA' ROTTA mNAV %s: mirror %.3f vs tool %.3f — bug "
                                "da investigare, numeri NON affidabili."
                                % (spec["ticker"], out["mnav"], tool_m))
            elif tool_m is None:
                # review V5 B3: mai un cross-check saltato in silenzio
                warnings.append("parita' mirror<->tool NON verificabile (derived assente "
                                "nel payload del tool): mirror non confermato, dichiarato.")

    else:  # cef_nav
        nav, px = spec["nav_per_share"], spec["price"]
        out["nav_per_share"] = round(nav, 2)
        out["discount_to_nav_pct"] = round((px / nav - 1.0) * 100, 2)
        tool_d = spec.get("tool_discount_pct")
        if tool_d is not None and abs(out["discount_to_nav_pct"] - tool_d) > 0.15:
            warnings.append("PARITA' ROTTA sconto NAV: mirror %.2f%% vs tool %.1f%% — bug "
                            "da investigare." % (out["discount_to_nav_pct"], tool_d))
        if tgt is not None:
            fv = nav * tgt

    if fv is not None:
        if fv <= 0:
            warnings.append("FV %.2f <= 0 dal target dichiarato: NON e' un numero — "
                            "FV n.d. dichiarato, rivedere target/input." % fv)
            out["fv_note"] = "FV n.d.: risultato non positivo col target dichiarato"
        else:
            out["fair_value_nav"] = round(fv, 2)
    if out.get("fair_value_nav") is None and "fv_note" not in out and spec.get("fv_note"):
        out["fv_note"] = spec["fv_note"]
    return out


def _mnav_sanity(spec, mv) -> Dict[str, Any]:
    """Sanity DEDICATA (D2): il FV di un veicolo e' 'upside alla convergenza del
    premio/sconto dichiarato', NON un target price — la distanza dal prezzo e'
    informazione, non un modello rotto: WARN solo su target fuori banda o parita'
    rotta, MAI BLOCK automatico su una view."""
    fv = mv.get("fair_value_nav")
    parity = [w for w in (mv.get("warnings") or []) if w.startswith("PARITA' ROTTA")]
    if fv is None:
        # review V5 M2: severity None, MAI "OK" — un badge OK verde su una scheda
        # senza fair value sarebbe un giudizio sanity che non esiste (F17 rende n.d.)
        return {"status": "n/d", "ratio": None, "severity": "WARN" if parity else None,
                "exclude_from_action_table": False,
                "headline": parity[0] if parity else (mv.get("fv_note")
                            or "FV n.d. dichiarato (nessun nav_target — D2)")}
    px = spec.get("price")
    ratio = round(fv / px, 2) if px else None
    sev, hl = "OK", None
    tgt = spec.get("nav_target")
    lo, hi = TARGET_WARN_BAND
    if tgt is not None and not lo <= tgt <= hi:
        sev = "WARN"
        hl = ("nav_target %.2f fuori banda di buon senso %.1f-%.1f (dichiarata): "
              "view forte da argomentare — mai BLOCK su una view (D2)." % (tgt, lo, hi))
    if parity:
        sev, hl = "WARN", parity[0]
    return {"status": "ok", "ratio": ratio,
            "upside_pct": round((ratio - 1) * 100, 1) if ratio else None,
            "reading": ("upside alla CONVERGENZA del premio/sconto dichiarato "
                        "dall'analista — NON un target price (D2)"),
            "severity": sev, "exclude_from_action_table": False, "headline": hl}


# ============================================================
# WORKBOOK (formule VIVE + valori; parita' col mirror misurata al collaudo)
# ============================================================
def _hdr(ws, spec, sub):
    from bellomberg.valuation.dcf_bank import GREYTX, L, T
    from datetime import datetime as _dt
    T(ws, "A1", f"{spec['company_name']} ({spec['ticker']}) - {sub}")
    L(ws, "A2", f"Motore mNAV V5 (audit/18) - {_dt.now():%d/%m/%Y %H:%M} - valuta {spec['currency']}"
                + (" - valori in $M / M azioni" if spec["kind"] == "dat_hype" else ""),
      italic=True, color=GREYTX)
    L(ws, "A4", "PERCHE' NON C'E' UN DCF: e' un VEICOLO — il valore e' il NAV del "
                "sottostante e il premio/sconto che il mercato gli riconosce, non flussi "
                "operativi da scontare (DCF vietato dal motore). Il canonico e' la FOTO "
                "della tesi: il tool live resta la fonte del giorno-per-giorno.", italic=True)
    L(ws, "A6", "TESI DELL'ANALISTA (variant view):", bold=True)
    L(ws, "A7", str(spec.get("variant_view") or "(nessuna variant view fornita)"))


def _fv_rows(ws, spec, mv, r, fv_formula, px_cell):
    """Blocco FAIR VALUE (D2): formula viva solo col target; n.d. dichiarato altrimenti.
    Ritorna la riga successiva."""
    from bellomberg.valuation.dcf_bank import GOLD, GREYTX, L, N
    L(ws, f"A{r}", "FAIR VALUE (semantica D2)", bold=True); r += 1
    if mv.get("fair_value_nav") is not None:
        L(ws, f"A{r}", "Target premio/sconto (ANALISTA)")
        N(ws, f"B{r}", spec["nav_target"], fmt="0.00")
        L(ws, f"C{r}", spec["_sources"].get("nav_target") or "ANALISTA", italic=True, color=GREYTX)
        t_cell = f"B{r}"; r += 1
        L(ws, f"A{r}", "FAIR VALUE per azione", bold=True)
        ws[f"B{r}"] = fv_formula(t_cell)
        ws[f"B{r}"].number_format = "#,##0.00"
        if Font is not None:
            ws[f"B{r}"].font = Font(bold=True, color=GOLD, size=11)
        fv_cell = f"B{r}"; r += 1
        L(ws, f"A{r}", "Upside alla convergenza")
        ws[f"B{r}"] = f"={fv_cell}/{px_cell}-1"
        ws[f"B{r}"].number_format = "0.0%"
        r += 1
        L(ws, f"A{r}", "NB: upside alla CONVERGENZA del premio/sconto dichiarato "
                       "dall'analista — NON un target price.", italic=True, color=GREYTX)
        r += 1
    else:
        L(ws, f"A{r}", "FV n.d. DICHIARATO: " + str(mv.get("fv_note") or
          "nessun nav_target dell'analista (D2)"), color=GREYTX)
        r += 1
    r += 1
    warns = mv.get("warnings") or []
    if warns:
        L(ws, f"A{r}", "ATTENZIONI (dal motore, da riportare nel report):", bold=True); r += 1
        for w in warns:
            L(ws, f"A{r}", "! " + w); r += 1
    return r


def _sheet_thesis_mstr(wb, spec, mv):
    from bellomberg.valuation.dcf_bank import GOLD, GREYTX, H, L, N
    ws = wb.create_sheet("Thesis & Assumptions", 0)
    _hdr(ws, spec, "Canonico mNAV DAT Bitcoin")
    H(ws, "A9", "INPUT UFFICIALI"); H(ws, "B9", "VALORE"); H(ws, "C9", "FONTE")
    L(ws, "A10", "BTC posseduti"); N(ws, "B10", spec["btc_holdings"], fmt="#,##0")
    L(ws, "C10", spec["_sources"]["record"], italic=True, color=GREYTX)
    L(ws, "A11", "Prezzo BTC (USD)"); N(ws, "B11", spec["px_btc"], fmt="#,##0")
    L(ws, "C11", spec["_sources"]["prezzi"], italic=True, color=GREYTX)
    L(ws, "A12", "BTC-NAV (USD)"); ws["B12"] = "=B10*B11"; ws["B12"].number_format = "#,##0"
    L(ws, "A13", "Azioni basic"); N(ws, "B13", spec["shares_basic"], fmt="#,##0")
    L(ws, "C13", "record della fonte (basic: convertibili/RSU esclusi, limite dichiarato)",
      italic=True, color=GREYTX)
    L(ws, "A14", "Prezzo azione (USD)"); N(ws, "B14", spec["price"], fmt="#,##0.00")
    L(ws, "A15", "Market cap basic (USD)"); ws["B15"] = "=B13*B14"; ws["B15"].number_format = "#,##0"
    if spec.get("ev_missing"):
        L(ws, "A16", "Debt / Preferred / Cassa: " + spec["_sources"]["ev_fields"], color=GREYTX)
    else:
        L(ws, "A16", "Debito (USD)"); N(ws, "B16", spec["debt"], fmt="#,##0")
        L(ws, "C16", spec["_sources"]["ev_fields"], italic=True, color=GREYTX)
        L(ws, "A17", "Preferred (USD)"); N(ws, "B17", spec["pref"], fmt="#,##0")
        L(ws, "A18", "Cassa (USD)"); N(ws, "B18", spec["cash"], fmt="#,##0")
    H(ws, "A20", "MISURE mNAV")
    L(ws, "A21", "mNAV equity basic = mktcap / BTC-NAV")
    ws["B21"] = "=B15/B12"; ws["B21"].number_format = "0.000"
    if spec.get("ev_missing"):
        L(ws, "A22", "mNAV EV: n.d. DICHIARATO (campi %s assenti — assente != zero)"
          % spec["ev_missing"], color=GREYTX)
        L(ws, "A23", "NAV equity per azione: n.d. DICHIARATO (stessi campi)", color=GREYTX)
    else:
        L(ws, "A22", "mNAV EV = (mktcap + debt + pref - cassa) / BTC-NAV", bold=True)
        ws["B22"] = "=(B15+B16+B17-B18)/B12"; ws["B22"].number_format = "0.000"
        if Font is not None:
            ws["B22"].font = Font(bold=True, color=GOLD, size=11)
        L(ws, "A23", "NAV equity per azione = (BTC-NAV - debt - pref + cassa) / azioni")
        ws["B23"] = "=(B12-B16-B17+B18)/B13"; ws["B23"].number_format = "#,##0.00"
    L(ws, "A25", "LEGGIMI (dal tool): " + str(spec.get("leggimi") or
      "per giudizi di sconto/premio usare il mNAV EV: davanti alle ordinarie ci sono "
      "debt+preferred."), italic=True, color=GREYTX)
    r = _fv_rows(ws, spec, mv, 27,
                 lambda t: f"=({t}*B12-B16-B17+B18)/B13", "B14")
    if mv.get("fair_value_nav") is not None:
        L(ws, f"A{r}", "NB target: inteso su mNAV EV (doctrine del tool) — FV = "
                       "(target x BTC-NAV - debt - pref + cassa) / azioni; a target 1,0 "
                       "coincide col NAV equity per azione.", italic=True, color=GREYTX)
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 80


def _sheet_thesis_purr(wb, spec, mv):
    from bellomberg.valuation.dcf_bank import GOLD, GREYTX, H, L, N
    ws = wb.create_sheet("Thesis & Assumptions", 0)
    _hdr(ws, spec, "Canonico mNAV DAT HYPE (mirror computeDashboard)")
    i = spec["inputs_musd"]
    H(ws, "A9", "INPUT UFFICIALI ($M salvo nota)"); H(ws, "B9", "VALORE"); H(ws, "C9", "FONTE")
    rows = (("Book NAV", "bookNAV"), ("Cash from ops (da sottrarre)", "cashFromOps"),
            ("Cash from financing", "cashFromFin"), ("Treasury deploy (da sottrarre)", "treasuryDeploy"),
            ("Digital assets a bilancio (da sottrarre)", "reportedDigital"),
            ("HYPE posseduti (M token)", "hypeHeld"), ("Tax rate (%)", "taxRate"),
            ("Tax basis", "taxBasis"), ("DTL a bilancio", "reportedDTL"),
            ("Azioni basic (M)", "basicShares"))
    r = 10
    cell = {}
    L(ws, "C10", spec["_sources"]["input_bilancio"], italic=True, color=GREYTX)
    for lbl, k in rows:
        L(ws, f"A{r}", lbl); N(ws, f"B{r}", i[k], fmt="#,##0.00")
        if k in (spec.get("inputs_zero_keys") or []):
            # la nota ASSENTE vince sulla nota fonte generale (riga 10 = bookNAV)
            L(ws, f"C{r}", "ASSENTE nel JSON IR -> 0 nella formula del sito (dichiarato)",
              italic=True, color=GREYTX)
        cell[k] = f"B{r}"; r += 1
    L(ws, f"A{r}", "Prezzo HYPE live (USD)"); N(ws, f"B{r}", spec["px_hype"], fmt="#,##0.00")
    L(ws, f"C{r}", spec["_sources"]["prezzi"], italic=True, color=GREYTX)
    hy = f"B{r}"; r += 1
    L(ws, f"A{r}", "Prezzo azione live (USD)"); N(ws, f"B{r}", spec["price"], fmt="#,##0.00")
    px = f"B{r}"; r += 2
    H(ws, f"A{r}", "WARRANT (treasury method, formula del sito)")
    L(ws, f"D{r}", spec["_sources"]["warrants"], italic=True, color=GREYTX)
    r += 1
    dil_cells = []
    if spec["warrants"]:
        H(ws, f"A{r}", "tranche"); H(ws, f"B{r}", "strike"); H(ws, f"C{r}", "amount (M)")
        H(ws, f"D{r}", "diluizione (M)")
        r += 1
        for n_, w in enumerate(spec["warrants"], start=1):
            L(ws, f"A{r}", f"warrant {n_}")
            N(ws, f"B{r}", w["strike"], fmt="#,##0.00")
            N(ws, f"C{r}", w["amount"], fmt="#,##0.00")
            ws[f"D{r}"] = f"=IF({px}>B{r},({px}-B{r})*C{r}/{px},0)"
            ws[f"D{r}"].number_format = "#,##0.000"
            dil_cells.append(f"D{r}")
            r += 1
    r += 1
    H(ws, f"A{r}", "FORMULA UFFICIALE (celle C7..C39 del sito)")
    r += 1
    L(ws, f"A{r}", "Valore HYPE corrente ($M)")
    ws[f"B{r}"] = f"={cell['hypeHeld']}*{hy}"; ws[f"B{r}"].number_format = "#,##0.0"
    hv = f"B{r}"; r += 1
    L(ws, f"A{r}", "Azioni FD (M) = basic + diluizione ITM")
    ws[f"B{r}"] = (f"={cell['basicShares']}+" + "+".join(dil_cells)) if dil_cells \
        else f"={cell['basicShares']}"
    ws[f"B{r}"].number_format = "#,##0.000"
    fd = f"B{r}"; r += 1
    L(ws, f"A{r}", "DTL teorica ($M, negativa se plusvalenza)")
    ws[f"B{r}"] = f"=-{cell['taxRate']}/100*({hv}-{cell['taxBasis']})"
    ws[f"B{r}"].number_format = "#,##0.0"
    dtl = f"B{r}"; r += 1
    L(ws, f"A{r}", "Delta DTL ($M) = DTL a bilancio - teorica")
    ws[f"B{r}"] = f"={cell['reportedDTL']}-{cell['taxRate']}/100*({hv}-{cell['taxBasis']})"
    ws[f"B{r}"].number_format = "#,##0.0"
    dch = f"B{r}"; r += 1
    L(ws, f"A{r}", "Adjusted NAV ($M)", bold=True)
    ws[f"B{r}"] = (f"={cell['bookNAV']}-{cell['cashFromOps']}+{cell['cashFromFin']}"
                   f"-{cell['treasuryDeploy']}-{cell['reportedDigital']}+{hv}+{dch}")
    ws[f"B{r}"].number_format = "#,##0.0"
    anav = f"B{r}"; r += 1
    L(ws, f"A{r}", "Adjusted NAV per azione FD (USD)")
    ws[f"B{r}"] = f"={anav}/{fd}"; ws[f"B{r}"].number_format = "#,##0.0000"
    anav_ps = f"B{r}"; r += 1
    L(ws, f"A{r}", "Adjusted NAV/FD con DTL add-back (riga C24 del sito)")
    ws[f"B{r}"] = f"=({anav}-{dtl})/{fd}"; ws[f"B{r}"].number_format = "#,##0.0000"
    anav_dtl = f"B{r}"; r += 1
    if mv.get("mnav") is None:
        # review V5 M3: caso degenerato (ANAV o FD non positivi) — il payload
        # dichiara n.d. e il foglio NON deve mostrare un rapporto negativo vivo
        # come headline: stessa resa del ramo dat_bitcoin ev_missing
        L(ws, f"A{r}", "mNAV: n.d. DICHIARATO (Adjusted NAV o azioni FD non positivi "
                       "— v. ATTENZIONI)", color=GREYTX); r += 1
        L(ws, f"A{r}", "mNAV con DTL add-back: n.d. DICHIARATO (stesso motivo)",
          color=GREYTX); r += 1
        r += 1
    else:
        L(ws, f"A{r}", "mNAV = prezzo azione / (Adjusted NAV / FD)", bold=True)
        ws[f"B{r}"] = f"={px}/{anav_ps}"; ws[f"B{r}"].number_format = "0.000"
        if Font is not None:
            ws[f"B{r}"].font = Font(bold=True, color=GOLD, size=11)
        r += 1
        L(ws, f"A{r}", "mNAV con DTL add-back")
        ws[f"B{r}"] = f"={px}/{anav_dtl}"; ws[f"B{r}"].number_format = "0.000"
        r += 2
    r = _fv_rows(ws, spec, mv, r, lambda t: f"={anav_ps}*{t}", px)
    ws.column_dimensions["A"].width = 52
    for c_ in "BCD":
        ws.column_dimensions[c_].width = 16
    ws.column_dimensions["C"].width = 70
    # riferimenti per il foglio Sensitivity (celle vive, mai valori copiati)
    spec["_cells"] = {"hy": hy, "px": px, "hv": hv, "fd": fd, "dtl": dtl,
                      "anav": anav, "anav_ps": anav_ps, "cell": cell}


def _sheet_thesis_cef(wb, spec, mv):
    from bellomberg.valuation.dcf_bank import GOLD, GREYTX, H, L, N
    ws = wb.create_sheet("Thesis & Assumptions", 0)
    _hdr(ws, spec, "Canonico NAV/sconto fondo chiuso")
    H(ws, "A9", "INPUT UFFICIALI"); H(ws, "B9", "VALORE"); H(ws, "C9", "FONTE")
    L(ws, "A10", "NAV per azione (USD)"); N(ws, "B10", spec["nav_per_share"], fmt="#,##0.00")
    L(ws, "C10", spec["_sources"]["nav"], italic=True, color=GREYTX)
    if spec.get("fx_rate"):
        L(ws, "A11", "Prezzo di mercato (GBp)"); N(ws, "B11", spec["price_quote"], fmt="#,##0.0")
        L(ws, "C11", spec["_sources"]["prezzo"], italic=True, color=GREYTX)
        L(ws, "A12", "Cambio GBP/USD (tasso del tool)"); N(ws, "B12", spec["fx_rate"], fmt="0.0000")
        L(ws, "A13", "Prezzo in USD = GBp/100 x cambio")
        ws["B13"] = "=B11/100*B12"; ws["B13"].number_format = "#,##0.00"
        px = "B13"; r0 = 14
    else:
        L(ws, "A11", "Prezzo di mercato (USD)"); N(ws, "B11", spec["price"], fmt="#,##0.00")
        L(ws, "C11", spec["_sources"]["prezzo"], italic=True, color=GREYTX)
        px = "B11"; r0 = 12
    L(ws, f"A{r0}", "Sconto/premio sul NAV = prezzo/NAV - 1", bold=True)
    ws[f"B{r0}"] = f"={px}/B10-1"; ws[f"B{r0}"].number_format = "0.0%"
    if Font is not None:
        ws[f"B{r0}"].font = Font(bold=True, color=GOLD, size=11)
    r = r0 + 1
    L(ws, f"A{r}", "NB: " + str(spec.get("discount_note") or
      "negativo = il mercato paga il fondo MENO dei suoi asset (sconto); il segnale "
      "e' lo sconto vs la sua storia (storico: voce dati futura dichiarata)."),
      italic=True, color=GREYTX)
    r += 2
    H(ws, f"A{r}", "RITORNI DEL NAV (dal sito ufficiale)")
    r += 1
    for k_ in ("MTD", "QTD", "YTD"):
        v_ = (spec.get("returns") or {}).get(k_)
        L(ws, f"A{r}", f"Ritorno {k_}")
        if v_ is None:
            L(ws, f"B{r}", "n.d.", color=GREYTX)
        else:
            N(ws, f"B{r}", v_ / 100.0, fmt="0.0%")
        r += 1
    r += 1
    _fv_rows(ws, spec, mv, r, lambda t: f"=B10*{t}", px)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 80
    spec["_cells"] = {"px": px}


def _sheet_fonti(wb, spec):
    from bellomberg.valuation.dcf_bank import GREYTX, H, L
    ws = wb.create_sheet("Fonti & Vintage")
    H(ws, "A1", "INPUT"); H(ws, "B1", "FONTE / VINTAGE")
    r = 2
    for k, v in (spec.get("_sources") or {}).items():
        L(ws, f"A{r}", k); L(ws, f"B{r}", str(v)); r += 1
    r += 1
    L(ws, f"A{r}", "VINTAGE", bold=True); r += 1
    for k, v in (spec.get("vintage") or {}).items():
        L(ws, f"A{r}", k)
        L(ws, f"B{r}", str(v) if v is not None else "n.d.",
          color=GREYTX if v is None else "1A1A1A")
        r += 1
    if spec.get("lag_note"):
        r += 1
        L(ws, f"A{r}", "lag della fonte (nota del sito):", bold=True); r += 1
        L(ws, f"B{r}", str(spec["lag_note"]), italic=True, color=GREYTX)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120


def _sheet_sens_mstr(wb, spec, mv):
    from bellomberg.valuation.dcf_bank import GREYTX, H, L, N
    ws = wb.create_sheet("Sensitivity")
    H(ws, "A1", "SENSITIVITY sul prezzo BTC (celle vive dal foglio Thesis)")
    ws_t = "'Thesis & Assumptions'!"
    cols = "BCDEF"
    H(ws, "A3", "shock BTC")
    for c_, s_ in zip(cols, SENS_SHOCKS):
        N(ws, f"{c_}3", s_, fmt="+0%;-0%;0%")
    L(ws, "A4", "Prezzo BTC (USD)")
    L(ws, "A5", "BTC-NAV (USD)")
    for c_ in cols:
        ws[f"{c_}4"] = f"={ws_t}B11*(1+{c_}3)"; ws[f"{c_}4"].number_format = "#,##0"
        ws[f"{c_}5"] = f"={ws_t}B10*{c_}4"; ws[f"{c_}5"].number_format = "#,##0"
    if spec.get("ev_missing"):
        L(ws, "A6", "mNAV EV / FV: n.d. DICHIARATO (campi %s assenti nel record)"
          % spec["ev_missing"], color=GREYTX)
    else:
        L(ws, "A6", "mNAV EV a prezzo azione corrente")
        for c_ in cols:
            ws[f"{c_}6"] = f"=({ws_t}B15+{ws_t}B16+{ws_t}B17-{ws_t}B18)/{c_}5"
            ws[f"{c_}6"].number_format = "0.000"
        L(ws, "A7", "NAV equity per azione (USD)")
        for c_ in cols:
            ws[f"{c_}7"] = f"=({c_}5-{ws_t}B16-{ws_t}B17+{ws_t}B18)/{ws_t}B13"
            ws[f"{c_}7"].number_format = "#,##0.00"
        if mv.get("fair_value_nav") is not None and spec.get("_target_cell"):
            t_cell = spec["_target_cell"]   # cella del target sul Thesis (riferimento VIVO)
            L(ws, "A8", "FV al target dichiarato (USD)")
            for c_ in cols:
                ws[f"{c_}8"] = f"=({ws_t}{t_cell}*{c_}5-{ws_t}B16-{ws_t}B17+{ws_t}B18)/{ws_t}B13"
                ws[f"{c_}8"].number_format = "#,##0.00"
    L(ws, "A10", "NB: banda di target dell'analista: modificare la cella target sul "
                 "Thesis (formule vive) — griglia dedicata non implementata, deviazione "
                 "dichiarata dal design V5.3.", italic=True)
    ws.column_dimensions["A"].width = 44
    for c_ in cols:
        ws.column_dimensions[c_].width = 15


def _sheet_sens_purr(wb, spec, mv):
    from bellomberg.valuation.dcf_bank import H, L, N
    ws = wb.create_sheet("Sensitivity")
    H(ws, "A1", "SENSITIVITY sul prezzo HYPE (celle vive dal foglio Thesis)")
    ws_t = "'Thesis & Assumptions'!"
    cc = spec.get("_cells") or {}
    cell, hy, px = cc.get("cell") or {}, cc.get("hy"), cc.get("px")
    fd = cc.get("fd")
    cols = "BCDEF"
    H(ws, "A3", "shock HYPE")
    for c_, s_ in zip(cols, SENS_SHOCKS):
        N(ws, f"{c_}3", s_, fmt="+0%;-0%;0%")
    rows = (
        ("Prezzo HYPE (USD)", "#,##0.00",
         lambda c: f"={ws_t}{hy}*(1+{c}3)"),
        ("Valore HYPE ($M)", "#,##0.0",
         lambda c: f"={ws_t}{cell['hypeHeld']}*{c}4"),
        ("Delta DTL ($M)", "#,##0.0",
         lambda c: f"={ws_t}{cell['reportedDTL']}-{ws_t}{cell['taxRate']}/100*({c}5-{ws_t}{cell['taxBasis']})"),
        ("Adjusted NAV ($M)", "#,##0.0",
         lambda c: (f"={ws_t}{cell['bookNAV']}-{ws_t}{cell['cashFromOps']}+{ws_t}{cell['cashFromFin']}"
                    f"-{ws_t}{cell['treasuryDeploy']}-{ws_t}{cell['reportedDigital']}+{c}5+{c}6")),
        ("Adjusted NAV per azione FD (USD)", "#,##0.0000",
         lambda c: f"={c}7/{ws_t}{fd}"),
        ("mNAV a prezzo azione corrente", "0.000",
         lambda c: f"={ws_t}{px}/{c}8"),
    )
    for j, (lbl, fmt, f_) in enumerate(rows, start=4):
        L(ws, f"A{j}", lbl)
        for c_ in cols:
            ws[f"{c_}{j}"] = f_(c_)
            ws[f"{c_}{j}"].number_format = fmt
    if mv.get("fair_value_nav") is not None and spec.get("_target_cell"):
        L(ws, "A10", "FV al target dichiarato (USD) = ANAV/FD x target")
        # la cella del target e' trovata sul Thesis dal builder (riferimento VIVO)
        t_cell = spec["_target_cell"]
        for c_ in cols:
            ws[f"{c_}10"] = f"={c_}8*{ws_t}{t_cell}"
            ws[f"{c_}10"].number_format = "#,##0.00"
    L(ws, "A12", "NB: diluizione warrant funzione del prezzo dell'azione (costante nello shock "
                 "HYPE); DTL teorica ricalcolata per shock come nella formula del sito. "
                 "Banda di target: modificare la cella target sul Thesis (formule vive) — "
                 "griglia dedicata non implementata, deviazione dichiarata dal design V5.3.",
      italic=True)
    ws.column_dimensions["A"].width = 46
    for c_ in cols:
        ws.column_dimensions[c_].width = 15


def _sheet_sens_cef(wb, spec, mv):
    from bellomberg.valuation.dcf_bank import H, L, N
    ws = wb.create_sheet("Sensitivity")
    H(ws, "A1", "GRIGLIA prezzo/NAV (celle vive dal foglio Thesis)")
    ws_t = "'Thesis & Assumptions'!"
    px = (spec.get("_cells") or {}).get("px") or "B11"
    cols = "BCDEF"
    H(ws, "A3", "prezzo/NAV")
    for c_, d_ in zip(cols, CEF_DISCOUNT_GRID):
        N(ws, f"{c_}3", d_, fmt="0.00")
    L(ws, "A4", "Prezzo implicito (USD) = NAV x rapporto")
    L(ws, "A5", "Upside vs prezzo corrente")
    for c_ in cols:
        ws[f"{c_}4"] = f"={ws_t}B10*{c_}3"; ws[f"{c_}4"].number_format = "#,##0.00"
        ws[f"{c_}5"] = f"={c_}4/{ws_t}{px}-1"; ws[f"{c_}5"].number_format = "0.0%"
    L(ws, "A7", "Lettura: ogni colonna e' un livello di sconto/premio a cui il mercato "
                "potrebbe prezzare il fondo — non una previsione. Il target dell'analista "
                "(se dichiarato) sta nel foglio Thesis.", italic=True)
    ws.column_dimensions["A"].width = 46
    for c_ in cols:
        ws.column_dimensions[c_].width = 14


def build_mnav_model(spec, output_path) -> Dict[str, Any]:
    """Workbook canonico mNAV (3 fogli: Thesis, Fonti & Vintage, Sensitivity) +
    payload dal mirror Python (fonte di verita'; i fogli riproducono le stesse
    formule VIVE — parita' misurata al collaudo live col bake)."""
    try:
        import openpyxl
    except ImportError:
        return {"ok": False, "error": "openpyxl non disponibile"}
    mv = compute_mnav_values(spec)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def _find_target_cell():
        # la cella del target sul Thesis (riferimento VIVO per la Sensitivity):
        # trovata per etichetta, mai una riga cablata
        ws_t = wb["Thesis & Assumptions"]
        for row in ws_t.iter_rows(min_col=1, max_col=1):
            for c in row:
                if c.value == "Target premio/sconto (ANALISTA)":
                    spec["_target_cell"] = "B%d" % c.row

    if spec["kind"] == "dat_bitcoin":
        _sheet_thesis_mstr(wb, spec, mv)
        _find_target_cell()
        _sheet_fonti(wb, spec)
        _sheet_sens_mstr(wb, spec, mv)
    elif spec["kind"] == "dat_hype":
        _sheet_thesis_purr(wb, spec, mv)
        _find_target_cell()
        _sheet_fonti(wb, spec)
        _sheet_sens_purr(wb, spec, mv)
    else:
        _sheet_thesis_cef(wb, spec, mv)
        _sheet_fonti(wb, spec)
        _sheet_sens_cef(wb, spec, mv)
    wb.calculation.fullCalcOnLoad = True   # cintura anti-anteprima-vuota (P0 17/07)
    try:
        wb.save(output_path)
    except Exception as e:
        return {"ok": False, "error": f"save: {e}"}
    from datetime import datetime as _dt
    out = {"ok": True, "path": output_path, "engine": "mnav",
           "payload_currency": spec.get("currency"),
           "method": _METHOD[spec["kind"]],
           "nav_target": spec.get("nav_target"),
           "price_quote": spec.get("price_quote"),
           "price_quote_currency": spec.get("price_quote_ccy"),
           "nav_vintage": spec.get("vintage"),
           "assumptions": {"nav_target": spec.get("nav_target"),
                           "kind": spec["kind"], "sources": spec.get("_sources")},
           "_timestamp": _dt.now().isoformat()}
    for k in ("nav_per_share", "nav_per_share_dtl_addback", "mnav_equity", "mnav_ev",
              "mnav", "mnav_dtl_addback", "discount_to_nav_pct", "btc_nav_usd",
              "adjusted_nav_musd", "fd_shares_m", "hype_value_musd",
              "fair_value_nav", "fv_note", "warnings"):
        if mv.get(k) is not None:
            out[k] = mv[k]
    if mv.get("fair_value_nav") is not None:
        # alias per la catena storica sanity/tesi (fair_value_base e' una delle
        # 4 chiavi che memo/persistenza gia' leggono) — stessa cifra, dichiarato
        out["fair_value_base"] = mv["fair_value_nav"]
        out["fv_alias_note"] = ("fair_value_base = fair_value_nav (alias per la catena "
                                "tesi/sanity; semantica: NAV x target dichiarato, NON "
                                "un target price — D2)")
    out["sanity"] = _mnav_sanity(spec, mv)
    return out
