"""
dcf_buyside_v3.py (#204b) - Generatore VAL v3 "template parity" col modello buy-side
del PM (Buy Side.xlsx, business game Rothschild/Luiss - DIGITRUST 9 fogli, ~900 formule).

Cosa replica del template:
  - STORICO accanto ai FORWARD (XBRL 10y quando c'e', via sec_xbrl)
  - ASSUMPTIONS stack per scenario con colonna COMMENTARY (scritta dall'AGENTE)
  - Pro-forma P&L driver-based con "o/w" splits, Cap R&D con roll di ammortamento 4y
  - UFCF build per scenario, DCF sul base con bridge to equity, sensitivity WACC x g
  - FAIR VALUE NUMERICI per scenario nel payload (mirror python delle formule)

Interfaccia: build_model_v3(spec, out_path, scenarios=None, history=None)
spec = quello di dcf_calibration.build_spec_from_ticker (riusato, non duplicato).
scenarios = dict bear/base/bull con driver paths + commentary (dall'agente via get_valuation).
Fallback: NESSUNO dal 16/07 (audit/13 §4 n.1): v3 KO = rifiuto dichiarato da dcf_engine;
il v2 dcf_buyside.py sta in attic/morti_20260902 dal 02/09 (0 importer misurati).
"""
from bellomberg.core.language import scoped_language, text as _lt
from bellomberg.reporting.i18n_excel import label as _xt
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPX_OK = True
except ImportError:
    OPX_OK = False

NAVY = "0B2545"; GOLD = "B08D2E"; WHITE = "FFFFFF"; GREYTX = "5A6472"; INK = "1A1A1A"
BLUE_INPUT = "1F4E9C"   # convenzione modelli: input blu, formule nere
N_FWD = 5
N_HIST = 4
DRIVERS = [
    ("revenue_growth", "Revenue Growth %", "0.0%"),
    ("gross_margin", "Gross Margin %", "0.0%"),
    ("rnd_pct", "R&D % of Revenues", "0.0%"),
    ("sga_pct", "SG&A (S&M+G&A) % of Rev", "0.0%"),
    ("capdev_pct", "Cap R&D/Dev % of Rev", "0.0%"),
    ("nwc_pct", "NWC % Rev", "0.0%"),
    ("capex_pct", "CapEx % Rev", "0.0%"),
    ("tax_rate", "Tax Rate", "0.0%"),
    # B11 (14/07, audit/09): promossi a driver espliciti (erano 2% / 60% / 30%
    # hardcoded nelle formule); stessi default, sovrascrivibili dall'agente
    ("da_tan_pct", "D&A tangibile % Rev", "0.0%"),
    ("personnel_pct", "o/w Personnel % of COGS", "0%"),
    ("services_pct", "o/w Services % of SG&A", "0%"),
]


def _safe(v, d=None):
    try:
        import math
        f = float(v)
        return f if math.isfinite(f) else d
    except (TypeError, ValueError):
        return d


def _hist_series(history, key, years):
    out = []
    items = ((history or {}).get("items") or {}).get(key) or {}
    for y in years:
        out.append(_safe(items.get(y)))
    return out


def _default_scenarios(spec, history):
    """Driver di default dai dati: base = spec growth/margini, bear/bull = +/- spread.
    B14 (14/07, audit/09): i margini bear/bull fanno FADE verso la mediana del
    settore (SECTOR_PRIORS di dcf_calibration) invece del solo drift meccanico:
    il drift pre-B14 (-1pp bear / +1.5pp bull a Y5) resta come FLOOR, cosi' il
    fade non puo' mai rendere i default MENO prudenti di prima; senza settore
    il comportamento e' identico al pre-B14."""
    g = list(spec.get("growth_path") or [0.08] * N_FWD)[:N_FWD]
    while len(g) < N_FWD:
        g.append(g[-1] * 0.85)
    gm = _safe(spec.get("gross_margin"), 0.45)
    rnd = _safe(spec.get("rnd_pct"), 0.05)
    sga = _safe(spec.get("sga_pct"), max(0.10, (1 - gm) * 0.35))
    capex = _safe(spec.get("capex_pct"), 0.04)
    nwc = _safe(spec.get("nwc_pct"), 0.02)
    tax = _safe(spec.get("tax_rate"), 0.25)
    sector_gm = None
    try:
        from bellomberg.valuation.dcf_calibration import _prior_for_sector, DEFAULT_PRIOR
        _p = _prior_for_sector(spec.get("sector") or "")
        # fade SOLO con un settore realmente matchato: il DEFAULT_PRIOR (gm 40%)
        # non e' una mediana di settore e farebbe crollare i margini a caso
        if _p is not DEFAULT_PRIOR:
            sector_gm = _safe(_p.get("gm"))
        if sector_gm is not None and not (0.05 <= sector_gm <= 0.95):
            sector_gm = None
    except Exception:
        sector_gm = None

    def gm_path(name):
        if name == "base":
            return [round(gm, 4)] * N_FWD
        if name == "bear":
            # fade a meta' strada verso il peggiore tra gm e mediana settore,
            # cap -6pp per prior anomali; floor -1pp = drift pre-B14
            fade = max((min(gm, sector_gm) - gm) * 0.5, -0.06) if sector_gm is not None else 0.0
            d5 = min(fade, -0.01)
        else:
            # bull: verso il migliore tra gm e mediana settore, cap +6pp; floor +1.5pp
            fade = min((max(gm, sector_gm) - gm) * 0.5, 0.06) if sector_gm is not None else 0.0
            d5 = max(fade, 0.015)
        return [round(min(max(gm + d5 * (i + 1) / N_FWD, 0.02), 0.95), 4) for i in range(N_FWD)]

    # M7 audit/13 (23/07): banda sanity sul margine base dei NON-ciclici — il
    # mid-cycle V2 normalizza i ciclici, ma per gli altri il gm TTM/spec entra nei
    # 3 scenari SENZA controlli: un one-off (impairment, COGS anomalo, dato Yahoo
    # sporco) diventerebbe la base di tutto il forecast. Si MISURA lo scarto vs
    # mediana storica XBRL (soglia 10pp, >=3 esercizi) o, senza storico, vs prior
    # di settore (15pp: banda larga, il prior non e' l'azienda) e si DICHIARA in
    # foglio+payload — mai sostituzioni zitte (regola 14/07). Soglie PM-regolabili.
    margin_sanity = None
    if not spec.get("_cyclical"):
        hist_gm, _n_gm = None, 0
        try:
            _items = (history or {}).get("items") or {}
            _revs = _items.get("revenue") or {}
            _gps = _items.get("gross_profit") or {}
            _cogs = _items.get("cost_of_revenue") or {}
            _ratios = []
            for _y, _rv in _revs.items():
                _rv = _safe(_rv)
                if not _rv:
                    continue
                _gp = _safe(_gps.get(_y))
                if _gp is None and _safe(_cogs.get(_y)) is not None:
                    _gp = _rv - _safe(_cogs.get(_y))
                if _gp is not None and -1.0 < _gp / _rv < 1.0:
                    _ratios.append(_gp / _rv)
            _n_gm = len(_ratios)
            if _n_gm >= 3:
                hist_gm = _median(_ratios)
        except Exception:
            hist_gm = None
        # review M7 (MEDIA-2): per i filer IFRS il V2.3 riaggiunge il D&A al gm del
        # modello (gm_model = gm_real + da_pct) mentre la mediana storica usa il
        # GrossProfit grezzo — confrontare gm_model produrrebbe un falso positivo
        # SISTEMATICO pari al re-add: lo scarto si misura sul gm REALE quando c'e'
        _gm_cmp = gm
        _gm_real = _safe(((spec.get("_calibration") or {}).get("gross_margin_real")))
        if _gm_real is not None:
            _gm_cmp = _gm_real
        _readd = ("" if _gm_cmp == gm else
                  " (scarto sul gm REALE %.0f%%; il base modello %.0f%% include il re-add D&A IFRS dichiarato)"
                  % (_gm_cmp * 100, gm * 100))
        if hist_gm is not None and abs(_gm_cmp - hist_gm) > 0.10:
            margin_sanity = ("ATTENZIONE M7: gross margin base %.0f%% devia %+.0fpp dalla "
                             "mediana storica %.0f%% su %d esercizi XBRL (possibile one-off "
                             "TTM): driver da validare con variant view%s"
                             % (_gm_cmp * 100, (_gm_cmp - hist_gm) * 100, hist_gm * 100,
                                _n_gm, _readd))
        elif hist_gm is None and sector_gm is not None and abs(_gm_cmp - sector_gm) > 0.15:
            # review M7 (BASSA-3): 1-2 esercizi non sono "storico n.d." — il buco
            # si dichiara com'e' (regola 14/07)
            _hole = ("storico %d esercizi < quorum 3" % _n_gm) if _n_gm else "storico n.d."
            margin_sanity = ("ATTENZIONE M7: gross margin base %.0f%% devia %+.0fpp dalla "
                             "mediana di settore %.0f%% (%s: banda larga 15pp): "
                             "driver da validare con variant view%s"
                             % (_gm_cmp * 100, (_gm_cmp - sector_gm) * 100, sector_gm * 100,
                                _hole, _readd))

    # M8 audit/13 (23/07): FONTE di ogni driver, riga per riga, nel blocco
    # ASSUMPTIONS — i driver_sources della calibrazione (REALE vs PROFILO vs
    # DEFAULT, gia' etichettati in dcf_calibration) + _growth/_margin_source
    # arrivavano al payload ma NON al foglio: l'analista vedeva i numeri senza
    # sapere cosa e' misurato e cosa e' convenzione. L'analista che fornisce un
    # driver sostituisce la fonte con "ANALISTA" (in _merge_scenarios).
    _cal_src = ((spec.get("_calibration") or {}).get("driver_sources") or {})
    # review M8 (MEDIA-1): con ebitda_margin_target dell'analista la calibrazione
    # NON tocca il gm (resta reale/prior) ma ricava l'OPEX per centrare il target:
    # l'etichetta "agent EBITDA-margin target" appartiene a R&D/SG&A, non al gm
    _ebt_target = (spec.get("_margin_source") == "agent EBITDA-margin target")
    # review M8 (BASSA-4): un growth_path corto viene esteso 0,85x oltre l'orizzonte
    _gp_in = spec.get("growth_path") or []
    _g_ext = (" | esteso 0,85x oltre l'orizzonte fornito"
              if (_gp_in and len(_gp_in) < N_FWD) else "")
    fonti = {
        "revenue_growth": ((str(spec.get("_growth_source") or "spec analista (growth_path fornito)")[:200]
                            + _g_ext) if _gp_in
                           else "DEFAULT 8% flat (growth_path n.d.)"),
        "gross_margin": (("calibrazione: gm reale/prior (il target EBITDA dell'analista "
                          "muove l'opex, NON il gm)") if _ebt_target
                         else (str(spec.get("_margin_source") or "spec/calibrazione (gm reale o modello)")[:200]
                               if spec.get("gross_margin") is not None
                               else "DEFAULT 45% (gross margin n.d.)")),
        "rnd_pct": (("derivato dal target EBITDA dell'analista (30% dell'opex "
                     "per centrare il target)") if _ebt_target
                    else ("calibrazione: split convenzionale dell'opex (30% R&D)"
                          if spec.get("rnd_pct") is not None else "DEFAULT 5% (R&D n.d.)")),
        "sga_pct": (("derivato dal target EBITDA dell'analista (45% S&M + 25% G&A "
                     "dell'opex per centrare il target)") if _ebt_target
                    else ("calibrazione: split convenzionale dell'opex (45% S&M + 25% G&A)"
                          if spec.get("sga_pct") is not None else "DEFAULT max(10%, (1-gm)*35%)")),
        "capex_pct": _cal_src.get("capex") or ("spec analista" if spec.get("capex_pct") is not None
                                               else "DEFAULT 4%"),
        "nwc_pct": _cal_src.get("nwc") or ("spec analista" if spec.get("nwc_pct") is not None
                                           else "DEFAULT 2%"),
        "tax_rate": _cal_src.get("tax") or ("spec analista" if spec.get("tax_rate") is not None
                                            else "DEFAULT 25%"),
        "da_tan_pct": _cal_src.get("da_tan") or ("spec analista" if spec.get("da_tan_pct") is not None
                                                 else "DEFAULT 2%"),
        "capdev_pct": "DEFAULT 0%: il v3 spesa l'R&D (capitalizzazione solo da analista)",
        "personnel_pct": "convenzione 60% (o/w cosmetico, non entra nell'UFCF)",
        "services_pct": "convenzione 30% (o/w cosmetico, non entra nell'UFCF)",
    }

    # audit/13 V2.4 (ok PM): spread ADDITIVI per profilo al posto dei moltiplicatori
    # x0.55/x1.30 — con growth negativa i moltiplicatori INVERTIVANO bear e bull.
    # Lo spread arriva dalla calibrazione (0.06 ciclici / 0.03 altri), dichiarato.
    spread = _safe(spec.get("scenario_spread"), 0.03)
    da_tan = _safe(spec.get("da_tan_pct"), 0.02)  # V2.3: D&A reale dal driver, non 2% fisso

    # V6 Lotto 2 (audit/19 V6.3, D3 PM 23/07): il range low/high della guidance che
    # ha GUIDATO il base (la calibrazione lo passa solo in quel caso) e' il bear/bull
    # VERO dichiarato dalla societa' sull'anno 1 — solo sul driver coperto
    # (revenue_growth); gli altri driver e gli anni 2+ restano allo spread meccanico.
    # Cintura B14: il bear da guidance non puo' MAI essere meno prudente del bear
    # meccanico — si prende il piu' prudente dei due, dichiarato in commentary.
    _grange = spec.get("_guidance_growth_range")
    if not isinstance(_grange, dict):
        _grange = None

    def stack(gshift, name):
        fade_note = (f"; margini con fade verso mediana settore {sector_gm:.0%}"
                     if (sector_gm is not None and name != "base") else "")
        cm_extra = {}
        if margin_sanity:
            cm_extra["_margin_sanity"] = margin_sanity  # M7: rimossa in _merge se l'analista fornisce il gm
        gpath = [round(x + gshift, 4) for x in g]
        if _grange and name == "bear" and _safe(_grange.get("low")) is not None:
            _mech = gpath[0]
            _low = _safe(_grange.get("low"))
            gpath[0] = round(min(_low, _mech), 4)
            cm_extra["_guidance_range"] = (
                "bear anno-1 %.1f%%: %s (range della guidance societaria %s [%s])"
                % (gpath[0] * 100,
                   ("low della guidance" if _low <= _mech else
                    "bear MECCANICO, piu' prudente del low %.1f%% della guidance (floor B14)" % (_low * 100)),
                   _grange.get("period"), _grange.get("src")))
        elif _grange and name == "bull" and _safe(_grange.get("high")) is not None:
            gpath[0] = round(_safe(_grange.get("high")), 4)
            cm_extra["_guidance_range"] = (
                "bull anno-1 %.1f%% = high del range della guidance societaria %s [%s]"
                % (gpath[0] * 100, _grange.get("period"), _grange.get("src")))
        return {
            "_driver_fonti": dict(fonti),  # M8: fonte per riga nel blocco ASSUMPTIONS
            "revenue_growth": gpath,
            "gross_margin": gm_path(name),
            "rnd_pct": [round(rnd, 4)] * N_FWD,
            "sga_pct": [round(sga, 4)] * N_FWD,
            "capdev_pct": [0.0] * N_FWD,
            "nwc_pct": [round(nwc, 4)] * N_FWD,
            "capex_pct": [round(capex, 4)] * N_FWD,
            "tax_rate": [round(tax, 4)] * N_FWD,
            # B11: default espliciti degli ex-hardcoded (splits 60/30); da_tan dal driver V2.3
            "da_tan_pct": [round(da_tan, 4)] * N_FWD,
            "personnel_pct": [0.60] * N_FWD,
            "services_pct": [0.30] * N_FWD,
            "commentary": {"_auto": f"scenario {name} generato dai dati (l'analista NON ha "
                                    f"fornito i driver: vedi nudge; spread additivo "
                                    f"{gshift:+.0%}){fade_note}", **cm_extra},
        }
    return {"bear": stack(-spread, "bear"), "base": stack(0.0, "base"),
            "bull": stack(+spread, "bull")}


def _merge_scenarios(agent, defaults):
    """Agent override sui default, driver per driver; commentary dell'agente preservata."""
    # driver SOSTANZIALI = muovono i numeri; i soli o/w splits (cosmetici, non
    # entrano nell'UFCF) non bastano a dichiarare lo scenario "dell'analista"
    substantive = tuple(k for k, _, _f in DRIVERS if k not in ("personnel_pct", "services_pct"))
    out = {}
    for sc in ("bear", "base", "bull"):
        d = dict(defaults[sc])
        a = (agent or {}).get(sc) or {}
        supplied = []
        for k, _, _f in DRIVERS:
            vals = a.get(k)
            if isinstance(vals, (int, float)):
                vals = [float(vals)] * N_FWD
            if isinstance(vals, list) and vals:
                clean = [_safe(x) for x in vals if _safe(x) is not None]
                if clean:
                    while len(clean) < N_FWD:
                        clean.append(clean[-1])
                    d[k] = [round(float(x), 4) for x in clean[:N_FWD]]
                    if k in substantive:
                        supplied.append(k)
        cm = dict(d.get("commentary") or {})
        # B14: scenario "dell'analista" solo con driver sostanziali; se parziale,
        # i default restanti restano DICHIARATI (fade incluso) — la discrezionalita'
        # si ordina, non si nasconde
        if supplied:
            # M8: il driver fornito ha come fonte l'analista, non piu' la calibrazione
            _f = dict(d.get("_driver_fonti") or {})
            for k in supplied:
                _f[k] = "ANALISTA (driver fornito in scenarios)"
            d["_driver_fonti"] = _f
            # M7: se l'analista ha fornito il margine, la sanity sul gm AUTO decade
            if "gross_margin" in supplied:
                cm.pop("_margin_sanity", None)
            # V6 Lotto 2: se l'analista fornisce il growth, la nota sul range della
            # guidance decade con lui (il suo path ha sostituito bear/bull anno-1)
            if "revenue_growth" in supplied:
                cm.pop("_guidance_range", None)
            auto_note = cm.pop("_auto", None)
            missing = [k for k in substantive if k not in supplied]
            if missing and auto_note:
                fade = ""
                if "; margini con fade" in auto_note and "gross_margin" in missing:
                    fade = "; " + auto_note.split("; ", 1)[1]
                cm["_parziale"] = ("driver dai DEFAULT (non forniti): " + ", ".join(missing) + fade)[:220]
        # review M7 (BASSA-7): niente chiavi riservate "_" dall'agente — un
        # commentary {"_margin_sanity": ...} contraffarebbe una nota del motore
        # nel foglio (rosso bold) e nel payload
        cm.update({k: str(v)[:220] for k, v in (a.get("commentary") or {}).items()
                   if not str(k).startswith("_")})
        d["commentary"] = cm
        out[sc] = d
    return out


# ---------- mirror numerico: UFCF + DCF per scenario ----------
def _scenario_numbers(spec, sc, last_rev, nwc0=None):
    """B12 (14/07): nwc0 = NWC storico reale da XBRL (CA-CL, in milioni) quando
    disponibile -> il delta NWC anno-1 parte dal bilancio vero invece del proxy
    rev x nwc_pct. B11: D&A tangibile dal driver esplicito (ex 2% hardcoded)."""
    rows = {"revenue": [], "ebitda": [], "ebit": [], "ufcf": []}
    rev = last_rev
    if nwc0 is not None:
        nwc_prev = nwc0
    else:
        # audit/11 §5: senza storico XBRL il FOGLIO usa l'hack =-NWC_y1*0.1 — il mirror
        # deve dare lo STESSO numero (prima: (rev1-rev0)*pct, divergenza sistematica
        # payload/workbook su ogni nome senza filing SEC). nwc_prev = 0.9*NWC_y1.
        nwc_prev = 0.9 * last_rev * (1 + sc["revenue_growth"][0]) * sc["nwc_pct"][0]
    documented = spec.get("documented_inputs") is True
    count = len(sc["revenue_growth"]) if documented else N_FWD
    da_tan_path = sc["da_tan_pct"] if documented else sc.get("da_tan_pct") or [0.02] * N_FWD
    amortization_years = spec["capdev_amortization_years"] if documented else 4
    capdev_hist = []
    for t in range(count):
        rev = rev * (1 + sc["revenue_growth"][t])
        gp = rev * sc["gross_margin"][t]
        opex = rev * (sc["rnd_pct"][t] + sc["sga_pct"][t])
        capdev = rev * sc["capdev_pct"][t]
        capdev_hist.append(capdev)
        ebitda = gp - opex + capdev
        da_tan = rev * da_tan_path[t]
        da_int = sum(capdev_hist[-amortization_years:]) / float(amortization_years)
        if documented:
            da_int += sc["opening_intangible_amortization"][t]
            rows.setdefault('research_amortization', []).append(da_int)
        ebit = ebitda - da_tan - da_int
        tax = max(ebit, 0) * sc["tax_rate"][t]
        capex = rev * sc["capex_pct"][t]
        nwc = rev * sc["nwc_pct"][t]
        d_nwc = nwc - nwc_prev
        nwc_prev = nwc
        ufcf = ebit - tax + da_tan + da_int - capex - capdev - d_nwc
        rows["revenue"].append(rev); rows["ebitda"].append(ebitda)
        rows["ebit"].append(ebit); rows["ufcf"].append(ufcf)
    return rows


def _median(vals):
    v = sorted(vals)
    if not v:
        return None
    m = len(v)
    return v[m // 2] if m % 2 else (v[m // 2 - 1] + v[m // 2]) / 2


def _comps_multiples(spec):
    """B8 (14/07, audit/09): FONTE UNICA dei multipli EV/EBITDA dei comps per
    mirror e foglio (mai due mediane divergenti). Esclusione outlier DICHIARATA
    come la media selettiva del template: multipli <=0 o >3x la mediana grezza.
    audit/12 V0.6: QUORUM n>=3 — con 1-2 multipli la 'mediana' e' il numero di un
    singolo peer (un large-cap deciso da un micro-cap a 2x; exit a 177x da n=2 dove
    l'outlier-check >3x mediana NON puo' scattare per costruzione): sotto quorum
    la mediana non si usa (None) e il chiamante ripiega sul proxy DICHIARANDO.
    Ritorna (rows [{name, m, excluded}], mediana sui tenuti | None, n_esclusi)."""
    rows = []
    for c in (spec.get("comps") or [])[:8]:
        ev, eb = _safe(c.get("ev")), _safe(c.get("ebitda"))
        m = _safe(c.get("ev_ebitda")) or ((ev / eb) if (ev is not None and eb and eb > 0) else None)
        rows.append({"name": str(c.get("name", "")), "m": m, "raw": c, "excluded": False})
    med0 = _median([r["m"] for r in rows if r["m"] and r["m"] > 0])
    for r in rows:
        r["excluded"] = bool(r["m"] is not None and (r["m"] <= 0 or (med0 and r["m"] > 3 * med0)))
    kept = [r["m"] for r in rows if r["m"] and not r["excluded"]]
    return rows, (_median(kept) if len(kept) >= 3 else None), sum(1 for r in rows if r["excluded"])


def _valid_precedents(spec):
    """B9: deal M&A validi (target + EV/EBITDA > 0), UNICA validazione per foglio
    Precedents e payload: mai una mediana nel JSON diversa da quella del foglio."""
    out = []
    for t in (spec.get("precedents") or [])[:12]:
        if not isinstance(t, dict):
            continue
        ev = _safe(t.get("ev_m")) or _safe(t.get("ev"))
        eb = _safe(t.get("ebitda_m")) or _safe(t.get("ebitda"))
        target = str(t.get("target") or "").strip()
        if ev is None or eb is None or not target:
            continue
        # fix review B: round PRIMA della validazione — un EBITDA 0.04 arrotondato
        # a 0.0 passava il check e poi divideva per zero (payload + #DIV/0! foglio)
        ev, eb = round(ev, 1), round(eb, 1)
        if ev > 0 and eb > 0:
            out.append({"target": target[:40], "acquirer": str(t.get("acquirer") or "")[:40],
                        "year": _safe(t.get("year")), "ev": ev, "ebitda": eb})
    return out


def _bridge_adjustments(spec):
    """Parita' A7: voci del bridge EV->Equity oltre il net debt (input agente/engine).
    Convenzione: value_m firmato, in milioni — negativo = debt-like che riduce l'equity
    (minorities, pension/TFR, earn-out, deferred, preferred), positivo = aggiunge
    (associates non consolidate). Lista vuota = comportamento pre-A7 invariato.
    Cap a 10 DOPO la validazione (8 voci agente + 2 auto engine): mai scartare in
    silenzio le auto-voci in coda, e le voci malformate non consumano slot."""
    out = []
    for a in (spec.get("equity_adjustments") or []):
        if not isinstance(a, dict):
            continue
        v = _safe(a.get("value_m"))
        lab = str(a.get("label") or "").strip()
        if v is None or not lab:
            continue
        out.append({"label": lab[:60], "value_m": v if spec.get("documented_inputs") is True else round(v, 1),
                    "commentary": str(a.get("commentary") or "")[:220]})
        if len(out) >= 10:
            break
    return out


def _shares_used(spec):
    """A7: fully diluted quando l'agente passa opzioni/SBC, altrimenti shares base."""
    return _safe(spec.get("diluted_shares_m")) or _safe(spec.get("shares_m"), 0) or 0


def _dcf_value(spec, ufcf, wacc, g, exit_multiple=None, ebitda_terminal=None,
               ebit_terminal=None, ronic=None, tax_term=None, *, terminal_value=None):
    """Fair value per share. Parita' A3+A5 (13/07, audit/09): mid-year convention
    parametrica (spec['mid_year'], default True = standard banking) e TV opzionale
    a exit multiple (EV/EBITDA x EBITDA terminale) in alternativa al Gordon.
    Parita' A7: equity = EV - net debt + somma adjustments; shares fully diluted.
    audit/13 V2.2.5: TV DISCIPLINATO (Damodaran) quando ebit_terminal/ronic/tax_term
    sono passati — FCFF terminale = EBIT_T x (1+g) x (1-tax) x (1-g/RONIC): la
    crescita perpetua COSTA reinvestimento, niente piu' Gordon sull'UFCF Y5 gratis.
    STESSA formula scritta viva nel foglio DCF (parita' B20 == mini base)."""
    from .cash_math import present_value
    if terminal_value is None and wacc <= g:
        return None
    off = 0.5 if spec.get("mid_year", True) else 0.0
    periods = spec["discount_periods"] if spec.get("documented_inputs") is True else [t + 1 - off for t in range(len(ufcf))]
    pv = present_value(ufcf,periods,wacc)
    if terminal_value is not None:
        tv = terminal_value
    elif exit_multiple and ebitda_terminal:
        tv = exit_multiple * ebitda_terminal
    elif ebit_terminal is not None and ronic and tax_term is not None:
        # review B1: parita' col foglio anche nel caso limite ronic<=g (TV=0 su
        # entrambi i lati, mai Gordon di nascosto solo nel mirror)
        tv = (ebit_terminal * (1 + g) * (1 - tax_term) * (1 - g / ronic) / (wacc - g)
              if ronic > g else 0.0)
    else:
        tv = ufcf[-1] * (1 + g) / (wacc - g)
    terminal_period = periods[-1] if spec.get("documented_inputs") is True else N_FWD - off
    ev = pv + tv / (1 + wacc) ** terminal_period
    net_debt = _safe(spec.get("net_debt"), 0) or 0
    eq = ev - net_debt + sum(a["value_m"] for a in _bridge_adjustments(spec))
    sh = _shares_used(spec)
    if sh <= 0:
        return None
    return eq / sh if spec.get('documented_inputs') is True else round(eq / sh, 2)


def _irr(flows):
    """G4 audit/14: IRR per bisezione (flows[0] = entry negativa, t in anni).
    None se il segno non cambia nel range [-95%, +1000%]: mai un numero inventato."""
    if any(cf != cf for cf in flows):  # review 17/07: NaN collasserebbe la bisezione su hi
        return None
    def npv(rate):
        return sum(cf / (1.0 + rate) ** t for t, cf in enumerate(flows))
    lo, hi = -0.95, 10.0
    f_lo, f_hi = npv(lo), npv(hi)
    if f_lo == 0:
        return lo
    if f_lo * f_hi > 0:
        return None
    mid = (lo + hi) / 2.0
    for _ in range(200):
        mid = (lo + hi) / 2.0
        v = npv(mid)
        if abs(v) < 1e-9:
            break
        if (v > 0) == (f_lo > 0):
            lo, f_lo = mid, v
        else:
            hi = mid
    return round(mid, 4)


def compute_fair_values_v3(spec, scenarios, last_rev, nwc0=None):
    wacc = _safe((spec.get("wacc_inputs") or {}).get("wacc")) or _safe(spec.get("wacc"), 0.09)
    # review B4: un terminal_g esplicito = 0.0 dell'analista si RISPETTA (l'or lo
    # trasformava in 2.5% zitto); il default 2.5% vale solo quando manca del tutto
    _tg = _safe(spec.get("terminal_g"))
    g = min(_tg if _tg is not None else 0.025,
            _safe((spec.get("wacc_inputs") or {}).get("rf"), 0.04) or 0.04)
    out = {"wacc_used": round(wacc, 4), "terminal_g_used": round(g, 4),
           "mid_year": bool(spec.get("mid_year", True))}
    # audit/13 V2.2.5: TERMINALE DISCIPLINATO — riattiva disciplined_terminal (era
    # codice morto): RONIC dall'ancora di profilo (ROC industry Damodaran, dichiarata)
    # con le regole di fade RONIC->WACC; senza ancora RONIC=WACC (reinvestment=g/WACC),
    # DICHIARATO. tax terminale = tax Y5 del base (unica per i 3 scenari, dichiarata).
    _ra = spec.get("ronic_anchor") or {}
    _tax_term = _safe((scenarios.get("base") or {}).get("tax_rate", [None] * N_FWD)[-1], 0.25)
    try:
        from bellomberg.market_data.damodaran_wacc import disciplined_terminal
        _rf_v = _safe((spec.get("wacc_inputs") or {}).get("rf"), 0.04) or 0.04
        _dt = disciplined_terminal(g, _rf_v, _safe(_ra.get("value")), wacc)
        _ronic = _dt["roic_terminal"]
        out["terminal_method"] = ("FCFF disciplinato: RONIC %.1f%% [%s], reinvestment g/RONIC "
                                  "= %.0f%%, tax term %.0f%%"
                                  % (_ronic * 100,
                                     _ra.get("source") or "ancora n.d. -> RONIC=WACC (dichiarato)",
                                     _dt["reinvestment_rate"] * 100, _tax_term * 100))
        if _dt.get("warnings"):
            out["terminal_warnings"] = _dt["warnings"]
    except Exception as _e:
        _ronic = None
        out["terminal_method"] = "Gordon su UFCF Y5 (disciplined_terminal KO: %s)" % _e
    out["ronic_used"] = round(_ronic, 4) if _ronic else None
    out["tax_terminal_used"] = round(_tax_term, 4)
    # B14: delta WACC opzionale per bear/bull (basis point, validati dall'engine);
    # il base resta SEMPRE sul WACC headline (B4 nel foglio, parita' B20==mini base).
    # fix review B: WACC di scenario CLAMPATO a g+25bp (un delta bull troppo negativo
    # mandava il mirror a None e il foglio a TV=0 con numeri diversi), e calcolato
    # UNA volta qui: il foglio legge fv['wacc_bear'/'wacc_bull'], mai ricalcolo suo.
    dbp = spec.get("wacc_delta_bp") if isinstance(spec.get("wacc_delta_bp"), dict) else {}

    def _wacc_sc(name):
        d = _safe(dbp.get(name)) if name in ("bear", "bull") else None
        if not d:
            return wacc
        w = wacc + d / 10000.0
        floor = g + 0.0025
        if w <= floor:
            out.setdefault("wacc_delta_notes", []).append(
                f"{name}: delta {d:+.0f}bp -> WACC {w:.2%} <= g {g:.2%}: clampato a {floor:.2%}")
            w = floor
        return round(w, 6)
    # A3 (13/07): exit multiple dai COMPS REALI; B8 (14/07): mediana dalla fonte
    # unica _comps_multiples, con esclusione outlier dichiarata. Mai il 10x cieco.
    comp_rows, comp_med, n_excl = _comps_multiples(spec)
    if comp_med:
        xm = comp_med
        n_used = sum(1 for r in comp_rows if r["m"] and not r["excluded"])
        xm_src = (f"mediana EV/EBITDA di {n_used} comps reali"
                  + (f" ({n_excl} outlier esclusi >3x mediana)" if n_excl else ""))
    else:
        xm = _safe(spec.get("exit_multiple")) or _safe((spec.get("terminal") or {}).get("exit_multiple"))
        _n_kept = sum(1 for r in comp_rows if r["m"] and not r["excluded"])
        _why = (f"comps sotto quorum: {_n_kept} multipli validi < 3" if _n_kept
                else "nessun comp con EV/EBITDA valido")
        xm_src = f"PROXY calibrazione di settore ({_why})" if xm else None
    out["exit_multiple_used"] = round(xm, 2) if xm else None
    out["exit_multiple_source"] = xm_src
    wsc = {n: _wacc_sc(n) for n in ("bear", "base", "bull")}
    vals = {}
    for sc_name in ("bear", "base", "bull"):
        nums = _scenario_numbers(spec, scenarios[sc_name], last_rev, nwc0=nwc0)
        w_s = wsc[sc_name]
        vals[sc_name] = _dcf_value(spec, nums["ufcf"], w_s, g,
                                   ebit_terminal=(nums["ebit"][-1] if _ronic else None),
                                   ronic=_ronic, tax_term=_tax_term)
        out[f"fair_value_{sc_name}"] = vals[sc_name]
        if xm and nums.get("ebitda"):
            out[f"fair_value_{sc_name}_exit"] = _dcf_value(
                spec, nums["ufcf"], w_s, g,
                exit_multiple=xm, ebitda_terminal=nums["ebitda"][-1])
        out[f"_{sc_name}_numbers"] = nums
    if any(_safe(dbp.get(k)) for k in ("bear", "bull")):
        out["wacc_delta_bp_used"] = {k: round(float(dbp[k])) for k in ("bear", "bull") if _safe(dbp.get(k))}
        # 6 decimali: STESSO numero che il foglio scrive nella cella WACC scenario
        out["wacc_bear"] = round(wsc["bear"], 6); out["wacc_bull"] = round(wsc["bull"], 6)
    # cross-check: multiplo EV/EBITDA implicito nel TV del base (disciplinato o Gordon)
    bn = out.get("_base_numbers") or {}
    if bn.get("ebitda") and bn["ebitda"][-1] and wacc > g and bn.get("ufcf"):
        if _ronic and bn.get("ebit"):
            tv_g = bn["ebit"][-1] * (1 + g) * (1 - _tax_term) * (1 - g / _ronic) / (wacc - g)
        else:
            tv_g = bn["ufcf"][-1] * (1 + g) / (wacc - g)
        out["implied_exit_multiple_gordon"] = round(tv_g / bn["ebitda"][-1], 2)
    pw = {"bear": 0.25, "base": 0.50, "bull": 0.25}
    if all(vals.get(k) for k in pw):
        out["fair_value_weighted"] = round(sum(vals[k] * w for k, w in pw.items()), 2)
        out["probability_weights"] = pw
    # B8: valutazione implicita dai comps = mediana EV/EBITDA x EBITDA fwd Y1 base,
    # stesso bridge e stesse shares del DCF (mirror del blocco nel foglio Comps)
    if comp_med and bn.get("ebitda"):
        nd = _safe(spec.get("net_debt"), 0) or 0
        eq_impl = (comp_med * bn["ebitda"][0] - nd
                   + sum(a["value_m"] for a in _bridge_adjustments(spec)))
        sh_u = _shares_used(spec)
        if sh_u > 0:
            out["fair_value_comps_implied"] = round(eq_impl / sh_u, 2)
    # B8: ponderazione parametrica dei METODI (default assente = 100% DCF, fair
    # value storici INVARIATI); blend dichiarato solo con pesi validi dall'engine
    mw = spec.get("method_weights") if isinstance(spec.get("method_weights"), dict) else None
    if mw and out.get("fair_value_weighted") and out.get("fair_value_comps_implied"):
        wd, wc = _safe(mw.get("dcf")), _safe(mw.get("comps"))
        if wd is not None and wc is not None and wd >= 0 and wc >= 0 and abs(wd + wc - 1) < 0.01:
            # pesi RAW (non arrotondati): la formula blend in Thesis usa questi stessi
            # valori — con round(2) un 1/3 diventava 0.33 e foglio/payload divergevano
            out["method_weights_used"] = {"dcf": wd, "comps": wc}
            out["fair_value_final"] = round(wd * out["fair_value_weighted"]
                                            + wc * out["fair_value_comps_implied"], 2)
    elif mw and out.get("fair_value_weighted"):
        # review V0: il degrade a 100% DCF era visibile solo per assenza — dichiarato
        out["method_weights_ignored"] = ("pesi metodi richiesti ma NON applicati: comps sotto "
                                         "quorum n>=3 o valutazione implicita non calcolabile — "
                                         "fair value 100% DCF")
    # G5 audit/14: delta DICHIARATO tra le due strade quando esistono entrambe
    # (stile "Delta vs top down" del modello Kairos) — in riga, non solo nel blend
    if out.get("fair_value_comps_implied") and out.get("fair_value_weighted"):
        out["methods_delta"] = round(out["fair_value_comps_implied"]
                                     / out["fair_value_weighted"] - 1.0, 4)
    # G4 audit/14 (Kairos, principio n.4): IRR di holding period PER SCENARIO —
    # entry al prezzo di oggi nella valuta dei FLUSSI (hp_entry_price, conversione
    # dichiarata da dcf_engine: mai unita' miste), dividendo corrente FLAT (nessun
    # modello dividendi operating: semplificazione dichiarata), exit a 3 anni al
    # multiplo dei comps (fonte dichiarata) su EBITDA Y3; net debt COSTANTE
    # (il funding plan e' il gap G8, residuo dichiarato).
    _p0 = _safe(spec.get("hp_entry_price"))
    if _p0 and _p0 > 0 and xm:
        _dps = _safe(spec.get("hp_dividend_ps"), 0) or 0
        _nd_hp = _safe(spec.get("net_debt"), 0) or 0
        _adj_hp = sum(a["value_m"] for a in _bridge_adjustments(spec))
        _shu = _shares_used(spec)
        _by_sc = {}
        _floored = []
        for _sc in ("bear", "base", "bull"):
            _eb = ((out.get(f"_{_sc}_numbers") or {}).get("ebitda")) or []
            # review 17/07: 'is not None' — EBITDA Y3 esattamente 0.0 non deve sparire zitto
            if len(_eb) >= 3 and _shu and _shu > 0 and _safe(_eb[2]) is not None:
                _exit_ps = (xm * _eb[2] - _nd_hp + _adj_hp) / _shu
                if _exit_ps < 0:
                    # review 17/07: limited liability — l'azionista non paga all'exit
                    _floored.append(_sc)
                    _exit_ps = 0.0
                _by_sc[_sc] = _irr([-_p0, _dps, _dps, _dps + _exit_ps])
        if _by_sc:
            _note = ((spec.get("hp_note") or "") + " | dividendo corrente flat e ND "
                     "costante (funding plan = G8): semplificazioni dichiarate")
            if _floored:
                _note += (" | equity exit NEGATIVA in %s -> floor 0 (limited liability, "
                          "dichiarato)" % ",".join(_floored))
            # review F4 finanza: dichiarare QUANDO l'ipotesi ND-costante morde — FCF
            # cumulato 3y (base, UFCF proxy di FCFE dichiarata) vs mkt cap di entry
            _ufcf3 = (((out.get("_base_numbers") or {}).get("ufcf")) or [])[:3]
            if len(_ufcf3) == 3 and all(_safe(x) is not None for x in _ufcf3):
                _fcf3 = sum(_ufcf3) - 3 * _dps * _shu
                _mcap0 = _p0 * _shu
                if _mcap0 > 0 and abs(_fcf3) > 0.15 * _mcap0:
                    _note += (" | ND costante ignora FCF cumulato 3y ~%+.0f%% del mkt cap: "
                              "IRR verosimilmente %s (dichiarato)"
                              % (100 * _fcf3 / _mcap0,
                                 "SOTTOSTIMATO" if _fcf3 > 0 else "SOVRASTIMATO"))
            _hi_out = {
                "years": 3, "by_scenario": _by_sc,
                "entry_price": _p0, "dividend_ps_flat": _dps,
                "exit_multiple": xm, "exit_multiple_source": xm_src,
                "note": _note}
            # review F3 finanza (G5): l'exit dei comps CONFRONTATA col multiplo implicito
            # nel terminale DCF (gia' calcolato, mai reso prima) — warning dichiarativo,
            # soglia 1.75x regolabile dal PM, MAI un cap silenzioso
            _ieg = _safe(out.get("implied_exit_multiple_gordon"))
            if _ieg and _ieg > 0:
                _gc = ("exit comps %.1fx vs multiplo implicito nel TV del DCF %.1fx (%+.0f%%)"
                       % (xm, _ieg, (xm / _ieg - 1) * 100))
                if xm / _ieg > 1.75:
                    _gc += (" — ATTENZIONE (>1.75x): per un ciclico verificare che i comps "
                            "siano mid-cycle e non di picco")
                _hi_out["gordon_check"] = _gc
            out["holding_irr"] = _hi_out
        else:
            out["holding_irr"] = {"years": 3, "by_scenario": {},
                                  "note": "IRR n.d.: EBITDA Y3 o shares mancanti (dichiarato)"}
    else:
        # review 17/07 (F2 finanza): con entry valida e xm assente la nota di successo
        # FX diventava il "motivo" del n.d. — precedenza corretta al quorum comps
        out["holding_irr"] = {"years": 3, "by_scenario": {},
                              "note": ("IRR n.d.: exit multiple dei comps assente (quorum)" if _p0
                                       else (spec.get("hp_note") or
                                             "IRR n.d.: prezzo di entry mancante (dichiarato)"))}
    # A7: bridge dichiarato nel payload (mirror di quel che il foglio DCF espone)
    adjs = _bridge_adjustments(spec)
    if adjs:
        out["equity_adjustments_used"] = adjs
        out["equity_adjustments_total"] = round(sum(a["value_m"] for a in adjs), 1)
    if spec.get("stance") in ("buy", "sell", "neutral"):
        out["stance"] = spec["stance"]
    if _safe(spec.get("diluted_shares_m")):
        out["shares_fully_diluted_m"] = round(float(spec["diluted_shares_m"]), 1)
    return out


# ---------- Excel helpers: CONVENZIONE BANCARIA (dal template del PM) ----------
# BLU = input/hardcoded (con fondo chiaro) | NERO = formula stesso foglio |
# ROSSO bold = link cross-foglio / derivate forward | azzurro DAEEF3 = righe chiave |
# grigio corsivo 8 = commentary | header anni: bianco su banda scura con bordi pieni.
INPUT_BLUE = "0000FF"; LINK_RED = "FF0000"; KEY_FILL = "DAEEF3"; INPUT_FILL = "EFF3FB"
SECTION_FILL = "2E5496"
_thin = Side(style="thin", color="9BA7B8")
BORDER_ALL = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)
BORDER_BOT = Border(bottom=_thin)


def TITLE(ws, c, t, ncols=10):
    ws[c] = t; ws[c].font = Font(bold=True, color=WHITE, size=12)
    row = ws[c].row
    for i in range(1, ncols + 1):
        ws.cell(row=row, column=i).fill = PatternFill("solid", fgColor=NAVY)


def H(ws, c, t, fill=SECTION_FILL, ncols=0):
    ws[c] = t; ws[c].font = Font(bold=True, color=WHITE, size=10)
    ws[c].fill = PatternFill("solid", fgColor=fill)
    if ncols:
        row = ws[c].row; col0 = ws[c].column
        for i in range(col0, col0 + ncols):
            ws.cell(row=row, column=i).fill = PatternFill("solid", fgColor=fill)


def YH(ws, c, t):
    """Header anno: bianco bold su navy con bordi pieni (come il template)."""
    ws[c] = t; ws[c].font = Font(bold=True, color=WHITE, size=9)
    ws[c].fill = PatternFill("solid", fgColor=NAVY); ws[c].border = BORDER_ALL
    ws[c].alignment = Alignment(horizontal="center")


def L(ws, c, t, bold=False, italic=False, color=INK, size=9):
    ws[c] = t; ws[c].font = Font(bold=bold, italic=italic, color=color, size=size)


def CMT(ws, c, t):
    ws[c] = t; ws[c].font = Font(italic=True, color="666666", size=8)


def N(ws, c, v, fmt="#,##0", bold=False, color=INK, is_input=False):
    ws[c] = v; ws[c].number_format = fmt
    if is_input:
        ws[c].font = Font(bold=bold, color=INPUT_BLUE, size=9)
        ws[c].fill = PatternFill("solid", fgColor=INPUT_FILL)
    else:
        ws[c].font = Font(bold=bold, color=color, size=9)


def F(ws, c, formula, fmt="#,##0", bold=False, color=INK, link=False):
    ws[c] = formula; ws[c].number_format = fmt
    ws[c].font = Font(bold=bold or link, color=LINK_RED if link else color, size=9)


def KEY(ws, row, col0=1, ncols=10):
    """Evidenzia una riga chiave (EBITDA/UFCF/Fair value): fill azzurro + bordo bottom."""
    for i in range(col0, col0 + ncols):
        cell = ws.cell(row=row, column=i)
        cell.fill = PatternFill("solid", fgColor=KEY_FILL)
        cell.border = BORDER_BOT
        rgb = None
        try:
            raw = cell.font.color.rgb if (cell.font and cell.font.color) else None
            raw = str(raw) if raw is not None else None
            if raw and len(raw) in (6, 8) and all(ch in "0123456789ABCDEFabcdef" for ch in raw):
                rgb = raw
        except Exception:
            rgb = None
        cell.font = Font(bold=True, color=rgb or INK, size=9)


def _sheet_thesis(wb, spec, fv, scenarios, history_ok, dcf_cells=None, comps_cells=None):
    """Parita' A6 (14/07, audit/09): il football field PUNTA alle celle vive del
    foglio DCF (link rossi, convenzione template) invece dei numeri Python congelati:
    toccando gli scenari, la Thesis si ricalcola. Fallback ai numeri del mirror se
    le celle non sono disponibili. B8: barra Comps nel football field + blend metodi
    parametrico. B14: banner quando TUTTI gli scenari sono _auto."""
    ws = wb.create_sheet("Thesis & Output", 0)
    dc = dcf_cells or {}
    TITLE(ws, "A1", _lt(f"{spec.get('company_name', spec.get('ticker'))} ({spec.get('ticker')}) - Modello buy-side v3",f"{spec.get('company_name', spec.get('ticker'))} ({spec.get('ticker')}) - Buy-side model v3"), ncols=4)
    L(ws, "A2", _lt(f"Generato {datetime.now():%d/%m/%Y %H:%M} - valuta {spec.get('currency', 'USD')} - storico XBRL: {('SI' if history_ok else 'NO (fallback yfinance)')}",f"Generated {datetime.now():%d/%m/%Y %H:%M} - currency {spec.get('currency', 'USD')} - XBRL history: {('SI' if history_ok else 'NO (fallback yfinance)')}"), italic=True, color=GREYTX)
    # B14: banner ben visibile quando NESSUNO scenario e' dell'analista
    if all((scenarios[s].get("commentary") or {}).get("_auto") for s in ("bear", "base", "bull")):
        ws["A3"] = _xt("ATTENZIONE: driver generati AUTOMATICAMENTE dai dati, non validati dall'analista (passa scenarios= in get_valuation)")
        ws["A3"].font = Font(bold=True, color="9C6500", size=9)
        for i in range(1, 5):
            ws.cell(row=3, column=i).fill = PatternFill("solid", fgColor="FFF2CC")
    H(ws, "A4", _xt("VARIANT VIEW DELL'ANALISTA"))
    L(ws, "A5", spec.get("_variant_view") or spec.get("variant_view") or _xt("(non fornita)"))
    H(ws, "A7", _xt("FOOTBALL FIELD (fair value per azione, link vivi al foglio DCF)")); H(ws, "B7", _xt("VALORE")); H(ws, "C7", _xt("PESO"))
    r = 8
    for sc, w in (("bear", "25%"), ("base", "50%"), ("bull", "25%")):
        L(ws, f"A{r}", f"DCF scenario {sc.upper()}")
        if dc.get(sc):
            F(ws, f"B{r}", f"=DCF!{dc[sc]}", "#,##0.00", bold=(sc == "base"), link=True)
        else:
            N(ws, f"B{r}", fv.get(f"fair_value_{sc}"), "#,##0.00", bold=(sc == "base"))
        L(ws, f"C{r}", w, color=GREYTX); r += 1
    w_row = r
    L(ws, f"A{r}", _xt("PONDERATO (25/50/25)"), bold=True)
    if dc.get("weighted"):
        F(ws, f"B{r}", f"=DCF!{dc['weighted']}", "#,##0.00", bold=True, link=True)
    else:
        N(ws, f"B{r}", fv.get("fair_value_weighted"), "#,##0.00", bold=True, color=GOLD)
    r += 1
    if dc.get("exit_base"):
        L(ws, f"A{r}", _xt("DCF TV exit multiple (base)"))
        F(ws, f"B{r}", f"=DCF!{dc['exit_base']}", "#,##0.00", link=True)
        L(ws, f"C{r}", "cross-check", color=GREYTX); r += 1
    # B8: barra Comps nel football field (link vivo) + blend metodi parametrico;
    # default senza method_weights = 100% DCF, fair value INVARIATO
    mw = fv.get("method_weights_used") or {}
    if comps_cells and comps_cells.get("per_share"):
        L(ws, f"A{r}", _xt("Comps EV/EBITDA implicito (mediana)"))
        F(ws, f"B{r}", f"=Comps!{comps_cells['per_share']}", "#,##0.00", link=True)
        L(ws, f"C{r}", f"{mw.get('comps', 0):.0%}" if mw else _xt("0% (informativo)"), color=GREYTX)
        comps_row = r; r += 1
        if mw and fv.get("fair_value_final") is not None:
            L(ws, f"A{r}", _lt(f"FAIR VALUE FINALE (combinazione {mw['dcf']:.0%} DCF / {mw['comps']:.0%} comps)",f"FINAL FAIR VALUE (blend {mw['dcf']:.0%} DCF / {mw['comps']:.0%} comps)"), bold=True)
            F(ws, f"B{r}", f"={mw['dcf']}*B{w_row}+{mw['comps']}*B{comps_row}", "#,##0.00", bold=True, color=GOLD)
            w_row = r  # upside calcolato sul blend dichiarato
            r += 1
    if spec.get("price"):
        # residuo (i) §9-quinquies n.5 (23/07, opzione A PM): con financialCurrency !=
        # quotazione (caso industriale: FV in USD, prezzo in EUR) la vecchia B{fv}/B{prezzo} mischiava
        # le valute (-29,5% al posto del vero -38,4%; payload/DB/F17 giusti). Il FV si
        # mostra in ENTRAMBE le valute su cella VIVA (tasso dichiarato, stesso del
        # payload) e l'upside punta SOLO alla cella convertita; cambio n.d. = upside
        # n.d. DICHIARATO, mai un numero a valute miste (regola 14/07).
        fxq = spec.get("fx_quote") or {}
        p_row = r
        L(ws, f"A{r}", _xt("Prezzo corrente") + (f" ({fxq['to']})" if fxq.get("to") else ""))
        N(ws, f"B{r}", spec["price"], "#,##0.00"); r += 1
        if fv.get("fair_value_weighted"):
            u_row = w_row
            if fxq.get("rate"):
                L(ws, f"A{r}", _lt(f"Fair value in {fxq['to']} (convertito @{fxq['rate']:.6g}, src {fxq.get('src', 'yfinance')})",f"Fair value in {fxq['to']} (convertito @{fxq['rate']:.6g}, src {fxq.get('src', 'yfinance')})"))
                F(ws, f"B{r}", f"=B{w_row}*{fxq['rate']}", "#,##0.00", bold=True)
                L(ws, f"C{r}", "cross-valuta", color=GREYTX)
                u_row = r; r += 1
            L(ws, f"A{r}", "Upside/Downside")
            if fxq.get("error"):
                L(ws, f"B{r}", _xt("n.d. — ") + fxq["error"], color=GREYTX); r += 1
            else:
                F(ws, f"B{r}", f"=B{u_row}/B{p_row}-1", "+0.0%;-0.0%", bold=True, color=GOLD); r += 1
    # G5 audit/14: delta tra metodi IN RIGA (stile "Delta vs top down" Kairos)
    if fv.get("methods_delta") is not None:
        L(ws, f"A{r}", _xt("Delta metodi: comps implicito vs DCF ponderato"))
        N(ws, f"B{r}", fv["methods_delta"], "+0.0%;-0.0%")
        L(ws, f"C{r}", "G5 audit/14", color=GREYTX); r += 1
    # G4 audit/14 (Kairos, principio n.4): IRR di holding period per scenario,
    # accanto all'upside — il tempo sta DENTRO la metrica
    _hi = fv.get("holding_irr") or {}
    _bs = _hi.get("by_scenario") or {}
    if _bs:
        r += 1
        H(ws, f"A{r}", _lt(f"IRR DI HOLDING PERIOD ({_hi.get('years', 3)} ANNI) PER SCENARIO — G4 audit/14",f"HOLDING PERIOD IRR ({_hi.get('years', 3)} YEARS) BY SCENARIO — G4 audit/14")); r += 1
        for _sc in ("bear", "base", "bull"):
            if _sc not in _bs:
                continue
            L(ws, f"A{r}", _lt(f"IRR annuo {_sc.upper()} (exit {_hi.get('exit_multiple')}x EV/EBITDA su Y3)",f"Annual IRR {_sc.upper()} (exit {_hi.get('exit_multiple')}x EV/EBITDA su Y3)"))
            if _bs[_sc] is None:
                L(ws, f"B{r}", _xt("n.d. (IRR fuori range [-95%,+1000%] o flussi senza cambio di segno)"), color=GREYTX)
            else:
                N(ws, f"B{r}", _bs[_sc], "+0.0%;-0.0%", bold=(_sc == "base"),
                  color=GOLD if _sc == "base" else INK)
            r += 1
        L(ws, f"A{r}", _lt(f"  entry {_hi.get('entry_price')} + div/az flat {_hi.get('dividend_ps_flat')} | exit: {_hi.get('exit_multiple_source')}",f"  entry {_hi.get('entry_price')} + flat dividend/share {_hi.get('dividend_ps_flat')} | exit: {_hi.get('exit_multiple_source')}"), italic=True, color=GREYTX); r += 1
        if _hi.get("gordon_check"):
            # review F3 finanza (G5): la strada comps-vs-DCF dichiarata in riga
            _gc_warn = _xt("ATTENZIONE") in _hi["gordon_check"]
            L(ws, f"A{r}", "  " + _hi["gordon_check"], italic=not _gc_warn, bold=_gc_warn,
              color="C00000" if _gc_warn else GREYTX); r += 1
        L(ws, f"A{r}", "  " + str(_hi.get("note") or ""), italic=True, color=GREYTX); r += 1
    elif _hi.get("note"):
        _nt = str(_hi["note"])
        L(ws, f"A{r}", _nt if _nt.startswith("IRR") else _xt("IRR holding n.d.: ") + _nt,
          italic=True, color=GREYTX); r += 1
    if fv.get("equity_adjustments_total") is not None:
        # il numero mostrato sono i SOLI adjustments: il totale bridge col net debt
        # sta nella riga TOTALE del foglio DCF (etichetta coerente, mai fuorviante)
        L(ws, f"A{r}", _lt(f"Ponte: {len(fv.get('equity_adjustments_used') or [])} adjustments oltre il net debt = {fv['equity_adjustments_total']:+,.1f}m",f"Bridge: {len(fv.get('equity_adjustments_used') or [])} adjustments beyond net debt = {fv['equity_adjustments_total']:+,.1f}m")
          + (f" - stance {fv['stance'].upper()}" if fv.get("stance") else "") + _xt(" (v. blocco BRIDGE nel foglio DCF)"), italic=True, color=GREYTX); r += 1
    elif fv.get("stance"):
        L(ws, f"A{r}", _lt(f"Stance {fv['stance'].upper()}: nessuna voce discrezionale inclusa nel bridge (solo net debt)",f"Stance {fv['stance'].upper()}: no discretionary items included in the bridge (net debt only)"), italic=True, color=GREYTX); r += 1
    r += 1
    H(ws, f"A{r}", _xt("PARAMETRI")); r += 1
    L(ws, f"A{r}", "WACC")
    if dc:
        F(ws, f"B{r}", "=DCF!B4", "0.00%", link=True)
    else:
        N(ws, f"B{r}", fv.get("wacc_used"), "0.00%")
    r += 1
    L(ws, f"A{r}", _xt("g terminale"))
    if dc:
        F(ws, f"B{r}", "=DCF!B5", "0.00%", link=True)
    else:
        N(ws, f"B{r}", fv.get("terminal_g_used"), "0.00%")
    r += 1
    if fv.get("wacc_delta_bp_used"):
        _d = fv["wacc_delta_bp_used"]
        L(ws, f"A{r}", "WACC delta bp (bear/bull)")
        L(ws, f"B{r}", " / ".join(f"{k} {v:+.0f}bp" for k, v in _d.items()), color=GREYTX)
        r += 1
    r += 1
    H(ws, f"A{r}", _xt("COMMENTARY DELL'ANALISTA (per scenario)")); r += 1
    for sc in ("bear", "base", "bull"):
        cm = scenarios[sc].get("commentary") or {}
        L(ws, f"A{r}", sc.upper(), bold=True); r += 1
        for k, txt in list(cm.items())[:8]:
            L(ws, f"A{r}", f"  {k}: {txt}", color=GREYTX); r += 1
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 14


def _sheet_historical(wb, spec, history):
    """B10 (14/07, audit/09): voci scalate in MILIONI (coerenti con gli scenari in
    {cur}m; heuristica come il builder: scala solo se i valori sono in unita');
    margini/derivate a FORMULE VIVE sulle righe base quando esistono, fallback ai
    valori Python del payload XBRL solo per le righe senza base nel foglio."""
    ws = wb.create_sheet("Historical 10Y")
    # V3 §9-sexies n.3: la FONTE del titolo e' quella vera (SEC o ESEF), mai hardcoded
    _src = str((history or {}).get("_source") or "SEC XBRL companyfacts")
    if not history or history.get("error"):
        TITLE(ws, "A1", _xt("Storico riga-per-riga (%s)") % _src, ncols=12)
        L(ws, "A3", _xt("Storico XBRL non disponibile per questo nome: ") + str((history or {}).get("error", "n/d")), italic=True, color=GREYTX)
        L(ws, "A4", _xt("(nomi senza filing SEC/ESEF: usare i 4 anni yfinance nel foglio scenario)"), italic=True, color=GREYTX)
        return []
    years = history.get("years") or []
    cols = [get_column_letter(2 + i) for i in range(len(years))]
    items = history.get("items") or {}
    rev_vals = [abs(_safe(v, 0) or 0) for v in (items.get("revenue") or {}).values()]
    if not rev_vals:
        # review 17/07 (ESEF): schemi senza riga revenue (banche Circ.262) — la scala
        # si decide sul massimo di TUTTE le voci, mai miliardi stampati in unita' piene
        rev_vals = [abs(_safe(v, 0) or 0) for s in items.values() for v in s.values()]
    scale = 1e6 if (rev_vals and max(rev_vals) > 1e7) else 1.0
    unit = ",".join(set((history.get("units") or {}).values()))
    TITLE(ws, "A1", _lt(f"Storico riga-per-riga ({_src}) - valori in {unit}{('m' if scale > 1 else '')}",f"Line-by-line history ({_src}) - values in {unit}{('m' if scale > 1 else '')}"), ncols=12)
    # review 17/07: copertura + BUCHI del repository + indice stantio, tutto in riga
    # nota (una sola riga disponibile: la 3 e' l'header anni)
    _notes = []
    if history.get("coverage_note"):
        _notes.append(str(history["coverage_note"]))
    if history.get("gaps"):
        _notes.append(_xt("BUCHI: ") + "; ".join(str(x) for x in history["gaps"][:3]))
    if history.get("index_note"):
        _notes.append(str(history["index_note"]))
    if history.get("accounting_basis"):
        _notes.append(str(history["accounting_basis"]))
    if _notes:
        L(ws, "A2", " | ".join(_notes), italic=True, color=GREYTX)
    YH(ws, "A3", _lt(f"VOCE ({unit}{('m' if scale > 1 else '')})",f"ITEM ({unit}{('m' if scale > 1 else '')})"))
    for i, y in enumerate(years):
        YH(ws, f"{cols[i]}3", str(y))
    ROWS = [("revenue", _xt("Net Revenues")), ("cost_of_revenue", "COGS"), ("gross_profit", _xt("Gross Profit")),
            ("rnd_expense", "R&D"), ("sga_expense", "SG&A"), ("operating_income", _xt("EBIT (Operating Income)")),
            ("interest_expense", _xt("Interest Expense")), ("pretax_income", _xt("Pre-tax Income")),
            ("tax_expense", _xt("Taxes")), ("net_income", _xt("Net Income")), ("eps_diluted", _xt("EPS diluted")),
            ("total_assets", _xt("Total Assets")), ("cash", _xt("Cash")), ("total_liabilities", _xt("Liabilities")),
            ("lt_debt", _xt("LT Debt")), ("equity", _xt("Equity")), ("cfo", "CFO"), ("capex", "CapEx"),
            ("dep_amort", "D&A"), ("dividends_paid", _xt("Dividends Paid")), ("buyback", "Buyback"), ("sbc", "SBC")]
    # review 17/07 (ESEF F4): le righe BANCARIE arrivavano al payload ma non al foglio
    for _bk, _bl in (("interest_revenue", _xt("Interest revenue (lordo)")),
                     ("fee_commission_net", _xt("Net fee & commission")),
                     ("fee_commission_income", _xt("Fee & commission income")),
                     ("impairment_ifrs9", _xt("Impairments IFRS9 (costo del rischio)")),
                     ("trading_income", _xt("Trading income")),
                     ("loans_to_customers", _xt("Loans to customers")),
                     ("loans_to_banks", _xt("Loans to banks")),
                     ("deposits_from_customers", _xt("Deposits from customers")),
                     ("deposits_from_banks", _xt("Deposits from banks"))):
        if _bk in items:
            ROWS.append((_bk, _bl))
    r = 4
    rowmap = {}
    for key, label in ROWS:
        if key not in items:
            continue
        L(ws, f"A{r}", label)
        for i, y in enumerate(years):
            v = _safe(items[key].get(y))
            if v is not None:
                # eps e' per-azione: mai scalato
                N(ws, f"{cols[i]}{r}", v if key == "eps_diluted" else v / scale,
                  "#,##0.00" if key == "eps_diluted" else "#,##0")
        rowmap[key] = r
        r += 1
    r += 1
    H(ws, f"A{r}", _xt("MARGINI E DERIVATE (formule vive sulle righe sopra)"))
    r += 1

    def _has(key, i):
        return key in rowmap and _safe(items[key].get(years[i])) is not None

    def _nonzero(key, i):
        v = _safe(items[key].get(years[i])) if key in rowmap else None
        return v is not None and abs(v) > 1e-9

    written = set()
    rev_row = rowmap.get("revenue")
    if rev_row:
        L(ws, f"A{r}", _xt("Revenue YoY %"))
        for i in range(1, len(years)):
            # denominatore MAI zero: il vecchio codice ometteva, mai #DIV/0!
            if _has("revenue", i) and _nonzero("revenue", i - 1):
                F(ws, f"{cols[i]}{r}", f"={cols[i]}{rev_row}/{cols[i-1]}{rev_row}-1", "0.0%")
        r += 1
    # margini a formula: (numeratore, label); payout = (div+bb)/NI con guardia perdite
    for num_key, label, dkey in (("gross_profit", _xt("Gross Margin %"), "gross_margin"),
                                 ("operating_income", _xt("EBIT Margin %"), "operating_margin"),
                                 ("net_income", _xt("Net Margin %"), "net_margin"),
                                 ("capex", _xt("CapEx % Rev"), "capex_pct_revenue")):
        if num_key not in rowmap or not rev_row:
            continue
        L(ws, f"A{r}", label)
        written.add(dkey)
        for i in range(len(years)):
            if _has(num_key, i) and _nonzero("revenue", i):
                F(ws, f"{cols[i]}{r}", f"={cols[i]}{rowmap[num_key]}/{cols[i]}{rev_row}", "0.0%")
        r += 1
    payout_parts = [k for k in ("dividends_paid", "buyback") if k in rowmap]
    if payout_parts and "net_income" in rowmap:
        L(ws, f"A{r}", _xt("Payout totale (div+bb)"))
        written.add("payout_total")
        ni = rowmap["net_income"]
        for i in range(len(years)):
            if _has("net_income", i) and any(_has(k, i) for k in payout_parts):
                num = "+".join(f"{cols[i]}{rowmap[k]}" for k in payout_parts)
                # in perdita il payout non e' significativo: "n/m" invece di un % negativo
                F(ws, f"{cols[i]}{r}", f"=IF({cols[i]}{ni}>0,({num})/{cols[i]}{ni},\"n/m\")", "0.0%")
        r += 1
    if all(k in rowmap for k in ("cfo", "capex")):
        L(ws, f"A{r}", "FCF (CFO-CapEx)")
        written.add("fcf")
        for i in range(len(years)):
            if _has("cfo", i) and _has("capex", i):
                F(ws, f"{cols[i]}{r}", f"={cols[i]}{rowmap['cfo']}-{cols[i]}{rowmap['capex']}", "#,##0")
        r += 1
    # fallback pre-B10: le derivate senza righe base nel foglio tornano dai valori
    # Python del payload XBRL (nessuna riga sparisce rispetto a prima)
    dd = history.get("derived") or {}
    for dkey, label, fmt, do_scale in (("gross_margin", _xt("Gross Margin %"), "0.0%", False),
                                       ("operating_margin", _xt("EBIT Margin %"), "0.0%", False),
                                       ("net_margin", _xt("Net Margin %"), "0.0%", False),
                                       ("capex_pct_revenue", _xt("CapEx % Rev"), "0.0%", False),
                                       ("payout_total", _xt("Payout totale (div+bb)"), "0.0%", False),
                                       ("fcf", "FCF (CFO-CapEx)", "#,##0", True)):
        if dkey in written or not (dd.get(dkey) or {}):
            continue
        L(ws, f"A{r}", label)
        for i, y in enumerate(years):
            v = _safe((dd.get(dkey) or {}).get(y))
            if v is not None:
                N(ws, f"{cols[i]}{r}", v / scale if do_scale else v, fmt)
        r += 1
    ws.column_dimensions["A"].width = 26
    for c in cols:
        ws.column_dimensions[c].width = 11
    return years


def _sheet_scenario(wb, spec, sc_name, sc, hist_rev, hist_years, hist_rows=None, nwc0=None,
                    last_rev=None):
    """B11: D&A tangibile e o/w splits presi dai driver ASSUMPTIONS (input blu)
    invece che hardcoded in formula. B12: delta NWC anno-1 dal NWC storico XBRL
    quando c'e' (nwc0, in milioni), fallback all'hack -NWC*0.1 con nota (stima).
    B13: storico COGS/GP/opex/EBITDA/EBIT compilato dalle serie XBRL (hist_rows,
    convenzione template: storici a valore, forecast a formula).
    Fix review B: storico allineato a DESTRA (l'ultima colonna storica e' sempre
    l'anno piu' recente: e' l'ancora della catena forecast E la colonna del nwc0);
    se il ricavo dell'ultimo anno manca, l'ancora viene garantita scrivendo
    last_rev (lo stesso numero usato dal mirror) come input dichiarato."""
    ws = wb.create_sheet(f"Scenario {sc_name.capitalize()}")
    cur = spec.get("currency", "USD")
    TITLE(ws, "A1", f"{spec.get('ticker')} - Pro-Forma {sc_name.upper()} ({cur}m)", ncols=11)
    y0 = datetime.now().year
    hy = [f"{y0 - N_HIST + i}A" for i in range(N_HIST)]
    fy = [f"{y0 + i}B" for i in range(N_FWD)]
    all_cols = [get_column_letter(2 + i) for i in range(N_HIST + N_FWD)]
    hcols, fcols = all_cols[:N_HIST], all_cols[N_HIST:]
    CM = get_column_letter(2 + N_HIST + N_FWD + 1)  # colonna commentary

    # ===== ASSUMPTIONS =====
    H(ws, "A3", _xt("ASSUMPTIONS (driver dell'analista)"), ncols=1)
    for i, y in enumerate(hy + fy):
        YH(ws, f"{all_cols[i]}3", y)
    YH(ws, f"{CM}3", _xt("COMMENTARY"))
    cm = sc.get("commentary") or {}
    fonti = sc.get("_driver_fonti") or {}
    arow = {}
    r = 4
    for key, label, fmt in DRIVERS:
        L(ws, f"A{r}", label)
        for i in range(N_FWD):
            N(ws, f"{fcols[i]}{r}", sc[key][i], fmt, is_input=True)
        if cm.get(key):
            CMT(ws, f"{CM}{r}", cm[key])
        elif fonti.get(key):
            # M8 audit/13: la FONTE del driver in riga quando l'analista non ha
            # commentato (il suo commento, se c'e', E' la storia della fonte)
            CMT(ws, f"{CM}{r}", _xt("fonte: ") + str(fonti[key]))
        arow[key] = r
        r += 1
    if cm.get("_auto"):
        L(ws, f"{CM}{r}", cm["_auto"], italic=True, color="C00000")
    elif cm.get("_parziale"):
        L(ws, f"{CM}{r}", cm["_parziale"], italic=True, color=GREYTX)
    if cm.get("_margin_sanity"):
        # M7 audit/13: banda sanity margini non-ciclici — visibile, mai zitta
        r += 1
        L(ws, f"{CM}{r}", cm["_margin_sanity"], bold=True, color="C00000")
    if cm.get("_guidance_range"):
        # V6 Lotto 2 (D3): bear/bull anno-1 dal range della guidance societaria,
        # scelta dichiarata in riga (floor B14 incluso)
        r += 1
        L(ws, f"{CM}{r}", cm["_guidance_range"], italic=True, color=GREYTX)
    r += 2

    # ===== PRO-FORMA P&L =====
    H(ws, f"A{r}", _xt("PRO-FORMA P&L")); pr = r + 1
    rows = {}
    def row(label, bold=False):
        nonlocal pr
        L(ws, f"A{pr}", label, bold=bold)
        rows[label] = pr
        pr += 1
        return pr - 1
    rev_r = row(_xt("Net Revenues"), bold=True)
    yoy_r = row(_xt("YoY %"))
    cogs_r = row("COGS")
    row(_xt("  o/w Personnel (driver)"))
    gp_r = row(_xt("Gross Profit"), bold=True)
    gm_r = row(_xt("Gross Margin %"))
    rnd_r = row("R&D")
    sga_r = row("SG&A (S&M + G&A)")
    row(_xt("  o/w Services (driver)"))
    opex_r = row(_xt("Total OpEx"), bold=True)
    capdev_r = row(_xt("Capitalized R&D/Dev"))
    ebitda_r = row("EBITDA", bold=True)
    ebitdam_r = row(_xt("EBITDA Margin %"))
    datan_r = row(_xt("D&A tangibile (driver % rev)"))
    daint_r = row(_xt("D&A intangibile (roll 4y Cap R&D)"))
    ebit_r = row("EBIT", bold=True)
    ebitm_r = row(_xt("EBIT Margin %"))

    # storici (valori) allineati a DESTRA + forward (formule)
    rev_vals = list(hist_rev)[-N_HIST:]
    off = N_HIST - len(rev_vals)
    hist_have = {"rev": [False] * N_HIST}
    for i, v in enumerate(rev_vals):
        v = _safe(v)
        if v is not None:
            # residuo (ii) §9-quinquies n.5 (23/07): hist_rev arriva GIA' in milioni
            # (scala decisa UNA volta in build_model_v3 sul massimo di TUTTE le serie,
            # review B 14/07; fallback revenue_hist_m = milioni per contratto) — il
            # vecchio riscaling per-valore (v/1e6 se >1e7) divideva una SECONDA volta
            # le mega-cap JPY/KRW con ricavi >1e7 mln: storico e catena forward
            # sballati di 1e6 nel foglio (payload giusto: usa last_rev, non la cella)
            N(ws, f"{hcols[off + i]}{rev_r}", v, "#,##0", is_input=True)
            # revenue e' il DENOMINATORE di margini/YoY: uno zero non e' una base valida
            hist_have["rev"][off + i] = abs(v) > 1e-9
    # ancora della catena forecast: l'ultima colonna storica DEVE avere il ricavo
    # (lo stesso last_rev del mirror), altrimenti Y1 = 0 in silenzio
    if not hist_have["rev"][N_HIST - 1] and _safe(last_rev):
        N(ws, f"{hcols[-1]}{rev_r}", float(last_rev), "#,##0", is_input=True)
        CMT(ws, f"{CM}{rev_r}", _xt("ultimo ricavo noto (ancora della catena forecast: stesso valore del mirror)"))
        hist_have["rev"][N_HIST - 1] = True
    # B13: storico completo dalle serie XBRL (gia' scalate in milioni dal builder);
    # COGS/R&D/SG&A col segno negativo di convenzione, EBIT/GP/EBITDA come sono
    # (mai abs: un EBIT storico negativo e' informazione, non errore)
    hr = hist_rows or {}
    for key, rr, neg in (("cogs", cogs_r, True), ("gross_profit", gp_r, False),
                         ("rnd", rnd_r, True), ("sga", sga_r, True),
                         ("ebitda", ebitda_r, False), ("ebit", ebit_r, False)):
        serie = list(hr.get(key) or [])[-N_HIST:]
        soff = N_HIST - len(serie)
        have = [False] * N_HIST
        for i, v in enumerate(serie):
            if v is not None:
                N(ws, f"{hcols[soff + i]}{rr}", -abs(v) if neg else v, "#,##0",
                  bold=(key in ("gross_profit", "ebitda", "ebit")), is_input=True)
                have[soff + i] = True
        hist_have[key] = have
    # margini, YoY e Total OpEx sugli anni storici, a formula, dove le basi esistono
    for i in range(N_HIST):
        c = hcols[i]
        if i > 0 and hist_have["rev"][i] and hist_have["rev"][i - 1]:
            F(ws, f"{c}{yoy_r}", f"={c}{rev_r}/{hcols[i-1]}{rev_r}-1", "0.0%")
        if hist_have.get("rnd", [False] * N_HIST)[i] and hist_have.get("sga", [False] * N_HIST)[i]:
            F(ws, f"{c}{opex_r}", f"={c}{rnd_r}+{c}{sga_r}", bold=True)
        if hist_have["rev"][i]:
            if hist_have.get("gross_profit", [False] * N_HIST)[i]:
                F(ws, f"{c}{gm_r}", f"={c}{gp_r}/{c}{rev_r}", "0.0%")
            if hist_have.get("ebitda", [False] * N_HIST)[i]:
                F(ws, f"{c}{ebitdam_r}", f"={c}{ebitda_r}/{c}{rev_r}", "0.0%")
            if hist_have.get("ebit", [False] * N_HIST)[i]:
                F(ws, f"{c}{ebitm_r}", f"={c}{ebit_r}/{c}{rev_r}", "0.0%")
    for i, c in enumerate(fcols):
        p = fcols[i - 1] if i > 0 else hcols[-1]
        F(ws, f"{c}{rev_r}", f"={p}{rev_r}*(1+{c}{arow['revenue_growth']})", bold=True)
        F(ws, f"{c}{yoy_r}", f"={c}{rev_r}/{p}{rev_r}-1", "0.0%")
        F(ws, f"{c}{cogs_r}", f"=-{c}{rev_r}*(1-{c}{arow['gross_margin']})")
        F(ws, f"{c}{cogs_r+1}", f"=ROUND({c}{cogs_r}*{c}{arow['personnel_pct']},1)")
        F(ws, f"{c}{gp_r}", f"={c}{rev_r}+{c}{cogs_r}", bold=True)
        F(ws, f"{c}{gm_r}", f"={c}{gp_r}/{c}{rev_r}", "0.0%")
        F(ws, f"{c}{rnd_r}", f"=-{c}{rev_r}*{c}{arow['rnd_pct']}")
        F(ws, f"{c}{sga_r}", f"=-{c}{rev_r}*{c}{arow['sga_pct']}")
        F(ws, f"{c}{sga_r+1}", f"=ROUND({c}{sga_r}*{c}{arow['services_pct']},1)")
        F(ws, f"{c}{opex_r}", f"={c}{rnd_r}+{c}{sga_r}", bold=True)
        F(ws, f"{c}{capdev_r}", f"={c}{rev_r}*{c}{arow['capdev_pct']}")
        F(ws, f"{c}{ebitda_r}", f"={c}{gp_r}+{c}{opex_r}+{c}{capdev_r}", bold=True)
        F(ws, f"{c}{ebitdam_r}", f"={c}{ebitda_r}/{c}{rev_r}", "0.0%")
        F(ws, f"{c}{datan_r}", f"=-{c}{rev_r}*{c}{arow['da_tan_pct']}")
        prevs = fcols[max(0, i - 3):i + 1]
        terms = "+".join(f"{cc}{capdev_r}" for cc in prevs)
        F(ws, f"{c}{daint_r}", f"=-({terms})/4")
        F(ws, f"{c}{ebit_r}", f"={c}{ebitda_r}+{c}{datan_r}+{c}{daint_r}", bold=True)
        F(ws, f"{c}{ebitm_r}", f"={c}{ebit_r}/{c}{rev_r}", "0.0%")
    # righe chiave evidenziate come nel template (EBITDA azzurra, EBIT bordata)
    KEY(ws, ebitda_r, 1, N_HIST + N_FWD + 1)
    for i in range(1, N_HIST + N_FWD + 2):
        ws.cell(row=ebit_r, column=i).border = BORDER_BOT
    pr += 1

    # ===== UFCF =====
    H(ws, f"A{pr}", _xt("UNLEVERED FREE CASH FLOW")); pr += 1
    tax_r = pr; L(ws, f"A{pr}", _xt("Taxes su EBIT")); pr += 1
    nwc_r = pr; L(ws, f"A{pr}", _xt("NWC (livello)")); pr += 1
    dnwc_r = pr; L(ws, f"A{pr}", _xt("Delta NWC")); pr += 1
    capex_r = pr; L(ws, f"A{pr}", "CapEx"); pr += 1
    ufcf_r = pr; L(ws, f"A{pr}", "UFCF", bold=True); pr += 1
    # B12: NWC di partenza dal bilancio XBRL (CA-CL) come input blu nell'ultima
    # colonna storica -> il delta Y1 diventa la STESSA formula degli altri anni;
    # senza storico resta l'hack -NWC*0.1, dichiarato '(stima)' in commentary
    if nwc0 is not None:
        N(ws, f"{hcols[-1]}{nwc_r}", nwc0, "#,##0", is_input=True)
        CMT(ws, f"{CM}{dnwc_r}", _xt("delta Y1 dal NWC OPERATIVO storico XBRL (CA - cassa - CL; "
                                 "debito a breve non separabile dai tag: proxy dichiarato)"))
    else:
        CMT(ws, f"{CM}{dnwc_r}", _xt("(stima) delta Y1 = -10% del NWC Y1: storico XBRL non disponibile"))
    for i, c in enumerate(fcols):
        p = fcols[i - 1] if i > 0 else (hcols[-1] if nwc0 is not None else None)
        F(ws, f"{c}{tax_r}", f"=-MAX({c}{ebit_r},0)*{c}{arow['tax_rate']}")
        F(ws, f"{c}{nwc_r}", f"={c}{rev_r}*{c}{arow['nwc_pct']}")
        F(ws, f"{c}{dnwc_r}", f"=-({c}{nwc_r}-{p}{nwc_r})" if p else f"=-{c}{nwc_r}*0.1")
        F(ws, f"{c}{capex_r}", f"=-{c}{rev_r}*{c}{arow['capex_pct']}")
        F(ws, f"{c}{ufcf_r}", f"={c}{ebit_r}+{c}{tax_r}-{c}{datan_r}-{c}{daint_r}+{c}{dnwc_r}+{c}{capex_r}-{c}{capdev_r}", bold=True)
    KEY(ws, ufcf_r, 1, N_HIST + N_FWD + 1)
    # nota: -datan/-daint con segno gia' negativo nelle righe = riaggiunta D&A
    ws.column_dimensions["A"].width = 30
    for c in all_cols + [CM]:
        ws.column_dimensions[c].width = 10
    ws.column_dimensions[CM].width = 60
    # audit/13 V2.2.5: ebit_row esposto — il TV disciplinato del foglio DCF parte
    # dall'EBIT terminale di scenario, non piu' dall'UFCF Y5
    return {"ufcf_row": ufcf_r, "ebitda_row": ebitda_r, "ebit_row": ebit_r,
            "fcols": fcols, "sheet": ws.title}


def _sheet_dcf(wb, spec, fv, refs):
    """Parita' A6+A7 (14/07, audit/09): oltre al DCF base, bridge to equity
    parametrico (adjustments input blu + stance) e mini-DCF a formule vive per
    bear/base/bull con ponderato 25/50/25. Ritorna le celle chiave per la Thesis."""
    base_ref = refs["base"]
    adjs = _bridge_adjustments(spec)
    diluted = _safe(spec.get("diluted_shares_m"))
    sh_ref = "$B$8" if diluted else "$B$7"
    ws = wb.create_sheet("DCF")
    TITLE(ws, "A1", _xt("DCF (formule vive): base + bridge to equity + mini-DCF per scenario"), ncols=7)
    H(ws, "A3", _xt("PARAMETRI"))
    L(ws, "A4", "WACC"); N(ws, "B4", fv.get("wacc_used"), "0.00%", is_input=True)
    L(ws, "A5", _xt("g terminale")); N(ws, "B5", fv.get("terminal_g_used"), "0.00%", is_input=True)
    L(ws, "A6", _xt("Net debt")); N(ws, "B6", _safe(spec.get("net_debt"), 0) or 0)
    L(ws, "A7", _xt("Shares (m)")); N(ws, "B7", _safe(spec.get("shares_m"), 0) or 0, "#,##0.0")
    if diluted:
        L(ws, "A8", _xt("Shares fully diluted (m)"))
        N(ws, "B8", diluted, "#,##0.0", is_input=True)
        CMT(ws, "C8", _xt("diluizione opzioni/RSU/SBC passata dall'analista: usata al posto delle shares base"))
    # parita' A3+A5 (13/07): parametri in colonna D per non slittare le righe
    L(ws, "D4", _xt("Mid-year offset (0,5=ON)")); N(ws, "E4", 0.5 if fv.get("mid_year", True) else 0.0, "0.0", is_input=True)
    L(ws, "D5", _xt("Exit multiple EV/EBITDA")); N(ws, "E5", _safe(fv.get("exit_multiple_used"), 0.0), "0.00", is_input=True)
    CMT(ws, "F5", fv.get("exit_multiple_source") or "")
    # audit/13 V2.2.5: parametri del TV DISCIPLINATO (stessa formula del mirror)
    ronic_v = _safe(fv.get("ronic_used"))
    taxt_v = _safe(fv.get("tax_terminal_used"), 0.25)
    if ronic_v:
        L(ws, "D6", _xt("RONIC terminale")); N(ws, "E6", ronic_v, "0.00%", is_input=True)
        L(ws, "D7", _xt("Tax terminale")); N(ws, "E7", taxt_v, "0.00%", is_input=True)
        CMT(ws, "F6", fv.get("terminal_method") or _xt("TV disciplinato: FCFF = EBIT_Y5 x (1+g) x (1-tax) x (1-g/RONIC)"))
    H(ws, "A9", _xt("ANNO"));
    cols = [get_column_letter(2 + i) for i in range(N_FWD)]
    sh = base_ref["sheet"]; ur = base_ref["ufcf_row"]
    for i, c in enumerate(cols):
        YH(ws, f"{c}9", f"Y{i+1}")
        F(ws, f"{c}10", f"='{sh}'!{base_ref['fcols'][i]}{ur}", link=True)
        F(ws, f"{c}11", f"=1/(1+$B$4)^({i + 1}-$E$4)", "0.000")
        F(ws, f"{c}12", f"={c}10*{c}11")
    L(ws, "A10", _xt("UFCF (da Scenario Base)"))
    L(ws, "A11", _xt("Discount factor (mid-year)"))
    L(ws, "A12", _xt("PV UFCF"))
    last = cols[-1]
    # A7: geometria del bridge calcolata PRIMA delle formule che la referenziano.
    # Blocco sotto la sensitivity (righe 23-28): header 30, -net debt 31, adjustments,
    # totale; il mini-DCF per scenario (A6) parte due righe sotto.
    br0 = 30
    br_total = br0 + 2 + len(adjs)
    bridge_ref = f"$B${br_total}" if adjs else None
    ms = (br_total + (3 if spec.get("stance") in ("buy", "sell", "neutral") else 2)) if adjs else br0
    L(ws, "A14", _xt("Sum PV UFCF")); F(ws, "B14", f"=SUM(B12:{last}12)", bold=True)
    # audit/13 V2.2.5: TV disciplinato (FCFF da EBIT terminale, reinvestment=g/RONIC)
    # quando il mirror ha un RONIC; altrimenti Gordon storico su UFCF Y5. STESSA
    # formula del mirror _dcf_value: la parita' B20 == payload resta verificabile.
    _er_ebit = base_ref.get("ebit_row")
    if ronic_v and _er_ebit:
        _ebit5 = f"'{sh}'!{base_ref['fcols'][-1]}{_er_ebit}"
        L(ws, "A15", _xt("Terminal Value (FCFF disciplinato)"))
        F(ws, "B15", f"=IF(AND($B$4>$B$5,$E$6>$B$5),{_ebit5}*(1+$B$5)*(1-$E$7)*(1-$B$5/$E$6)/($B$4-$B$5),0)")
        CMT(ws, "C15", _xt("Damodaran: FCFF terminale = EBIT_Y5 x (1+g) x (1-tax) x (1-g/RONIC) — la crescita perpetua costa reinvestimento"))
    else:
        L(ws, "A15", _xt("Terminal Value (Gordon)")); F(ws, "B15", f"=IF($B$4>$B$5,{last}10*(1+$B$5)/($B$4-$B$5),0)")
    L(ws, "A16", _xt("PV Terminal")); F(ws, "B16", f"=B15*{last}11")
    L(ws, "A17", _xt("Enterprise Value"), bold=True); F(ws, "B17", "=B14+B16", bold=True)
    if adjs:
        L(ws, "A18", _lt(f'(-) Debito netto + {len(adjs)} adjustments (bridge, riga {br_total})',f'(-) Net Debt + {len(adjs)} adjustments (bridge, row {br_total})'))
        F(ws, "B18", f"={bridge_ref}")
    else:
        L(ws, "A18", _xt("(-) Net Debt")); F(ws, "B18", "=-B6")
        if spec.get("stance") in ("buy", "sell", "neutral"):
            # stance dichiarata anche SENZA adjustments: il criterio "nessuna voce
            # discrezionale inclusa" e' esso stesso una scelta da documentare
            CMT(ws, "C18", _lt(f"Stance dell'analista: {spec['stance'].upper()} — nessuna voce discrezionale inclusa nel bridge",f"Analyst stance: {spec['stance'].upper()} — no discretionary items included in the bridge"))
    L(ws, "A19", _xt("Equity Value"), bold=True); F(ws, "B19", "=B17+B18", bold=True)
    L(ws, "A20", _xt("FAIR VALUE / SHARE (Gordon)"), bold=True); F(ws, "B20", f"=B19/{sh_ref}", "#,##0.00", bold=True, color=GOLD)
    KEY(ws, 20, 1, 2)
    # A3: secondo TV a exit multiple + cross-check del multiplo implicito (colonne D/E)
    er = base_ref.get("ebitda_row")
    ebitda_y5 = f"'{sh}'!{base_ref['fcols'][-1]}{er}" if er else None
    if ebitda_y5:
        L(ws, "D15", "TV exit (=E5 x EBITDA Y5)"); F(ws, "E15", f"={ebitda_y5}*$E$5")
        L(ws, "D16", _xt("PV TV exit")); F(ws, "E16", f"=E15*{last}11")
        L(ws, "D17", _xt("EV (exit)")); F(ws, "E17", "=B14+E16")
        L(ws, "D19", _xt("Equity (exit)")); F(ws, "E19", f"=E17+{bridge_ref}" if adjs else "=E17-B6")
        L(ws, "D20", _xt("FAIR VALUE / SHARE (exit)"), bold=True); F(ws, "E20", f"=E19/{sh_ref}", "#,##0.00", bold=True, color=GOLD)
        L(ws, "D21", _xt("Multiplo implicito TV Gordon")); F(ws, "E21", f"=IF({ebitda_y5}<>0,B15/{ebitda_y5},0)", "0.00")
        CMT(ws, "F21", _xt("cross-check: se lontano dal multiplo comps, il Gordon e' aggressivo o conservativo"))
    # A4: sensitivity WACC x g a FORMULE VIVE (header numerici modificabili)
    H(ws, "A23", _xt("SENSITIVITY fair value Gordon (VIVA): WACC \\ g"))
    base_w = fv.get("wacc_used") or 0.09; base_g = fv.get("terminal_g_used") or 0.02
    gs = [round(base_g + d, 4) for d in (-0.01, -0.005, 0, 0.005, 0.01)]
    wsz = [round(base_w + d, 4) for d in (-0.015, -0.0075, 0, 0.0075, 0.015)]
    ufcf_refs = [f"'{sh}'!{base_ref['fcols'][t]}{ur}" for t in range(N_FWD)]
    for j, gg in enumerate(gs):
        N(ws, f"{get_column_letter(2 + j)}23", gg, "0.00%", is_input=True)
    for i, ww in enumerate(wsz):
        N(ws, f"A{24 + i}", ww, "0.00%", bold=(ww == base_w), is_input=True)
        wc = f"$A${24 + i}"
        for j in range(len(gs)):
            c = get_column_letter(2 + j)
            gc = f"{c}$23"
            terms = "+".join(f"{ref}/(1+{wc})^({t + 1}-$E$4)" for t, ref in enumerate(ufcf_refs))
            if ronic_v and _er_ebit:
                # V2.2.5: nella sensitivity il reinvestimento SEGUE la g della griglia
                # (g/RONIC): piu' crescita perpetua = piu' reinvestimento, vivo in cella
                tv = f"{_ebit5}*(1+{gc})*(1-$E$7)*(1-{gc}/$E$6)/({wc}-{gc})/(1+{wc})^({N_FWD}-$E$4)"
                _cond = f"AND({wc}>{gc},$E$6>{gc})"  # review B1: g sopra il RONIC = n/a, non un TV negativo "valido"
            else:
                tv = f"{ufcf_refs[-1]}*(1+{gc})/({wc}-{gc})/(1+{wc})^({N_FWD}-$E$4)"
                _cond = f"{wc}>{gc}"
            eq_adj = f"+{bridge_ref}" if adjs else "-$B$6"
            F(ws, f"{c}{24 + i}", f"=IF({_cond},({terms}+{tv}{eq_adj})/{sh_ref},\"n/a\")", "#,##0.00",
              bold=(i == 2 and j == 2), color=GOLD if (i == 2 and j == 2) else INK)

    # ===== A7: BRIDGE TO EQUITY parametrico (solo se l'agente/engine passa voci) =====
    if adjs:
        H(ws, f"A{br0}", _xt("BRIDGE TO EQUITY (EV -> Equity): net debt + adjustments dichiarati"), ncols=3)
        L(ws, f"A{br0 + 1}", _xt("(-) Net Debt")); F(ws, f"B{br0 + 1}", "=-$B$6")
        for k, a in enumerate(adjs):
            rr = br0 + 2 + k
            L(ws, f"A{rr}", a["label"])
            N(ws, f"B{rr}", a["value_m"], "#,##0.0", is_input=True)
            if a.get("commentary"):
                CMT(ws, f"C{rr}", a["commentary"])
        L(ws, f"A{br_total}", _xt("TOTALE bridge (Net debt + adjustments)"), bold=True)
        F(ws, f"B{br_total}", f"=SUM(B{br0 + 1}:B{br_total - 1})", bold=True)
        KEY(ws, br_total, 1, 2)
        if spec.get("stance") in ("buy", "sell", "neutral"):
            L(ws, f"A{br_total + 1}", _lt(f"Stance dell'analista: {spec['stance'].upper()}",f"Analyst stance: {spec['stance'].upper()}"), italic=True, color=GREYTX)
            CMT(ws, f"B{br_total + 1}", _xt("criterio dichiarato di inclusione delle voci discrezionali (lezione Luiss: la discrezionalita' si ordina, non si nasconde)"))

    # ===== A6: MINI-DCF PER SCENARIO a formule vive (Gordon; B5/E4 condivisi) =====
    # B14: riga WACC per scenario — di default =$B$4 (formula, si ricalcola col
    # parametro), input blu solo quando l'agente passa wacc_delta_bp bear/bull;
    # il BASE resta sempre =$B$4 (parita' B20 == colonna base garantita)
    H(ws, f"A{ms}", _xt("MINI-DCF PER SCENARIO (Gordon; g B5, mid-year E4 condivisi; WACC per scenario)"), ncols=4)
    dbp = spec.get("wacc_delta_bp") if isinstance(spec.get("wacc_delta_bp"), dict) else {}
    base_w = _safe(fv.get("wacc_used"), 0.09)
    sc_cols = {"bear": "B", "base": "C", "bull": "D"}
    for name, c in sc_cols.items():
        YH(ws, f"{c}{ms + 1}", name.upper())
    L(ws, f"A{ms + 2}", _xt("WACC scenario"))
    L(ws, f"A{ms + 3}", _xt("Sum PV UFCF"))
    L(ws, f"A{ms + 4}", _xt("PV Terminal"))
    L(ws, f"A{ms + 5}", _xt("Enterprise Value"))
    L(ws, f"A{ms + 6}", _xt("Equity Value"))
    L(ws, f"A{ms + 7}", _xt("FAIR VALUE / SHARE"), bold=True)
    for name, c in sc_cols.items():
        # fonte unica: il WACC di scenario (clampato a g+25bp) arriva dal mirror
        # (fv['wacc_bear'/'wacc_bull'], 6 decimali) — mai ricalcolato qui
        w_cell = _safe(fv.get(f"wacc_{name}")) if name in ("bear", "bull") else None
        if w_cell and abs(w_cell - base_w) > 1e-9:
            N(ws, f"{c}{ms + 2}", w_cell, "0.00%", is_input=True)
        else:
            F(ws, f"{c}{ms + 2}", "=$B$4", "0.00%")
        wc = f"{c}${ms + 2}"
        ref = refs[name]
        urefs = [f"'{ref['sheet']}'!{ref['fcols'][t]}{ref['ufcf_row']}" for t in range(N_FWD)]
        terms = "+".join(f"{u}/(1+{wc})^({t + 1}-$E$4)" for t, u in enumerate(urefs))
        F(ws, f"{c}{ms + 3}", f"={terms}")
        if ronic_v and ref.get("ebit_row"):
            # V2.2.5: stesso TV disciplinato del blocco base, con l'EBIT Y5 dello scenario
            _e5s = f"'{ref['sheet']}'!{ref['fcols'][-1]}{ref['ebit_row']}"
            F(ws, f"{c}{ms + 4}",
              f"=IF(AND({wc}>$B$5,$E$6>$B$5),{_e5s}*(1+$B$5)*(1-$E$7)*(1-$B$5/$E$6)/({wc}-$B$5)/(1+{wc})^({N_FWD}-$E$4),0)")
        else:
            F(ws, f"{c}{ms + 4}", f"=IF({wc}>$B$5,{urefs[-1]}*(1+$B$5)/({wc}-$B$5)/(1+{wc})^({N_FWD}-$E$4),0)")
        F(ws, f"{c}{ms + 5}", f"={c}{ms + 3}+{c}{ms + 4}", bold=True)
        F(ws, f"{c}{ms + 6}", f"={c}{ms + 5}+{bridge_ref}" if adjs else f"={c}{ms + 5}-$B$6")
        F(ws, f"{c}{ms + 7}", f"={c}{ms + 6}/{sh_ref}", "#,##0.00", bold=True, color=GOLD)
    if any(_safe(dbp.get(k)) for k in ("bear", "bull")):
        CMT(ws, f"E{ms + 2}", _xt("wacc_delta_bp dell'analista: bear/bull scontati a WACC diverso dal base (input blu)"))
    KEY(ws, ms + 7, 1, 4)
    L(ws, f"A{ms + 8}", _xt("PONDERATO (25/50/25)"), bold=True)
    F(ws, f"B{ms + 8}", f"=0.25*B{ms + 7}+0.5*C{ms + 7}+0.25*D{ms + 7}", "#,##0.00", bold=True, color=GOLD)
    KEY(ws, ms + 8, 1, 2)
    CMT(ws, f"C{ms + 8}", _xt("stessa ponderazione di probability_weights nel payload; la colonna BASE deve coincidere con B20"))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["D"].width = 26
    # exit_base esposto alla Thesis SOLO con un multiplo reale: con E5=0 il fair
    # value exit e' un numero fuorviante (TV azzerato), meglio non linkarlo.
    # sh_ref/bridge_total/net_debt servono al foglio Comps (B8) per la
    # valutazione implicita con lo STESSO bridge del DCF.
    return {"bear": f"B{ms + 7}", "base": f"C{ms + 7}", "bull": f"D{ms + 7}",
            "weighted": f"B{ms + 8}", "gordon_base": "B20",
            "exit_base": "E20" if (ebitda_y5 and _safe(fv.get("exit_multiple_used"))) else None,
            "sh_ref": sh_ref, "bridge_total": (f"$B${br_total}" if adjs else None),
            "net_debt": "$B$6"}


def _sheet_wacc(wb, spec):
    """Parita' A1+A2 (13/07, audit/09): il WACC deve avere il suo BUILD visibile e
    challengabile (foglio 'Wacc' del template Buy Side; adattamento di
    dcf_buyside.sheet_wacc, v2 oggi in attic). Input BLU modificabili, build CAPM+Hamada a formule
    vive, riconciliazione col WACC usato nel DCF, disclosure dei proxy del motore
    Damodaran, sensitivity Beta x ERP viva."""
    ws = wb.create_sheet("WACC")
    wi = spec.get("wacc_inputs") or {}
    dw = spec.get("_damodaran_wacc") or {}
    _tax = _safe(wi.get("tax"), 0.25)
    kd_pre = (dw.get("kd_aftertax") / (1 - _tax)
              if (_safe(dw.get("kd_aftertax")) and _tax < 1) else wi.get("kd", 0.055))
    TITLE(ws, "A1", _xt("WACC — build CAPM (Hamada), input blu modificabili"), ncols=8)
    L(ws, "A3", _xt("Input"), bold=True)
    rows = [
        # review pre-commit (M4): niente "live" fisso in etichetta — il tier vero
        # (live/LKG/static, con data) sta nella fonte src_rf qui sotto
        (_xt("Risk-free (10Y)"), wi.get("rf", 0.04), "0.00%"),
        (_xt("Equity Risk Premium"), wi.get("erp", 0.05), "0.00%"),
        (_xt("Country Risk Premium"), wi.get("crp", 0.0), "0.00%"),
        (_xt("Beta unlevered"), wi.get("beta_u", 1.0), "0.00"),
        ("D/E", dw.get("de_used") if dw.get("de_used") is not None else wi.get("de", 0.3), "0.00"),
        (_xt("Tax rate"), wi.get("tax", 0.25), "0.00%"),
        (_xt("Kd pre-tax"), kd_pre, "0.00%"),
    ]
    for i, (lab, val, fmt) in enumerate(rows, start=4):
        L(ws, f"A{i}", lab)
        N(ws, f"B{i}", _safe(val, 0.0), fmt, is_input=True)
    L(ws, "A12", _xt("Build (formule vive)"), bold=True)
    L(ws, "A13", _xt("Beta levered (Hamada)")); F(ws, "B13", "=B7*(1+B8*(1-B9))", "0.00")
    L(ws, "A14", _xt("Cost of equity (CAPM)")); F(ws, "B14", "=B4+B13*B5+B6", "0.00%", bold=True)
    L(ws, "A15", _xt("Kd after-tax")); F(ws, "B15", "=B10*(1-B9)", "0.00%")
    L(ws, "A16", _xt("Peso equity")); F(ws, "B16", "=1/(1+B8)", "0.0%")
    L(ws, "A17", _xt("Peso debito")); F(ws, "B17", "=B8/(1+B8)", "0.0%")
    L(ws, "A18", _xt("WACC (build)"), bold=True); F(ws, "B18", "=B16*B14+B17*B15", "0.00%", bold=True)
    KEY(ws, 18, ncols=4)
    L(ws, "A20", _xt("WACC usato nel DCF"), bold=True)
    N(ws, "B20", _safe(wi.get("wacc"), 0.0), "0.00%", bold=True)
    CMT(ws, "C20", _xt("se diverso dal build: numero del motore Damodaran (bottom-up beta peer + synthetic rating)"))
    if dw:
        L(ws, "A22", _xt("Motore Damodaran"), bold=True)
        L(ws, "A23", _xt("Rating sintetico")); L(ws, "B23", str(dw.get("synthetic_rating") or "n/d"))
        L(ws, "A24", _xt("Interest coverage")); N(ws, "B24", _safe(dw.get("interest_coverage_used"), 0.0), "0.00")
        L(ws, "A25", _xt("Default spread")); N(ws, "B25", _safe(dw.get("default_spread"), 0.0), "0.00%")
    notes = dw.get("inputs_note") or {}
    L(ws, "A27", _xt("Fonti degli input (ogni proxy e' DICHIARATO)"), bold=True)
    r = 28
    if notes:
        for k, v in notes.items():
            L(ws, f"A{r}", k)
            CMT(ws, f"B{r}", str(v))
            r += 1
    else:
        # review pre-commit (M3): con motore Damodaran KO le fonti V1 della
        # calibrazione (rf/erp/crp/tax/de + eventuale 'wacc' di riserva) si mostrano
        # comunque — prima usciva un messaggio generico e ormai falso
        _wsrc = wi.get("sources") or {}
        if _wsrc:
            for k, v in _wsrc.items():
                L(ws, f"A{r}", "src_" + str(k))
                CMT(ws, f"B{r}", str(v))
                r += 1
        else:
            CMT(ws, f"A{r}", _xt("motore Damodaran KO e nessuna fonte in wacc_inputs: input dal profilo, NON etichettati"))
            r += 1
    rr0 = r + 2
    L(ws, f"A{rr0}", _xt("Sensitivity WACC: Beta unlevered x ERP (formule vive)"), bold=True)
    betas = [round(_safe(wi.get("beta_u"), 1.0) * f, 2) for f in (0.85, 1.0, 1.15, 1.3)]
    erps = [0.045, 0.05, 0.055, 0.06]
    YH(ws, f"A{rr0 + 1}", "Beta\\ERP")
    for j, e in enumerate(erps):
        YH(ws, f"{get_column_letter(2 + j)}{rr0 + 1}", f"{e:.1%}")
    for k, bta in enumerate(betas):
        rr = rr0 + 2 + k
        N(ws, f"A{rr}", bta, "0.00")
        for j, e in enumerate(erps):
            c = get_column_letter(2 + j)
            F(ws, f"{c}{rr}", f"=$B$16*($B$4+{bta}*(1+$B$8*(1-$B$9))*{e}+$B$6)+$B$17*$B$15", "0.00%")
    ws.column_dimensions["A"].width = 30
    for col in "BCDE":
        ws.column_dimensions[col].width = 13


def _sheet_comps(wb, spec, base_ref=None, dcf_cells=None):
    """B8 (14/07, audit/09): oltre alla tabella multipli, VALUTAZIONE IMPLICITA
    viva: EV = mediana EV/EBITDA (outlier esclusi CON nota, come la media
    selettiva del template) x EBITDA fwd Y1 dello Scenario Base (link rosso),
    equity con lo STESSO bridge del DCF, per-share e premio/sconto vs DCF.
    I multipli mancanti vengono derivati da ev/sales/ebitda grezzi dei peer.
    Ritorna le celle chiave per il football field della Thesis."""
    ws = wb.create_sheet("Comps")
    TITLE(ws, "A1", _xt("Trading comps (peer del sub-settore)"), ncols=6)
    comps = spec.get("comps") or []
    if spec.get("_peer_note"):
        # 16/07: la PROVENIENZA dei peer va detta nel foglio (auto dal sub-settore
        # via yfinance, con l'elenco dichiarato di chi non aveva dati)
        L(ws, "A2", str(spec["_peer_note"]), italic=True, color=GREYTX)
    rows, med, n_excl = _comps_multiples(spec)
    YH(ws, "A3", "Peer"); YH(ws, "B3", "EV/Sales"); YH(ws, "C3", "EV/EBITDA"); YH(ws, "D3", "P/E")
    r = 4
    med_cells = []
    for cdat, rr in zip(comps[:8], rows):
        L(ws, f"A{r}", str(cdat.get("name", "")))
        ev, sales = _safe(cdat.get("ev")), _safe(cdat.get("sales"))
        ev_sales = _safe(cdat.get("ev_sales")) or ((ev / sales) if (ev is not None and sales and sales > 0) else None)
        if ev_sales is not None:
            N(ws, f"B{r}", ev_sales, "0.0\"x\"")
        if rr["m"] is not None:
            N(ws, f"C{r}", rr["m"], "0.0\"x\"", color=GREYTX if rr["excluded"] else INK)
            if rr["excluded"]:
                CMT(ws, f"E{r}", _xt("escluso dalla mediana: outlier (>3x mediana grezza o <=0)"))
            else:
                med_cells.append(f"C{r}")
        pe = _safe(cdat.get("pe"))
        if pe is not None:
            N(ws, f"D{r}", pe, "0.0\"x\"")
        r += 1
    if not comps:
        # 16/07 (feedback PM su un industriale: "non trova comp e peer"): il buco si dichiara PER
        # ESTESO dentro il foglio. NB review: il primo fix scriveva la nota PRIMA di
        # questo early-return preesistente, che la sovrascriveva con la riga corta.
        L(ws, "A4", _xt("Peer NON trovati per questo nome: ne' passati dall'analista nella "
                    "chiamata (peers=[...]) ne' disponibili dalla lista auto del "
                    "sub-settore. Il fair value NON usa il metodo comps: buco dichiarato."),
          italic=True, color=GREYTX)
        ws.column_dimensions["A"].width = 24
        return None
    med_row = r
    L(ws, f"A{r}", _xt("MEDIANA") + (_lt(f' ({n_excl} outlier esclusi)',f' ({n_excl} outliers excluded)') if n_excl else "")
      + ("" if med else _xt(" — SOTTO QUORUM (n<3): informativa, NON usata nel fair value")), bold=True)
    F(ws, f"B{r}", f"=MEDIAN(B4:B{r-1})", "0.0\"x\"", bold=True)
    if med_cells:
        F(ws, f"C{r}", f"=MEDIAN({','.join(med_cells)})", "0.0\"x\"", bold=True)
    F(ws, f"D{r}", f"=MEDIAN(D4:D{r-1})", "0.0\"x\"", bold=True)
    KEY(ws, r, 1, 4)
    r += 2
    # valutazione implicita viva (solo con mediana valida — quorum n>=3, audit/12 V0.6 —
    # e ref del DCF/base): il gate usa la STESSA fonte del payload, mai due verita'
    if not (med and med_cells and base_ref and dcf_cells):
        L(ws, f"A{r}", _xt("(valutazione implicita non calcolabile: multipli validi sotto il "
                       "quorum n>=3 o riferimenti DCF assenti — il fair value usa il proxy "
                       "di settore, dichiarato nel payload)"), italic=True, color=GREYTX)
        ws.column_dimensions["A"].width = 30
        return None
    H(ws, f"A{r}", _xt("VALUTAZIONE IMPLICITA (mediana x EBITDA fwd Y1 Base, bridge del DCF)"), ncols=3)
    ebitda_ref = f"'{base_ref['sheet']}'!{base_ref['fcols'][0]}{base_ref['ebitda_row']}"
    r += 1
    L(ws, f"A{r}", _xt("EBITDA fwd Y1 (Scenario Base)")); F(ws, f"B{r}", f"={ebitda_ref}", link=True); eb_row = r; r += 1
    L(ws, f"A{r}", _xt("EV implicito")); F(ws, f"B{r}", f"=C{med_row}*B{eb_row}"); ev_row = r; r += 1
    bridge = dcf_cells.get("bridge_total")
    if bridge:
        L(ws, f"A{r}", _xt("(-) Net debt + adjustments (bridge DCF)"))
        F(ws, f"B{r}", f"=DCF!{bridge}", link=True)
    else:
        L(ws, f"A{r}", _xt("(-) Net Debt (da DCF)"))
        F(ws, f"B{r}", f"=-DCF!{dcf_cells.get('net_debt', '$B$6')}", link=True)
    r += 1
    L(ws, f"A{r}", _xt("Equity implicito"), bold=True); F(ws, f"B{r}", f"=B{ev_row}+B{r-1}", bold=True); r += 1
    L(ws, f"A{r}", _xt("PER SHARE implicito (comps)"), bold=True)
    F(ws, f"B{r}", f"=B{r-1}/DCF!{dcf_cells.get('sh_ref', '$B$7')}", "#,##0.00", bold=True, color=GOLD)
    KEY(ws, r, 1, 2)
    ps_row = r; r += 1
    if dcf_cells.get("weighted"):
        L(ws, f"A{r}", _xt("Premio/sconto vs DCF ponderato"))
        F(ws, f"B{r}", f"=B{ps_row}/DCF!{dcf_cells['weighted']}-1", "+0.0%;-0.0%")
        r += 1
    ws.column_dimensions["A"].width = 34
    return {"per_share": f"B{ps_row}", "median": f"C{med_row}"}


def _sheet_precedents(wb, spec):
    """B9 (14/07, audit/09): porting del foglio Precedents v2, dati SOLO
    dall'agente via get_valuation (deal M&A con fonte [src: tool] obbligatoria
    nella provenienza). Validazione _safe: EV/EBITDA non numerici o <=0 scartati
    CON conteggio — i dati sporchi del template (celle stringa, EV=0, multipli
    negativi) non si replicano. MEDIAN viva; senza dati, nota esplicita."""
    ws = wb.create_sheet("Precedents")
    TITLE(ws, "A1", _xt("Precedent transactions (M&A, dati dell'analista)"), ncols=6)
    heads = (_xt("Target"), _xt("Acquirer"), _xt("Year"), "EV (m)", "EBITDA (m)", "EV/EBITDA")
    for i, h in enumerate(heads):
        YH(ws, f"{get_column_letter(1 + i)}3", h)
    raw = spec.get("precedents") or []
    deals = _valid_precedents(spec)
    skipped = min(len(raw), 12) - len(deals)
    r = 4
    for t in deals:
        L(ws, f"A{r}", t["target"])
        L(ws, f"B{r}", t["acquirer"])
        if t.get("year"):
            N(ws, f"C{r}", int(t["year"]), "0", is_input=True)
        N(ws, f"D{r}", t["ev"], "#,##0", is_input=True)
        N(ws, f"E{r}", t["ebitda"], "#,##0", is_input=True)
        F(ws, f"F{r}", f"=D{r}/E{r}", "0.0\"x\"")
        r += 1
    if r > 4:
        L(ws, f"A{r}", _xt("MEDIANA"), bold=True)
        F(ws, f"F{r}", f"=MEDIAN(F4:F{r-1})", "0.0\"x\"", bold=True)
        KEY(ws, r, 1, 6)
        if skipped:
            CMT(ws, f"A{r + 1}", _lt(f'{skipped} deal scartati in validazione (EV/EBITDA mancanti o <=0)',f'{skipped} deals rejected during validation (EV/EBITDA missing or <=0)'))
    else:
        L(ws, "A4", _xt("(nessuna transazione fornita dall'analista: passa precedents=[{target, acquirer, "
                    "year, ev_m, ebitda_m}] in get_valuation, cifre dai tool con [src: tool])"), italic=True, color=GREYTX)
        if skipped:
            CMT(ws, "A5", _lt(f'{skipped} deal ricevuti ma scartati in validazione (EV/EBITDA mancanti o <=0)',f'{skipped} deals received but rejected during validation (EV/EBITDA missing or <=0)'))
    for col in "ABCDEF":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["A"].width = 24


@scoped_language
def build_model_v3(spec, out_path, scenarios=None, history=None):
    """Workbook v3 template-parity. Ritorna fair value numerici per scenario + ponderato."""
    if not OPX_OK:
        return {"ok": False, "error": _xt("openpyxl non disponibile")}
    defaults = _default_scenarios(spec, history)
    sc = _merge_scenarios(scenarios, defaults)
    agent_provided = bool(scenarios)

    # ultimo ricavo: da storico XBRL (scalato a milioni se serve) o dallo spec.
    # fix review B (14/07): la SCALA si decide sul massimo di TUTTE le serie della
    # finestra, non solo sull'ultimo ricavo — sec_xbrl fa l'UNIONE degli anni di
    # tutti gli item, quindi l'ultimo anno puo' avere CA/CL ma non revenue; con la
    # vecchia euristica nwc0/hist_rows restavano in unita' piene (fair value garbage).
    years_hist = (history or {}).get("years") or []
    yrs_win = years_hist[-N_HIST:]
    _scan_keys = ("revenue", "cost_of_revenue", "gross_profit", "rnd_expense", "sga_expense",
                  "operating_income", "dep_amort", "current_assets", "current_liabilities")
    _all_raw = ([abs(v) for k in _scan_keys for v in _hist_series(history, k, yrs_win)
                 if v is not None] if years_hist else [])
    scale = 1e6 if (_all_raw and max(_all_raw) > 1e7) else 1.0
    rev_hist_raw = _hist_series(history, "revenue", yrs_win) if years_hist else []
    rev_hist_raw = [(float(v) / scale if v is not None else None) for v in rev_hist_raw]
    last_rev = rev_hist_raw[-1] if (rev_hist_raw and _safe(rev_hist_raw[-1])) else None
    if not last_rev:
        last_rev = _safe(spec.get("revenue_m")) or _safe(spec.get("last_revenue"), 100.0)
        if not any(_safe(x) is not None for x in rev_hist_raw):
            rev_hist_raw = spec.get("revenue_hist_m") or []

    # B13: serie storiche (stessa scala dei ricavi) per i fogli Scenario;
    # B12: NWC reale = current assets - current liabilities dell'ultimo anno XBRL
    hist_rows = {}
    nwc0 = None
    if years_hist:
        for k, src in (("cogs", "cost_of_revenue"), ("gross_profit", "gross_profit"),
                       ("rnd", "rnd_expense"), ("sga", "sga_expense"), ("ebit", "operating_income")):
            serie = _hist_series(history, src, yrs_win)
            if any(v is not None for v in serie):
                hist_rows[k] = [(v / scale if v is not None else None) for v in serie]
        _ebit_s = _hist_series(history, "operating_income", yrs_win)
        _da_s = _hist_series(history, "dep_amort", yrs_win)
        _ebd = [((_ebit_s[i] + _da_s[i]) / scale if (_ebit_s[i] is not None and _da_s[i] is not None) else None)
                for i in range(len(yrs_win))]
        if any(v is not None for v in _ebd):
            hist_rows["ebitda"] = _ebd
        # fix review B: NWC OPERATIVO, non CA-CL totali — la cassa non e' circolante
        # (il debito a breve non e' separabile dai tag XBRL correnti: proxy dichiarato
        # nella nota del foglio). Senza cassa disponibile, fallback CA-CL dichiarato.
        _ca = _hist_series(history, "current_assets", yrs_win)
        _cl = _hist_series(history, "current_liabilities", yrs_win)
        _cash = _hist_series(history, "cash", yrs_win)
        if _ca and _cl and _ca[-1] is not None and _cl[-1] is not None:
            _cash0 = _cash[-1] if (_cash and _cash[-1] is not None) else 0.0
            nwc0 = round((_ca[-1] - _cash0 - _cl[-1]) / scale, 1)

    fv = compute_fair_values_v3(spec, sc, last_rev, nwc0=nwc0)
    fv["_base_numbers"] = fv.pop("_base_numbers", None) or fv.get("_base_numbers")
    # DCF-Q1: evidence describes ACTUAL resolved drivers, never changes the DCF.
    from bellomberg.valuation.dcf_quality import assess_quality
    from bellomberg.valuation.dcf_quality_sheet import append_quality_sheet
    quality_spec = dict(spec, _forecast_years=[datetime.now().year + i for i in range(N_FWD)])
    quality_spec["_valuation_controls"] = {key: fv.get(key) for key in (
        "wacc_used", "terminal_g_used", "ronic_used", "terminal_method", "terminal_warnings",
        "mid_year", "wacc_delta_bp_used", "wacc_bear", "wacc_bull", "equity_adjustments_used",
        "equity_adjustments_total", "method_weights_used")}
    quality = assess_quality(quality_spec, scenarios, sc, spec.get("_analysis_context"),
                             spec.get("_previous_valuation_snapshot"), last_rev=last_rev,
                             nwc0=nwc0, compute_fv=compute_fair_values_v3)
    if spec.get("_previous_snapshot_note"):
        quality["previous_snapshot_note"] = spec["_previous_snapshot_note"]

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    _sheet_historical(wb, spec, history)
    # A6: ref di TUTTI e 3 gli scenari al foglio DCF (mini-DCF bear/base/bull vivi);
    # Comps/Precedents PRIMA della Thesis (che li linka), Thesis sempre a indice 0.
    refs = {}
    for name in ("bear", "base", "bull"):
        refs[name] = _sheet_scenario(wb, spec, name, sc[name], rev_hist_raw, yrs_win,
                                     hist_rows=hist_rows, nwc0=nwc0, last_rev=last_rev)
    dcf_cells = _sheet_dcf(wb, spec, fv, refs)
    _sheet_wacc(wb, spec)
    comps_cells = _sheet_comps(wb, spec, base_ref=refs["base"], dcf_cells=dcf_cells)
    _sheet_precedents(wb, spec)
    _sheet_thesis(wb, spec, fv, sc, history_ok=bool(years_hist),
                  dcf_cells=dcf_cells, comps_cells=comps_cells)
    append_quality_sheet(wb, quality)
    wb["Thesis & Output"]["A3"] = (
        _lt(f"QUALITA: {quality['status']} — vedi Qualita e revisioni; completezza documentale, non validazione economica",f"QUALITY: {quality['status']} — see Qualita e revisioni; documentary completeness, not economic validation"))
    wb["Thesis & Output"]["A3"].font = Font(bold=True, color="9C6500", size=9)
    # P0 17/07: cintura — se il bake COM di dcf_engine fallisse, Excel ricalcola
    # comunque all'apertura (openpyxl non scrive i valori cached delle formule)
    wb.calculation.fullCalcOnLoad = True
    try:
        wb.save(out_path)
    except Exception as e:
        return {"ok": False, "error": f"save: {e}"}
    out = {"ok": True, "path": out_path, "engine": "operating_v3",
           "sheets": [s for s in wb.sheetnames],
           "agent_scenarios": agent_provided,
           "analytical_quality": quality,
           "_timestamp": datetime.now().isoformat(timespec="seconds")}
    out.update({k: v for k, v in fv.items() if not k.startswith("_")})
    # M7 audit/13: la banda sanity margini (non-ciclici) arriva anche al payload —
    # post-merge, quindi gia' decaduta se l'analista ha fornito il gross_margin
    _ms = (sc["base"].get("commentary") or {}).get("_margin_sanity")
    if _ms:
        out["margin_sanity"] = _ms
    # B9: mediana precedents nel payload (stessa validazione del foglio)
    deals = _valid_precedents(spec)
    if deals:
        out["precedents_n"] = len(deals)
        out["precedents_median_ev_ebitda"] = round(_median([d["ev"] / d["ebitda"] for d in deals]), 2)
    if not agent_provided:
        out["_analyst_note"] = (_xt("ATTENZIONE: scenari generati dai dati. Da analista buy-side dovresti passare "
                                "scenarios={bear/base/bull: {revenue_growth:[...], gross_margin:[...], ..., "
                                "commentary:{driver: 'perche''}}} - il modello e' la TUA tesi, commentary inclusa."))
    return out


if __name__ == "__main__":
    import json
    spec = {"ticker": "TEST", "company_name": "TestCo", "currency": "USD", "price": 50.0,
            "shares_m": 100.0, "net_debt": 200.0, "gross_margin": 0.62, "rnd_pct": 0.12,
            "sga_pct": 0.25, "capex_pct": 0.03, "nwc_pct": 0.03, "tax_rate": 0.25,
            "growth_path": [0.15, 0.13, 0.11, 0.09, 0.07],
            "wacc_inputs": {"wacc": 0.092, "rf": 0.043}, "terminal_g": 0.025,
            "_variant_view": "test variant view", "comps": [{"name": "PEER1", "ev_sales": 5.1, "ev_ebitda": 14.2, "pe": 22.0}]}
    history = {"years": [2022, 2023, 2024, 2025],
               "items": {"revenue": {2022: 800e6, 2023: 900e6, 2024: 1020e6, 2025: 1150e6},
                         "net_income": {2022: 60e6, 2023: 75e6, 2024: 90e6, 2025: 110e6}},
               "derived": {"gross_margin": {2024: 0.61, 2025: 0.62}}, "units": {"revenue": "USD"}}
    scen = {"base": {"revenue_growth": [0.16, 0.14, 0.12, 0.10, 0.08],
                     "commentary": {"revenue_growth": "guidance 18%, modello 16% per execution risk"}}}
    # A7: bridge parametrico + stance + fully diluted nel test
    spec["equity_adjustments"] = [
        {"label": "Minority interest", "value_m": -80.0, "commentary": "[src: test] quota terzi"},
        {"label": "Associates (non consolidate)", "value_m": 40.0, "commentary": "[src: test] partecipazione 30%"}]
    spec["stance"] = "buy"
    spec["diluted_shares_m"] = 104.0
    import os, tempfile
    r = build_model_v3(spec, os.path.join(tempfile.gettempdir(), "VAL_V3_TEST.xlsx"), scenarios=scen, history=history)
    print(json.dumps({k: v for k, v in r.items() if not k.startswith("_")}, indent=1, ensure_ascii=False))
