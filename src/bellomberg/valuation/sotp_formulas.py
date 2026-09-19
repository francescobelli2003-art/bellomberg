"""Self-contained SOTP with live child models, ownership and parent-only claims."""
from copy import copy, deepcopy
from openpyxl import Workbook
from openpyxl.formula import Tokenizer
from openpyxl.cell.cell import MergedCell
from bellomberg.core.language import text as tr
from .linked_model import LinkedModel, SCENARIOS, ready, total
from .documented_presentation import _finish, _line, PERCENT


def _translate(expression, names):
    """Rewrite authored sheet references only; never touch literal source strings."""
    tokens = Tokenizer('='+expression).items
    for token in tokens:
        if token.type == 'OPERAND' and token.subtype == 'RANGE' and '!' in token.value:
            sheet, cell = token.value.rsplit('!', 1)
            name = sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet
            if name not in names:
                raise ValueError('Unresolved child formula sheet: '+name)
            token.value = "'"+names[name].replace("'", "''")+"'!"+cell
    return ''.join(t.value for t in tokens)


def _embed(wb, child, index):
    from .linked_dispatch import apply_linked_model
    tmp = Workbook()
    tmp.active.title = 'Valuation'
    if not apply_linked_model(tmp, child) or not hasattr(tmp, '_model_link'):
        tmp.close()
        raise ValueError('Child linked model unavailable: '+str(child.get('method')))
    names = {ws.title: f'C{index}-'+ws.title[:27-len(str(index))] for ws in tmp}
    if len(set(names.values())) != len(names):
        raise ValueError('Child sheet name collision')
    for source in tmp:
        target = wb.create_sheet(names[source.title])
        for row in source:
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                dest = target.cell(cell.row, cell.column, cell.value)
                dest.data_type = cell.data_type
                if cell.data_type == 'f':
                    dest.value = '='+_translate(cell.value[1:], names)
                if cell.has_style:
                    for attr in ('font','fill','border','alignment','protection'):
                        setattr(dest, attr, copy(getattr(cell, attr)))
                    dest.number_format = cell.number_format
                if cell.comment:
                    dest.comment = copy(cell.comment)
        for area in source.merged_cells.ranges:
            target.merge_cells(str(area))
        for key, dimension in source.row_dimensions.items():
            target.row_dimensions[key] = copy(dimension)
        for key, dimension in source.column_dimensions.items():
            target.column_dimensions[key] = copy(dimension)
        for attr in ('sheet_format','sheet_properties','page_setup','page_margins','print_options','views'):
            setattr(target, attr, deepcopy(getattr(source, attr)))
        target.print_area = source.print_area.split('!', 1)[-1] if source.print_area else None
        target.print_title_rows = source.print_title_rows
    link = tmp._model_link
    translated = {kind: {key: [_translate(v, names) for v in value] if isinstance(value, list)
                        else _translate(value, names) for key, value in values.items()} for kind, values in link.items()}
    tmp.close()
    return translated


def apply_sotp_formulas(wb, payload):
    if not ready(payload, {'mixed_business_sotp'}):
        return False
    children = payload.get('child_valuations', {})
    m = LinkedModel(wb, payload, skip_drivers={'children'})
    m.quotation_checks()
    links = {name: _embed(wb, child, i) for i, (name, child) in enumerate(children.items(), 1)}
    shares = m.ref('model','shares')
    parent = {k: m.ref('model','parent_balance',k) for k in ('cash','debt','preferred','other_liabilities','other_assets','common_equity')}
    for s in SCENARIOS:
        ws = m.sheet('SOTP '+s, tr('Somma delle parti — ', 'Sum of the parts — ')+s, [])
        m.check(s, 'Parent shares', shares+'>0')
        for k in ('cash','debt','preferred','other_liabilities'):
            m.check(s, 'Parent '+k, parent[k]+'>=0')
        m.check(s, 'No unsupported parent assets', parent['other_assets']+'=0')
        book = total(m.ref('model','parent_balance','investments_book',name) for name in children)
        m.equal(s, 'Standalone parent accounting', parent['common_equity'],
                f'{book}+{parent["cash"]}-{parent["debt"]}-{parent["preferred"]}-{parent["other_liabilities"]}')
        owned, row = [], 9
        for name, child in children.items():
            link = links[name]
            held = m.ref('model','segments',name,'held_common_shares_m')
            fx = m.ref(s,'segment_fx',name,'to_group_rate')
            denominator, raw = link['shares'][s], link['raw'][s]
            m.check(s, name+' input validity', link['guards'][s]+'="OK"')
            m.check(s, name+' ownership', f'AND({held}>0,{denominator}>0,{held}<={denominator},{raw}>0,{fx}>0)')
            if child['method']=='digital_asset_nav':
                m.check(s, name+' common issued shares', held+'<='+link['refs'][('model','capitalization','basic_shares')])
            if child['financial_currency'] == payload['financial_currency']:
                m.check(s, name+' same currency FX', fx+'=1')
            for call in link['calls'][s]:
                m.check(s, name+' no parent funding call', '('+call+')<=0.000000001')
            _line(ws, row, name, 6, True)
            m.calc(ws,row+1,4,tr('Valore child per azione', 'Child value per share'),raw)
            m.calc(ws,row+2,4,tr('Quota posseduta', 'Ownership'),held+'/'+denominator,PERCENT)
            owned.append(m.calc(ws,row+3,4,tr('Valore della partecipazione', 'Owned equity value'),f'{raw}*{held}*{fx}'))
            row += 6
        # Single claims follow their owner; split instrument inventories must reconcile.
        groups = {}
        parent_debt_claims = []
        for i, claim in enumerate(m.value(s,'claims')['items']):
            if claim['allocation']=='PARENT':
                if claim['role']=='debt':
                    amount=m.ref(s,'claims','items',i,'amount')
                    m.check(s, claim['claim_id']+' settlement',amount+'>=0')
                    parent_debt_claims.append(amount)
                    continue
                ref = parent[claim['role']]
            else:
                target = claim['record_ref']
                ref = '0' if target['driver']=='debt_schedule' and target['path']==['no_debt'] else links[claim['allocation']]['refs'][
                    (target['scenario'],target['driver'],*target['path'])]
            key = (claim['allocation'], ref)
            groups.setdefault(key, []).append((i, claim))
        for (allocation, ref), items in groups.items():
            if len(items)==1:
                m.bind(s,'claims',('items',items[0][0],'amount'),ref)
            else:
                net = allocation!='PARENT' and items[0][1]['record_ref']['driver']=='net_debt'
                terms=[]
                for i, claim in items:
                    amount=m.ref(s,'claims','items',i,'amount')
                    if claim['role']!='equity_adjustments':
                        m.check(s, claim['claim_id']+' amount',amount+'>=0')
                    terms.append(('-' if net and claim['role']=='cash' else '')+amount)
                m.equal(s, allocation+' claim inventory',total(terms),ref)
        settlement_debt = total(parent_debt_claims)
        m.check(s, 'Parent debt settlement covers book', settlement_debt+'>='+parent['debt']+'-0.000000001')
        rate, growth = m.ref(s,'central_discount_rate'), m.ref(s,'central_growth')
        m.check(s, 'Central rate / growth', f'AND({rate}>MAX(0,{growth}),{growth}>-1)')
        costs = []
        for i in range(m.n):
            cost = m.ref(s,'central_cash_cost',i)
            m.check(s, f'Central cash cost {i+1}', cost+'>=0')
            costs.append(f'{cost}/(1+{rate})^{m.period(i)}')
        terminal = f'{m.ref(s,"central_cash_cost",m.n-1)}*(1+{growth})/({rate}-{growth})/(1+{rate})^{m.period(m.n-1)}'
        cost_pv = m.calc(ws,row,4,tr('Valore attuale costi centrali', 'PV central costs'),total([*costs,terminal]))
        equity = m.calc(ws,row+2,4,tr('Capitale azionario del gruppo', 'Group equity value'),
                        f'{total(owned)}+{parent["cash"]}-{settlement_debt}-{parent["preferred"]}-{parent["other_liabilities"]}-{cost_pv}')
        m.results[s] = m.calc(ws,row+3,4,tr('Valore per azione', 'Value per share'),equity+'/'+shares)
        _finish(ws,row+6,6)
    m.finish()
    return True
