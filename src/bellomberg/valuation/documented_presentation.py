"""Reader-facing presentation of an approved FCFF snapshot; no valuation engine."""
from math import ceil, isfinite

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.views import Selection
from openpyxl.utils import get_column_letter as col

from bellomberg.core.language import text as tr
from .market_quote import market_quote_view

NAVY, BLUE, PALE = '17324D', '315C85', 'EDF2F7'
AMOUNT = '#,##0.0;(#,##0.0);"—"'
PRICE = '#,##0.00;(#,##0.00);"—"'
PERCENT = '0.0%;(0.0%);"—"'
LABELS = {
    'revenue': ('Ricavi', 'Revenue'), 'ebitda': ('EBITDA', 'EBITDA'),
    'ebit': ('Risultato operativo (EBIT)', 'Operating profit (EBIT)'),
    'ufcf': ('Flusso di cassa disponibile (FCFF)', 'Unlevered free cash flow (FCFF)'),
    'research_amortization': ('Ammortamento sviluppo capitalizzato', 'Capitalized development amortization'),
    'revenue_growth': ('Crescita dei ricavi', 'Revenue growth'),
    'gross_margin': ('Margine lordo', 'Gross margin'), 'rnd_pct': ('Ricerca / ricavi', 'R&D / revenue'),
    'sga_pct': ('Costi commerciali e generali / ricavi', 'SG&A / revenue'),
    'capdev_pct': ('Sviluppo capitalizzato / ricavi', 'Capitalized development / revenue'),
    'da_tan_pct': ('Ammortamento materiale / ricavi', 'Tangible depreciation / revenue'),
    'tax_rate': ('Aliquota fiscale', 'Tax rate'), 'capex_pct': ('Investimenti / ricavi', 'Capex / revenue'),
    'nwc_pct': ('Capitale circolante / ricavi', 'Working capital / revenue'),
    'wacc': ('Costo del capitale (WACC)', 'Cost of capital (WACC)'),
    'terminal_growth': ('Crescita perpetua', 'Perpetual growth'),
    'terminal_ronic': ('Rendimento dei nuovi investimenti', 'Return on new invested capital'),
    'net_debt': ('Debito netto', 'Net debt'), 'equity_adjustments': ('Rettifiche al capitale azionario', 'Equity adjustments'),
    'shares': ('Azioni diluite', 'Diluted shares'), 'opening_nwc': ('Capitale circolante iniziale', 'Opening working capital'),
    'normalized_ebit': ('EBIT normalizzato', 'Normalized EBIT'),
    'capitalized_research_adjustment': ('Rettifica ricerca capitalizzata', 'Capitalized research adjustment'),
    'cycle_adjustment': ('Rettifica di ciclo', 'Cycle adjustment'),
    'expiring_product_loss': ('Utili persi per prodotti in scadenza', 'Expiring product earnings loss'),
    'replacement_product_income': ('Utili dei prodotti sostitutivi', 'Replacement product earnings'),
    'other_adjustment': ('Altre rettifiche', 'Other adjustments'),
}


def _label(key):
    return tr(*LABELS[key]) if key in LABELS else key


def _number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and isfinite(value)


def _put(ws, row, column, value, fmt=None, bold=False):
    cell = ws.cell(row, column, value if value is not None else tr('n.d.', 'n.a.'))
    if isinstance(cell.value, str):
        cell.data_type = 's'  # Evidence and company names are never executable formulas.
    cell.font = Font(name='Arial', size=10, color=NAVY, bold=bold)
    cell.alignment = Alignment(vertical='center', horizontal='right' if _number(value) else 'left')
    if fmt:
        cell.number_format = fmt
    return cell


def _line(ws, row, text, end=8, section=False, height=28):
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=end)
    c = _put(ws, row, 2, text, bold=section)
    c.alignment = Alignment(vertical='center', wrap_text=True)
    if section:
        for n in range(2, end + 1):
            ws.cell(row, n).fill = PatternFill('solid', fgColor=PALE)
    ws.row_dimensions[row].height = height


def _sheet(wb, name, title, end=8):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 42
    ws.column_dimensions['C'].width = 12
    for n in range(4, end + 1):
        ws.column_dimensions[col(n)].width = 17
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.outlinePr.summaryRight = False
    _line(ws, 2, title, end)
    ws['B2'].font = Font(name='Arial', size=16, color=NAVY, bold=True)
    for n in range(2, end + 1):
        ws.cell(3, n).border = Border(bottom=Side(style='thin', color=BLUE))
    ws.freeze_panes = 'D8'
    ws.print_title_rows = '7:7'
    return ws


def _header(ws, row, values, start=2):
    for n, value in enumerate(values, start):
        c = _put(ws, row, n, value, bold=True)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 32


def _finish(ws, last, end=8):
    for r in range(4, last + 1):
        if ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = 21
    ws.print_area = f'B2:{col(end)}{last}'


def _values(ws, row, driver, unit, values, fmt, count=None):
    bold = driver in ('revenue', 'ufcf', 'normalized_ebit')
    _put(ws, row, 2, _label(driver), bold=bold)
    _put(ws, row, 3, unit)
    for i in range(len(values) if count is None else count):
        _put(ws, row, i + 4, values[i] if i < len(values) else None, fmt)


def _detail(ws, widths):
    """Keep the source evidence verbatim, with enough room for wrapped text."""
    ws.sheet_view.showGridLines = False
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws:
        lines = 1
        for cell in row:
            cell.font = Font(name='Arial', size=10, color=NAVY)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if _number(cell.value):
                cell.number_format = PRICE
            width = ws.column_dimensions[cell.column_letter].width
            lines = max(lines, sum(max(1, ceil(len(s) / max(1, width - 2)))
                                   for s in str(cell.value or '').split('\n')))
        ws.row_dimensions[row[0].row].height = min(409, 15 * lines + 8)
    ws.print_area = f'A1:{col(ws.max_column)}{ws.max_row}'


def present_operating(wb, payload):
    """Keep raw evidence/quote cells and make dated, printable analytical views."""
    evidence = payload['analytical_quality'].get('rows', [])
    inputs = {(r.get('scenario'), r.get('driver')): r.get('values') for r in evidence}
    periods = payload['analytical_quality'].get('snapshot', {}).get('forecast_years', [])
    years = [str(p) for p in periods]
    end = max(8, len(years) + 3)
    scenarios = payload['calculation_details']['scenarios']
    names = dict(zip(('bear', 'base', 'bull'), (tr('Ribassista', 'Bear'), 'Base', tr('Rialzista', 'Bull'))))
    money = tr('mln ', 'million ') + str(payload.get('financial_currency') or tr('n.d.', 'n.a.'))
    usable = payload['valuation_usability']['usable'] is True
    quote = market_quote_view(payload.get('market_quote'),
        as_of=(payload.get('market_quote') or {}).get('acquired_as_of'), usable=usable)
    static = tr('Snapshot statico. Ipotesi diverse richiedono una nuova valutazione.',
                'Static snapshot. Changed assumptions require a new valuation.')
    for key in ('bear', 'base', 'bull'):
        if key in wb:
            # Retain every calculation field, including future/unknown extensions.
            for label, value in wb[key].values:
                wb['Valuation'].append([key + '.' + str(label), value])
            del wb[key]
        ws = _sheet(wb, key, payload['ticker'] + ' — ' + names[key], end)
        _line(ws, 4, static, end, height=32)
        _line(ws, 5, tr('Previsioni annuali; date di fine periodo.', 'Annual forecasts; period-end dates.'), end)
        _header(ws, 7, [tr('Previsioni operative', 'Operating forecast'), tr('Unita', 'Units'), *years])
        data = scenarios.get(key, {})
        for row, driver in enumerate(('revenue', 'ebitda', 'ebit', 'ufcf', 'research_amortization'), 8):
            _values(ws, row, driver, money, data.get('rows', {}).get(driver, []), AMOUNT, len(years))
        _line(ws, 15, tr('Normalizzazione del risultato terminale', 'Terminal earnings normalization'), end, True)
        for row, driver in enumerate(('capitalized_research_adjustment', 'cycle_adjustment',
                'expiring_product_loss', 'replacement_product_income', 'other_adjustment', 'normalized_ebit'), 17):
            _values(ws, row, driver, money, [data.get('terminal_bridge', {}).get(driver)], AMOUNT)
        _line(ws, 25, tr('Risultato dello scenario', 'Scenario result'), end, True)
        _put(ws, 27, 2, tr('Valore per azione in valuta di bilancio', 'Value per share in financial currency'))
        _put(ws, 27, 4, data.get('fair_value_per_share'), PRICE)
        _line(ws, 29, tr('BOZZA: fair value non utilizzabile.', 'DRAFT: fair value not usable.') if not usable else
              tr('Valuta di quotazione: vedi sintesi.', 'Quote currency: see Summary.'), end, height=34)
        _finish(ws, 30, end)

    assumptions = _sheet(wb, 'Assumptions', tr('Ipotesi documentate', 'Documented assumptions'), end)
    _line(assumptions, 4, static, end, height=32)
    row = 7
    drivers = ('revenue_growth', 'gross_margin', 'rnd_pct', 'sga_pct', 'capdev_pct', 'da_tan_pct', 'tax_rate', 'capex_pct', 'nwc_pct')
    for key in ('base', 'bear', 'bull'):
        _header(assumptions, row, [names[key], tr('Unita', 'Units'), *years])
        for driver in drivers:
            row += 1
            _values(assumptions, row, driver, '%', inputs.get((key, driver)) or [], PERCENT, len(years))
        row += 3
    _header(assumptions, row, [tr('Capitale e sconto', 'Capital and discounting'), tr('Unita', 'Units'), *names.values()])
    for driver in ('wacc', 'terminal_growth', 'terminal_ronic', 'net_debt', 'equity_adjustments', 'shares', 'opening_nwc'):
        row += 1
        ratio = driver in ('wacc', 'terminal_growth', 'terminal_ronic')
        values = [inputs.get(('model' if driver in ('shares', 'opening_nwc') else key, driver)) for key in names]
        unit = '%' if ratio else tr('mln azioni', 'million shares') if driver == 'shares' else money
        _values(assumptions, row, driver, unit, values, PERCENT if ratio else PRICE if driver == 'shares' else AMOUNT)
    _line(assumptions, row + 3, tr('Fonti e motivazioni: Qualita e revisioni.',
                                 'Sources and rationales: Qualita e revisioni.'), end, height=32)
    _finish(assumptions, row + 4, end)

    segment_data = inputs.get(('base', 'revenue_build')) or {}
    if segment_data.get('basis') == 'segment_guidance':
        seg = _sheet(wb, 'Segments', tr('Segmenti e ipotesi di crescita', 'Segments and growth assumptions'), end)
        _line(seg, 4, tr('Basi storiche in ', 'Historical bases in ') + money + '. ' + static, end, height=34)
        row = 7
        for key in ('base', 'bear', 'bull'):
            _header(seg, row, [names[key], tr('Base storica', 'Historical base'), *years])
            for item in (inputs.get((key, 'revenue_build')) or {}).get('segments', []):
                row += 1
                _put(seg, row, 2, item.get('label')); _put(seg, row, 3, item.get('base'), AMOUNT)
                for i, value in enumerate(item.get('growth', []), 4):
                    _put(seg, row, i, value, PERCENT)
            row += 3
        _line(seg, row, tr('Previsioni in %. Consolidamento con segno originale.',
                           'Forecasts in %. Consolidation retains its original sign.'), end, height=32)
        _finish(seg, row + 1, end)

    summary = _sheet(wb, 'Summary', payload['ticker'] + ' — ' + tr('Valutazione FCFF', 'FCFF valuation'), 6)
    summary.freeze_panes = None
    summary.sheet_view.selection = [Selection(activeCell='B2', sqref='B2')]
    entity = (inputs.get(('model', 'perimeter')) or {}).get('entity') or payload.get('company') or payload['ticker']
    _line(summary, 4, entity, 6, height=30)
    _line(summary, 5, tr('Fair value utilizzabile', 'Fair value usable') if usable else
          tr('BOZZA — fair value non utilizzabile', 'DRAFT — fair value not usable'), 6, section=True)
    _header(summary, 7, [tr('Scenari di valutazione', 'Valuation scenarios'), tr('Unita', 'Units'), *names.values()])
    for row, label, field, unit, fmt in (
        (8, 'Fair value', 'fair_value_', payload.get('currency'), PRICE),
        (9, tr('Upside sul prezzo osservato', 'Upside to observed price'), 'upside_', '%', PERCENT),
        (11, _label('wacc'), 'wacc', '%', PERCENT), (12, _label('terminal_growth'), 'terminal_growth', '%', PERCENT)):
        _put(summary, row, 2, label, bold=row == 8); _put(summary, row, 3, unit)
        for c, key in enumerate(('bear', 'base', 'bull'), 4):
            value = payload.get(field + key) if field == 'fair_value_' else quote.get(field + key + '_pct') if field == 'upside_' else inputs.get((key, field))
            if field in ('fair_value_', 'upside_') and not usable:
                value = None
            if field == 'upside_' and _number(value):
                value /= 100
            _put(summary, row, c, value, fmt, bold=row == 8)
    _line(summary, 15, tr('Date e riferimenti di mercato', 'Dates and market references'), 6, section=True)
    fields = [(tr('Data della valutazione', 'Model valuation date'), payload.get('valuation_date')),
              (tr('Prezzo del modello', 'Model price'), payload.get('price')),
              (tr('Prezzo osservato', 'Observed price'), quote.get('price')),
              (tr('Ora del prezzo osservato (UTC)', 'Observed quote time (UTC)'), quote.get('observed_at')),
              (tr('Fonte / mercato', 'Source / exchange'), ' / '.join(str(quote.get(k) or tr('n.d.', 'n.a.')) for k in ('source_id', 'exchange')))]
    for row, (caption, value) in enumerate(fields, 17):
        summary.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        _put(summary, row, 2, caption)
        summary.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        _put(summary, row, 5, value, PRICE if _number(value) else None)
    _line(summary, 23, tr('Fair value alla data del modello, non rivalutato alla quotazione osservata. ',
                          'Fair value at the model date, not rolled forward to the observed quote. ') + static, 6, height=46)
    reasons = payload['valuation_usability'].get('reasons') or []
    checks = (payload.get('sanity') or {}).get('scenario_checks', {})
    warnings = [names[k] + ': ' + str(v.get('headline') or v.get('reading') or v.get('severity'))
                for k, v in checks.items() if k in names and v.get('severity') != 'OK']
    _line(summary, 26, tr('Attenzione: ', 'Attention: ') + '; '.join([*reasons, *warnings]) if reasons or warnings else
          tr('Dati e controlli nei fogli finali.', 'Data and checks in the final tabs.'), 6, height=55)
    if usable and all(_number(payload.get('fair_value_' + k)) for k in names):
        chart = BarChart(); chart.type = 'bar'; chart.style = 13
        chart.add_data(Reference(summary, min_col=4, max_col=6, min_row=8, max_row=8), from_rows=True)
        chart.set_categories(Reference(summary, min_col=4, max_col=6, min_row=7, max_row=7))
        chart.title = tr('Valore per azione', 'Value per share') + ' (' + str(payload.get('currency') or '') + ')'
        chart.series[0].graphicalProperties.solidFill = BLUE
        chart.series[0].graphicalProperties.line.solidFill = BLUE
        chart.legend = None; chart.height = 6.5; chart.width = 19
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = True
        chart.dLbls.showCatName = True
        chart.dLbls.showSerName = False
        chart.dLbls.showLegendKey = False
        chart.dLbls.separator = ': '
        chart.dLbls.dLblPos = 'outEnd'
        summary.add_chart(chart, 'B29')
    _finish(summary, 40, 6)
    summary.page_setup.orientation = 'portrait'; summary.page_setup.fitToHeight = 1
    _detail(wb['Valuation'], {'A': 48, 'B': 100})
    _detail(wb['Metodo e dati'], {'A': 32, 'B': 35, 'C': 90})
    _detail(wb['Qualita e revisioni'], {'A': 22, 'B': 65, 'F': 65})
    ordered = ['Summary', 'base', 'bear', 'bull', 'Assumptions', 'Segments', 'Valuation', 'Metodo e dati', 'Qualita e revisioni']
    for index, name in enumerate(n for n in ordered if n in wb):
        wb.move_sheet(wb[name], index - wb.index(wb[name]))
    wb.active = 0
