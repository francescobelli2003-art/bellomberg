"""
Memory layer per il consigliere multi-agent.

Componenti:
- SQLite: dati strutturati (portfolio, trade history, memos, decisions, feedback, specialist reports)
- ChromaDB: semantic search su memo, decision rationale, PM feedback

Tabelle SQLite principali:
  positions          - portafoglio LIVE (sostituisce Excel)
  trade_history      - storico ordini (buy/sell)
  position_prices    - snapshot prezzi (per P/L live)
  memos              - memo settimanali archivio
  decisions          - raccomandazioni del Capo con tracking esito
  specialist_reports - tutti i report degli specialisti per memo (6 dal 15/07/2026, Event Desk = News+Politics fusi)
  llm_usage          - token/costo per agente+round di ogni run (costo EUR, FX e stato: n.d. se non prezzabile)
  pm_feedback        - feedback dell'utente
  themes_tracked     - temi macro evoluti nel tempo
  chat_sessions      - placeholder per UI futuro
  chat_messages      - placeholder per UI futuro

ChromaDB collections:
  memo_chunks        - chunks dei memo storici per semantic search
  decision_rationale - rationale ogni decisione (perche')
  pm_feedback_emb    - feedback PM embedded per semantic search
"""
from bellomberg.core.language import text as _storage_text
from bellomberg.core.presentation import message as _message, join_messages
from bellomberg.core.paths import (CHROMA_PATH as _CHROMA_PATH, DATA_DIR as _DATA_DIR,
                                   PROJECT_ROOT, REPORT_DIR as _REPORT_DIR,
                                   RESEARCH_NOTES_DIR as _RESEARCH_NOTES_DIR,
                                   SQLITE_PATH as _SQLITE_PATH)
import os
import sqlite3
import json
import math
import re
from datetime import date, datetime, timedelta
from contextlib import contextmanager

try:
    import chromadb
    from chromadb.config import Settings
    CHROMA_AVAILABLE = True
except ImportError:
    CHROMA_AVAILABLE = False


# 02/09 (pubblicazione B4): percorso dati ANCORATO a questo file, o a
# BELLOMBERG_DATA_DIR. Prima `DB_DIR = "data"` era relativo alla cwd:
# `python /altra/cartella/bellomberg_api.py` creava un DB VUOTO altrove senza
# dirlo (lezione CLAUDE.md "0 posizioni = path sbagliato"). La junction
# data/ resta lo stesso archivio: core.paths ne risolve la destinazione reale.
REPO_DIR = str(PROJECT_ROOT)
try:  # il .env lo carica config.py, ma price_updater/recover_db/regenerate_memo
      # importano memory_db PRIMA di config (review 02/09, F2): qui non si
      # sovrascrive una variabile gia' presente nell'ambiente (override=False)
    from dotenv import load_dotenv as _load_dotenv
    _load_dotenv(PROJECT_ROOT / ".env")
except ImportError:
    pass
DB_DIR = str(_DATA_DIR)
SQLITE_PATH = str(_SQLITE_PATH)
CHROMA_PATH = str(_CHROMA_PATH)
REPORT_DIR = str(_REPORT_DIR)
RESEARCH_NOTES_DIR = str(_RESEARCH_NOTES_DIR)


# ============================================================
# POLICY: LE PAROLE DEL PM NON SI TAGLIANO A UN LETTERALE (21/08, audit/25)
# ============================================================
# Il testo del PM viveva tagliato a [:120] in quattro punti (piu' [:100], [:300]
# e [:200] altrove) dentro blocchi intitolati "PAROLE DIRETTE DEL PM (VINCOLANTI)".
# Misurato sul DB vero: 28 righe, 2.442 char in tutto, 563 persi (23%) su 8 righe;
# tutte e 28 INTERE costano 3.007 char contro un margine libero di 4.126
# (specialisti) e 6.974 (Capo). Il taglio non serviva a niente e costava il senso:
# I due esempi qui sotto sono INVENTATI (03/09, P4): fino a oggi erano due frasi
# del PM copiate verbatim dal DB, con il numero della decisione e il ticker. La
# forma del guasto e' quella vera, ed e' l'unica cosa che serve capire.
# Una riga perdeva "oppure SE LA CONVINZIONE E' MASSIMA" (un vincolo condizionale
# tagliato prima della condizione diventa un divieto secco), un'altra perdeva
# "quell'area geografica mi piace molto", cioe' l'istruzione di ricerca.
# Peggio del taglio: la riga faceva `[:120] + "\""`, rimettendo la virgoletta di
# CHIUSURA dopo il taglio — il troncamento si travestiva da citazione completa, e
# il Capo nella run non ha tool per recuperare la coda.
MAX_CHAR_FEEDBACK_PM = 2000   # stessa policy di current_facts.MAX_CHAR_TESI: sono le stesse parole
MAX_RIGHE_FEEDBACK_PM = 60    # tetto sulle RIGHE, non sul testo. Oggi ne esistono 28.
# Audit 11/09 (Fable 5.1): una run RIPETUTA (10/09 22:45, memo #54, sette ore dopo la run
# 15:30 del memo #53) si archivia in modo reversibile marcando `memos.notes` e
# `decisions.outcome_notes` con questo prefisso: le sue decisioni escono dalla memoria del
# Capo (non sono ne' pendenti ne' decadute: il PM non le ha mai valutate) e dal track record
# dello scorekeeper (non doppiano le call della stessa settimana). Tutto resta nel DB.
MARCATORE_DUPLICATO = "[DUPLICATO"


def e_duplicato(riga):
    """True se la riga (memo o decisione) porta il marcatore della run ripetuta."""
    if not isinstance(riga, dict):
        return False
    for k in ("notes", "outcome_notes"):
        if str(riga.get(k) or "").lstrip().startswith(MARCATORE_DUPLICATO):
            return True
    return False

_MANDATO_MEMO_RE = re.compile(
    r"^\[MANDATO PM: impronta ([0-9a-f]{64}); origine [^\]\r\n]+; "
    r"dichiarato_il [^\]\r\n]+\]$",
    re.MULTILINE,
)


def _fingerprint_memo(testo):
    """Legge solo il marker canonico del memo, mai occorrenze narrative di «impronta»."""
    match = _MANDATO_MEMO_RE.search(testo or "")
    return match.group(1) if match else None


def _fingerprint_mandato_corrente():
    """Snapshot del mandato corrente per una singola costruzione della memoria."""
    from bellomberg.core import mandato_pm
    try:
        return mandato_pm.impronta(mandato_pm.carica())
    except mandato_pm.MandatoMancante:
        return None


def _etichetta_mandato_storico(testo_memo, fingerprint_corrente):
    storico = _fingerprint_memo(testo_memo)
    if storico is None:
        return ("[MANDATO STORICO: provenienza mandato non disponibile; "
                "non trattare il testo come preferenza PM corrente]")
    if fingerprint_corrente is None:
        return ("[MANDATO STORICO: impronta presente ma mandato corrente non dichiarato; "
                "non trattare il testo come preferenza PM corrente]")
    if storico != fingerprint_corrente:
        return ("[MANDATO STORICO: impronta diversa dal mandato corrente; "
                "non trattare il testo come preferenza PM corrente]")
    return "[MANDATO STORICO: stessa impronta del mandato corrente]"


def _marcatore_troncamento(parts, testo, max_chars):
    """Il marcatore del taglio al cap, che NOMINA le sezioni cadute.

    «coda persa — dichiarato» diceva CHE aveva tagliato ma non COSA, e la prima
    sezione a cadere e' il TRACK RECORD (l'hit-rate delle call dell'agente):
    l'agente perdeva la misura di quanto ha avuto ragione e nessuno poteva
    accorgersene, perche' il marcatore non finisce in nessun log (21/08, audit/25)."""
    taglio = max_chars - 230
    perse = []
    for p in parts:
        if not isinstance(p, str):
            continue
        testa = p.strip()
        if not testa.startswith("---"):
            continue
        if testo.find(testa) >= taglio:
            perse.append(testa.strip("- ").split("\n")[0][:70])
    return ("\n[MEMORIA TRONCATA al cap %d su %d caratteri veri — dichiarato. "
            "SEZIONI PERSE: %s. NON dedurne che non esistano]"
            % (max_chars, len(testo),
               "; ".join(perse) if perse else "coda dell'ultima sezione"))


def pm_verbatim(testo, ident="", virgolette=False, prefisso="",
                fonte="decisions.pm_feedback"):
    """Il testo del PM, intero. Se supera la policy, il taglio si DICHIARA coi
    numeri — e la virgoletta di chiusura NON viene rimessa: un'apertura senza
    chiusura e' il segnale che la citazione non e' finita.

    `prefisso` sta DENTRO le virgolette (serve alla riga del veto, che apre la
    citazione prima del testo): va passato qui e non concatenato fuori, altrimenti
    il chiamante si ritrova a chiudere lui le virgolette su un testo tagliato —
    che e' esattamente il difetto curato."""
    s = (testo or "").strip()
    if len(s) <= MAX_CHAR_FEEDBACK_PM:
        return ('"%s%s"' % (prefisso, s)) if virgolette else (prefisso + s)
    corpo = prefisso + s[:MAX_CHAR_FEEDBACK_PM]
    if virgolette:
        corpo = '"' + corpo
    return (corpo + " [...TAGLIATO QUI: di questo commento del PM arrivano %d caratteri "
            "su %d%s. Il resto e' nel DB (%s) e NON e' stato letto: "
            "non dedurne che il PM non abbia detto altro]"
            % (MAX_CHAR_FEEDBACK_PM, len(s),
               (", decisione #%s" % ident) if ident else "", fonte))


# ============================================================
# HARDENING #32 (audit 01 F-01): helper unico per le connessioni SQLite.
# 5+ processi scrivono lo stesso DB (API, price updater, news feed, briefing,
# consigliere): WAL evita "database is locked" e riduce la finestra di corruzione.
# - busy_timeout=5000 + synchronous=NORMAL ad OGNI connessione
# - journal_mode=WAL una tantum per path/processo (persistente nel file .db)
# NB: WAL crea file -wal/-shm accanto al .db -> junction data/ fuori OneDrive
# resta raccomandata (vedi changelog 32).
# ============================================================
_WAL_DONE: set = set()


def connect_sqlite(path: str = SQLITE_PATH, **kwargs) -> sqlite3.Connection:
    """Apre una connessione sqlite3 con i PRAGMA di robustezza del progetto.
    Da usare al posto di sqlite3.connect() in tutti i siti raw del repo."""
    conn = sqlite3.connect(path, **kwargs)
    try:
        conn.execute("PRAGMA busy_timeout=5000")
        conn.execute("PRAGMA synchronous=NORMAL")
        # riallineamento 23/07 (audit/20): prima le FK erano attive solo in
        # MemoryDB._conn — i writer raw (news, briefing, twr, favorites)
        # giravano senza enforcement referenziale.
        conn.execute("PRAGMA foreign_keys=ON")
        ap = os.path.abspath(path)
        if ap not in _WAL_DONE:
            conn.execute("PRAGMA journal_mode=WAL")
            _WAL_DONE.add(ap)
    except Exception as e:
        # mai bloccare l'apertura per un PRAGMA fallito — ma dichiarato nel log
        # (quick-win audit/21 §6 n.16): un DB senza WAL/FK deve lasciare traccia
        try:
            print(f"[DB] PRAGMA fallito su {os.path.basename(path)}: "
                  f"{type(e).__name__}: {e}", flush=True)
        except OSError:
            pass
    return conn


# ============================================================
# SCHEMAS
# ============================================================

SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS positions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL UNIQUE,
    nome TEXT,
    quantita REAL NOT NULL DEFAULT 0,
    prezzo_medio REAL,
    valuta TEXT DEFAULT 'EUR',
    data_apertura TEXT,
    tesi TEXT,
    temi_monitoraggio TEXT,
    note TEXT,
    last_updated TEXT,
    is_active INTEGER DEFAULT 1
);

CREATE TABLE IF NOT EXISTS trade_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL,
    action TEXT NOT NULL CHECK(action IN ('BUY','SELL','TRIM','ADD','DIVIDEND')),
    quantita REAL NOT NULL,
    prezzo REAL NOT NULL,
    valuta TEXT DEFAULT 'EUR',
    data TEXT NOT NULL,
    note TEXT,
    pm_rationale TEXT,
    linked_decision_id INTEGER,
    created_at TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS position_prices (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL,
    prezzo REAL NOT NULL,
    valuta TEXT DEFAULT 'USD',
    source TEXT,
    timestamp TEXT DEFAULT (datetime('now'))
);
CREATE INDEX IF NOT EXISTS idx_prices_ticker_time ON position_prices(ticker, timestamp);

CREATE TABLE IF NOT EXISTS memos (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    timestamp TEXT NOT NULL,
    title TEXT,
    full_markdown TEXT,
    pdf_path TEXT,
    appendix_path TEXT,
    dcf_files TEXT,
    capo_tokens_in INTEGER,
    capo_tokens_out INTEGER,
    portfolio_nav_eur REAL,
    notes TEXT
);

CREATE TABLE IF NOT EXISTS decisions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    memo_id INTEGER,
    timestamp TEXT NOT NULL,
    action TEXT NOT NULL,
    ticker TEXT,
    eur_amount REAL,
    timing TEXT,
    confidence TEXT,
    rationale TEXT,
    status TEXT DEFAULT 'PENDING' CHECK(status IN ('PENDING','EXECUTED','SKIPPED','EXPIRED','PARTIAL')),
    pm_feedback TEXT,
    outcome_pct REAL,
    outcome_eur REAL,
    outcome_notes TEXT,
    closed_at TEXT,
    FOREIGN KEY(memo_id) REFERENCES memos(id)
);
CREATE INDEX IF NOT EXISTS idx_decisions_status ON decisions(status, timestamp);
CREATE INDEX IF NOT EXISTS idx_decisions_ticker ON decisions(ticker);

CREATE TABLE IF NOT EXISTS specialist_reports (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    memo_id INTEGER,
    specialist TEXT NOT NULL,
    round_n INTEGER NOT NULL,
    content TEXT,
    timestamp TEXT DEFAULT (datetime('now')),
    FOREIGN KEY(memo_id) REFERENCES memos(id)
);
CREATE INDEX IF NOT EXISTS idx_reports_specialist ON specialist_reports(specialist, timestamp DESC);

CREATE TABLE IF NOT EXISTS llm_usage (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    memo_id INTEGER,
    agent TEXT NOT NULL,
    round_n INTEGER,
    model TEXT NOT NULL,
    tokens_in INTEGER DEFAULT 0,
    tokens_out INTEGER DEFAULT 0,
    cache_read INTEGER DEFAULT 0,
    cache_write INTEGER DEFAULT 0,
    api_calls INTEGER DEFAULT 0,
    duration_s REAL,
    cost_eur REAL,
    fx_rate REAL,
    fx_source TEXT,
    cost_status TEXT,
    cache_ttl TEXT,
    timestamp TEXT DEFAULT (datetime('now')),
    FOREIGN KEY(memo_id) REFERENCES memos(id)
);
CREATE INDEX IF NOT EXISTS idx_llm_usage_memo ON llm_usage(memo_id);

CREATE TABLE IF NOT EXISTS pm_feedback (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    memo_id INTEGER,
    decision_id INTEGER,
    specialist TEXT,
    feedback_text TEXT NOT NULL,
    sentiment TEXT CHECK(sentiment IN ('POSITIVE','NEGATIVE','NEUTRAL','SUGGESTION')),
    timestamp TEXT DEFAULT (datetime('now')),
    FOREIGN KEY(memo_id) REFERENCES memos(id),
    FOREIGN KEY(decision_id) REFERENCES decisions(id)
);

CREATE TABLE IF NOT EXISTS themes_tracked (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    theme TEXT NOT NULL,
    first_mentioned TEXT,
    last_mentioned TEXT,
    conviction TEXT,
    status TEXT,
    notes TEXT
);

CREATE TABLE IF NOT EXISTS chat_sessions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    specialist TEXT,
    title TEXT,
    started_at TEXT DEFAULT (datetime('now')),
    last_activity TEXT
);

CREATE TABLE IF NOT EXISTS chat_messages (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    session_id INTEGER NOT NULL,
    role TEXT NOT NULL CHECK(role IN ('user','assistant','system')),
    content TEXT NOT NULL,
    tokens_in INTEGER,
    tokens_out INTEGER,
    timestamp TEXT DEFAULT (datetime('now')),
    FOREIGN KEY(session_id) REFERENCES chat_sessions(id)
);

CREATE TABLE IF NOT EXISTS news_feed (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT NOT NULL,
    snippet TEXT,
    source TEXT,
    url TEXT UNIQUE,
    published_at TEXT,
    pulled_at TEXT DEFAULT (datetime('now')),
    ticker_mentioned TEXT,
    theme TEXT,
    provider TEXT,
    sentiment TEXT,
    sentiment_score REAL,
    relevance INTEGER,
    notified INTEGER DEFAULT 0,
    headline_it TEXT,
    why_matters TEXT
);

CREATE INDEX IF NOT EXISTS idx_news_pulled ON news_feed(pulled_at DESC);
CREATE INDEX IF NOT EXISTS idx_news_ticker ON news_feed(ticker_mentioned);
CREATE INDEX IF NOT EXISTS idx_news_relevance ON news_feed(relevance DESC);
CREATE INDEX IF NOT EXISTS idx_news_alert ON news_feed(notified, relevance);

CREATE TABLE IF NOT EXISTS decision_notes (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    decision_id INTEGER NOT NULL,
    autore TEXT NOT NULL CHECK(autore IN ('PM','AI')),
    testo TEXT NOT NULL,
    timestamp TEXT DEFAULT (datetime('now','localtime')),
    FOREIGN KEY(decision_id) REFERENCES decisions(id)
);
CREATE INDEX IF NOT EXISTS idx_decision_notes ON decision_notes(decision_id, timestamp);

-- V6 guidance pipeline (audit/19, decisioni PM D1-D4 23/07): registro delle
-- guidance SOCIETARIE con fonte OBBLIGATORIA. status in DB: active|superseded;
-- la staleness (D4: oltre valid_until) si CALCOLA in lettura, mai un cron.
CREATE TABLE IF NOT EXISTS company_guidance (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL,
    metric TEXT NOT NULL CHECK(metric IN ('revenue_growth','revenue_abs','eps',
        'ebitda_margin','gross_margin','capex_pct','other')),
    period TEXT NOT NULL,             -- es. FY2026, H2-2026, Q3-2026
    value_low REAL,
    value_mid REAL NOT NULL,
    value_high REAL,
    unit TEXT NOT NULL,               -- pct(frazione) | musd | eps | other-dichiarata
    source_doc TEXT NOT NULL,         -- documento+riferimento (fonte OBBLIGATORIA, D1)
    source_date TEXT NOT NULL,        -- data del documento
    effective_date TEXT NOT NULL,     -- data di registrazione/decorrenza
    valid_until TEXT NOT NULL,        -- D4: prossima trimestrale dal calendar o +120g
    valid_until_source TEXT,          -- da dove viene la scadenza (calendar|fallback)
    status TEXT NOT NULL DEFAULT 'active',   -- active | superseded
    superseded_by INTEGER,            -- id della guidance che l'ha sostituita
    entered_by TEXT,                  -- PM | specialista | chat (D1)
    note TEXT,
    created_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_guidance ON company_guidance(ticker, metric, period, status);

CREATE TABLE IF NOT EXISTS valuation_theses (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL,
    date TEXT NOT NULL,
    price_at_thesis REAL,
    fair_value REAL,
    growth_path TEXT,
    ebitda_margin_target REAL,
    terminal_growth REAL,
    variant_view TEXT,
    engine TEXT,
    subsector TEXT,
    memo_id INTEGER,
    sanity_severity TEXT,   -- audit/12 V0.3: OK/WARN del motore (BLOCK non si salva piu')
    sanity_headline TEXT,   -- il PERCHE' del giudizio sanity, mostrato in F17
    profile_key TEXT        -- audit/13 V1.7a: chiave PULITA del profilo sub-settore
);
CREATE INDEX IF NOT EXISTS idx_theses_ticker ON valuation_theses(ticker, date DESC);

-- Riallineamento 23/07 (audit/20, migrazione 4): tabelle che PRIMA vivevano solo
-- fuori dallo schema canonico — favorite_companies era DDL a runtime in
-- bellomberg_api._fav_db (che resta come cintura idempotente), cash_movements e
-- nav_snapshots nascevano solo da tools/migrations/setup_twr_tables.py. Un DB ricreato da
-- zero ora parte completo. Stessi statement, idempotenti.
CREATE TABLE IF NOT EXISTS favorite_companies (
    ticker TEXT PRIMARY KEY,
    name TEXT,
    sector TEXT,
    industry TEXT,
    note TEXT,
    added_at TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS cash_movements (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    type TEXT NOT NULL CHECK(type IN ('DEPOSIT','WITHDRAWAL')),
    amount_eur REAL NOT NULL CHECK(amount_eur > 0),
    note TEXT,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE INDEX IF NOT EXISTS idx_cash_movements_date ON cash_movements(date);

CREATE TABLE IF NOT EXISTS cash_state (
    singleton_id INTEGER PRIMARY KEY CHECK(singleton_id = 1),
    balance_cents INTEGER NOT NULL,
    updated_at TEXT NOT NULL,
    source TEXT NOT NULL,
    version INTEGER NOT NULL DEFAULT 1
);

CREATE TABLE IF NOT EXISTS nav_snapshots (
    date TEXT PRIMARY KEY,
    nav_total_eur REAL NOT NULL,
    invested_eur REAL,
    cash_eur REAL,
    source TEXT,
    created_at TEXT DEFAULT (datetime('now'))
);

-- IV History (voce quant P1, ok PM 25/07 sera): un'istantanea AL GIORNO per
-- ticker della term structure IV (ATM/RR25/BF25 per expiry, da vol_surface su
-- Polygon) -> base dell'IV Rank vero (percentile storico, iv_history.py).
-- UNIQUE = idempotenza del collector (INSERT OR IGNORE: prima foto del giorno).
-- Stessi statement, idempotenti, nella migrazione 6 per i DB esistenti.
CREATE TABLE IF NOT EXISTS iv_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL,
    snap_date TEXT NOT NULL,
    expiry TEXT NOT NULL,
    days INTEGER NOT NULL,
    atm_iv REAL NOT NULL,
    rr25 REAL,
    bf25 REAL,
    spot REAL,
    spot_source TEXT,
    source TEXT NOT NULL,
    created_at TEXT DEFAULT (datetime('now')),
    UNIQUE(ticker, snap_date, expiry)
);
CREATE INDEX IF NOT EXISTS idx_iv_history_ticker_date ON iv_history(ticker, snap_date);
"""

# ============================================================
# MIGRAZIONI VERSIONATE (DB igiene §9-bis n.6, 21/07)
# ============================================================
# Lista ordinata (version, description, [statements]). Applicate UNA volta da
# _apply_migrations e registrate in schema_version. Regole: version crescente,
# MAI riusare/modificare una migrazione gia' spedita (aggiungerne una nuova);
# statements possibilmente idempotenti (IF NOT EXISTS) per robustezza.
MIGRATIONS = [
    (1, "indici da query reali: trade_history(ticker,data) [WHERE ticker ORDER BY data, "
        "MAX(data) GROUP BY ticker] + pm_feedback(decision_id)/(memo_id) [canale veti 21/07]",
     ["CREATE INDEX IF NOT EXISTS idx_trade_ticker_data ON trade_history(ticker, data)",
      "CREATE INDEX IF NOT EXISTS idx_pm_feedback_decision ON pm_feedback(decision_id)",
      "CREATE INDEX IF NOT EXISTS idx_pm_feedback_memo ON pm_feedback(memo_id)"]),
    (2, "F10 flag veto (opzione A scelta dal PM 22/07 sui mockup): veto ETERNO su una "
        "decisione, motivo obbligatorio, revocabile; il canale della run lo legge per sempre",
     ["ALTER TABLE decisions ADD COLUMN veto INTEGER DEFAULT 0",
      "ALTER TABLE decisions ADD COLUMN veto_reason TEXT",
      "ALTER TABLE decisions ADD COLUMN veto_at TEXT",
      "ALTER TABLE decisions ADD COLUMN veto_revoked_at TEXT"]),
    # V6 guidance (audit/19, ok PM 23/07): la CREATE vive nello SCHEMA_SQL (IF NOT
    # EXISTS, i DB nuovi la prendono da li'); la voce qui DOCUMENTA il cambio in
    # schema_version sui DB esistenti — stessi statement, idempotenti per regola.
    (3, "V6 guidance pipeline (audit/19, decisioni PM D1-D4 23/07): registro "
        "company_guidance con fonte obbligatoria + indice",
     ["""CREATE TABLE IF NOT EXISTS company_guidance (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL,
    metric TEXT NOT NULL CHECK(metric IN ('revenue_growth','revenue_abs','eps',
        'ebitda_margin','gross_margin','capex_pct','other')),
    period TEXT NOT NULL,
    value_low REAL,
    value_mid REAL NOT NULL,
    value_high REAL,
    unit TEXT NOT NULL,
    source_doc TEXT NOT NULL,
    source_date TEXT NOT NULL,
    effective_date TEXT NOT NULL,
    valid_until TEXT NOT NULL,
    valid_until_source TEXT,
    status TEXT NOT NULL DEFAULT 'active',
    superseded_by INTEGER,
    entered_by TEXT,
    note TEXT,
    created_at TEXT
)""",
      "CREATE INDEX IF NOT EXISTS idx_guidance ON company_guidance(ticker, metric, period, status)"]),
    # Riallineamento 23/07 (audit/20, pacchetto "verita' dei numeri" Lotto A):
    # tabelle fuori-schema portate nel canonico + indice query alert. Tutti gli
    # statement sono idempotenti (safe su DB vivo, fresco e backup vecchi); le
    # colonne news_feed.headline_it/why_matters NON stanno qui perche' ALTER non
    # e' idempotente e il DB vivo le ha gia' -> _migrate_news_i18n (legacy-style,
    # guard su PRAGMA table_info) copre i DB vecchi, SCHEMA_SQL i DB nuovi.
    (4, "riallineamento 23/07 (audit/20): favorite_companies + cash_movements + "
        "nav_snapshots nello schema canonico (prima solo _fav_db runtime / "
        "setup_twr_tables.py) + indice news_feed(notified, relevance)",
     ["""CREATE TABLE IF NOT EXISTS favorite_companies (
    ticker TEXT PRIMARY KEY,
    name TEXT,
    sector TEXT,
    industry TEXT,
    note TEXT,
    added_at TEXT DEFAULT (datetime('now'))
)""",
      """CREATE TABLE IF NOT EXISTS cash_movements (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    type TEXT NOT NULL CHECK(type IN ('DEPOSIT','WITHDRAWAL')),
    amount_eur REAL NOT NULL CHECK(amount_eur > 0),
    note TEXT,
    created_at TEXT DEFAULT (datetime('now'))
)""",
      "CREATE INDEX IF NOT EXISTS idx_cash_movements_date ON cash_movements(date)",
      """CREATE TABLE IF NOT EXISTS nav_snapshots (
    date TEXT PRIMARY KEY,
    nav_total_eur REAL NOT NULL,
    invested_eur REAL,
    cash_eur REAL,
    source TEXT,
    created_at TEXT DEFAULT (datetime('now'))
)""",
      "CREATE INDEX IF NOT EXISTS idx_news_alert ON news_feed(notified, relevance)"]),
    # F6 (pacchetto verita' dei numeri Lotto B, 23/07): realized P/L PERSISTITO
    # per trade al momento di SELL/TRIM — prima ricalcolato a parte a ogni giro
    # (portfolio_analytics), zero base durevole per attribution/outcome loop.
    # Pattern migrazione 2 (colonne SOLO qui, non in SCHEMA_SQL: ALTER non
    # idempotente — i DB freschi le prendono da questa migrazione).
    # realized_local = valuta di quotazione; realized_eur = con FX del giorno del
    # trade (n.d. = NULL DICHIARATO, mai un cambio inventato). Backfill storico:
    # tools/migrations/backfill_realized.py (dry-run di default).
    (5, "F6 riallineamento 23/07 (audit/20): realized P/L per trade persistito "
        "(realized_local valuta di quotazione + realized_eur a FX del giorno; "
        "NULL = n.d. dichiarato; backfill via tools/migrations/backfill_realized.py)",
     ["ALTER TABLE trade_history ADD COLUMN realized_local REAL",
      "ALTER TABLE trade_history ADD COLUMN realized_eur REAL"]),
    # Pattern migrazione 3: la CREATE vive anche nello SCHEMA_SQL (IF NOT EXISTS,
    # i DB nuovi la prendono da li'); qui gli stessi statement idempotenti per i
    # DB esistenti + registrazione in schema_version.
    (6, "IV History (voce quant P1, ok PM 25/07 sera): snapshot giornaliero "
        "ATM IV/RR25/BF25 per expiry da vol_surface (Polygon), collector in "
        "iv_history.py agganciato al task PriceUpdater; base dell'IV Rank vero",
     ["""CREATE TABLE IF NOT EXISTS iv_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ticker TEXT NOT NULL,
    snap_date TEXT NOT NULL,
    expiry TEXT NOT NULL,
    days INTEGER NOT NULL,
    atm_iv REAL NOT NULL,
    rr25 REAL,
    bf25 REAL,
    spot REAL,
    spot_source TEXT,
    source TEXT NOT NULL,
    created_at TEXT DEFAULT (datetime('now')),
    UNIQUE(ticker, snap_date, expiry)
)""",
      "CREATE INDEX IF NOT EXISTS idx_iv_history_ticker_date "
      "ON iv_history(ticker, snap_date)"]),
    (7, "cassa operativa atomica: saldo singleton SQLite in centesimi; "
        "bootstrap dati separato e controllato da portfolio.json",
     ["""CREATE TABLE IF NOT EXISTS cash_state (
    singleton_id INTEGER PRIMARY KEY CHECK(singleton_id = 1),
    balance_cents INTEGER NOT NULL,
    updated_at TEXT NOT NULL,
    source TEXT NOT NULL,
    version INTEGER NOT NULL DEFAULT 1
)"""]),
]


# ============================================================
# DB CLASS
# ============================================================

from bellomberg.storage.journal import JOURNAL_MIGRATION
from bellomberg.agents.score_history import SCORE_HISTORY_MIGRATION

MIGRATIONS.extend([JOURNAL_MIGRATION, SCORE_HISTORY_MIGRATION])

# Lotto trade_date (12/09, Fable 5.1): la data del trade la sceglie il PM. Tre colonne
# ADDITIVE su trade_history, pattern migrazione 5 (ALTER solo qui, i DB freschi le
# prendono da questa voce). NULL = riga LEGACY (scritta prima della migrazione), mai
# un default zitto:
#   ora_convenzionale  1 = il PM ha dato la SOLA data e l'ora 12:00:00 e' la convenzione
#                        (la stessa dei trade storici importati senza ora); 0 = ora misurata.
#   link_origin        'explicit' = linked_decision_id scelto dal PM; 'none' = il PM ha
#                        dichiarato «nessuna decisione» (l'inferenza lo salta);
#                        'unknown' = «non so» (l'inferenza resta ammessa).
#   fx_fonte           cambio usato per la cassa: 'storico' (serie daily del giorno del
#                        trade), 'corrente' (live al momento della registrazione),
#                        'identity' (trade in EUR).
TRADE_DATE_MIGRATION = (
    10, "data del trade scelta dal PM (lotto 12/09): ora_convenzionale, link_origin, "
        "fx_fonte su trade_history (NULL = legacy)",
    ["ALTER TABLE trade_history ADD COLUMN ora_convenzionale INTEGER",
     "ALTER TABLE trade_history ADD COLUMN link_origin TEXT",
     "ALTER TABLE trade_history ADD COLUMN fx_fonte TEXT"])
MIGRATIONS.append(TRADE_DATE_MIGRATION)

# A documented holding balance is not a purchase or an external cash flow.
POSITION_OPENINGS_MIGRATION = (
    11, "documented opening holdings, separate from trades and cash",
    ["""CREATE TABLE IF NOT EXISTS position_openings (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ticker TEXT NOT NULL UNIQUE,
        quantita REAL NOT NULL CHECK(quantita > 0),
        prezzo_medio REAL NOT NULL CHECK(prezzo_medio >= 0),
        valuta TEXT NOT NULL,
        as_of TEXT NOT NULL,
        precisione_data TEXT NOT NULL CHECK(precisione_data IN ('day','second')),
        provenienza TEXT NOT NULL,
        nota TEXT,
        created_at TEXT NOT NULL,
        FOREIGN KEY(ticker) REFERENCES positions(ticker)
    )"""])
MIGRATIONS.append(POSITION_OPENINGS_MIGRATION)

OUTPUT_LANGUAGE_MIGRATION = (
    12, "language of new generated output; legacy metadata remains NULL",
    [f"ALTER TABLE {table} ADD COLUMN output_language TEXT CHECK(output_language IN ('it','en'))"
     for table in ("memos", "specialist_reports", "llm_usage", "chat_sessions", "chat_messages")])
MIGRATIONS.append(OUTPUT_LANGUAGE_MIGRATION)


class DataTradeNonValida(ValueError):
    """La data di un trade non e' ISO, e' nel futuro o e' prima del 2000."""


class RicalcoloImpossibile(RuntimeError):
    """Il replay per data dei trade di un ticker non torna (dati incoerenti): niente
    scritture a meta', la transazione va annullata."""


DATA_TRADE_MIN_ANNO = 2000
ORA_CONVENZIONALE = "12:00:00"


def normalizza_data_trade(data, oggi=None):
    """(data_iso_con_ora, ora_convenzionale) per `trade_history.data`.

    Accetta None (= adesso, ora misurata), `YYYY-MM-DD` (= giorno scelto dal PM, ora
    12:00:00 PER CONVENZIONE -> ora_convenzionale 1), un ISO con l'ora (`T` o spazio,
    con o senza secondi -> ora_convenzionale 0) o un datetime. Rifiuta con
    DataTradeNonValida: non ISO, giorno nel FUTURO (rispetto a `oggi`, giorno locale),
    anno prima del 2000. E' il parse che `log_trade` fa SEMPRE (recon 12/09 §4.2: prima
    il DB accettava qualunque stringa e i consumer si difendevano a valle)."""
    oggi = oggi or date.today()
    if data is None:
        return datetime.now().isoformat(timespec="seconds"), 0
    if isinstance(data, datetime):
        dt, sola_data = data, False
    elif isinstance(data, date):
        dt, sola_data = datetime(data.year, data.month, data.day), True
    else:
        s = str(data).strip()
        if not s:
            return datetime.now().isoformat(timespec="seconds"), 0
        try:
            dt = datetime.fromisoformat(s)
        except ValueError:
            raise DataTradeNonValida(
                _storage_text(f"data del trade '{s}' non e' ISO: usa YYYY-MM-DD (ora 12:00 per convenzione) oppure YYYY-MM-DDTHH:MM:SS", f"Trade date '{s}' is not ISO: use YYYY-MM-DD (12:00 by convention) or YYYY-MM-DDTHH:MM:SS"))
        sola_data = len(s) == 10
    if dt.tzinfo is not None:
        dt = dt.astimezone().replace(tzinfo=None)
    if dt.year < DATA_TRADE_MIN_ANNO:
        raise DataTradeNonValida(
            _storage_text(f"data del trade {dt.date().isoformat()} prima del {DATA_TRADE_MIN_ANNO}: refuso sull'anno?", f'Trade date {dt.date().isoformat()} before {DATA_TRADE_MIN_ANNO}: year typo?'))
    if dt.date() > oggi:
        raise DataTradeNonValida(
            _storage_text(f'data del trade {dt.date().isoformat()} nel futuro (oggi {oggi.isoformat()}): un trade si registra dopo averlo eseguito', f'Trade date {dt.date().isoformat()} is in the future (today {oggi.isoformat()}): record a trade after execution'))
    if sola_data:
        return dt.strftime("%Y-%m-%d") + "T" + ORA_CONVENZIONALE, 1
    return dt.isoformat(timespec="seconds"), 0


# Manual override per match nome Excel -> ticker quando fuzzy non basta
NAME_TO_TICKER_OVERRIDES = {
    "hyperliquidstrategies": "PURR",
    "hyperliquid": "PURR",
    "purr": "PURR",
}

def _parse_eur_amount(s):
    """Estrae un importo EUR da testo libero dell'ACTION TABLE.
    Gestisce: suffissi k/mila (x1000), m/mln/milioni (x1e6), mld (x1e9);
    formato europeo (60.000 / 1,3) e US (60,000 / 1.3); range (50-70k -> media);
    simboli e segni. Ritorna float o None.
    """
    if not s:
        return None
    import re as _re
    txt = str(s).strip().lower()
    txt = txt.replace("\u2212", "-")  # minus unicode
    txt = txt.replace("eur", "").replace("\u20ac", "").replace("$", "")
    # moltiplicatore
    mult = 1.0
    if "mld" in txt or "miliard" in txt or _re.search(r"\b(?:billion|bn)\b", txt):
        mult = 1e9
    elif "mln" in txt or "milion" in txt or _re.search(r"\b(?:million|mm)\b", txt):
        mult = 1e6
    elif "mila" in txt or _re.search(r"\bthousand\b", txt) or _re.search(r"\dk\b", txt) or _re.search(r"\dk", txt) or txt.rstrip().endswith("k"):
        mult = 1e3
    elif _re.search(r"\d\s*m\b", txt):
        mult = 1e6
    def _to_float(raw):
        raw = raw.strip()
        neg = raw.startswith("-")
        raw = raw.lstrip("+-")
        if "." in raw and "," in raw:
            if raw.rfind(",") > raw.rfind("."):
                raw = raw.replace(".", "").replace(",", ".")
            else:
                raw = raw.replace(",", "")
        elif "," in raw:
            dec = raw.split(",")[-1]
            raw = raw.replace(",", ".") if len(dec) <= 2 else raw.replace(",", "")
        elif "." in raw:
            dec = raw.split(".")[-1]
            if len(dec) == 3 and mult == 1.0:
                raw = raw.replace(".", "")
        try:
            v = float(raw)
            return -v if neg else v
        except ValueError:
            return None
    # RANGE "50-70k" -> media dei due estremi (il '-' e' separatore, non segno)
    range_m = _re.search(r"(\d[\d.,]*)\s*[-\u2013]\s*(\d[\d.,]*)", txt)
    if range_m:
        a = _to_float(range_m.group(1)); b = _to_float(range_m.group(2))
        if a is not None and b is not None:
            return round((a + b) / 2 * mult, 2)
    # token numerici singoli
    nums = _re.findall(r"[-+]?\d[\d.,]*", txt)
    if not nums:
        return None
    vals = [v for v in (_to_float(n) for n in nums[:2]) if v is not None]
    if not vals:
        return None
    base = sum(vals) / len(vals) if len(vals) > 1 else vals[0]
    return round(base * mult, 2)




# ============================================================
# CASSA OPERATIVA: fonte runtime unica SQLite
# ============================================================
# Il file vive nella radice del repo (runtime, non committato). Prima tre
# lettori lo aprivano con path RELATIVO ("portfolio.json"): due con
# `except: pass` (qui e in portfolio_analytics — col cwd sbagliato la cassa
# valeva 0 zitta e `nav_total_eur` collassava su `nav_eur` senza che nessun
# campo lo dicesse; descritto dalla chat frontend leggendo
# portfolio_analytics.py:577-583, ponte F43 (1)), il terzo (agent_tools)
# tornava `{"error"}`. Il percorso e' ANCORATO a questo file, non al cwd, e lo
# usano anche gli SCRITTORI (trade, movimento di cassa) in bellomberg_api.
PORTFOLIO_JSON_PATH = str(PROJECT_ROOT / "portfolio.json")



class CashNotInitialized(RuntimeError):
    """Il ledger non ha ancora un saldo iniziale misurato."""


class CashConfirmationRequired(ValueError):
    """Existing cash guards with language-independent identifiers for clients."""

    def __init__(self, code, message):
        self.code = code
        super().__init__(message)


def _cash_state_from_conn(conn):
    row = conn.execute(
        "SELECT balance_cents FROM cash_state WHERE singleton_id=1"
    ).fetchone()
    if row is None:
        return {"cash_eur": 0.0, "cash_source": None,
                "cash_source_note": _storage_text('cash_state SQLite non inizializzata: cassa 0 NON misurata', 'SQLite cash_state is not initialized: zero cash has NOT been measured')}
    return {"cash_eur": row[0] / 100.0, "cash_source": "sqlite:cash_state",
            "cash_source_note": None}


def read_cash_state(path=None):
    """Legge la fonte unica SQLite; non usa mai portfolio.json come fallback."""
    with connect_sqlite(path or SQLITE_PATH) as conn:
        return _cash_state_from_conn(conn)


def _cash_balance_cents_from_conn(conn):
    row = conn.execute(
        "SELECT balance_cents FROM cash_state WHERE singleton_id=1"
    ).fetchone()
    if row is None:
        raise CashNotInitialized(
            _storage_text('cash_state SQLite non inizializzata: cassa 0 NON misurata', 'SQLite cash_state is not initialized: zero cash has NOT been measured'))
    return row[0]


def leggi_cassa_portfolio(path=None):
    """Alias transitorio per i consumer esistenti; non legge più JSON."""
    return read_cash_state(path)


# S2: new databases include these append-only tables. Existing databases use the
# separate dry-run/backup migration; saving a thesis never performs DDL.
VALUATION_SNAPSHOT_STATEMENTS = [
    """CREATE TABLE IF NOT EXISTS valuation_snapshots (
        snapshot_id TEXT NOT NULL, generation_id TEXT NOT NULL, ticker TEXT NOT NULL,
        metadata_version INTEGER NOT NULL, payload_json TEXT NOT NULL,
        payload_sha256 TEXT NOT NULL, created_at TEXT NOT NULL,
        PRIMARY KEY(snapshot_id, generation_id))""",
    "CREATE INDEX IF NOT EXISTS idx_valuation_snapshots_ticker ON valuation_snapshots(ticker, created_at)",
    """CREATE TABLE IF NOT EXISTS valuation_snapshot_links (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        snapshot_id TEXT NOT NULL, generation_id TEXT NOT NULL,
        thesis_id INTEGER, decision_id INTEGER, created_at TEXT NOT NULL,
        CHECK ((thesis_id IS NOT NULL) != (decision_id IS NOT NULL)),
        FOREIGN KEY(snapshot_id,generation_id) REFERENCES valuation_snapshots(snapshot_id,generation_id),
        FOREIGN KEY(thesis_id) REFERENCES valuation_theses(id),
        FOREIGN KEY(decision_id) REFERENCES decisions(id))""",
    "CREATE UNIQUE INDEX IF NOT EXISTS idx_valuation_snapshot_thesis ON valuation_snapshot_links(thesis_id)",
    "CREATE INDEX IF NOT EXISTS idx_valuation_snapshot_decision ON valuation_snapshot_links(decision_id, id)",
]
for _snapshot_table in ("valuation_snapshots", "valuation_snapshot_links"):
    for _snapshot_operation in ("UPDATE", "DELETE"):
        VALUATION_SNAPSHOT_STATEMENTS.append(
            f"CREATE TRIGGER IF NOT EXISTS {_snapshot_table}_{_snapshot_operation.lower()}_immutable "
            f"BEFORE {_snapshot_operation} ON {_snapshot_table} "
            "BEGIN SELECT RAISE(ABORT, 'valuation snapshot is immutable'); END")


class MemoryDB:
    """Persistent memory layer per il consigliere."""

    def __init__(self, db_path=None, chroma_path=None):
        # ⚠️ 21/08: qui i default erano `db_path=SQLITE_PATH, chroma_path=CHROMA_PATH`,
        # cioe' LEGATI ALLA DEFINIZIONE della classe. Effetto misurato: il presidio
        # `(a-ter)` di tests/conftest.py, che dal 26/07 redirige `CHROMA_PATH` in
        # tmp, NON funzionava — `__init__.__defaults__` restava `data\chroma` — e
        # ogni giro di suite riscriveva lo store vettoriale di PRODUZIONE del PM
        # (`data/chroma/chroma.sqlite3`, 6,4 MB: sha 67a062c2… -> 1e154d2b… dopo
        # `pytest tests/test_prezzi_freschezza.py`). Cadeva nell'uso NORMALE — chi
        # passa solo `db_path` — non in quello patologico, quindi il tripwire sul DB
        # non lo vedeva. Risolvendoli alla CHIAMATA, le due costanti di modulo
        # tornano a essere la fonte di verita' per chiunque le rediriga.
        self.db_path = SQLITE_PATH if db_path is None else db_path
        self.chroma_path = CHROMA_PATH if chroma_path is None else chroma_path
        db_path, chroma_path = self.db_path, self.chroma_path
        os.makedirs(os.path.dirname(db_path), exist_ok=True)
        os.makedirs(chroma_path, exist_ok=True)
        self._init_sqlite()
        self._init_chroma()

    def _init_sqlite(self):
        with self._conn() as conn:
            fresh = conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' LIMIT 1").fetchone() is None
            conn.executescript(SCHEMA_SQL)
            if fresh:
                for statement in VALUATION_SNAPSHOT_STATEMENTS:
                    conn.execute(statement)
                # D1A (13/09, Claude Opus 5): archivio append-only dei record documentati,
                # stesso patto degli snapshot: i DB nuovi lo ricevono qui, gli esistenti
                # SOLO da tools/migrations/migra_method_records.py (dry-run, backup).
                from bellomberg.storage.method_records_store import crea_tabelle as _crea_archivio_record
                _crea_archivio_record(conn)
                from bellomberg.storage.filing_store import ensure_schema as _crea_archivio_filing
                _crea_archivio_filing(conn)
                conn.commit()
            self._migrate_trade_history_check(conn)
            self._migrate_archive_override(conn)
            self._migrate_news_i18n(conn)
            self._apply_migrations(conn)

    def _apply_migrations(self, conn):
        """DB igiene §9-bis n.6 (21/07): migrazioni VERSIONATE minime. Ogni voce di
        MIGRATIONS gira UNA volta sola e viene registrata in schema_version (audit:
        quando e cosa). Le due _migrate_* legacy sopra restano com'erano (idempotenti
        ad ogni init); da qui in avanti i cambi schema passano da questa lista.
        Una migrazione fallita NON viene registrata e BLOCCA le successive (ordine
        garantito), con errore dichiarato a log — mai marcata applicata a vuoto."""
        conn.execute("""CREATE TABLE IF NOT EXISTS schema_version (
            version INTEGER PRIMARY KEY,
            applied_at TEXT NOT NULL,
            description TEXT NOT NULL
        )""")
        applied = {r[0] for r in conn.execute("SELECT version FROM schema_version")}
        for version, description, statements in MIGRATIONS:
            if version in applied:
                continue
            try:
                # review 23/07 (F10, finding MEDIA-4): transazione ESPLICITA — in
                # autocommit i DDL committavano subito e un crash tra due ALTER
                # lasciava la migrazione a meta' (riavvio: "duplicate column" =
                # init brickato). BEGIN rende atomico statements+registrazione.
                conn.execute("BEGIN IMMEDIATE")
                for stmt in statements:
                    conn.execute(stmt)
                conn.execute(
                    "INSERT INTO schema_version (version, applied_at, description) VALUES (?, ?, ?)",
                    (version, datetime.now().isoformat(timespec="seconds"), description))
                conn.execute("COMMIT")
                print(f"[MEMORY_DB] migrazione {version} applicata: {description}")
            except Exception as e:
                try:
                    conn.execute("ROLLBACK")
                except Exception:
                    pass
                print(f"[MEMORY_DB] ERRORE migrazione {version} ({description}): {e} — "
                      "migrazioni successive SOSPESE (ordine garantito)")
                raise

    def _migrate_archive_override(self, conn):
        """F10-C (PM 17/07): colonna per l'override MANUALE dell'archivio —
        NULL = regole automatiche, 1 = archiviata dal PM, 0 = riportata in pagina
        dal PM (l'override vince sempre sull'automatico). Idempotente, si crea
        da sola al riavvio del backend (stesso rituale di decision_notes)."""
        try:
            cols = {r[1] for r in conn.execute("PRAGMA table_info(decisions)")}
            if cols and "archive_override" not in cols:
                conn.execute("ALTER TABLE decisions ADD COLUMN archive_override INTEGER")
                print("[memory_db] added decisions.archive_override (F10-C)")
        except Exception as e:
            print(f"[memory_db] archive_override migration failed: {e}")

    def _migrate_news_i18n(self, conn):
        """Riallineamento 23/07 (audit/20, schema drift F-04): headline_it e
        why_matters erano scritte da auto_pull_feed ma NON stavano in SCHEMA_SQL —
        un DB ricreato da zero rompeva il feed (le aggiungeva solo il one-shot
        migrate_news_v3, oggi in attic). Ora SCHEMA_SQL le ha nella CREATE (DB
        nuovi); questa guard idempotente copre i DB VECCHI dove news_feed esiste
        gia' senza le colonne. Non sta in MIGRATIONS perche' ALTER non e'
        idempotente e il DB vivo le ha gia' (stesso rituale di archive_override)."""
        try:
            cols = {r[1] for r in conn.execute("PRAGMA table_info(news_feed)")}
            for col in ("headline_it", "why_matters"):
                if cols and col not in cols:
                    conn.execute(f"ALTER TABLE news_feed ADD COLUMN {col} TEXT")
                    print(f"[memory_db] added news_feed.{col} (riallineamento F-04)")
        except Exception as e:
            print(f"[memory_db] news_i18n migration failed: {e}")

    def set_decision_archive(self, decision_id, override):
        """F10-C: imposta l'override archivio (True/False) o lo azzera (None ->
        tornano le regole automatiche). Ritorna False se la decisione non esiste."""
        with self._conn() as conn:
            row = conn.execute("SELECT id FROM decisions WHERE id=?", (decision_id,)).fetchone()
            if not row:
                return False
            val = None if override is None else (1 if override else 0)
            conn.execute("UPDATE decisions SET archive_override=? WHERE id=?", (val, decision_id))
            return True

    def _migrate_trade_history_check(self, conn):
        """Migration: se trade_history esiste con il vecchio CHECK constraint
        senza 'DIVIDEND', rebuild la tabella per supportare DIVIDEND."""
        try:
            row = conn.execute(
                "SELECT sql FROM sqlite_master WHERE type='table' AND name='trade_history'"
            ).fetchone()
            if not row or not row[0]:
                return
            ddl = row[0]
            # If old schema (without DIVIDEND in CHECK), rebuild
            if "CHECK(action IN ('BUY','SELL','TRIM','ADD'))" in ddl and "DIVIDEND" not in ddl:
                conn.executescript("""
                    BEGIN;
                    CREATE TABLE trade_history_new (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        ticker TEXT NOT NULL,
                        action TEXT NOT NULL CHECK(action IN ('BUY','SELL','TRIM','ADD','DIVIDEND')),
                        quantita REAL NOT NULL,
                        prezzo REAL NOT NULL,
                        valuta TEXT DEFAULT 'EUR',
                        data TEXT NOT NULL,
                        note TEXT,
                        pm_rationale TEXT,
                        linked_decision_id INTEGER,
                        created_at TEXT DEFAULT (datetime('now'))
                    );
                    INSERT INTO trade_history_new
                        SELECT id, ticker, action, quantita, prezzo, valuta, data,
                               note, pm_rationale, linked_decision_id, created_at
                        FROM trade_history;
                    DROP TABLE trade_history;
                    ALTER TABLE trade_history_new RENAME TO trade_history;
                    COMMIT;
                """)
                print("[memory_db] migrated trade_history to allow DIVIDEND action")
        except Exception as e:
            print(f"[memory_db] migration check failed: {e}")

    def _init_chroma(self):
        if not CHROMA_AVAILABLE:
            self.chroma_client = None
            self.col_memos = None
            self.col_decisions = None
            self.col_feedback = None
            return
        try:
            self.chroma_client = chromadb.PersistentClient(
                path=self.chroma_path,
                settings=Settings(anonymized_telemetry=False, allow_reset=False),
            )
            self.col_memos = self.chroma_client.get_or_create_collection("memo_chunks")
            self.col_decisions = self.chroma_client.get_or_create_collection("decision_rationale")
            self.col_feedback = self.chroma_client.get_or_create_collection("pm_feedback_emb")
        except Exception as e:
            print("[MemoryDB] ChromaDB init failed: " + str(e))
            self.chroma_client = None
            self.col_memos = None
            self.col_decisions = None
            self.col_feedback = None

    @contextmanager
    def _conn(self):
        conn = connect_sqlite(self.db_path)  # hardening #32: WAL + busy_timeout
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys = ON")
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    # ========================================================
    # PORTFOLIO MANAGEMENT
    # ========================================================

    def add_or_update_position(self, ticker, nome=None, quantita=None, prezzo_medio=None,
                                valuta=None, data_apertura=None, tesi=None,
                                temi_monitoraggio=None, note=None):
        """Crea o aggiorna una posizione. quantita=0 chiude la posizione (soft delete).

        20/08 (Opus 5): `valuta` era `"EUR"` di DEFAULT e veniva riscritta a ogni
        chiamata, non solo quando la si passava — `manage_portfolio.py update --tesi` (oggi in attic)
        e `close` la chiamano senza valuta, quindi aggiornare una tesi portava una quota in GBX
        da GBX a EUR (un carico in pence letto come se fosse in euro, circa cento
        volte il vero) e valeva per 9 posizioni su 28.
        Ora `None` = "non toccare"; sull'INSERT il default dichiarato resta EUR, perche'
        una posizione con valuta NULL verrebbe sommata nel NAV come se fosse in euro
        senza dirlo (regola 14/07)."""
        with self._conn() as conn:
            existing = conn.execute(
                "SELECT * FROM positions WHERE ticker = ?", (ticker,)).fetchone()
            now = datetime.now().isoformat(timespec="seconds")
            if existing:
                fields = []
                values = []
                if nome is not None: fields.append("nome=?"); values.append(nome)
                if quantita is not None: fields.append("quantita=?"); values.append(quantita)
                if prezzo_medio is not None: fields.append("prezzo_medio=?"); values.append(prezzo_medio)
                if valuta: fields.append("valuta=?"); values.append(valuta)
                if data_apertura: fields.append("data_apertura=?"); values.append(data_apertura)
                if tesi is not None: fields.append("tesi=?"); values.append(tesi)
                if temi_monitoraggio is not None:
                    fields.append("temi_monitoraggio=?")
                    values.append(json.dumps(temi_monitoraggio) if isinstance(temi_monitoraggio, list) else temi_monitoraggio)
                if note is not None: fields.append("note=?"); values.append(note)
                fields.append("last_updated=?"); values.append(now)
                if quantita is not None:
                    fields.append("is_active=?"); values.append(1 if quantita > 0 else 0)
                values.append(ticker)
                conn.execute("UPDATE positions SET " + ", ".join(fields) + " WHERE ticker=?", values)
                return existing["id"]
            else:
                cur = conn.execute("""
                    INSERT INTO positions (ticker, nome, quantita, prezzo_medio, valuta,
                                          data_apertura, tesi, temi_monitoraggio, note, last_updated, is_active)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?)
                """, (ticker, nome, quantita or 0, prezzo_medio, valuta or "EUR", data_apertura or now,
                      tesi,
                      json.dumps(temi_monitoraggio) if isinstance(temi_monitoraggio, list) else temi_monitoraggio,
                      note, now, 1 if (quantita or 0) > 0 else 0))
                return cur.lastrowid

    def get_opening_positions(self, ticker=None):
        """Documented balances; their as_of is not an acquisition date."""
        with self._conn() as conn:
            rows = conn.execute("SELECT * FROM position_openings"
                                + (" WHERE ticker=?" if ticker is not None else "")
                                + " ORDER BY created_at,id",
                                ((str(ticker).strip().upper(),) if ticker is not None else ())).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def _normalize_opening(ticker, quantita, prezzo_medio, valuta, as_of, provenienza,
                           nome=None, nota=None):
        for label, value in (("ticker", ticker), ("valuta", valuta), ("provenienza", provenienza), ("as_of", as_of)):
            if not isinstance(value, str) or not value.strip():
                raise ValueError(_storage_text(f'{label} obbligatorio per il saldo iniziale documentato', f'{label} required for the documented opening balance'))
        ticker, valuta = ticker.strip().upper(), valuta.strip().upper()
        if not re.fullmatch(r"[A-Z]{3}", valuta):
            raise ValueError(_storage_text('valuta di quotazione non valida', 'Invalid quote currency'))
        for label, value, positive in (("quantita", quantita, True), ("prezzo_medio", prezzo_medio, False)):
            if (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
                    or value < 0 or (positive and value == 0)):
                raise ValueError(_storage_text(f'{label} non valido per il saldo iniziale', f'{label} invalid for the opening balance'))
        if not math.isfinite(quantita * prezzo_medio):
            raise ValueError(_storage_text('controvalore nativo del saldo iniziale non finito', 'Opening balance native amount is not finite'))
        stamp = as_of.strip()
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[T ].+)?", stamp):
            raise ValueError(_storage_text('as_of deve essere una data ISO esplicita', 'as_of must be an explicit ISO date'))
        when, conventional = normalizza_data_trade(stamp)
        if not conventional and datetime.fromisoformat(when) > datetime.now():
            raise ValueError(_storage_text('as_of nel futuro: il saldo deve essere gia noto', 'as_of is in the future: the balance must already be known'))
        return {"ticker": ticker, "quantita": float(quantita), "prezzo_medio": float(prezzo_medio),
                "valuta": valuta, "as_of": when[:10] if conventional else when,
                "precisione_data": "day" if conventional else "second", "provenienza": provenienza.strip(),
                "nome": nome, "nota": nota}

    def prepare_position_opening(self, **values):
        opening = self._normalize_opening(**values)
        context = self.trade_context(opening["ticker"])
        if context["position"] is not None or context["trades"] or context["opening"] is not None:
            raise RicalcoloImpossibile(_storage_text('posizione, storico o saldo iniziale gia presenti: nessun retrofit implicito', 'Position, history or opening balance already present: no implicit retrofit'))
        return {"opening": opening, "expected_context": context["fingerprint"],
                "position": {k: opening[k] for k in ("ticker", "nome", "quantita", "prezzo_medio", "valuta")}
                            | {"data_apertura": None},
                "cash_delta_eur": 0.0,
                "cash_disponibile_eur": context["cash"]["balance_cents"] / 100 if context["cash"] else None,
                "performance_note": _storage_text('Saldo noto alla data indicata, acquisti precedenti non documentati. Nessun trade o movimento cassa creato; performance da snapshot successivi alla registrazione completa.', 'Balance known at the stated date; earlier purchases are undocumented. No trade or cash movement created; performance starts from snapshots after complete registration.')}

    def create_position_opening(self, *, expected_context, **values):
        from datetime import timezone
        opening = self._normalize_opening(**values)
        with self._conn() as conn:
            conn.execute("BEGIN IMMEDIATE")
            context = self.trade_context(opening["ticker"], _conn=conn)
            if (context["fingerprint"] != expected_context or context["position"] is not None
                    or context["trades"] or context["opening"] is not None):
                raise RicalcoloImpossibile(_storage_text('anteprima del saldo iniziale cambiata: ripeti la conferma', 'Opening balance preview changed: preview and confirm again'))
            created = datetime.now(timezone.utc).isoformat()
            conn.execute("INSERT INTO positions(ticker,nome,quantita,prezzo_medio,valuta,data_apertura,last_updated,is_active) "
                         "VALUES (?,?,?,?,?,NULL,?,1)",
                         (opening["ticker"], opening["nome"], opening["quantita"], opening["prezzo_medio"], opening["valuta"], created))
            cur = conn.execute("INSERT INTO position_openings(ticker,quantita,prezzo_medio,valuta,as_of,precisione_data,provenienza,nota,created_at) "
                               "VALUES (?,?,?,?,?,?,?,?,?)", tuple(opening[k] for k in
                               ("ticker", "quantita", "prezzo_medio", "valuta", "as_of", "precisione_data", "provenienza", "nota")) + (created,))
            return dict(conn.execute("SELECT * FROM position_openings WHERE id=?", (cur.lastrowid,)).fetchone())

    @staticmethod
    def _assert_trade_after_opening(trade, opening):
        if not opening:
            return
        day, start_day = trade["data"][:10], opening["as_of"][:10]
        ambiguous = (opening["precisione_data"] == "day" or trade.get("ora_convenzionale"))
        if day < start_day or (day == start_day and (ambiguous or trade["data"] <= opening["as_of"])):
            raise RicalcoloImpossibile(_storage_text('trade precedente al saldo iniziale o ordine intraday non documentato: as_of descrive il saldo noto, non un acquisto', 'Trade predates the opening balance or intraday order is undocumented: as_of describes the known balance, not a purchase'))

    @staticmethod
    def _apply_position_replay(conn, ticker, trade_id, plan):
        after = plan["dopo"]
        conn.execute("UPDATE positions SET quantita=?,prezzo_medio=?,data_apertura=?,is_active=?,last_updated=? WHERE ticker=?",
                     (after["quantita"], after["prezzo_medio"], after["data_apertura"],
                      int(after["quantita"] > 0), datetime.now().isoformat(timespec="seconds"), ticker))
        for row in plan["updates"]:
            conn.execute("UPDATE trade_history SET realized_local=?,realized_eur=? WHERE id=?",
                         (row["realized_local"], row["realized_eur"], row["id"] if row["id"] is not None else trade_id))

    # soglia di accorciamento oltre la quale si chiede conferma: sotto meta' del
    # testo precedente non e' una correzione, e' una riscrittura (POLICY, PM)
    TESI_TAGLIO_SOSPETTO = 0.5

    def _tesi_history_path(self):
        """Storico delle tesi ACCANTO al DB, non dentro: una tabella nuova sarebbe
        una migrazione, e i task applicherebbero da soli qualunque MIGRATIONS al
        primo tick (decisione PM sulla lettera A/B ancora aperta)."""
        return os.path.join(os.path.dirname(os.path.abspath(self.db_path)),
                            "tesi_history.json")

    def get_tesi(self, ticker):
        """Tesi corrente + storico delle versioni precedenti (piu' recente prima)."""
        t = str(ticker or "").upper().strip()
        with self._conn() as conn:
            row = conn.execute("SELECT tesi, last_updated, is_active FROM positions "
                               "WHERE ticker=?", (t,)).fetchone()
        if row is None:
            return {"error": _storage_text('posizione %s non trovata', 'Position %s not found') % t,
                    "code": "thesis_position_missing"}
        storico = []
        try:
            with open(self._tesi_history_path(), "r", encoding="utf-8") as f:
                storico = json.load(f).get(t, [])
        except FileNotFoundError:
            pass
        except Exception as e:
            # regola 14/07: uno storico illeggibile si DICHIARA, non si finge vuoto
            return {"ticker": t, "tesi": row["tesi"], "last_updated": row["last_updated"],
                    "is_active": bool(row["is_active"]), "storico": None,
                    "storico_errore": "%s: %s" % (type(e).__name__, e)}
        return {"ticker": t, "tesi": row["tesi"], "last_updated": row["last_updated"],
                "is_active": bool(row["is_active"]),
                "storico": list(reversed(storico)), "versioni": len(storico)}

    def update_tesi(self, ticker, tesi, conferma=False, autore=None):
        """Riscrive la TESI di una posizione, e SOLO quella (20/08, richiesta PM:
        "vorrei vedere e modificare le tesi dall'app").

        Non passa da `add_or_update_position` di proposito: quella funzione tocca
        piu' campi e fino a oggi riscriveva la valuta. Qui l'UPDATE nomina il solo
        campo `tesi` (+ last_updated).

        Guardie PRIMA della scrittura — rifiuto = zero scritture:
        · posizione inesistente -> errore
        · tesi non stringa -> errore
        · svuotare una tesi che c'era, o accorciarla sotto meta', chiede `conferma`
          esplicita: e' la cintura contro il salvataggio distratto che cancella gli
          addendum (il PM li modifica a mano dall'app)
        · testo identico -> nessuna scrittura, dichiarato (idempotente)
        La versione precedente finisce SEMPRE nello storico prima di essere
        sostituita: senza, un errore di battitura sarebbe irreversibile.
        """
        t = str(ticker or "").upper().strip()
        if not isinstance(tesi, str):
            return {"error": _storage_text('tesi deve essere testo, non %s', 'Thesis must be text, not %s') % type(tesi).__name__}
        nuova = tesi.strip()
        with self._conn() as conn:
            row = conn.execute("SELECT tesi FROM positions WHERE ticker=? AND is_active=1",
                               (t,)).fetchone()
            if row is None:
                return {"error": _storage_text('posizione %s non trovata o non attiva', 'Position %s not found or inactive') % t,
                        "code": "thesis_position_missing"}
            vecchia = (row["tesi"] or "").strip()

            if nuova == vecchia:
                return {"ok": True, "ticker": t, "invariata": True, "scritture": 0,
                        "nota": _storage_text("testo identico a quello a registro: nessuna scrittura", "Text identical to the stored version: no writes")}
            if vecchia and not nuova and not conferma:
                return {"code": "thesis_empty_confirmation", "error": (_storage_text("GUARDIA TESI: stai SVUOTANDO una tesi di %d caratteri. Se e' voluto rimanda con conferma=true (la versione precedente resta comunque nello storico).", 'THESIS GUARD: you are EMPTYING a thesis of %d characters. If intended, resubmit with conferma=true (the previous version remains in history).') % len(vecchia))}
            if (vecchia and nuova
                    and len(nuova) < len(vecchia) * self.TESI_TAGLIO_SOSPETTO
                    and not conferma):
                return {"code": "thesis_shortening_confirmation", "error": (_storage_text("GUARDIA TESI: la tesi passerebbe da %d a %d caratteri "
                                  "(-%.0f%%). Se stai riscrivendo apposta rimanda con "
                                  "conferma=true; se invece hai perso un addendum, "
                                  "recuperalo da GET /positions/%s/tesi.",
                                  "THESIS GUARD: the thesis would change from %d to %d characters "
                                  "(-%.0f%%). If you are intentionally rewriting it, resubmit with "
                                  "conferma=true; if an addendum was lost, recover it from GET /positions/%s/tesi.")
                                  % (len(vecchia), len(nuova),
                                     100.0 * (1 - len(nuova) / len(vecchia)), t))}

            # storico PRIMA della scrittura: se fallisce, non si sovrascrive nulla
            versione = None
            if vecchia:
                try:
                    p = self._tesi_history_path()
                    try:
                        with open(p, "r", encoding="utf-8") as f:
                            tutto = json.load(f)
                    except FileNotFoundError:
                        tutto = {}
                    tutto.setdefault(t, []).append({
                        "tesi": row["tesi"],
                        "sostituita_il": datetime.now().isoformat(timespec="seconds"),
                        "da": autore or "n.d.",
                        "caratteri": len(row["tesi"] or ""),
                    })
                    os.makedirs(os.path.dirname(p), exist_ok=True)
                    with open(p, "w", encoding="utf-8") as f:
                        json.dump(tutto, f, ensure_ascii=False, indent=1)
                    versione = len(tutto[t])
                except Exception as e:
                    return {"error": (_storage_text('storico tesi NON scritto (%s: %s): la tesi precedente sarebbe andata persa, non sovrascrivo', 'Thesis history NOT saved (%s: %s): the previous thesis would be lost, so it was not overwritten')
                                      % (type(e).__name__, e))}

            now = datetime.now().isoformat(timespec="seconds")
            n = conn.execute("UPDATE positions SET tesi=?, last_updated=? WHERE ticker=?",
                             (nuova, now, t)).rowcount
        out = {"ok": True, "ticker": t, "scritture": n, "caratteri": len(nuova),
               "caratteri_precedenti": len(vecchia), "versione_storico": versione,
               "last_updated": now,
               "nota": (_storage_text("tesi aggiornata; la precedente e' la versione %s dello storico", 'Thesis updated; the previous text is version %s in history')
                        % versione) if versione else _storage_text("tesi scritta (non ce n'era una)", 'Thesis saved (none existed previously)')}
        # 20/08: durante il collaudo ho corrotto io una tesi passandola per la shell
        # ('è' -> 'Ã¨', UTF-8 riletto come latin-1). Il testo era gia' scritto e nessuno
        # l'avrebbe detto: il campo e' libero e il mojibake e' testo valido. Non si
        # RIFIUTA (potrebbe essere una citazione legittima), si DICHIARA — e lo storico
        # tiene la versione buona, quindi e' recuperabile.
        sospette = [s for s in ("Ã¨", "Ã©", "Ã ", "Ã²", "Ã¹", "Ã¬", "â€™", "â€œ", "Ã ")
                    if s in nuova]
        if sospette:
            out["avviso_encoding"] = (
                _storage_text("il testo contiene sequenze tipiche di UTF-8 riletto come latin-1 (%s): probabile testo passato per un canale con la codifica sbagliata. Scritto lo stesso, ma controllalo — la versione precedente e' nello storico.", 'The text contains sequences typical of UTF-8 read as Latin-1 (%s): it likely passed through a channel with the wrong encoding. Saved, but check it — the previous version is in history.') % ", ".join(repr(s) for s in sospette))
        return out

    def log_trade(self, ticker, action, quantita, prezzo, valuta="EUR",
                  data=None, note=None, pm_rationale=None, linked_decision_id=None,
                  ora_convenzionale=None, link_origin=None, fx_fonte=None,
                  _conn=None, _realized_fx="legacy"):
        """Registra storico + posizione senza cambiare il saldo corrente.

        E' il percorso intenzionale per importazioni storiche e fixture. Le
        operazioni correnti devono usare `execute_trade`, che include cash_state
        nella stessa transazione (e, se il trade e' retrodatato, la ricostruzione
        della posizione per data).

        `data`: None = adesso; `YYYY-MM-DD` = giorno scelto dal PM con ora 12:00:00
        PER CONVENZIONE (ora_convenzionale=1, la stessa dei trade storici importati
        senza ora); ISO con ora = ora misurata. Si PARSA SEMPRE (normalizza_data_trade):
        non ISO / futuro / prima del 2000 -> DataTradeNonValida, nessuna scrittura.
        `ora_convenzionale`/`link_origin`/`fx_fonte`: colonne della migrazione 10;
        se non passate, l'ora la decide il parse e le altre restano NULL (legacy).

        Actions:
          BUY/ADD  - increase position, recompute avg cost
          SELL/TRIM - reduce position, deactivate if qty=0
          DIVIDEND - record cash income only (qty stays, no avg cost change).
                     Convention: quantita = num shares that paid, prezzo = EUR per share,
                     valuta = 'EUR'. NAV market value automatically drops via ex-div price (yfinance).
        """
        action = action.upper()
        if action not in ("BUY", "SELL", "TRIM", "ADD", "DIVIDEND"):
            raise ValueError(_storage_text('action deve essere BUY/SELL/TRIM/ADD/DIVIDEND', 'action must be BUY/SELL/TRIM/ADD/DIVIDEND'))
        # Normalizzazione (review 03/08): canonici QUI, non solo all'endpoint —
        # "mstr"/" MSTR" con la SELECT esatta bypassavano la guardia valute e
        # aprivano una riga positions PARALLELA; "usd" passava ma finiva raw in
        # trade_history. DB gia' canonico (0 righe non-upper, misurato 03/08).
        ticker = (ticker or "").strip().upper()
        if not ticker:
            raise ValueError(_storage_text('ticker vuoto', 'Empty ticker'))
        valuta = (valuta or "EUR").strip().upper()
        data, _ora_conv = normalizza_data_trade(data)
        if ora_convenzionale is None:
            ora_convenzionale = _ora_conv
        from contextlib import nullcontext
        with (nullcontext(_conn) if _conn is not None else self._conn()) as conn:
            if _conn is None:
                conn.execute("BEGIN IMMEDIATE")  # standalone import shares the replay's locked state
            baseline_plan = None
            baseline_row = conn.execute("SELECT * FROM position_openings WHERE ticker=?", (ticker,)).fetchone()
            if baseline_row is not None:
                baseline = dict(baseline_row)
                event = {"ticker": ticker, "action": action, "quantita": quantita, "prezzo": prezzo,
                         "valuta": valuta, "data": data, "ora_convenzionale": ora_convenzionale}
                self._assert_trade_after_opening(event, baseline)
                if _conn is None:
                    if _realized_fx == "legacy":
                        _realized_fx = None  # An undocumented acquisition FX is never a live-FX proxy.
                    baseline_plan = self._trade_replay_plan(event, self.trade_context(ticker, _conn=conn),
                                                          {(data[:10], valuta): _realized_fx})
            # GUARDIA VALUTE (blocco cassa/valute 03/08, ponte F22): la guardia
            # dell'endpoint copre solo POST /trade — da qui passano TUTTI i
            # chiamanti (script, import, automazioni) ed e' da qui che il
            # carico MSTR #46/#47 si e' contaminato (un prezzo in EUR dentro un
            # book USD).
            # Confronto contro la STESSA riga su cui il blend scrive, anche
            # inattiva: la riapertura riusa la riga senza aggiornarne la
            # valuta, e un discorde lascerebbe pm e label in valute diverse
            # (il NAV convertirebbe col cambio sbagliato). DIVIDEND esente
            # come in guardia_prezzi (dividendo per azione; un fondo paga USD su
            # quotazione GBX). Prima dell'INSERT: il rifiuto non scrive NULLA.
            pos = conn.execute("SELECT * FROM positions WHERE ticker=?", (ticker,)).fetchone()
            if action != "DIVIDEND" and pos is not None:
                _v_pos = (pos["valuta"] or "").strip().upper()
                _v_trd = valuta
                if _v_pos and _v_trd and _v_trd != _v_pos:
                    raise ValueError(
                        _storage_text(f'GUARDIA VALUTE: valuta del trade ({_v_trd}) diversa da quella della posizione {ticker} ({_v_pos}): rifiutato senza conversione automatica (classe del carico contaminato #46/#47). Reinserisci nella valuta di quotazione della posizione.', f'CURRENCY GUARD: trade currency ({_v_trd}) differs from position {ticker} ({_v_pos}): refused without automatic conversion (contaminated cost basis class #46/#47). Re-enter in the position quote currency.'))
            cur = conn.execute("""
                INSERT INTO trade_history (ticker, action, quantita, prezzo, valuta, data,
                                            note, pm_rationale, linked_decision_id,
                                            ora_convenzionale, link_origin, fx_fonte)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
            """, (ticker, action, quantita, prezzo, valuta, data, note, pm_rationale, linked_decision_id,
                  int(ora_convenzionale), link_origin, fx_fonte))
            trade_id = cur.lastrowid

            # DIVIDEND: only logs income, NO position update
            if action == "DIVIDEND":
                return trade_id

            # Aggiorna position con average-cost basis (BUY/SELL/TRIM/ADD)
            if action in ("BUY", "ADD"):
                if pos:
                    old_q = pos["quantita"] or 0
                    old_pm = pos["prezzo_medio"] or 0
                    new_q = old_q + quantita
                    new_pm = ((old_q * old_pm) + (quantita * prezzo)) / new_q if new_q > 0 else prezzo
                    conn.execute("UPDATE positions SET quantita=?, prezzo_medio=?, last_updated=?, is_active=1 WHERE id=?",
                                  (new_q, new_pm, data, pos["id"]))
                    # 28/08 (review): una riga CHIUSA che si riapre e' una posizione
                    # NUOVA: data_apertura = data del trade. Prima restava quella della
                    # detenzione vecchia per sempre, e la regola «P&L GG dal carico»
                    # (che si fida di data_apertura) non sarebbe mai scattata su un
                    # ri-acquisto. Su una riga APERTA (ADD) la data non si tocca.
                    if old_q <= 0 or not pos["is_active"]:
                        conn.execute("UPDATE positions SET data_apertura=? WHERE id=?",
                                     (data, pos["id"]))
                else:
                    conn.execute("""INSERT INTO positions (ticker, quantita, prezzo_medio, valuta, data_apertura, last_updated, is_active)
                                    VALUES (?,?,?,?,?,?,1)""",
                                  (ticker, quantita, prezzo, valuta, data, data))
            elif action in ("SELL", "TRIM"):
                if pos:
                    old_q = pos["quantita"] or 0
                    new_q = old_q - quantita
                    conn.execute("UPDATE positions SET quantita=?, last_updated=?, is_active=? WHERE id=?",
                                  (max(new_q, 0), data, 1 if new_q > 0 else 0, pos["id"]))
                    # F6 (riallineamento 23/07): realized P/L persistito al momento
                    # del trade — STESSA FORMULA del metodo #164 ma avg dal DB
                    # (positions.prezzo_medio), che puo' differire di centesimi
                    # dall'avg replay di portfolio_analytics/backfill (misurato:
                    # 10/27 ticker, seed/arrotondamenti storici — dichiarato).
                    # Solo la qty realmente posseduta genera realized (excess =
                    # refuso ticker, nessun realized — coerente col replay).
                    avg = pos["prezzo_medio"]
                    if avg is None or old_q <= 0:
                        print(f"[MEMORY_DB] realized n.d. per trade {trade_id} "
                              f"({ticker}): prezzo_medio assente o qty 0 — dichiarato")
                    if avg is not None and old_q > 0:
                        sold = min(quantita, old_q)
                        realized_local = round(sold * (prezzo - avg), 6)
                        realized_eur = None
                        ccy = (valuta or "EUR").upper()
                        if ccy == "EUR":
                            realized_eur = realized_local
                        elif _realized_fx != "legacy":
                            # execute_trade supplies the observed FX before taking the
                            # write lock. None declares a missing historical conversion.
                            if (_realized_fx is not None and math.isfinite(_realized_fx)
                                    and _realized_fx > 0):
                                realized_eur = round(realized_local * _realized_fx, 6)
                        else:
                            try:
                                # FX CORRENTE = FX del giorno del trade nel flusso
                                # normale (data=adesso). Limite dichiarato: un trade
                                # RETRODATATO qui prende l'FX di oggi — per lo
                                # storico vero c'e' tools/migrations/backfill_realized.py
                                # (FX del giorno effettivo del trade).
                                from bellomberg.cli.price_updater import get_fx_to_eur
                                fx = get_fx_to_eur(ccy)
                                if fx:
                                    realized_eur = round(realized_local * fx, 6)
                                # fx assente -> realized_eur resta NULL = n.d.
                                # DICHIARATO (regola 14/07), mai un cambio inventato
                            except Exception:
                                pass
                        try:
                            conn.execute(
                                "UPDATE trade_history SET realized_local=?, realized_eur=? WHERE id=?",
                                (realized_local, realized_eur, trade_id))
                        except Exception as e:
                            if _conn is not None:
                                raise  # atomic trade/cash execution cannot hide a realized write failure
                            # trade valido comunque, buco dichiarato a log con la
                            # causa VERA (review Lotto B: "migrazione mancante?" su
                            # un lock transitorio era una diagnosi fuorviante)
                            _cause = ("migrazione 5 non applicata su questo DB"
                                      if "no column" in str(e).lower() or "has no column" in str(e).lower()
                                      else type(e).__name__ + ": " + str(e)[:120])
                            print(f"[MEMORY_DB] realized non persistito per trade {trade_id}: {_cause}")
            if baseline_plan:
                self._apply_position_replay(conn, ticker, trade_id, baseline_plan)
            return trade_id

    def trade_context(self, ticker, linked_decision_id=None, _conn=None):
        """Read the rows that a trade may change; no price/provider calls or writes."""
        from contextlib import nullcontext
        from hashlib import sha256
        with (nullcontext(_conn) if _conn is not None else self._conn()) as conn:
            row = conn.execute("SELECT * FROM positions WHERE ticker=?", (ticker,)).fetchone()
            cash = conn.execute("SELECT * FROM cash_state WHERE singleton_id=1").fetchone()
            decision = conn.execute("SELECT * FROM decisions WHERE id=?", (linked_decision_id,)).fetchone()
            opening = conn.execute("SELECT * FROM position_openings WHERE ticker=?", (ticker,)).fetchone()
            context = {"position": dict(row) if row else None, "cash": dict(cash) if cash else None,
                       "opening": dict(opening) if opening else None,
                       "decision": dict(decision) if decision else None,
                       "trades": [dict(r) for r in conn.execute(
                           "SELECT * FROM trade_history WHERE ticker=? ORDER BY data,id", (ticker,))]}
        context["fingerprint"] = sha256(json.dumps(context, sort_keys=True, allow_nan=False,
                                                   separators=(",", ":")).encode()).hexdigest()
        return context

    @classmethod
    def _validate_trade_decision(cls, trade, context):
        did, origin = trade.get("linked_decision_id"), trade.get("link_origin")
        if origin not in (None, "explicit", "none", "unknown"):
            raise ValueError(_storage_text('link_origin non valido', 'Invalid link_origin'))
        if did is None:
            if origin == "explicit":
                raise ValueError(_storage_text('legame esplicito senza decisione', 'Explicit link without a decision'))
            return None
        if origin in ("none", "unknown"):
            raise ValueError(_storage_text('decisione esplicita e senza decisione/non so incompatibili', 'Explicit decision is incompatible with no decision/unknown'))
        decision = context["decision"]
        if not decision:
            raise ValueError(_storage_text(f'decisione {did} inesistente', f'Decision {did} does not exist'))
        if str(decision["ticker"]).strip().upper() != trade["ticker"]:
            raise ValueError(_storage_text(f"decisione {did} riguarda {decision['ticker']}, non {trade['ticker']}", f"Decision {did} concerns {decision['ticker']}, not {trade['ticker']}"))
        verso = cls._VERSO.get(trade["action"])
        if verso is None or verso != cls._VERSO.get(decision["action"]):
            raise ValueError(_storage_text(f'decisione {did}: verso incompatibile con il trade', f'Decision {did}: direction incompatible with the trade'))
        if decision["status"] not in ("PENDING", "PARTIAL", "EXECUTED") or decision.get("veto"):
            raise ValueError(_storage_text(f"decisione {did}: stato {decision['status']} o veto incompatibile", f"Decision {did}: status {decision['status']} or veto incompatible"))
        try:
            t_dec = datetime.fromisoformat(decision["timestamp"])
            t_trade = datetime.fromisoformat(trade["data"])
        except (ValueError, TypeError):
            raise ValueError(_storage_text(f'decisione {did}: data non valida, legame non verificabile', f'Decision {did}: invalid date, link cannot be verified'))
        # With a date-only trade the intraday ordering is not known, not noon.
        before = (t_trade.date() < t_dec.date() if trade.get("ora_convenzionale")
                  else t_trade < t_dec)
        if before:
            raise ValueError(_storage_text(f'trade precedente alla decisione {did}: legame rifiutato', f'Trade predates decision {did}: link refused'))
        return {"id": did, "status": decision["status"],
                "nota": (_storage_text('Decisione PENDING: il trade non cambia lo stato scelto dal PM.', 'PENDING decision: the trade does not change the status chosen by the PM.')
                         if decision["status"] == "PENDING" else None)}

    @classmethod
    def _trade_replay_plan(cls, trade, context, realized_fx=None):
        """Strict chronological average cost for backdating or documented balances.

        No opening position is inferred from an incomplete trade history. FX is supplied
        by the caller; absent historical FX leaves realized_eur unavailable.
        """
        opening = context.get("opening")
        cls._assert_trade_after_opening(trade, opening)
        old = []
        for row in context["trades"]:
            try:
                if not isinstance(row["data"], str) or not row["data"].strip():
                    raise DataTradeNonValida(_storage_text('data del trade storico assente', 'Historical trade date missing'))
                when, _ = normalizza_data_trade(row["data"])
            except (ValueError, TypeError) as exc:
                raise RicalcoloImpossibile(_storage_text(f'replay impossibile: data legacy non valida ({exc})', f'Replay impossible: invalid legacy date ({exc})')) from exc
            old.append(dict(row, data=when))
            cls._assert_trade_after_opening(old[-1], opening)
        following = [r for r in old if r["data"] > trade["data"]]
        if (not following and not opening) or trade["action"] == "DIVIDEND":
            return None
        position = context["position"]
        qty_before = (opening["quantita"] if opening else 0) + sum(r["quantita"] * (1 if r["action"] in ("BUY", "ADD") else -1)
                         for r in old if r["action"] != "DIVIDEND")
        if not position or not math.isclose(qty_before, position["quantita"], abs_tol=1e-7):
            raise RicalcoloImpossibile(_storage_text('replay impossibile: storico non riconciliato con la posizione; manca una base iniziale documentata', 'Replay impossible: history does not reconcile with the position; a documented opening balance is missing'))
        qty, avg, opened = (opening["quantita"], opening["prezzo_medio"], None) if opening else (0.0, 0.0, None)
        if opening:
            check_qty, check_avg = qty, avg
            for row in old:
                if row["action"] == "DIVIDEND":
                    continue
                q, price = row["quantita"], row["prezzo"]
                if row["valuta"] != opening["valuta"] or not math.isfinite(q) or q <= 0 or not math.isfinite(price) or price <= 0:
                    raise RicalcoloImpossibile(_storage_text('storico del saldo iniziale non valido', 'Invalid opening balance history'))
                if row["action"] in ("BUY", "ADD"):
                    check_avg = (check_qty * check_avg + q * price) / (check_qty + q)
                    check_qty += q
                elif row["action"] in ("SELL", "TRIM") and q <= check_qty + 1e-9:
                    check_qty = max(0.0, check_qty - q)
                else:
                    raise RicalcoloImpossibile(_storage_text('storico del saldo iniziale non riconciliabile', 'Opening balance history cannot be reconciled'))
            if (position["valuta"] != opening["valuta"] or trade["valuta"] != opening["valuta"]
                    or position["prezzo_medio"] is None or not math.isclose(check_avg, position["prezzo_medio"], abs_tol=1e-7)):
                raise RicalcoloImpossibile(_storage_text('costo/valuta della posizione divergono dal saldo iniziale documentato e dai trade', 'Position cost/currency diverge from the documented opening balance and trades'))
        updates, notes = [], []
        unknown_opening_fx = bool(opening and opening["valuta"] != "EUR")
        current = dict(trade, id=None)
        all_rows = sorted([*old, current], key=lambda r: (r["data"], r["id"] if r["id"] is not None else float("inf")))
        for row in all_rows:
            action, q, price = row["action"], row["quantita"], row["prezzo"]
            if action == "DIVIDEND":
                continue
            if (not math.isfinite(q) or q <= 0 or not math.isfinite(price) or price <= 0
                    or row["valuta"] != trade["valuta"]):
                raise RicalcoloImpossibile(_storage_text('replay impossibile: quantita/prezzo/valuta dello storico non validi', 'Replay impossible: invalid historical quantity/price/currency'))
            if action in ("BUY", "ADD"):
                if row["id"] is None and action == "ADD" and qty <= 1e-9:
                    raise RicalcoloImpossibile(_storage_text('replay impossibile: ADD senza posizione alla data del trade; usa BUY', 'Replay impossible: ADD without a position on the trade date; use BUY'))
                if qty <= 1e-9:
                    opened = row["data"]
                    avg, qty = 0.0, 0.0
                avg = (qty * avg + q * price) / (qty + q)
                qty += q
            elif action in ("SELL", "TRIM"):
                if q > qty + 1e-9:
                    raise RicalcoloImpossibile(
                        _storage_text(f"replay impossibile: vendita {row['data']} di {q:g} oltre {qty:g} detenute", f"Replay impossible: sale {row['data']} of {q:g} exceeds {qty:g} held"))
                local = round(q * (price - avg), 6)
                if row["id"] is None or row in following:
                    rate = 1.0 if row["valuta"] == "EUR" else (realized_fx or {}).get(
                        (row["data"][:10], row["valuta"]))
                    eur = round(local * rate, 6) if rate is not None else None
                    if unknown_opening_fx:
                        eur = None
                        notes.append(_storage_text('Saldo iniziale estero: FX di acquisto ignoto, realized EUR n.d.; realized locale noto.', 'Foreign opening balance: acquisition FX unknown, realized EUR unavailable; local realized P&L known.'))
                    elif eur is None:
                        notes.append(_storage_text(
                            f"FX storico {row['valuta']} del {row['data'][:10]} n.d.: realized EUR n.d.",
                            f"Historical {row['valuta']} FX for {row['data'][:10]} unavailable: EUR realized P&L unavailable."))
                    updates.append({"id": row["id"], "realized_local": local, "realized_eur": eur})
                qty = max(0.0, qty - q)
                if qty <= 1e-9:
                    unknown_opening_fx = False
            else:
                raise RicalcoloImpossibile(_storage_text('replay impossibile: azione storica sconosciuta', 'Replay impossible: unknown historical action'))
        realized_before = [r["realized_local"] for r in old if r["action"] in ("SELL", "TRIM")]
        changed = {r["id"]: r["realized_local"] for r in updates}
        realized_after = [changed.get(r["id"], r["realized_local"]) for r in old
                          if r["action"] in ("SELL", "TRIM")]
        if None in changed:
            realized_after.append(changed[None])
        def total(values):
            return round(sum(values), 6) if all(v is not None for v in values) else None
        return {"valuta": trade["valuta"],
                "prima": {"quantita": position["quantita"], "prezzo_medio": position["prezzo_medio"],
                          "data_apertura": position["data_apertura"], "realized": total(realized_before)},
                "dopo": {"quantita": qty, "prezzo_medio": avg, "data_apertura": opened,
                         "realized": total(realized_after)},
                "trade_successivi": [r["id"] for r in following], "note": list(dict.fromkeys(notes)),
                "baseline": opening,
                "updates": updates}

    def execute_trade(self, *, cash_delta_cents, expected_context=None, realized_fx=None, **trade):
        """Registra trade, posizione e saldo cassa nella stessa transazione.

        `log_trade` resta il percorso esplicito per importazioni storiche e
        fixture: registra lo storico senza modificare il saldo corrente.
        """
        if isinstance(cash_delta_cents, bool) or not isinstance(cash_delta_cents, int):
            raise ValueError(_storage_text('cash_delta_cents deve essere un intero', 'cash_delta_cents must be an integer'))
        with self._conn() as conn:
            conn.execute("BEGIN IMMEDIATE")
            state = _cash_state_from_conn(conn)
            if state["cash_source"] is None:
                raise CashNotInitialized(state["cash_source_note"])
            old_cents = _cash_balance_cents_from_conn(conn)
            trade["ticker"] = trade["ticker"].strip().upper()
            trade["action"] = trade["action"].upper()
            trade["valuta"] = trade.get("valuta", "EUR").strip().upper()
            trade["data"], conventional = normalizza_data_trade(trade.get("data"))
            trade.setdefault("ora_convenzionale", conventional)
            context = self.trade_context(trade["ticker"], trade.get("linked_decision_id"), _conn=conn)
            if expected_context is not None and context["fingerprint"] != expected_context:
                raise RicalcoloImpossibile(_storage_text('anteprima cambiata: cassa, posizione, trade o decisione modificati; ripeti la conferma', 'Preview changed: cash, position, trade or decision changed; preview and confirm again'))
            decision = self._validate_trade_decision(trade, context)
            plan = self._trade_replay_plan(trade, context, realized_fx)
            trade_id = self.log_trade(_conn=conn, **trade)
            if plan:
                self._apply_position_replay(conn, trade["ticker"], trade_id, plan)
                plan.pop("updates")
            new_cents = old_cents + cash_delta_cents
            n = conn.execute(
                "UPDATE cash_state SET balance_cents=?, updated_at=?, version=version+1 "
                "WHERE singleton_id=1 AND balance_cents=?",
                (new_cents, datetime.now().isoformat(timespec="seconds"), old_cents),
            ).rowcount
            if n != 1:
                raise RuntimeError(_storage_text('saldo cassa cambiato durante il trade: rollback', 'Cash balance changed during the trade: rolled back'))
            return {"trade_id": trade_id, "cash_eur": new_cents / 100.0,
                    "ricalcolo": plan, "decisione": decision}

    # Soglia di conferma sui movimenti cassa (delega PM 13/08 "come meglio
    # credi", finding 5 review): sopra questo importo serve conferma=True.
    # E' POLICY, non fisica: il PM puo' cambiarla — un numero solo, qui.
    CASH_SOGLIA_CONFERMA_EUR = 100_000

    # Oltre questi minuti dall'ultimo snapshot il prezzo e' dichiarato VECCHIO
    # (19/08, Opus 5). L'updater scrive ogni ~15 min 24/7: un buco di 2 ore
    # significa che qualcosa si e' rotto, non che la borsa e' chiusa. POLICY,
    # non fisica: un numero solo, qui.
    PRICE_STALE_AFTER_MIN = 120

    def log_cash_movement(self, tipo, importo_eur, data=None, nota=None,
                          cassa_disponibile=None, conferma=False,
                          conferma_soglia=False, conferma_duplicato=False,
                          _conn=None):
        """Registra un movimento di cassa nel registro cash_movements
        (voce PM 12/08: versamenti/prelievi "come se fosse un trade").

        Percorso storico: non modifica cash_state. Le operazioni correnti usano
        apply_cash_movement, che registra movimento e saldo insieme.
        Guardie PRIMA dell'INSERT — rifiuto = zero scritture:
        tipo solo DEPOSIT/WITHDRAWAL · importo > 0 · prelievo oltre
        `cassa_disponibile` (se il chiamante la passa) rifiutato coi due
        numeri · data ISO o rifiuto (mai una data inventata) · duplicato
        stessa terna (data,tipo,importo) e importo sopra soglia rifiutati
        salvo conferma esplicita (delega PM 13/08: doppio click e
        fat-finger non passano zitti).

        F43(5) 31/08 (F37 del frontend, decisione PM 25/08): OGNI guardia ha
        la SUA chiave — `conferma_soglia` spegne solo la soglia importo,
        `conferma_duplicato` solo il doppio-click; prima `conferma` le
        spegneva entrambe insieme (confermare un 150k per la soglia disarmava
        anche la guardia doppio invio). `conferma=True` resta come ALIAS di
        entrambe per i chiamanti esistenti.
        """
        tipo = (tipo or "").strip().upper()
        if tipo not in ("DEPOSIT", "WITHDRAWAL"):
            raise ValueError(_storage_text(f'tipo deve essere DEPOSIT o WITHDRAWAL, non {tipo!r}', f'tipo must be DEPOSIT or WITHDRAWAL, not {tipo!r}'))
        try:
            importo = float(importo_eur)
        except (TypeError, ValueError):
            raise ValueError(_storage_text(f'importo_eur non numerico: {importo_eur!r}', f'importo_eur is not numeric: {importo_eur!r}'))
        # isfinite: json.loads accetta Infinity/NaN (lenienza python) e un inf
        # passerebbe il CHECK > 0 fino a corrompere portfolio.json a valle
        if not math.isfinite(importo) or not importo > 0:
            raise ValueError(_storage_text(f'importo_eur deve essere positivo e finito: {importo}', f'importo_eur must be positive and finite: {importo}'))
        if (tipo == "WITHDRAWAL" and cassa_disponibile is not None
                and importo > float(cassa_disponibile)):
            raise ValueError(
                _storage_text(f'GUARDIA CASSA: prelievo {importo:.2f} EUR oltre la cassa disponibile {float(cassa_disponibile):.2f} EUR: rifiutato.', f'CASH GUARD: withdrawal {importo:.2f} EUR exceeds available cash of {float(cassa_disponibile):.2f} EUR: refused.'))
        if data is not None:
            try:
                # si salva il PARSATO, non il raw (review 12/08): "…T23:59:59"
                # o "20260812" verbatim in tabella spostano il flusso di giorno
                # nel raggruppamento lessicografico del TWR e l'IRR lo droppa
                data = datetime.fromisoformat(str(data)).date().isoformat()
            except ValueError:
                raise ValueError(_storage_text(f'data non ISO (YYYY-MM-DD): {data!r}', f'Date is not ISO (YYYY-MM-DD): {data!r}'))
        data = data or datetime.now().strftime("%Y-%m-%d")
        # 2 decimali all'INSERT: la somma del registro e la cassa (round a 2
        # nel chiamante) non devono divergere di sub-cent
        importo = round(importo, 2)
        # la coda «(conferma=true ...)» nei DUE rifiuti e' PORTANTE, non
        # ridondante: la regex di app/src/lib/cassa.ts:603 accende il bottone
        # «CONFERMO, E' VOLUTO» di F7 sul letterale `conferma=true` — in
        # «conferma_soglia=true» dopo «conferma» c'e' un underscore e NON matcha
        if importo > self.CASH_SOGLIA_CONFERMA_EUR and not (conferma or conferma_soglia):
            raise CashConfirmationRequired("cash_threshold",
                _storage_text(f"GUARDIA IMPORTO: {importo:.2f} EUR sopra la soglia di {self.CASH_SOGLIA_CONFERMA_EUR} EUR — possibile errore di battitura. Se e' voluto, ripeti con conferma_soglia=true (conferma=true vale per ENTRAMBE le guardie).", f'AMOUNT GUARD: {importo:.2f} EUR above the threshold of {self.CASH_SOGLIA_CONFERMA_EUR} EUR — possible typo. If intended, repeat with conferma_soglia=true (conferma=true applies to BOTH guards).'))
        from contextlib import nullcontext
        with (nullcontext(_conn) if _conn is not None else self._conn()) as conn:
            if not (conferma or conferma_duplicato):
                dup = conn.execute(
                    "SELECT id FROM cash_movements WHERE date=? AND type=? "
                    "AND amount_eur=?", (data, tipo, importo)).fetchone()
                if dup:
                    raise CashConfirmationRequired("cash_duplicate",
                        _storage_text(f"GUARDIA DUPLICATO: esiste gia' un {tipo} di {importo:.2f} EUR in data {data} (id={dup[0]}) — doppio click? Se e' un secondo movimento vero, ripeti con conferma_duplicato=true (conferma=true vale per ENTRAMBE le guardie).", f'DUPLICATE GUARD: an existing {tipo} of {importo:.2f} EUR on date {data} (id={dup[0]}) — double click? If this is a separate actual movement, repeat with conferma_duplicato=true (conferma=true applies to BOTH guards).'))
            cur = conn.execute(
                "INSERT INTO cash_movements (date, type, amount_eur, note) "
                "VALUES (?,?,?,?)", (data, tipo, importo, nota))
            return cur.lastrowid

    def apply_cash_movement(self, tipo, importo_eur, data=None, nota=None,
                            conferma=False, conferma_soglia=False,
                            conferma_duplicato=False):
        """Registra movimento e saldo in un'unica transazione SQLite.

        Il primo DEPOSIT esplicito inizializza il saldo da zero. Ogni altra
        operazione richiede un saldo già misurato.
        """
        from decimal import Decimal, InvalidOperation
        tipo_norm = (tipo or "").strip().upper()
        if isinstance(importo_eur, bool):
            raise ValueError(_storage_text('importo_eur non numerico', 'importo_eur is not numeric'))
        try:
            amount_dec = Decimal(str(importo_eur))
        except InvalidOperation as exc:
            raise ValueError(_storage_text('importo_eur non numerico', 'importo_eur is not numeric')) from exc
        if (not amount_dec.is_finite() or amount_dec <= 0
                or amount_dec.as_tuple().exponent < -2):
            raise ValueError(_storage_text('importo_eur deve essere positivo, finito e avere massimo 2 decimali', 'importo_eur must be positive, finite and have at most 2 decimal places'))
        amount_cents = int(amount_dec * 100)
        with self._conn() as conn:
            conn.execute("BEGIN IMMEDIATE")
            state = _cash_state_from_conn(conn)
            inizializza = state["cash_source"] is None
            if inizializza and tipo_norm != "DEPOSIT":
                raise CashNotInitialized(state["cash_source_note"])
            old_cents = 0 if inizializza else _cash_balance_cents_from_conn(conn)
            mov_id = self.log_cash_movement(
                tipo, amount_cents / 100.0, data=data, nota=nota,
                cassa_disponibile=old_cents / 100.0,
                conferma=conferma, conferma_soglia=conferma_soglia,
                conferma_duplicato=conferma_duplicato, _conn=conn,
            )
            new_cents = old_cents + (amount_cents if tipo_norm == "DEPOSIT" else -amount_cents)
            now = datetime.now().isoformat(timespec="seconds")
            if inizializza:
                conn.execute(
                    "INSERT INTO cash_state(singleton_id,balance_cents,updated_at,source) "
                    "VALUES(1,?,?,?)", (new_cents, now, "explicit:first-deposit"))
            else:
                n = conn.execute(
                    "UPDATE cash_state SET balance_cents=?, updated_at=?, version=version+1 "
                    "WHERE singleton_id=1 AND balance_cents=?",
                    (new_cents, now, old_cents),
                ).rowcount
                if n != 1:
                    raise RuntimeError(_storage_text('saldo cassa cambiato durante il movimento: rollback', 'Cash balance changed during the movement: rolled back'))
            return {"movement_id": mov_id, "cash_eur": new_cents / 100.0}

    def get_cash_movements(self, limit=100):
        """Registro movimenti di cassa, piu' recenti prima (per F7/F14)."""
        # clamp: in sqlite LIMIT negativo = TUTTO, un limit=-1 diventerebbe
        # un dump involontario; tetto 1000 come cintura simmetrica
        limit = max(1, min(int(limit), 1000))
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT id, date, type, amount_eur, note, created_at "
                "FROM cash_movements ORDER BY date DESC, id DESC LIMIT ?",
                (limit,)).fetchall()
            return [dict(r) for r in rows]

    def get_portfolio(self, active_only=True, now=None):
        """Restituisce il portafoglio LIVE come list di dict (drop-in replacement per get_portfolio_live).

        `now` (ISO o datetime) e' iniettabile per collaudi deterministici sulla
        freschezza dei prezzi; di default e' l'orologio di sistema.
        """
        from datetime import datetime as _dt
        # ⚠️ `position_prices.timestamp` e' scritto in UTC NAIVE (misurato
        # 19/08: la riga delle 23:38:52 locali sta a disco come 21:38:53).
        # Confrontarlo con l'ora LOCALE dava +120 min fissi in Italia e
        # marchiava VECCHIE 21 posizioni su 28 appena aggiornate — un allarme
        # sempre acceso e' peggio di nessun allarme.
        if now is None:
            _adesso = _dt.utcnow()
        elif isinstance(now, str):
            _adesso = _dt.fromisoformat(now.replace("Z", "")[:19])
        else:
            _adesso = now
        # 20/08: gli orari di borsa vogliono un istante CONSAPEVOLE del fuso.
        # `_adesso` e' UTC naive (v. sopra): lo si etichetta come UTC senza
        # spostarlo, cosi' `now` resta iniettabile e i collaudi restano
        # deterministici anche sulla domanda "il mercato era aperto?".
        from datetime import timezone as _tzmod
        _adesso_locale = (_adesso if _adesso.tzinfo is not None
                          else _adesso.replace(tzinfo=_tzmod.utc))
        with self._conn() as conn:
            query = "SELECT * FROM positions"
            if active_only:
                query += " WHERE is_active=1 AND quantita > 0"
            query += " ORDER BY quantita * COALESCE(prezzo_medio, 0) DESC"
            rows = conn.execute(query).fetchall()

            positions = []
            for r in rows:
                # Get latest price snapshot
                last_price_row = conn.execute(
                    "SELECT prezzo, timestamp FROM position_prices WHERE ticker=? "
                    "ORDER BY timestamp DESC LIMIT 1",
                    (r["ticker"],)).fetchone()
                last_price = last_price_row["prezzo"] if last_price_row else None
                # ETA' del prezzo: una MISURA, non una promessa (19/08). Prima
                # `price_stale` era `last_price is None`, cioe' guardava solo
                # l'ESISTENZA: un prezzo fermo da ore usciva "fresco" — ed e'
                # esattamente come i 5 ticker polygon hanno mostrato per un
                # giorno intero la chiusura di ieri senza che nulla lo dicesse.
                price_age_min = None
                if last_price_row is not None and last_price_row["timestamp"]:
                    try:
                        _ts = _dt.fromisoformat(str(last_price_row["timestamp"])[:19])
                        price_age_min = int((_adesso - _ts).total_seconds() // 60)
                    except (ValueError, TypeError):
                        price_age_min = None   # timestamp illeggibile: n.d. dichiarato
                _vecchio = (price_age_min is not None
                            and price_age_min > self.PRICE_STALE_AFTER_MIN)
                # 20/08 (ok PM "rendila sensibile"): la soglia a tempo fisso non
                # distingueva "mercato chiuso" da "updater morto". Il lunedi' alle
                # 07:00 TUTTO il book risultava vecchio, e un allarme sempre acceso
                # e' peggio di nessun allarme — la stessa lezione che il fix del
                # 19/08 si era gia' portato dietro una volta (21 posizioni su 28
                # marchiate per un fuso). A mercato CHIUSO un prezzo fermo e' il
                # prezzo GIUSTO: l'eta' resta esposta, l'allarme no.
                # `aperto is None` = non determinabile: si tiene l'allarme. Un falso
                # allarme e' visibile e si corregge; un guasto nascosto no.
                _mkt_aperto = None
                try:
                    from bellomberg.market_data.mercati import aperto as _mkt_open
                    _mkt_aperto = _mkt_open(r["ticker"], _adesso_locale)
                except Exception:
                    _mkt_aperto = None      # modulo assente/rotto: allarme conservativo
                _vecchio = _vecchio and (_mkt_aperto is not False)
                # prev_close per il P&L GIORNALIERO (richiesta chat frontend 23/07,
                # ok PM): ultimo snapshot del GIORNO precedente all'ultimo giorno
                # con prezzi — festivi/weekend compresi (il "giorno prima" e' il
                # giorno di borsa precedente, non ieri calendario). Fonte
                # DICHIARATA: position_prices locale (updater ~15min: l'ultimo
                # snapshot del giorno ~ chiusura), MAI yfinance qui. Assente =
                # None = n.d. dichiarato, il frontend non mostra un GG% finto.
                prev_row = conn.execute(
                    "SELECT prezzo, timestamp FROM position_prices WHERE ticker=? "
                    "AND date(timestamp) < (SELECT date(MAX(timestamp)) FROM "
                    "position_prices WHERE ticker=?) "
                    "ORDER BY timestamp DESC LIMIT 1",
                    (r["ticker"], r["ticker"])).fetchone()
                prev_close = prev_row["prezzo"] if prev_row else None
                prev_close_ts = prev_row["timestamp"] if prev_row else None
                prev_close_source = "position_prices" if prev_row else None
                # 28/08 (decisione PM 27/08 sera, «P&L GG dal carico»): una posizione
                # APERTA il giorno dell'ultimo prezzo non ha una chiusura precedente
                # (l'updater fotografa solo le posizioni in book): il suo daily e' il
                # P&L dal prezzo di carico, DICHIARATO da `prev_close_source`. Solo
                # quel giorno: aperta ieri senza snapshot resta n.d. (voce futura:
                # chiusura dallo storico, fonte dichiarata). Nel SUO giorno di
                # apertura il carico vince anche su uno snapshot precedente (review
                # 28/08): un ticker chiuso e ricomprato ha in tabella gli snapshot
                # intraday della detenzione vecchia, che farebbero da falsa chiusura.
                # `data_apertura` e' in ora locale, gli snapshot in UTC naive: si
                # confrontano le DATE. Fra le 00:00 e le 02:00 locali le due date
                # divergono (l'updater fotografa 24/7): un acquisto in quella finestra
                # (after-hours USA) resta n.d. fino alle 02:00, poi prende lo snapshot
                # delle 01:50 come «chiusura» — dichiarato, raro, non muto.
                if last_price_row is not None and r["prezzo_medio"]:
                    _giorno_ultimo_prezzo = str(last_price_row["timestamp"] or "")[:10]
                    if _giorno_ultimo_prezzo and str(r["data_apertura"] or "")[:10] == _giorno_ultimo_prezzo:
                        prev_close = float(r["prezzo_medio"])
                        prev_close_ts = r["data_apertura"]
                        prev_close_source = "carico"
                val_mercato = (r["quantita"] or 0) * (last_price or r["prezzo_medio"] or 0)
                pl_eur = ((last_price - r["prezzo_medio"]) * r["quantita"]) if (last_price is not None and r["prezzo_medio"] is not None) else None
                pl_pct = ((last_price - r["prezzo_medio"]) / r["prezzo_medio"] * 100) if (last_price and r["prezzo_medio"]) else None
                positions.append({
                    "ticker": r["ticker"],
                    "nome": r["nome"],
                    "quantita": r["quantita"],
                    "prezzo_medio": r["prezzo_medio"],
                    "prezzo_live": last_price,
                    "valuta": r["valuta"],
                    "valore_mercato": val_mercato,
                    "pl_eur": pl_eur,
                    "pl_pct": pl_pct,
                    # P&L daily (chat frontend 23/07): prev_close in valuta di
                    # quotazione + timestamp della fonte; GG% = live/prev - 1.
                    "prev_close": prev_close,
                    "prev_close_ts": prev_close_ts,
                    # 28/08: da dove viene prev_close — "position_prices" | "carico" | None
                    "prev_close_source": prev_close_source,
                    # F-CONT-4 (riallineamento 23/07, regola no-fallback 14/07): senza
                    # snapshot prezzo la posizione vale il COSTO (P&L n.d.) — prima
                    # accadeva IN SILENZIO, ora e' un flag che UI/agenti possono dire.
                    # STALE = prezzo assente OPPURE piu' vecchio della soglia.
                    # Un prezzo vecchio resta MIGLIORE del costo e continua a
                    # valutare la posizione (il NAV non salta a ogni buco
                    # dell'updater), ma da qui in poi lo DICHIARA.
                    # Dal 20/08 la soglia e' SENSIBILE AL MERCATO: un prezzo fermo
                    # a borsa chiusa non e' un guasto, e non accende piu' l'allarme.
                    # L'eta' resta sempre esposta, e `mercato_aperto` dice perche'.
                    "price_stale": (last_price is None) or _vecchio,
                    "price_age_minutes": price_age_min,
                    "mercato_aperto": _mkt_aperto,
                    "price_source": (
                        _message('prezzo_medio (fallback DICHIARATO: P&L n.d.)', 'prezzo_medio (DECLARED fallback: P&L n/a)') if last_price is None
                        else _message('snapshot VECCHIO ({v0} min, soglia {v1}, mercato APERTO)', 'STALE snapshot ({v0} min, threshold {v1}, market OPEN)', v0=price_age_min, v1=self.PRICE_STALE_AFTER_MIN) if _vecchio
                        else _message("snapshot di {v0} min fa, oltre la soglia {v1} ma a mercato CHIUSO: e' l'ultimo prezzo valido, non un buco", 'snapshot {v0} min ago, beyond threshold {v1} but market CLOSED: the latest valid price, not a data gap', v0=price_age_min, v1=self.PRICE_STALE_AFTER_MIN)
                        if (price_age_min is not None
                            and price_age_min > self.PRICE_STALE_AFTER_MIN
                            and _mkt_aperto is False)
                        else "snapshot"),
                    "tesi": r["tesi"],
                    "data_apertura": r["data_apertura"],
                })
            return positions

    def get_portfolio_summary(self):
        """Stato aggregato con valori convertiti in EUR per multi-currency."""
        positions = self.get_portfolio()
        # audit/11 §2: conversione EUR per-posizione con TRACCIAMENTO — prima un errore a
        # meta' loop (o una valuta senza FX, es. CAD) lasciava valori nativi sommati come
        # EUR in silenzio, col summary etichettato 'normalized to EUR'.
        try:
            from bellomberg.cli.price_updater import get_fx_to_eur, get_fx_sources
        except Exception:
            get_fx_to_eur = None
            get_fx_sources = lambda: {}
        fx_incomplete = []
        fx_sources = {}
        for p in positions:
            cur = (p.get("valuta") or "EUR").upper()
            if cur == "EUR":
                p["fx_to_eur"] = 1.0
                p["fx_source"] = "identity"
                continue
            fx, fx_source = None, "assente"
            if get_fx_to_eur is not None:
                try:
                    fx = get_fx_to_eur(cur)
                    fx_source = get_fx_sources().get(
                        "GBP" if cur == "GBX" else cur, "n.d." if fx is not None else "assente")
                except Exception:
                    fx, fx_source = None, "assente"
            fx_sources[cur] = fx_source
            p["fx_source"] = fx_source
            if fx is None:
                fx_incomplete.append(f"{p.get('ticker')}:{cur}")
                continue
            # fx esposto per posizione (chat frontend 23/07): serve al P&L daily
            # in EUR = qty*(live-prev_close)*fx_to_eur — mai un cambio inventato
            # lato UI. Per le posizioni EUR resta il default 1.0 impostato sotto.
            p["fx_to_eur"] = fx
            if p.get("valore_mercato"):
                p["valore_mercato_eur"] = p["valore_mercato"] * fx
                p["valore_mercato"] = p["valore_mercato_eur"]  # override
            if p.get("pl_eur"):
                p["pl_eur"] = p["pl_eur"] * fx

        # pl_eur_fx (28/08, decisione PM): il P&L col cambio STORICO del costo,
        # chiavi ADDITIVE per posizione + totali SOLO a copertura piena (mai un
        # parziale spacciato per completo); ogni n.d. porta il suo motivo.
        _openings = []
        try:
            from bellomberg.portfolio.portfolio_analytics import pl_fx_per_posizione
            with self._conn() as conn:
                _trades = [dict(r) for r in conn.execute(
                    "SELECT id, ticker, action, quantita, prezzo, valuta, data "
                    "FROM trade_history ORDER BY data ASC, id ASC")]
            _openings = self.get_opening_positions()
            _st = pl_fx_per_posizione(_trades, datetime.now().strftime("%Y-%m-%d"),
                                     **({"openings": _openings} if _openings else {}))
        except Exception as _e:
            _st = {"error": _message('costo storico n.d. ({v0}: {v1})', 'historical cost n/a ({v0}: {v1})', v0=type(_e).__name__, v1=_e)}
        pl_fx_nd = []
        for p in positions:
            p["position_opening"] = next((row for row in _openings if row["ticker"] == p["ticker"]), None)
            p["costo_eur_storico"] = None
            p["pl_eur_fx"] = None
            p["pl_pct_fx"] = None
            p["fx_pl_eur"] = None
            p["fx_pl_note"] = None
            _cur_label = _classifica_valuta(p.get("ticker") or "?", p.get("valuta"))
            _cur = _cur_label.valore
            if _cur is None:
                p["fx_pl_note"] = _cur_label.dichiarazione
            elif "error" in _st:
                p["fx_pl_note"] = _st["error"]
            else:
                h = _st["per_ticker"].get(p["ticker"])
                if h is None:
                    p["fx_pl_note"] = _message('nessun trade in trade_history: costo storico n.d.', 'no trade in trade_history: historical cost n/a')
                elif h["costo_eur_storico"] is None:
                    p["fx_pl_note"] = join_messages("; ", h["note"]) or _message('FX storico n.d.', 'Historical FX n/a')
                elif abs(h["qty"] - float(p.get("quantita") or 0)) > 1e-6:
                    p["fx_pl_note"] = (_message('replay dei trade {v0:g} vs posizione {v1:g}: costo storico n.d.', 'trade replay {v0:g} vs position {v1:g}: historical cost n/a', v0=h['qty'], v1=float(p.get('quantita') or 0)))
                elif _cur != "EUR" and p.get("fx_to_eur") is None:
                    # review 28/08 (ALTO): senza FX di oggi valore_mercato e pl_eur
                    # sono rimasti in valuta NATIVA (fx_incomplete): un numero in
                    # valuta mista spacciato per EUR e' la classe vietata il 14/07
                    p["fx_pl_note"] = (_message('FX di oggi per {v0} n.d.: valore di mercato in valuta nativa, pl_eur_fx n.d.', 'Current FX for {v0} n/a: market value in native currency, pl_eur_fx n/a', v0=_cur))
                elif p.get("pl_eur") is None or not p.get("valore_mercato"):
                    p["fx_pl_note"] = _message('prezzo live o valore di mercato n.d.', 'Live price or market value n/a')
                else:
                    # review 28/08 (MEDIO): la componente cambio e' PURA — pool in
                    # valuta x FX di oggi meno pool in EUR — cosi' e' zero per
                    # costruzione sulle righe EUR (prima il residuo fra prezzo_medio
                    # del DB e costo medio del replay, centesimi, usciva come «cambio»:
                    # pochi euro su una riga EUR). pl_eur_fx resta ancorato al prezzo_medio del DB,
                    # il libro di F1: pl_eur + componente cambio.
                    _fx_oggi = 1.0 if _cur == "EUR" else float(p["fx_to_eur"])
                    costo = h["costo_eur_storico"]
                    fx_comp = round(h["costo_nativo"] * _fx_oggi - costo, 2)
                    p["fx_pl_eur"] = fx_comp
                    p["pl_eur_fx"] = round(float(p["pl_eur"]) + fx_comp, 2)
                    p["costo_eur_storico"] = round(float(p["valore_mercato"]) - p["pl_eur_fx"], 2)
                    if p["costo_eur_storico"]:
                        p["pl_pct_fx"] = round(p["pl_eur_fx"] / p["costo_eur_storico"] * 100, 2)
                    else:
                        p["fx_pl_note"] = _message('costo storico 0: pl_pct_fx n.d.', 'Historical cost 0: pl_pct_fx n/a')
                    if h["note"]:
                        p["fx_pl_note"] = join_messages("; ", h["note"]) if not p["fx_pl_note"] \
                            else join_messages('; ', [p['fx_pl_note'], join_messages('; ', h['note'])])
            if p["pl_eur_fx"] is None:
                pl_fx_nd.append(p["ticker"])
        fx_pl_basis = (_st.get("fx_basis") if "error" not in _st else None)

        # Voce 10 §9-quattuortrigies (01/08): l'ORDER BY di get_portfolio ordina
        # sul costo in valuta NATIVA (Londra in GBX entra x100: il libro usciva
        # `quota-1.L, quota-2.L, ...` in testa a prescindere dal valore vero). Qui i
        # valori sono appena stati convertiti in EUR (salvo fx_incomplete,
        # DICHIARATE nel source): il libro si riordina sul numero vero.
        positions.sort(key=lambda p: p.get("valore_mercato") or 0, reverse=True)

        # Un importo rimasto nella valuta nativa non e' un totale EUR. Prima il
        # payload dichiarava fx_incomplete ma pubblicava e persisteva comunque
        # la somma mista sotto chiavi *_eur.
        total_value = (sum(p.get("valore_mercato") or 0 for p in positions)
                       if not fx_incomplete else None)
        for p in positions:
            p["peso_pct"] = ((p["valore_mercato"] / total_value * 100)
                             if total_value is not None and total_value > 0 else None)

        # Saldo operativo dalla stessa fonte SQLite transazionale dei movimenti.
        _cassa = read_cash_state(self.db_path)
        cash_eur = _cassa["cash_eur"]

        nav_total = total_value + cash_eur if total_value is not None else None

        # --- totale «come il broker» (31/08, decisione PM: proposta 1 + «devi
        # pure mettere i dividendi nel calcolo») = aperto + realizzato + dividendi,
        # ogni componente dichiarata, MAI un parziale zitto. Il «P&L» del sito
        # del broker i dividendi li ESCLUDE (misura 31/08): la nota dice come
        # tornare al suo numero.
        _note_tot = []
        if _openings:
            _note_tot.append(_message("Saldi iniziali: realizzato e dividendi coprono solo i trade registrati; gli incassi precedenti ai saldi documentati sono ignoti. Nessun totale dell'intera storia del broker.", 'Opening balances: realized P&L and dividends cover recorded trades only; receipts before documented balances are unknown. No total for the entire broker history.'))
        try:
            with self._conn() as conn:
                _vn, _vnull, _vsum = conn.execute(
                    "SELECT COUNT(*), COALESCE(SUM(realized_eur IS NULL),0), "
                    "COALESCE(SUM(realized_eur),0) FROM trade_history "
                    "WHERE action IN ('SELL','TRIM')").fetchone()
                _div_rows = conn.execute(
                    "SELECT ticker, quantita, prezzo, valuta FROM trade_history "
                    "WHERE action='DIVIDEND'").fetchall()
        except Exception as _e:
            # lettura fallita = n.d. DICHIARATO (regola 14/07), mai un crash del
            # summary intero ne' uno zero spacciato per misura
            _vn = realizzato_vendite = dividendi_eur = None
            _note_tot.append(_message('lettura trade_history fallita ({v0}: {v1}): realizzato e dividendi n.d.', 'trade_history read failed ({v0}: {v1}): realized P&L and dividends n/a', v0=type(_e).__name__, v1=_e))
        else:
            if _vnull:
                realizzato_vendite = None
                _note_tot.append(_message('{v0} vendite su {v1} senza realized_eur: realizzato n.d.', '{v0} sales out of {v1} without realized_eur: realized P&L n/a', v0=_vnull, v1=_vn))
            else:
                realizzato_vendite = round(_vsum, 2)
            _div_sum, _div_nd = 0.0, []
            for _r in _div_rows:
                if (_r["valuta"] or "EUR").upper() == "EUR":
                    _div_sum += (_r["quantita"] or 0) * (_r["prezzo"] or 0)
                else:
                    # un importo nativo sommato come EUR e' la classe vietata il 14/07;
                    # la conversione storica dei dividendi non-EUR e' una voce futura
                    _div_nd.append(f"{_r['ticker']} in {_r['valuta']}")
            if _div_nd:
                dividendi_eur = None
                _note_tot.append(_message('dividendi senza conversione storica ({tickers}): dividendi_eur n.d.', 'Dividends lack historical conversion ({tickers}): dividendi_eur n/a', tickers=', '.join(dict.fromkeys(_div_nd))
))
            else:
                dividendi_eur = round(_div_sum, 2)
        # aperto: 0 e' una MISURA solo a book vuoto; con posizioni scoperte
        # (pl_fx_nd) resta n.d. — stessa regola di totale_pl_eur_fx
        _aperto = (0.0 if not positions else
                   (round(sum(p["pl_eur_fx"] for p in positions), 2) if not pl_fx_nd else None))
        if positions and _aperto is None:
            _note_tot.append(_message('aperto n.d. (v. pl_fx_nd): totale n.d.', 'open P&L n/a (see pl_fx_nd): total n/a'))
        totale_broker = (round(_aperto + realizzato_vendite + dividendi_eur, 2)
                         if None not in (_aperto, realizzato_vendite, dividendi_eur) else None)
        _nota_broker = (_message("aperto (totale_pl_eur_fx; 0 a book vuoto) + realizzato vendite SELL/TRIM (realized_eur) + dividendi (quantita' x prezzo, EUR). Il P&L del sito del broker ESCLUDE i dividendi (misura 31/08): per quel confronto sottrarre dividendi_eur.{warning}", 'open P&L (totale_pl_eur_fx; 0 for an empty book) + realized SELL/TRIM P&L (realized_eur) + dividends (quantity x price, EUR). The broker website P&L EXCLUDES dividends (measured 31/08): subtract dividendi_eur for that comparison.{warning}', warning=_message(' ATTENZIONE: {notes}', ' WARNING: {notes}', notes=join_messages('; ', _note_tot)) if _note_tot else ''



))

        return {
            "source": (_message('Database SQLite (live, multivaluta normalizzata in EUR)', 'SQLite database (live, multi-currency normalized to EUR)')
                       if not fx_incomplete else
                       _message('SQLite database (live) — ATTENZIONE: FX non disponibile per {currencies}: totali EUR e NAV n.d.; valori per posizione restano in valuta nativa', 'SQLite database (live) — WARNING: FX unavailable for {currencies}: EUR totals and NAV n/a; position values remain in native currency', currencies=', '.join(fx_incomplete)

)),
            "fx_incomplete": fx_incomplete or None,
            "fx_sources": fx_sources or None,
            # F-CONT-4: ticker valorizzati al COSTO per assenza di snapshot prezzo
            # (P&L n.d. su quei nomi) — buco DICHIARATO, mai piu' zitto.
            "stale_positions": [p["ticker"] for p in positions if p.get("price_stale")] or None,
            "n_positions": len(positions),
            "positions": positions,
            "totale_valore_mercato_eur": (round(total_value, 2)
                                            if total_value is not None else None),
            "cash_disponibile_eur": round(cash_eur, 2),
            # Fonte SQLite quando inizializzata; None quando lo zero e' un buco.
            "cash_source": _cassa["cash_source"],
            "cash_source_note": _cassa["cash_source_note"],
            "nav_total_eur": (round(nav_total, 2) if nav_total is not None else None),
            "totale_pl_eur": (round(sum(p.get("pl_eur") or 0 for p in positions), 2)
                               if not fx_incomplete else None),
            # pl_eur_fx (28/08): totali SOLO se ogni posizione ha il suo numero;
            # altrimenti None e l'elenco dei n.d. (il motivo sta in fx_pl_note)
            "totale_pl_eur_fx": (round(sum(p["pl_eur_fx"] for p in positions), 2)
                                 if positions and not pl_fx_nd else None),
            "totale_fx_pl_eur": (round(sum(p["fx_pl_eur"] for p in positions), 2)
                                 if positions and not pl_fx_nd else None),
            "pl_fx_nd": pl_fx_nd or None,
            # il metodo, dichiarato in testa (review 28/08): FX di oggi = fx_to_eur
            # per riga (price_updater.get_fx_to_eur), costo storico = pool al FX
            # daily del giorno di ogni acquisto, e la divergenza da nav_history
            "fx_pl_basis": (_message('{basis}; FX di oggi = fx_to_eur per riga (price_updater)', '{basis}; current FX = fx_to_eur per row (price_updater)', basis=fx_pl_basis)
                            if fx_pl_basis else None),
            # totale «come il broker» (31/08): componenti sempre dichiarate,
            # totale SOLO a componenti piene (il buco sta nella nota)
            "realizzato_eur_vendite": realizzato_vendite,
            "realizzato_vendite_n": _vn,
            "dividendi_eur": dividendi_eur,
            "totale_aperto_piu_realizzato_eur": totale_broker,
            "totale_aperto_piu_realizzato_note": _nota_broker,
            "timestamp": datetime.now().isoformat(timespec="seconds"),
        }

    def update_price(self, ticker, prezzo, valuta="USD", source="manual"):
        """Salva snapshot prezzo."""
        with self._conn() as conn:
            conn.execute("INSERT INTO position_prices (ticker, prezzo, valuta, source) VALUES (?,?,?,?)",
                          (ticker, prezzo, valuta, source))

    # ========================================================
    # IMPORT FROM EXCEL (one-shot migration)
    # ========================================================

    def import_from_excel(self, excel_path):
        """Importa il portafoglio dall'Excel attuale. Idempotente."""
        try:
            import openpyxl
        except ImportError:
            return {"error": "openpyxl required"}

        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
        except Exception as e:
            return {"error": "Open failed: " + str(e)}

        # Cerca il sheet 'Holdings'
        target_sheet = None
        for s in wb.sheetnames:
            if s.lower() in ("holdings", "portafoglio", "positions"):
                target_sheet = s
                break
        if not target_sheet:
            target_sheet = wb.sheetnames[0]

        ws = wb[target_sheet]
        # Cerca header
        HEADER_KEYWORDS = ["ticker", "nome", "q.t", "quantit", "prezzo", "val. mercato"]
        header_row = None
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if not row: continue
            txt = " ".join(str(c).lower() for c in row if c)
            if sum(1 for kw in HEADER_KEYWORDS if kw in txt) >= 2:
                header_row = row_idx
                break
        if not header_row:
            return {"error": "Header non trovato"}

        headers = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        col_map = {}
        for i, h in enumerate(headers):
            if not h: continue
            hl = str(h).lower().strip()
            if "ticker" in hl or "simbolo" in hl: col_map["ticker"] = i
            elif "nome" in hl: col_map["nome"] = i
            elif "q.t" in hl or "quantit" in hl: col_map["quantita"] = i
            elif "prezzo medio" in hl or "avg" in hl: col_map["prezzo_medio"] = i
            elif "valuta" in hl or "currency" in hl: col_map["valuta"] = i

        imported = 0
        errors = []
        for row in ws.iter_rows(min_row=header_row+1, values_only=True):
            if not row: continue
            ticker_raw = row[col_map.get("ticker", 0)] if "ticker" in col_map else None
            ticker = str(ticker_raw).strip() if ticker_raw else ""
            if not ticker or ticker.upper() in ("TOTALE", "TOTAL", "TOT") or ticker.startswith("#"):
                # Fallback su nome
                if "nome" in col_map:
                    nome_raw = row[col_map["nome"]]
                    if nome_raw:
                        ticker = str(nome_raw).strip()[:20]
                if not ticker: continue
            try:
                q = float(row[col_map["quantita"]]) if "quantita" in col_map and row[col_map["quantita"]] else 0
                pm = float(row[col_map["prezzo_medio"]]) if "prezzo_medio" in col_map and row[col_map["prezzo_medio"]] else None
                nome = str(row[col_map["nome"]]) if "nome" in col_map and row[col_map["nome"]] else ticker
                valuta = str(row[col_map.get("valuta", -1)]) if col_map.get("valuta") is not None and row[col_map.get("valuta")] else "EUR"
                self.add_or_update_position(ticker=ticker, nome=nome, quantita=q,
                                              prezzo_medio=pm, valuta=valuta)
                imported += 1
            except Exception as e:
                errors.append(ticker + ": " + str(e))

        return {"imported": imported, "errors": errors, "sheet": target_sheet}

    # ========================================================
    # MEMOS & DECISIONS
    # ========================================================

    def save_memo(self, full_markdown, pdf_path=None, appendix_path=None,
                  dcf_files=None, capo_tokens_in=0, capo_tokens_out=0,
                  portfolio_nav_eur=None, title=None, output_language=None):
        """Salva memo + chunks in ChromaDB. Ritorna memo_id."""
        from bellomberg.core.language import capture_language
        selected = capture_language(output_language)
        ts = datetime.now().isoformat(timespec="seconds")
        title = title or _storage_text("Nota di ricerca settimanale - ", "Weekly Research Note - ", language=selected) + datetime.now().strftime("%d/%m/%Y")
        with self._conn() as conn:
            cur = conn.execute("""INSERT INTO memos (timestamp, title, full_markdown, pdf_path,
                                                     appendix_path, dcf_files, capo_tokens_in,
                                                     capo_tokens_out, portfolio_nav_eur, output_language)
                                  VALUES (?,?,?,?,?,?,?,?,?,?)""",
                                (ts, title, full_markdown, pdf_path, appendix_path,
                                 json.dumps(dcf_files or []), capo_tokens_in, capo_tokens_out,
                                 portfolio_nav_eur, selected))
            memo_id = cur.lastrowid

        # Chunk + embed in ChromaDB
        if self.col_memos and full_markdown:
            try:
                # Chunking semplice per sezione (## headers)
                chunks = self._chunk_markdown(full_markdown)
                if chunks:
                    self.col_memos.add(
                        documents=chunks,
                        metadatas=[{"memo_id": memo_id, "timestamp": ts, "chunk_idx": i}
                                    for i in range(len(chunks))],
                        ids=["memo_" + str(memo_id) + "_chunk_" + str(i) for i in range(len(chunks))]
                    )
            except Exception as e:
                print("[MemoryDB] ChromaDB embed memo error: " + str(e))

        return memo_id

    def _chunk_markdown(self, text, max_chunk_chars=1200):
        """Split per sezioni ##, poi rispetta max_chunk_chars."""
        sections = re.split(r"\n(?=## )", text)
        chunks = []
        for sec in sections:
            sec = sec.strip()
            if not sec: continue
            if len(sec) <= max_chunk_chars:
                chunks.append(sec)
            else:
                # Split sub-paragrafi
                paras = sec.split("\n\n")
                buf = ""
                for p in paras:
                    if len(buf) + len(p) + 2 < max_chunk_chars:
                        buf += "\n\n" + p if buf else p
                    else:
                        if buf: chunks.append(buf)
                        buf = p
                if buf: chunks.append(buf)
        return chunks

    def search_memos_semantic(self, query, n_results=5):
        """Semantic search sui memo passati. Ritorna list di dict {memo_id, content, distance}."""
        if not self.col_memos:
            return []
        try:
            results = self.col_memos.query(query_texts=[query], n_results=n_results)
            out = []
            fingerprint_corrente = _fingerprint_mandato_corrente()
            ids = (results.get("ids") or [[]])[0]
            metadati = (results.get("metadatas") or [[]])[0]
            documenti = (results.get("documents") or [[]])[0]
            distanze = (results.get("distances") or [[]])[0]
            if not ids:
                return []
            memo_ids = [m.get("memo_id") for m in metadati]
            testi_memo = {}
            errore_lookup = None
            try:
                ids_validi = [int(i) for i in memo_ids if i is not None]
                if ids_validi:
                    ph = ",".join("?" * len(ids_validi))
                    with self._conn() as conn:
                        righe = conn.execute(
                            "SELECT id, full_markdown FROM memos WHERE id IN (" + ph + ")",
                            ids_validi,
                        ).fetchall()
                    testi_memo = {int(r["id"]): r["full_markdown"] for r in righe}
            except Exception as e:
                errore_lookup = type(e).__name__ + ": " + str(e)[:160]
            for i, chunk_id in enumerate(ids):
                meta = metadati[i] if i < len(metadati) else {}
                memo_id = meta.get("memo_id")
                try:
                    chiave = int(memo_id)
                except (TypeError, ValueError):
                    chiave = None
                if errore_lookup:
                    etichetta = ("[MANDATO STORICO: provenienza mandato non disponibile "
                                 "(lookup memo fallito: " + errore_lookup + "); "
                                 "non trattare il testo come preferenza PM corrente]")
                else:
                    etichetta = _etichetta_mandato_storico(
                        testi_memo.get(chiave), fingerprint_corrente)
                out.append({
                    "chunk_id": chunk_id,
                    "memo_id": memo_id,
                    "content": etichetta + "\n" + (documenti[i] if i < len(documenti) else ""),
                    "distance": distanze[i] if i < len(distanze) else None,
                })
            return out
        except Exception as e:
            print("[MemoryDB] search error: " + str(e))
            return []

    def get_recent_memos(self, n=3):
        """Ultimi N memo."""
        with self._conn() as conn:
            rows = conn.execute("SELECT * FROM memos ORDER BY id DESC LIMIT ?", (n,)).fetchall()
            return [dict(r) for r in rows]

    def save_specialist_report(self, memo_id, specialist, round_n, content, output_language=None):
        from bellomberg.core.language import capture_language
        selected = capture_language(output_language)
        with self._conn() as conn:
            conn.execute("INSERT INTO specialist_reports (memo_id, specialist, round_n, content, output_language) VALUES (?,?,?,?,?)",
                          (memo_id, specialist, round_n, content, selected))

    def save_llm_usage(self, memo_id, usage_log, output_language=None):
        """Persiste il consumo LLM della run (una riga per agente+round).
        usage_log = lista di dict prodotta da Blackboard.record_usage.
        Ritorna quante righe ha inserito.
        NB: cost_eur resta NULL se il round non e' prezzabile (modello fuori listino
        o FX assente): il buco si dichiara con cost_status, non si azzera."""
        if not usage_log:
            return 0
        from bellomberg.core.language import capture_language
        selected = capture_language(output_language)
        rows = []
        for e in usage_log:
            rows.append((
                memo_id,
                e.get("agent"),
                e.get("round"),
                e.get("model") or "n.d.",   # model e' NOT NULL: modello ignoto = dichiarato, non inventato
                e.get("in", 0) or 0,
                e.get("out", 0) or 0,
                e.get("cache_read", 0) or 0,
                e.get("cache_write", 0) or 0,
                e.get("api_calls", 0) or 0,
                e.get("duration_s"),
                e.get("cost_eur"),
                e.get("fx_rate"),
                e.get("fx_source"),
                e.get("status"),
                e.get("cache_ttl"),
                selected,
            ))
        with self._conn() as conn:
            conn.executemany("""INSERT INTO llm_usage
                                (memo_id, agent, round_n, model, tokens_in, tokens_out,
                                 cache_read, cache_write, api_calls, duration_s, cost_eur,
                                 fx_rate, fx_source, cost_status, cache_ttl, output_language)
                                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", rows)
        return len(rows)

    # Fusione 15/07/2026: eventdesk EREDITA a lettura la storia firmata da
    # news+politics (report, feedback PM). Alias read-time, nessuna migrazione:
    # i dati nel DB restano coi nomi originali, reversibile togliendo la riga.
    _LEGACY_SPECIALIST_ALIASES = {"eventdesk": ("news", "politics")}

    def get_recent_specialist_reports(self, specialist, n=3):
        """Ultimi N report di uno specialista (round 2, finalizzato).
        eventdesk eredita anche i report legacy news/politics (alias fusione 15/07)."""
        names = (specialist,) + self._LEGACY_SPECIALIST_ALIASES.get(specialist, ())
        ph = ",".join("?" * len(names))
        with self._conn() as conn:
            rows = conn.execute("""SELECT sr.*, m.full_markdown AS memo_full_markdown
                                    FROM specialist_reports sr
                                    LEFT JOIN memos m ON m.id=sr.memo_id
                                    WHERE sr.specialist IN (""" + ph + """) AND sr.round_n=2
                                    ORDER BY sr.id DESC LIMIT ?""", (*names, n)).fetchall()
            return [dict(r) for r in rows]

    # ========================================================
    # DECISIONS (extracted from Capo memo ACTION TABLE)
    # ========================================================

    def extract_and_save_decisions(self, memo_id, memo_markdown, usage_out=None):
        """Parse ACTION TABLE dal memo del Capo. Estrae righe e salva in decisions.
        #200c (15/07): prima via = estrazione STRUTTURATA (Sonnet, tool forzato,
        celle verbatim); la regex storica resta come FALLBACK DICHIARATO. L'EUR
        lo converte in entrambi i casi _parse_eur_amount (contratto invariato).

        usage_out (opzionale, #44/finding 4): dict passato PARI PARI a
        extract_rows_structured, che lo riempie col consumo della chiamata Sonnet
        (il chiamante la registra su blackboard.record_usage). Firma retro-compatibile:
        regenerate_memo.py e gli altri chiamanti restano invariati."""
        ts = datetime.now().isoformat(timespec="seconds")
        rows = []
        structured_ok = False
        try:
            from bellomberg.agents.action_table_extract import extract_rows_structured
            sr = extract_rows_structured(memo_markdown, usage_out=usage_out)
            if sr.get("error"):
                print("[MemoryDB] #200c estrazione strutturata KO ("
                      + str(sr["error"])[:150] + "): fallback parser regex")
            else:
                rows = [(r["action"], r["ticker"], _parse_eur_amount(r["size_raw"]),
                         r["timing"], r["confidence"]) for r in sr["rows"]]
                structured_ok = True
                print("[MemoryDB] #200c ACTION TABLE strutturata: "
                      + str(len(rows)) + " righe (tool forzato)")
                # Fable 5 16/07 (residuo del bug run #45): 0 righe SENZA errore con il
                # marker '## ACTION TABLE' PRESENTE nel memo = risposta sospetta di
                # Sonnet, non prova di assenza -> si attiva il fallback regex invece
                # di salvare una pagina decisioni vuota in silenzio.
                if not rows and re.search(r"##\s*ACTION TABLE", memo_markdown or "",
                                          re.IGNORECASE):
                    structured_ok = False
                    print("[MemoryDB] #200c CROSS-CHECK: 0 righe ma il marker ACTION "
                          "TABLE esiste nel memo -> fallback parser regex (dichiarato)")
        except Exception as e:
            print("[MemoryDB] #200c estrazione strutturata exception ("
                  + str(e)[:120] + "): fallback parser regex")

        if not structured_ok:
            # --- FALLBACK: parser regex storico (invariato) ---
            action_table_match = re.search(
                r"##\s*ACTION TABLE.*?\n(\|.*?\|.*?\n)+",
                memo_markdown, re.IGNORECASE | re.DOTALL)
            if not action_table_match:
                return []
            table_text = action_table_match.group(0)
            for line in table_text.split("\n"):
                if not line.startswith("|"): continue
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                # Skip header e separator
                from bellomberg.core.language import ACTION_TABLE_HEADERS
                if not cells or "---" in cells[0] or cells[0].lower() in ACTION_TABLE_HEADERS:
                    continue
                if len(cells) < 5:
                    continue
                action = cells[0]
                ticker = cells[1]
                eur_str = cells[2]
                timing = cells[3]
                confidence = cells[4]
                # Parse EUR amount (robusto: gestisce k/mila/mln, formato europeo, range)
                eur = _parse_eur_amount(eur_str)
                rows.append((action, ticker, eur, timing, confidence))
        if not rows:
            return []

        # ---- GUARDIA BOOK-AWARE (#31, deterministica) ----
        # Riconcilia OGNI decisione parsata con il book reale PRIMA del salvataggio:
        # - ticker GIA' in portafoglio -> nota nel rationale 'POSIZIONE ESISTENTE: peso X% dal GG/MM'
        #   e action BUY normalizzata in ADD (vietato BUY su posseduto);
        # - ticker con un trade negli ultimi 7 giorni -> prefisso '[!] ACQUISTATO IL GG/MM:'
        #   (RIDOTTO/VENDUTO per TRIM/SELL) nel testo della decisione.
        # Il contratto del parser (5 colonne -> action/ticker/eur/timing/confidence, status
        # PENDING, ritorno lista id) resta INVARIATO; rationale era una colonna gia'
        # esistente e mai valorizzata (il FE F10 la mostra nell'expand della riga).
        def _ddmm(s):
            s = str(s or "")
            return (s[8:10] + "/" + s[5:7]) if len(s) >= 10 else "?"

        guard_pos = {}
        try:
            for p in (self.get_portfolio_summary().get("positions") or []):
                t = (p.get("ticker") or "").upper()
                if t:
                    guard_pos[t] = p
        except Exception as e:
            print("[MemoryDB] book-guard: portafoglio non disponibile (" + str(e) + ")")

        decision_ids = []
        with self._conn() as conn:
            guard_trades = {}
            try:
                cutoff = (datetime.now() - timedelta(days=7)).isoformat(timespec="seconds")
                for tr in conn.execute(
                        "SELECT ticker, action, quantita, data FROM trade_history "
                        "WHERE data >= ? AND action != 'DIVIDEND' ORDER BY data DESC",
                        (cutoff,)).fetchall():
                    tt = (tr["ticker"] or "").upper()
                    if tt and tt not in guard_trades:
                        guard_trades[tt] = tr
            except Exception as e:
                print("[MemoryDB] book-guard: trade_history non disponibile (" + str(e) + ")")

            for (action, ticker, eur, timing, confidence) in rows:
                rationale = None
                try:
                    tk = (ticker or "").upper()
                    note_parts = []
                    tr = guard_trades.get(tk)
                    if tr:
                        verbo = {"BUY": "ACQUISTATO", "ADD": "ACQUISTATO",
                                 "TRIM": "RIDOTTO", "SELL": "VENDUTO"}.get(
                                     (tr["action"] or "").upper(), "TRADATO")
                        note_parts.append("[!] " + verbo + " IL " + _ddmm(tr["data"]) + ":")
                    pos = guard_pos.get(tk)
                    if pos:
                        peso = pos.get("peso_pct")
                        peso_s = ("%.1f%%" % peso) if isinstance(peso, (int, float)) else "n/d"
                        note_parts.append("POSIZIONE ESISTENTE: peso " + peso_s +
                                          " dal " + _ddmm(pos.get("data_apertura")))
                        if (action or "").strip().upper() == "BUY":
                            action = "ADD"
                            note_parts.append("(BUY normalizzato in ADD dalla guardia "
                                              "book-aware: titolo gia' in portafoglio)")
                    if note_parts:
                        rationale = " ".join(note_parts)
                except Exception as e:
                    print("[MemoryDB] book-guard skip su " + str(ticker) + ": " + str(e))
                cur = conn.execute("""INSERT INTO decisions (memo_id, timestamp, action, ticker,
                                                              eur_amount, timing, confidence,
                                                              rationale, status)
                                       VALUES (?,?,?,?,?,?,?,?,'PENDING')""",
                                    (memo_id, ts, action, ticker, eur, timing, confidence, rationale))
                decision_ids.append(cur.lastrowid)
        return decision_ids

    @staticmethod
    def _require_valuation_snapshot_schema(conn):
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if not {"valuation_snapshots", "valuation_snapshot_links"}.issubset(tables):
            raise RuntimeError(_storage_text('valuation snapshot schema assente: eseguire prima il dry-run di tools/migrations/migra_valuation_metadata.py', 'Valuation snapshot schema missing: first run the dry-run of tools/migrations/migra_valuation_metadata.py'))

    def _save_valuation_snapshot(self, conn, ticker, payload):
        import hashlib
        import uuid
        self._require_valuation_snapshot_schema(conn)
        if not isinstance(payload, dict):
            raise ValueError(_storage_text('valuation payload deve essere un oggetto', 'Valuation payload must be an object'))
        ticker = str(ticker).strip().upper()
        if not re.fullmatch(r"[A-Z0-9^][A-Z0-9.^=_/-]*", ticker):
            raise ValueError(_storage_text('snapshot ticker mancante o invalido', 'Snapshot ticker missing or invalid'))
        if str(payload.get("ticker", "")).strip().upper() != ticker:
            raise ValueError(_storage_text('snapshot ticker discordante', 'Snapshot ticker mismatch'))
        snapshot_id, generation_id = payload.get("snapshot_id"), payload.get("generation_id")
        if not isinstance(snapshot_id, str) or not re.fullmatch(r"[0-9a-f]{64}", snapshot_id):
            raise ValueError(_storage_text('snapshot_id SHA256 mancante o invalido', 'snapshot_id SHA256 missing or invalid'))
        try:
            uuid.UUID(str(generation_id))
        except (ValueError, AttributeError):
            raise ValueError(_storage_text('generation_id UUID mancante o invalido', 'generation_id UUID missing or invalid')) from None
        encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)
        existing = conn.execute("SELECT payload_json FROM valuation_snapshots WHERE snapshot_id=? AND generation_id=?",
                                (snapshot_id, generation_id)).fetchone()
        if existing:
            if existing[0] != encoded:
                raise ValueError(_storage_text('snapshot immutabile: la generazione esiste con contenuto diverso', 'Immutable snapshot: the generation exists with different content'))
            return snapshot_id
        conn.execute("INSERT INTO valuation_snapshots VALUES (?,?,?,?,?,?,?)",
                     (snapshot_id, generation_id, ticker, 1, encoded,
                      hashlib.sha256(encoded.encode("utf-8")).hexdigest(), datetime.now().isoformat(timespec="microseconds")))
        return snapshot_id

    def _link_valuation_snapshot(self, conn, snapshot_id, generation_id, thesis_id=None, decision_id=None):
        self._require_valuation_snapshot_schema(conn)
        if (thesis_id is None) == (decision_id is None):
            raise ValueError(_storage_text('specificare thesis_id oppure decision_id', 'Specify thesis_id or decision_id'))
        snapshot = conn.execute("SELECT ticker FROM valuation_snapshots WHERE snapshot_id=? AND generation_id=?",
                                (snapshot_id, generation_id)).fetchone()
        column, table, row_id = ("thesis_id", "valuation_theses", thesis_id) if thesis_id is not None else ("decision_id", "decisions", decision_id)
        row = conn.execute(f"SELECT ticker FROM {table} WHERE id=?", (row_id,)).fetchone()
        if snapshot is None or row is None or str(snapshot[0]).upper() != str(row[0]).upper():
            raise ValueError(_storage_text('snapshot/riferimento assente o ticker discordante', 'Snapshot/reference missing or ticker mismatch'))
        if conn.execute(f"SELECT 1 FROM valuation_snapshot_links WHERE snapshot_id=? AND generation_id=? AND {column}=?",
                        (snapshot_id, generation_id, row_id)).fetchone():
            return
        conn.execute("INSERT INTO valuation_snapshot_links(snapshot_id,generation_id,thesis_id,decision_id,created_at) VALUES (?,?,?,?,?)",
                     (snapshot_id, generation_id, thesis_id, decision_id, datetime.now().isoformat(timespec="microseconds")))

    def save_valuation_snapshot(self, ticker, payload, *, thesis_id=None, decision_id=None):
        """Append the exact result, including incomplete research and its acquisition bundle."""
        with self._conn() as conn:
            snapshot_id = self._save_valuation_snapshot(conn, ticker, payload)
            if thesis_id is not None or decision_id is not None:
                self._link_valuation_snapshot(conn, snapshot_id, payload["generation_id"], thesis_id, decision_id)
            return snapshot_id

    def link_valuation_snapshot(self, snapshot_id, *, generation_id, thesis_id=None, decision_id=None):
        with self._conn() as conn:
            self._link_valuation_snapshot(conn, snapshot_id, generation_id, thesis_id, decision_id)

    def get_valuation_snapshot(self, snapshot_id, *, generation_id=None):
        """Read through the common gate; a reused input hash requires an exact generation."""
        from bellomberg.valuation.dcf_quality import normalize_valuation_payload
        import hashlib
        with self._conn() as conn:
            self._require_valuation_snapshot_schema(conn)
            query = "SELECT payload_json,payload_sha256 FROM valuation_snapshots WHERE snapshot_id=?"
            args = [snapshot_id]
            if generation_id is not None:
                query += " AND generation_id=?"
                args.append(generation_id)
            rows = conn.execute(query, args).fetchall()
        if not rows:
            return None
        if len(rows) != 1:
            raise ValueError(_storage_text('snapshot ambiguo: specificare generation_id', 'Ambiguous snapshot: specify generation_id'))
        encoded, digest = rows[0]
        if hashlib.sha256(encoded.encode("utf-8")).hexdigest() != digest:
            raise ValueError(_storage_text('snapshot integrity SHA256 discordante', 'Snapshot integrity SHA256 mismatch'))
        return normalize_valuation_payload(json.loads(encoded))

    def get_latest_valuation_snapshots(self):
        """Latest generation per canonical ticker, including cases with no workbook."""
        with self._conn() as conn:
            self._require_valuation_snapshot_schema(conn)
            rows = conn.execute("SELECT snapshot_id,generation_id,ticker,created_at FROM valuation_snapshots ORDER BY created_at DESC,rowid DESC").fetchall()
        latest = {}
        for row in rows:
            if row["ticker"] not in latest:
                latest[row["ticker"]] = {"created_at": row["created_at"], "payload": self.get_valuation_snapshot(
                    row["snapshot_id"], generation_id=row["generation_id"])}
        return latest

    def save_valuation_thesis(self, ticker, variant_view=None, growth_path=None, price=None,
                              fair_value=None, ebitda_margin_target=None, terminal_growth=None,
                              engine=None, subsector=None, memo_id=None,
                              sanity_severity=None, sanity_headline=None, profile_key=None,
                              valuation_payload=None, reuse_generation=False):
        """#198 p.3: salva la tesi di valutazione dell'agente (variant view + growth) per
        verificarla week-on-week vs quello che la societa' consegna.
        audit/12 V0.3: porta anche il giudizio sanity del motore (OK/WARN + motivo), cosi'
        S2: anche le analisi incomplete/BLOCK vengono archiviate con snapshot;
        il fair value della tesi resta NULL quando il controllo comune lo rifiuta.
        audit/13 V1.7a: profile_key = chiave PULITA del profilo sub-settore (la colonna
        subsector resta la traccia legacy _matched, forme miste); colonna nuova via
        tools/migrations/migra_profile_key.py — su DB non migrato il buco si DICHIARA."""
        import json as _j
        from datetime import datetime as _dt
        try:
            with self._conn() as conn:
                if reuse_generation:
                    if valuation_payload is None:
                        raise ValueError("reuse_generation requires an exact valuation payload")
                    # A resumed worker must not duplicate the thesis after a
                    # crash between registration and its next checkpoint.
                    conn.execute("BEGIN IMMEDIATE")
                if valuation_payload is not None:
                    from bellomberg.valuation.dcf_quality import normalize_valuation_payload
                    normalized = normalize_valuation_payload(valuation_payload)
                    fair_value = next((normalized[key] for key in (
                        "fair_value_final", "fair_value_weighted", "fair_value_blend", "fair_value_base",
                        "fair_value_nav", "fair_value") if normalized.get(key) is not None), None
                    ) if normalized["valuation_usability"]["usable"] else None
                    self._save_valuation_snapshot(conn, ticker, valuation_payload)
                    if reuse_generation:
                        saved = conn.execute("SELECT thesis_id FROM valuation_snapshot_links "
                            "WHERE snapshot_id=? AND generation_id=? AND thesis_id IS NOT NULL ORDER BY id LIMIT 1",
                            (valuation_payload["snapshot_id"], valuation_payload["generation_id"])).fetchone()
                        if saved:
                            return saved[0]
                cols = {row[1] for row in conn.execute("PRAGMA table_info(valuation_theses)")}
                if profile_key and "profile_key" not in cols:
                    print("[memory_db] colonna profile_key assente in valuation_theses: "
                          "NON salvata — lancia tools/migrations/migra_profile_key.py")
                if "sanity_severity" in cols and "profile_key" in cols:
                    cur = conn.execute(
                        "INSERT INTO valuation_theses (ticker, date, price_at_thesis, fair_value, "
                        "growth_path, ebitda_margin_target, terminal_growth, variant_view, engine, "
                        "subsector, memo_id, sanity_severity, sanity_headline, profile_key) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (ticker.upper(), _dt.now().isoformat(timespec="seconds"), price, fair_value,
                         _j.dumps(growth_path) if growth_path else None, ebitda_margin_target,
                         terminal_growth, variant_view, engine, subsector, memo_id,
                         sanity_severity, sanity_headline, profile_key))
                elif "sanity_severity" in cols:
                    cur = conn.execute(
                        "INSERT INTO valuation_theses (ticker, date, price_at_thesis, fair_value, "
                        "growth_path, ebitda_margin_target, terminal_growth, variant_view, engine, "
                        "subsector, memo_id, sanity_severity, sanity_headline) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (ticker.upper(), _dt.now().isoformat(timespec="seconds"), price, fair_value,
                         _j.dumps(growth_path) if growth_path else None, ebitda_margin_target,
                         terminal_growth, variant_view, engine, subsector, memo_id,
                         sanity_severity, sanity_headline))
                else:
                    # DB non ancora migrato: la tesi si salva comunque, il buco si DICHIARA
                    if sanity_severity or sanity_headline:
                        print("[memory_db] colonne sanity assenti in valuation_theses: severity/"
                              "headline NON salvati — lancia tools/migrations/migra_tesi_sanity.py")
                    cur = conn.execute(
                        "INSERT INTO valuation_theses (ticker, date, price_at_thesis, fair_value, "
                        "growth_path, ebitda_margin_target, terminal_growth, variant_view, engine, "
                        "subsector, memo_id) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                        (ticker.upper(), _dt.now().isoformat(timespec="seconds"), price, fair_value,
                         _j.dumps(growth_path) if growth_path else None, ebitda_margin_target,
                         terminal_growth, variant_view, engine, subsector, memo_id))
                thesis_id = cur.lastrowid
                if valuation_payload is not None:
                    self._link_valuation_snapshot(conn, valuation_payload["snapshot_id"],
                                                  valuation_payload["generation_id"], thesis_id=thesis_id)
                return thesis_id
        except Exception as e:
            print("[memory_db] save_valuation_thesis failed: " + str(e))
            return None

    def get_valuation_history(self, ticker, n=10):
        """Storico tesi di valutazione su un ticker (per verificare se reggono nel tempo)."""
        from bellomberg.valuation.dcf_quality import assess_valuation_usability
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT * FROM valuation_theses WHERE UPPER(ticker)=? ORDER BY id DESC LIMIT ?",
                (ticker.upper(), n)).fetchall()
            history = [dict(r) for r in rows]
            has_snapshots = conn.execute("SELECT 1 FROM sqlite_master WHERE name='valuation_snapshot_links'").fetchone()
            for item in history:
                link = conn.execute("SELECT snapshot_id,generation_id FROM valuation_snapshot_links WHERE thesis_id=?",
                                    (item["id"],)).fetchone() if has_snapshots else None
                if link:
                    item["valuation_snapshot_id"] = item["snapshot_id"] = link["snapshot_id"]
                    item["generation_id"] = link["generation_id"]
                    item["valuation_payload"] = self.get_valuation_snapshot(link["snapshot_id"], generation_id=link["generation_id"])
                    item["valuation_usability"] = item["valuation_payload"]["valuation_usability"]
                    if not item["valuation_usability"]["usable"]:
                        item["fair_value"] = None
                else:
                    item["valuation_payload"] = None
                    item["valuation_usability"] = assess_valuation_usability(item)
            return history

    # ============================================================
    # V6 GUIDANCE (audit/19, decisioni PM D1-D4 23/07)
    # ============================================================
    # bande di plausibilita' per metric (D1, DICHIARATE): valori come FRAZIONI per
    # pct; fuori banda = rifiuto spiegato, mai un numero assurdo nel registro
    # review V6 M2: capex/ricavi oltre 1 e' REALE nei capital-intensive in ciclo di
    # investimento (Terna ~0.8): banda larga 0..2, il registro deve poterli contenere
    GUIDANCE_BANDS = {"revenue_growth": (-0.90, 3.00), "ebitda_margin": (0.0, 1.0),
                      "gross_margin": (0.0, 1.0), "capex_pct": (0.0, 2.00),
                      "revenue_abs": (0.0, 5e7), "eps": (-1e4, 1e5), "other": None}
    GUIDANCE_UNITS = {"revenue_growth": "pct", "ebitda_margin": "pct", "gross_margin": "pct",
                      "capex_pct": "pct", "revenue_abs": "musd", "eps": "eps"}
    GUIDANCE_FALLBACK_DAYS = 120   # D4: senza data earnings dal calendar

    @staticmethod
    def _chiave_unit(unit):
        """Identita' di una GRANDEZZA a partire dall'etichetta `unit`.

        Per metric='other' l'unit e' l'unica cosa che dice CHE COSA e' il numero
        ('meur (EBITA FY)' vs 'meur (FOCF FY)'), quindi decide chi sostituisce chi.
        Due etichette che l'occhio legge uguali devono dare la stessa chiave: qui
        si normalizzano forme unicode compatibili (NFKC), spazi non-ASCII (NBSP,
        spazio sottile: invisibili a schermo), caratteri a larghezza zero, spazi
        ripetuti, e maiuscole con casefold (che, a differenza di lower() di SQLite,
        copre anche gli accenti).
        NON normalizza il CONTENUTO: 'meur (EBITA FY)' e 'meur (EBITA FY26)' restano
        due chiavi diverse, e infatti add_guidance le dichiara in 'convivono'.
        """
        import unicodedata as _ud
        s = _ud.normalize("NFKC", str(unit or ""))
        s = "".join(" " if _ud.category(ch) == "Zs" else ch
                    for ch in s if ch not in "​‌‍⁠﻿")
        return " ".join(s.split()).casefold()

    def add_guidance(self, ticker, metric, period, value_mid, value_low=None,
                     value_high=None, unit=None, source_doc=None, source_date=None,
                     valid_until=None, valid_until_source=None, entered_by=None,
                     note=None, today=None):
        """Registra una guidance SOCIETARIA (D1: fonte OBBLIGATORIA — source_doc +
        source_date; senza, rifiuto spiegato). Supersede automatico della riga attiva
        stesso ticker+metric+period (storico conservato). valid_until: la passa il
        chiamante (prossima trimestrale dal calendar); assente = effective+120g
        DICHIARATO (D4). Ritorna dict con esito, mai eccezioni verso il chiamante."""
        import re as _re
        from datetime import date, timedelta
        try:
            t = str(ticker or "").upper().strip()
            m = str(metric or "").strip().lower()
            # review V6 M3: normalizzazione LEGGERA del period — "FY26"/"FY 2026"/
            # "Q3 2026" collassano su un formato solo (FY2026 / Q3-2026), altrimenti
            # il supersede su match esatto salta e restano due righe attive per la
            # stessa guidance vera. Formati fuori pattern restano com'erano.
            p = str(period or "").strip().upper().replace(" ", "-")
            _pm = _re.fullmatch(r"(FY|Q[1-4]|H[12])-?(\d{2}|\d{4})", p)
            if _pm:
                _yr = _pm.group(2) if len(_pm.group(2)) == 4 else "20" + _pm.group(2)
                p = ("FY" + _yr) if _pm.group(1) == "FY" else f"{_pm.group(1)}-{_yr}"
            if not t or not p:
                return {"error": "ticker e period obbligatori"}
            if m not in self.GUIDANCE_BANDS:
                return {"error": "metric '%s' non valida (attese: %s)"
                                 % (metric, ", ".join(sorted(self.GUIDANCE_BANDS)))}
            if not (source_doc and str(source_doc).strip()) or not (source_date and str(source_date).strip()):
                return {"error": ("fonte OBBLIGATORIA (D1): servono source_doc (documento/"
                                  "riferimento, es. 'press release Q2 FY26, sito IR') E "
                                  "source_date — una guidance senza fonte NON entra nel registro")}
            vals = {}
            for k, v in (("value_mid", value_mid), ("value_low", value_low),
                         ("value_high", value_high)):
                if v is None and k != "value_mid":
                    vals[k] = None
                    continue
                try:
                    vals[k] = float(v)
                except (TypeError, ValueError):
                    return {"error": f"{k} non numerico: {v!r}"}
            band = self.GUIDANCE_BANDS[m]
            if band:
                lo_b, hi_b = band
                for k, v in vals.items():
                    if v is not None and not lo_b <= v <= hi_b:
                        return {"error": ("%s %.4g fuori dalla banda di plausibilita' "
                                          "%.4g..%.4g per metric '%s' (dichiarata): input "
                                          "rotto o unita' sbagliata (pct = FRAZIONE, es. "
                                          "0.18 per 18%%)" % (k, v, lo_b, hi_b, m))}
            lo_, mid_, hi_ = vals["value_low"], vals["value_mid"], vals["value_high"]
            if (lo_ is not None and lo_ > mid_) or (hi_ is not None and hi_ < mid_):
                return {"error": "range incoerente: serve value_low <= value_mid <= value_high"}
            # review V6 B2: l'unit e' FISSA per metrica (la banda la presume) — un
            # override zitto su revenue_abs ("miliardi") renderebbe la banda cieca;
            # unita' custom = metric 'other' dichiarata
            _default_u = self.GUIDANCE_UNITS.get(m)
            if unit and _default_u and str(unit).strip() != _default_u:
                return {"error": ("unit '%s' non ammessa per metric '%s' (fissa: '%s', la "
                                  "banda di plausibilita' la presume) — per unita' custom "
                                  "usa metric 'other' con unit dichiarata" % (unit, m, _default_u))}
            u = str(unit).strip() if unit else _default_u
            if not u:
                return {"error": "metric 'other' richiede unit ESPLICITA (dichiarata)"}
            # review V6 M1/B1: le date si PARSANO (ISO) — un valid_until non-ISO
            # renderebbe la guidance immortale in silenzio (confronto stringhe)
            try:
                sd = date.fromisoformat(str(source_date)[:10])
            except (TypeError, ValueError):
                return {"error": ("source_date '%s' non ISO (atteso YYYY-MM-DD): la fonte "
                                  "D1 richiede una data parsabile" % source_date)}
            _today = today or date.today()
            eff = _today.isoformat()
            if valid_until:
                try:
                    vu = date.fromisoformat(str(valid_until)[:10]).isoformat()
                except (TypeError, ValueError):
                    return {"error": ("valid_until '%s' non ISO (atteso YYYY-MM-DD): una "
                                      "scadenza non parsabile = guidance immortale in "
                                      "silenzio, rifiutata" % valid_until)}
                vus = valid_until_source or "chiamante"
            else:
                vu = (_today + timedelta(days=self.GUIDANCE_FALLBACK_DAYS)).isoformat()
                vus = "fallback effective+%dg (calendar n.d., D4)" % self.GUIDANCE_FALLBACK_DAYS
            with self._conn() as conn:
                cur = conn.execute(
                    "INSERT INTO company_guidance (ticker, metric, period, value_low, "
                    "value_mid, value_high, unit, source_doc, source_date, effective_date, "
                    "valid_until, valid_until_source, status, entered_by, note, created_at) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?, 'active', ?, ?, ?)",
                    (t, m, p, lo_, mid_, hi_, u, str(source_doc).strip(),
                     str(source_date)[:10], eff, vu, vus, entered_by, note,
                     datetime.now().isoformat(timespec="seconds")))
                new_id = cur.lastrowid
                # 20/08 (gate trimestrali arretrate, Opus 5): il supersede sostituisce
                # la STESSA grandezza, e per metric='other' la grandezza la dice `unit`
                # — i 5 target FY26 di Leonardo sono tutti in euro (revenue_abs ha unit
                # fissa 'musd'), quindi tutti 'other': senza `unit` nel confronto l'EBITA
                # veniva dichiarato "sostituito" dal FOCF. Misurato su copia del DB:
                # 3 righe inserite, 1 sola attiva.
                # Il confronto si fa in PYTHON, non con lower()/trim() di SQLite:
                # lower() di SQLite e' solo-ASCII ('UTILITÀ' -> 'utilitÀ') e trim()
                # non tocca gli spazi INTERNI, quindi due etichette identiche a schermo
                # (una con NBSP, una con spazio normale) restavano due righe attive con
                # due numeri diversi per la stessa grandezza — il rovescio del rischio,
                # trovato dalla review avversariale del 20/08.
                # Per le metriche a unit fissa (pct/musd/eps) il comportamento e'
                # identico a prima, perche' l'unit e' imposta da GUIDANCE_UNITS e non
                # puo' variare (override rifiutato sopra).
                _chiave = self._chiave_unit(u)
                gemelle, accanto = [], []
                for _r in conn.execute(
                        "SELECT id, unit, value_mid FROM company_guidance WHERE ticker=? "
                        "AND metric=? AND period=? AND status='active' AND id != ?",
                        (t, m, p, new_id)).fetchall():
                    (gemelle if self._chiave_unit(_r["unit"]) == _chiave
                     else accanto).append(_r)
                n_sup = 0
                for _r in gemelle:
                    n_sup += conn.execute(
                        "UPDATE company_guidance SET status='superseded', superseded_by=? "
                        "WHERE id=?", (new_id, _r["id"])).rowcount
            out = {"ok": True, "id": new_id, "ticker": t, "metric": m, "period": p,
                   "valid_until": vu, "valid_until_source": vus,
                   "superseded": n_sup,
                   "nota": ("guidance registrata con fonte; sostituite %d righe attive "
                            "precedenti (storico conservato)" % n_sup) if n_sup else
                           "guidance registrata con fonte"}
            # regola 14/07 applicata al rovescio del supersede: `superseded: 0` da solo
            # e' AMBIGUO — vale sia "prima registrazione" (legittimo) sia "ho lasciato
            # attiva una riga che forse e' la stessa grandezza con un'altra etichetta".
            # Le righe che restano accanto si DICHIARANO, con i loro numeri, cosi' il
            # chiamante (agente o PM) puo' accorgersi del doppione.
            if accanto:
                out["convivono"] = [{"id": _r["id"], "unit": _r["unit"],
                                     "value_mid": _r["value_mid"]} for _r in accanto]
                out["nota"] += ("; RESTANO ATTIVE %d righe con stesso ticker/metric/period "
                                "e unit DIVERSA (in 'convivono'): se una di quelle e' la "
                                "STESSA grandezza scritta con un'altra etichetta, il "
                                "registro ora afferma due numeri per la stessa cosa — "
                                "riscrivila con la stessa `unit` per sostituirla"
                                % len(accanto))
            # review V6 B6: backfill di un documento vecchio col fallback 120g dalla
            # REGISTRAZIONE = "fresca" oltre il dovuto — dichiarato, mai zitto
            if "fallback" in vus and (_today - sd).days > 90:
                out["warning"] = ("source_date %s e' di %d giorni fa e la scadenza e' il "
                                  "fallback +120g dalla REGISTRAZIONE: se nel frattempo e' "
                                  "uscita una trimestrale nuova, registra QUELLA guidance "
                                  "(questa verrebbe superseded)" % (sd.isoformat(),
                                                                    (_today - sd).days))
            return out
        except Exception as e:
            return {"error": f"add_guidance: {type(e).__name__}: {e}"}

    def get_guidance(self, ticker, include_history=False, today=None):
        """Guidance del ticker: attive con STALENESS calcolata in lettura (D4: oltre
        valid_until = stale DICHIARATA, la colonna status non si tocca) + storiche a
        richiesta. Mai eccezioni verso il chiamante."""
        from datetime import date
        try:
            t = str(ticker or "").upper().strip()
            if not t:
                return {"error": "ticker mancante"}
            _today = (today or date.today()).isoformat()
            with self._conn() as conn:
                act = [dict(r) for r in conn.execute(
                    "SELECT * FROM company_guidance WHERE ticker=? AND status='active' "
                    "ORDER BY metric, period", (t,))]
                out = {"ticker": t, "active": act, "as_of": _today}
                for g in act:
                    g["stale"] = bool(g.get("valid_until") and _today > g["valid_until"])
                    if g["stale"]:
                        g["stale_note"] = ("STALE: scaduta il %s (D4) — NON guida piu' i "
                                           "default del modello, aggiornarla dalla "
                                           "trimestrale nuova" % g["valid_until"])
                n_stale = sum(1 for g in act if g["stale"])
                if n_stale:
                    out["nota"] = "%d guidance ATTIVE ma STALE (dichiarate riga per riga)" % n_stale
                if include_history:
                    out["history"] = [dict(r) for r in conn.execute(
                        "SELECT * FROM company_guidance WHERE ticker=? AND status != 'active' "
                        "ORDER BY id DESC LIMIT 50", (t,))]
                    if len(out["history"]) == 50:
                        # review V6 B3: troncamento DICHIARATO, mai zitto
                        out["history_truncated"] = "storico troncato alle 50 righe piu' recenti"
            return out
        except Exception as e:
            return {"error": f"get_guidance: {type(e).__name__}: {e}"}

    def get_recent_trades(self, n=15):
        """Ultimi N trade ESEGUITI dal PM con note e rationale (per la memoria del consigliere)."""
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT id, ticker, action, quantita, prezzo, valuta, data, note, pm_rationale, "
                "linked_decision_id, created_at, ora_convenzionale, link_origin, fx_fonte, "
                "realized_local, realized_eur FROM trade_history ORDER BY data DESC, id DESC LIMIT ?", (n,)).fetchall()
            return [dict(r) for r in rows]

    # verso di una decisione/trade: BUY e ADD comprano, TRIM e SELL vendono
    _VERSO = {"BUY": "LONG", "ADD": "LONG", "TRIM": "SHORT", "SELL": "SHORT"}

    def esecuzioni_delle_decisioni(self, decisioni, giorni=7):
        """{decision_id: {trades, trade_ids, eur, pct, inferito, data}} per le decisioni
        EXECUTED: i trade che le hanno eseguite. Legame ESPLICITO (linked_decision_id)
        oppure INFERITO: stesso ticker, stesso verso, trade nei `giorni` DOPO la decisione,
        trade non legato ad altro. `eur` = somma quantita' x prezzo SOLO se tutti i trade
        sono in EUR (altrimenti None: niente cambio inventato); `pct` = eur / proposto.

        Audit 11/09 (Fable 5.1, run 10/09): il blotter non invia mai linked_decision_id
        (nessun trade legato), quindi «esegue decisione #N» non compariva MAI e il Capo
        leggeva di una decisione eseguita il solo importo PROPOSTO, non quello ESEGUITO.
        Da qui una decisione gia' eseguita tornava in tavola come ancora da completare."""
        out = {}
        requested = [d for d in (decisioni or [])
               if str(d.get("status") or "").upper() in ("PENDING", "EXECUTED", "PARTIAL")
               and d.get("ticker") and d.get("id") is not None]
        if not requested:
            return out
        with self._conn() as conn:
            righe = conn.execute(
                "SELECT id, ticker, action, quantita, prezzo, valuta, data, linked_decision_id, link_origin, ora_convenzionale "
                "FROM trade_history WHERE action != 'DIVIDEND' ORDER BY data,id").fetchall()
            # Candidate assignment is global: asking for one decision must not
            # steal a trade from a closer decision omitted from that view.
            exe = [dict(r) for r in conn.execute(
                "SELECT id,ticker,action,timestamp,status FROM decisions WHERE status IN ('EXECUTED','PARTIAL')")]
        trades = [dict(r) for r in righe]
        explicit_ids = {t["linked_decision_id"] for t in trades if t.get("linked_decision_id") is not None}
        assignments = {}
        for trade in trades:
            if (trade.get("linked_decision_id") is not None or trade.get("link_origin") == "none"
                    or trade.get("ora_convenzionale") == 1):
                continue
            eligible = []
            for d in exe:
                if d["id"] in explicit_ids:
                    continue  # explicit execution remains authoritative for that decision
                if (str(d["ticker"]).strip().upper() != str(trade["ticker"]).strip().upper()
                        or self._VERSO.get(d["action"]) is None
                        or self._VERSO.get(d["action"]) != self._VERSO.get(trade["action"])):
                    continue
                try:
                    at, start = datetime.fromisoformat(trade["data"]), datetime.fromisoformat(d["timestamp"])
                    if start <= at <= start + timedelta(days=int(giorni)):
                        eligible.append((start, d["id"]))
                except (ValueError, TypeError):
                    continue
            if eligible:
                latest = max(start for start, _ in eligible)
                nearest = [did for start, did in eligible if start == latest]
                if len(nearest) == 1:
                    assignments[trade["id"]] = nearest[0]
        for d in requested:
            did = d["id"]
            espliciti = [t for t in trades if t.get("linked_decision_id") == did]
            if espliciti:
                scelti, inferito = espliciti, False
            else:
                scelti = [t for t in trades if assignments.get(t["id"]) == did]
                inferito = True
            if not scelti:
                continue
            in_eur = all((t.get("valuta") or "EUR").upper() == "EUR" for t in scelti)
            eur = (sum(float(t["quantita"]) * float(t["prezzo"]) for t in scelti)
                   if in_eur else None)
            prop = d.get("eur_amount")
            pct = None
            if eur is not None and isinstance(prop, (int, float)) and prop:
                pct = int(round(eur / float(prop) * 100))
            out[did] = {"trades": scelti, "trade_ids": [t["id"] for t in scelti], "eur": eur,
                        "pct": pct, "inferito": inferito,
                        "data": max(str(t.get("data") or "")[:10] for t in scelti)}
        return out

    def stato_decisione(self, d, esecuzioni=None):
        """Lo STATO di una decisione come fatto del DB, in una frase: 'ESEGUITA il GG/MM per X
        EUR (...)' / 'ESEGUITA (trade non identificato nel registro)' / 'SKIPPED (non eseguita
        questa volta)' / 'EXPIRED (decaduta senza esecuzione)' / 'PARTIAL' / 'PENDING'.
        Audit 11/09 (Fable 5.1): il feedback del PM era reso in tre punti (blocco BINDING dei
        desk, PAROLE DIRETTE del Capo, piede del validator) e in NESSUNO c'era lo stato: il
        Capo ha letto un commento su una decisione SKIPPED come una richiesta da soddisfare.
        Un formatter solo per i tre punti."""
        st = str(d.get("status") or "?").upper()
        esecuzioni = esecuzioni or {}
        if d.get("id") in esecuzioni:
            return self._frase_esecuzione(esecuzioni[d["id"]], d.get("eur_amount"))
        if st == "EXECUTED":
            return "ESEGUITA (trade non identificato nel registro)"
        if st == "PARTIAL":
            return "PARTIAL (eseguita in parte, trade non identificato nel registro)"
        if st == "SKIPPED":
            return "SKIPPED (non eseguita questa volta)"
        if st == "EXPIRED":
            return "EXPIRED (decaduta senza esecuzione)"
        return st

    @staticmethod
    def _frase_esecuzione(info, proposto):
        """'ESEGUITA il GG/MM per X EUR (Y% dei Z proposti; legame inferito)'."""
        dt = str(info.get("data") or "")
        ddmm = (dt[8:10] + "/" + dt[5:7]) if len(dt) >= 10 else "?"
        s = "ESEGUITA il " + ddmm
        if info.get("eur") is not None:
            s += " per {:,.0f} EUR".format(info["eur"])
            dett = []
            if info.get("pct") is not None and isinstance(proposto, (int, float)) and proposto:
                dett.append("{}% dei {:,.0f} proposti".format(info["pct"], float(proposto)))
            if info.get("inferito"):
                dett.append("legame inferito")
            if dett:
                s += " (" + "; ".join(dett) + ")"
        else:
            q = ", ".join("%s x %s %s" % (t.get("quantita"), t.get("prezzo"), t.get("valuta") or "?")
                          for t in (info.get("trades") or [])[:3])
            s += " (" + q + ("; legame inferito" if info.get("inferito") else "") + ")"
        return s

    def get_decisions_with_pm_feedback(self, n=10):
        """Decisioni con feedback testuale del PM, le piu' rilevanti per i prompt.
        21/07 (memo #46): query DEDICATA perche' i veti non devono scadere con la
        finestra 'ultime N decisioni' — con ~11 decisioni per memo, un veto della
        settimana prima esce dalla top-10 in una run sola. Review 21/07: PRIORITA'
        ai feedback su decisioni NON eseguite (SKIPPED/EXPIRED = disaccordi/veti):
        le note di esecuzione routinarie ('comprati 5000...') non devono spingere
        fuori un 'basta proporla' (misurato: ~2,8 feedback nuovi a memo).
        F10 (23/07): le righe con VETO attivo sono SEMPRE incluse e vengono PRIMA,
        FUORI dalla finestra n — un veto e' eterno finche' il PM non lo revoca."""
        with self._conn() as conn:
            try:
                veto_rows = conn.execute(
                    "SELECT * FROM decisions WHERE COALESCE(veto,0)=1 "
                    "ORDER BY id DESC").fetchall()
            except sqlite3.OperationalError as e:
                # SOLO colonna mancante = migrazione 2 non ancora applicata (nessun
                # veto puo' esistere in quel file): legacy DICHIARATO. Qualsiasi
                # altro OperationalError (locked/disk) DEVE propagare — un fallback
                # zitto farebbe evaporare i veti (review 23/07, finding ALTA-2).
                if "no such column" not in str(e).lower():
                    raise
                print("[MEMORY_DB] canale veti: colonna 'veto' assente (migrazione 2 "
                      "al riavvio backend) — canale in modalita' legacy senza veti")
                veto_rows = []
            veto_ids = {r["id"] for r in veto_rows}
            rows = conn.execute(
                "SELECT * FROM decisions WHERE pm_feedback IS NOT NULL "
                "AND TRIM(pm_feedback) != '' "
                "ORDER BY CASE WHEN UPPER(COALESCE(status,'')) IN ('SKIPPED','EXPIRED') "
                "THEN 0 ELSE 1 END, id DESC LIMIT ?", (n,)).fetchall()
            return ([dict(r) for r in veto_rows]
                    + [dict(r) for r in rows if r["id"] not in veto_ids])

    def ids_decisioni_con_feedback_pm(self):
        """TUTTI gli id con un commento del PM, per poter DIRE quanti restano
        fuori dalla finestra (21/08): il taglio sulle righe esisteva gia' ma era
        muto, e con 11 SKIPPED su 28 il bucket saturava LIMIT 5 e LIMIT 10 —
        nessun commento su una decisione EXECUTED/PARTIAL poteva entrare."""
        with self._conn() as conn:
            return [r[0] for r in conn.execute(
                "SELECT id FROM decisions WHERE pm_feedback IS NOT NULL "
                "AND TRIM(pm_feedback) != '' ORDER BY id DESC")]

    def _fuori_finestra_pm(self, dentro):
        """Riga di dichiarazione delle righe rimaste fuori. Vuota se non ne resta
        nessuna: un blocco non deve annunciare buchi che non ha."""
        try:
            ids = self.ids_decisioni_con_feedback_pm()
        except Exception as e:
            return (" [quante righe restino fuori NON e' verificabile in questa run: "
                    + type(e).__name__ + "]")
        presenti = {d.get("id") for d in dentro}
        fuori = [i for i in ids if i not in presenti]
        if not fuori:
            return ""
        # ⚠️ NIENTE «si recuperano dal DB»: questa riga ora viaggia anche nei
        # prompt di R1/R2 e del red team (voce (2b)), e NESSUN lettore in run
        # ha un tool che legga `decisions` — il red team ha 3 tool read-only
        # di portafoglio, il Capo non ha tool affatto. Una via di recupero
        # promessa a chi non puo' percorrerla e' la bugia del 21/08 (review
        # 22/08 sera-2, difetto dormiente: oggi la finestra non taglia).
        return (" [ATTENZIONE: %d commenti del PM su %d restano FUORI da questo blocco "
                "(id %s): NON sono revocati e valgono ancora — in questa run non hai "
                "modo di leggerli, trattali come vincoli non visti, non come inesistenti]"
                % (len(fuori), len(ids), ", ".join("#%s" % i for i in fuori[:12])
                   + (" ..." if len(fuori) > 12 else "")))

    def get_recent_decisions(self, n=10, status_filter=None):
        """Ultime N decisioni con status. status_filter='PENDING' per pending only."""
        with self._conn() as conn:
            if status_filter:
                rows = conn.execute("SELECT * FROM decisions WHERE status=? ORDER BY id DESC LIMIT ?",
                                     (status_filter, n)).fetchall()
            else:
                rows = conn.execute("SELECT * FROM decisions ORDER BY id DESC LIMIT ?", (n,)).fetchall()
            return [dict(r) for r in rows]

    def update_decision(self, decision_id, status=None, pm_feedback=None,
                        outcome_pct=None, outcome_eur=None, outcome_notes=None):
        """Aggiorna status/feedback/outcome di una decisione."""
        fields = []
        values = []
        if status is not None: fields.append("status=?"); values.append(status)
        if pm_feedback is not None: fields.append("pm_feedback=?"); values.append(pm_feedback)
        if outcome_pct is not None: fields.append("outcome_pct=?"); values.append(outcome_pct)
        if outcome_eur is not None: fields.append("outcome_eur=?"); values.append(outcome_eur)
        if outcome_notes is not None: fields.append("outcome_notes=?"); values.append(outcome_notes)
        if status in ("EXECUTED", "SKIPPED", "EXPIRED"):
            fields.append("closed_at=?"); values.append(datetime.now().isoformat(timespec="seconds"))
        if not fields: return False
        values.append(decision_id)
        with self._conn() as conn:
            conn.execute("UPDATE decisions SET " + ", ".join(fields) + " WHERE id=?", values)
        return True

    def set_decision_veto(self, decision_id, reason):
        """F10 opzione A (PM 22/07): VETO ETERNO su una decisione, motivo OBBLIGATORIO.
        Imposta SKIPPED + flag veto: resta nel canale della run per sempre, finche'
        il PM non lo revoca. Solleva OperationalError se la migrazione 2 non e'
        ancora applicata (l'endpoint la dichiara come 503, pattern F10-C)."""
        reason = (reason or "").strip()
        if not reason:
            return None
        now = datetime.now().isoformat(timespec="seconds")
        with self._conn() as conn:
            # review 23/07 (finding MEDIA-6): SKIPPED solo sulle PENDING — vetare una
            # decisione gia' EXECUTED non deve riscrivere la storia dell'esecuzione
            cur = conn.execute(
                "UPDATE decisions SET veto=1, veto_reason=?, veto_at=?, veto_revoked_at=NULL, "
                "status=CASE WHEN UPPER(COALESCE(status,''))='PENDING' THEN 'SKIPPED' "
                "ELSE status END, "
                "closed_at=COALESCE(closed_at, ?) WHERE id=?",
                (reason, now, now, decision_id))
            if not cur.rowcount:
                return None
            row = conn.execute("SELECT * FROM decisions WHERE id=?", (decision_id,)).fetchone()
            return dict(row)

    def revoke_decision_veto(self, decision_id):
        """Revoca del veto (la UI chiede conferma): il flag si spegne, motivo e date
        restano in storia (veto_reason/veto_at + veto_revoked_at)."""
        now = datetime.now().isoformat(timespec="seconds")
        with self._conn() as conn:
            cur = conn.execute(
                "UPDATE decisions SET veto=0, veto_revoked_at=? WHERE id=? AND COALESCE(veto,0)=1",
                (now, decision_id))
            if not cur.rowcount:
                return None
            row = conn.execute("SELECT * FROM decisions WHERE id=?", (decision_id,)).fetchone()
            return dict(row)

    def auto_expire_stale_decisions(self, days: int = 7) -> int:
        """16/07 (richiesta PM): le proposte OPERATIVE (non RESEARCH) rimaste PENDING
        oltre `days` giorni si archiviano DA SOLE come EXPIRED — il memo e' settimanale,
        una proposta non eseguita in una settimana o viene riproposta dal Capo nel memo
        nuovo, o e' decaduta. Tutto RESTA nel DB (status/feedback/outcome ricordati);
        la RESEARCH non si tocca: la archivia il PM (bottone) o la promuove la run.
        Ritorna il numero di righe archiviate (0 = nulla da fare)."""
        now = datetime.now().isoformat(timespec="seconds")
        with self._conn() as conn:
            # review 16/07: i timestamp in DB sono ISO LOCALI con 'T'; datetime('now') di
            # sqlite e' UTC con spazio -> il confronto slittava di ~1 giorno. strftime con
            # 'localtime' produce lo STESSO formato e fuso dei dati.
            # decisione PM 17/07 (opzione A): le righe FISSATE ATTIVE dal PM
            # (archive_override=0) l'automatismo NON le tocca MAI — il suo click
            # vince; restano PENDING finche' non decide lui.
            cur = conn.execute(
                "UPDATE decisions SET status='EXPIRED', closed_at=?, "
                "outcome_notes=COALESCE(outcome_notes,'') || ' [auto-archiviata: PENDING da oltre "
                + str(int(days)) + " giorni]' "
                "WHERE status='PENDING' AND upper(action) <> 'RESEARCH' "
                "AND COALESCE(archive_override, 1) <> 0 "
                "AND timestamp < strftime('%Y-%m-%dT%H:%M:%S', 'now', 'localtime', ?)",
                (now, "-%d days" % int(days)))
            n = cur.rowcount or 0
            # F10 v3 (mandato PM 16/07 notte, mockup A approvato): anche la RESEARCH
            # ha una finestra — piu' lunga (30 giorni) perche' la ricerca ha piu'
            # tempo, ma non infinita: senza avanzamento si auto-archivia DICHIARANDO.
            # Il bottone ARCHIVIA del PM e la promozione della run restano le vie maestre.
            cur2 = conn.execute(
                "UPDATE decisions SET status='EXPIRED', closed_at=?, "
                "outcome_notes=COALESCE(outcome_notes,'') || ' [auto-archiviata: research senza "
                "avanzamento da oltre 30 giorni]' "
                "WHERE status='PENDING' AND upper(action) = 'RESEARCH' "
                "AND COALESCE(archive_override, 1) <> 0 "  # decisione PM 17/07 (A): pin attivo intoccabile
                "AND timestamp < strftime('%Y-%m-%dT%H:%M:%S', 'now', 'localtime', '-30 days')",
                (now,))
            return n + (cur2.rowcount or 0)

    # ========================================================
    # DECISION NOTES (F10 v3: thread PM<->AI sulle research)
    # ========================================================

    def add_decision_note(self, decision_id, autore, testo):
        """F10 v3 (mandato PM 16/07 notte): nota sul filo di una decisione RESEARCH.
        autore in ('PM','AI'). Valida che la decisione esista e sia RESEARCH —
        il thread e' il canale di interazione della pipeline research, non un
        commento generico. Ritorna l'id nota o None (con motivo a log)."""
        autore = str(autore or "").upper()
        testo = str(testo or "").strip()
        if autore not in ("PM", "AI") or not testo:
            print("[memory_db] add_decision_note rifiutata: autore/testo non validi")
            return None
        try:
            with self._conn() as conn:
                row = conn.execute("SELECT upper(action) FROM decisions WHERE id=?",
                                   (int(decision_id),)).fetchone()
                if not row:
                    print(f"[memory_db] add_decision_note: decisione {decision_id} inesistente")
                    return None
                if row[0] != "RESEARCH":
                    print(f"[memory_db] add_decision_note: decisione {decision_id} non e' RESEARCH")
                    return None
                cols = {r[1] for r in conn.execute("PRAGMA table_info(decision_notes)")}
                if not cols:
                    print("[memory_db] tabella decision_notes assente: lancia "
                          "tools/migrations/migra_decision_notes.py")
                    return None
                cur = conn.execute(
                    "INSERT INTO decision_notes (decision_id, autore, testo) VALUES (?,?,?)",
                    (int(decision_id), autore, testo[:2000]))
                return cur.lastrowid
        except Exception as e:
            print("[memory_db] add_decision_note failed: " + str(e))
            return None

    def get_decision_notes(self, decision_id, n=20):
        """Thread note di una decisione (vecchie -> nuove). [] se tabella assente
        (DB non migrato: buco gia' dichiarato in scrittura)."""
        try:
            with self._conn() as conn:
                rows = conn.execute(
                    "SELECT id, autore, testo, timestamp FROM decision_notes "
                    "WHERE decision_id=? ORDER BY timestamp ASC, id ASC LIMIT ?",
                    (int(decision_id), int(n))).fetchall()
                return [dict(r) for r in rows]
        except Exception:
            return []

    # ========================================================
    # PM FEEDBACK
    # ========================================================

    def add_pm_feedback(self, feedback_text, sentiment="NEUTRAL", memo_id=None,
                       decision_id=None, specialist=None):
        """Registra feedback del PM. Anche embedded in ChromaDB per semantic search."""
        with self._conn() as conn:
            cur = conn.execute("""INSERT INTO pm_feedback (memo_id, decision_id, specialist,
                                                             feedback_text, sentiment)
                                   VALUES (?,?,?,?,?)""",
                                (memo_id, decision_id, specialist, feedback_text, sentiment))
            feedback_id = cur.lastrowid

        if self.col_feedback:
            try:
                self.col_feedback.add(
                    documents=[feedback_text],
                    metadatas=[{"feedback_id": feedback_id, "sentiment": sentiment,
                                "specialist": specialist or "", "memo_id": memo_id or 0,
                                "decision_id": decision_id or 0,
                                "timestamp": datetime.now().isoformat()}],
                    ids=["feedback_" + str(feedback_id)]
                )
            except Exception as e:
                print("[MemoryDB] feedback embed error: " + str(e))
        return feedback_id

    def get_recent_pm_feedback(self, n=5, specialist=None):
        """Ultimi N feedback PM, filtrabili per specialista.
        eventdesk eredita anche i feedback legacy news/politics (alias fusione 15/07)."""
        with self._conn() as conn:
            if specialist:
                names = (specialist,) + self._LEGACY_SPECIALIST_ALIASES.get(specialist, ())
                ph = ",".join("?" * len(names))
                rows = conn.execute("""SELECT * FROM pm_feedback WHERE specialist IN (""" + ph + """) OR specialist IS NULL
                                        ORDER BY id DESC LIMIT ?""", (*names, n)).fetchall()
            else:
                rows = conn.execute("SELECT * FROM pm_feedback ORDER BY id DESC LIMIT ?", (n,)).fetchall()
            return [dict(r) for r in rows]

    def search_feedback_semantic(self, query, n_results=3):
        """Cerca feedback per query semantica."""
        if not self.col_feedback:
            return []
        try:
            r = self.col_feedback.query(query_texts=[query], n_results=n_results)
            out = []
            if r.get("ids") and r["ids"][0]:
                for i, fb_id in enumerate(r["ids"][0]):
                    out.append({
                        "feedback_id": fb_id,
                        "content": r["documents"][0][i],
                        "metadata": r["metadatas"][0][i],
                    })
            return out
        except Exception:
            return []

    # ========================================================
    # CONTEXT BUILDERS PER SPECIALISTI
    # ========================================================

    def _pm_binding_parts(self):
        """Le righe del blocco «PM feedback on past decisions (BINDING)» e gli id
        le cui PAROLE sono davvero rese (le righe con veto rendono `veto_reason`,
        non `pm_feedback`: servono a `build_specialist_memory_context` per non
        affermare «il testo e' nel blocco sopra» quando non c'e').

        21/07 (memo #46): blocco dedicato con query propria — senza, il feedback
        era decorativo e la covered call vietata dal PM (#186) e' stata
        riproposta. Review 23/07 (F10, ALTA-1): il blocco sta IN TESTA, prima
        dei report — il troncamento a max_chars del caller mangiava la coda, e
        la coda erano i VETI. 22/08 sera-2 (voce (2b), ok PM): estratto qui
        perche' ora lo leggono ANCHE i Round 1/2 e il red team — un solo testo,
        piu' lettori. R0 e' rimasto byte-identico al pre-refactor: MISURATO
        una tantum contro HEAD dalla review 22/08 sera-2 (837/837 byte); il
        test della batteria pinna l'AUTO-CONSISTENZA fra i lettori (stesso
        testo per tutti), non il confronto storico.
        """
        fb_dec = self.get_decisions_with_pm_feedback(n=MAX_RIGHE_FEEDBACK_PM)
        _pm_resi = {d.get("id") for d in fb_dec if not d.get("veto")}
        # Audit 11/09 (Fable 5.1, run 10/09): la riga «#N ADD TICKER | PM: eseguito...»
        # arrivava ai desk SENZA stato ne' importo eseguito e tre desk l'hanno letta come un
        # ordine da eseguire («ADD TICKER per soddisfare #N»). Lo stato entra nella riga:
        # ESEGUITA (con quanto) / SKIPPED / EXPIRED / PENDING.
        try:
            _esec = self.esecuzioni_delle_decisioni(fb_dec)
        except Exception as _ee:
            print("[MEMORY_DB] esecuzioni delle decisioni non lette (dichiarato): " + str(_ee))
            _esec = {}
        parts = []
        if fb_dec:
            parts.append("\n--- PM feedback on past decisions (BINDING) ---"
                         + self._fuori_finestra_pm(fb_dec))
            for d in fb_dec:
                if d.get("veto"):
                    parts.append("#" + str(d.get("id")) + " " + str(d.get("action")) + " "
                                 + str(d.get("ticker")) + " | ⛔ VETO ATTIVO dal "
                                 + str(d.get("veto_at") or "")[:10] + " (ETERNO finche' il PM "
                                 "non lo revoca): "
                                 + pm_verbatim(d.get("veto_reason"), d.get("id")))
                else:
                    parts.append("#" + str(d.get("id")) + " " + str(d.get("action")) + " "
                                 + str(d.get("ticker")) + " | " + self.stato_decisione(d, _esec)
                                 + " | PM: " + pm_verbatim(d.get("pm_feedback"), d.get("id")))
            parts.append("Le righe '| PM:' sono la voce DIRETTA del PM e sono VINCOLANTI: un "
                         "rifiuto (es. 'basta proporla') vieta di riproporre quell'idea in ogni "
                         "variante senza fatti NUOVI dichiarati che citino il rifiuto. Le righe "
                         "⛔ VETO sono divieti PERMANENTI: mai riproporre (nemmeno varianti) "
                         "senza fatti nuovi dichiarati CHE CITINO il veto per numero. "
                         "Lo STATO fra le barre e' un fatto del registro: una riga ESEGUITA e' "
                         "un'operazione GIA' FATTA dal PM (il commento non e' un ordine da "
                         "rieseguire); una riga SKIPPED/EXPIRED commenta una NON-esecuzione "
                         "(leggila come motivo o condizione del mancato via, mai come richiesta "
                         "di eseguire quella stessa idea; se il PM indica un'alternativa, quella "
                         "e' la richiesta).")
        return parts, _pm_resi

    def build_pm_binding_block(self):
        """Il SOLO blocco BINDING (feedback del PM sulle decisioni + veti) come
        testo a se', per chi deve vederlo FUORI dalla memoria di Round 0: i
        Round 1/2 degli specialisti e il red team (22/08 sera-2, voce (2b)).
        "" se a registro non c'e' nulla — e' il chiamante a dichiararlo."""
        parts, _ = self._pm_binding_parts()
        return "\n".join(parts)

    def build_specialist_memory_context(self, specialist_name, max_chars=3000):
        """Costruisce blocco di memoria da iniettare nel prompt del specialista."""
        parts = ["=== YOUR MEMORY (from previous runs) ==="]

        _vincoli, _pm_resi = self._pm_binding_parts()
        parts.extend(_vincoli)

        # Ultimi 3 tuoi report
        reports = self.get_recent_specialist_reports(specialist_name, n=3)
        if reports:
            fingerprint_corrente = _fingerprint_mandato_corrente()
            parts.append("\n--- Your last 3 reports (round 2 final, most recent first) ---")
            for r in reports:
                snippet = (r.get("content") or "")[:800]
                etichetta = _etichetta_mandato_storico(
                    r.get("memo_full_markdown"), fingerprint_corrente)
                parts.append("[" + r.get("timestamp", "")[:10] + "] "
                             + etichetta + "\n" + snippet)
        else:
            parts.append("(No previous reports - this is your first run)")

        # Ultimi 5 feedback PM
        feedbacks = self.get_recent_pm_feedback(n=5, specialist=specialist_name)
        if feedbacks:
            parts.append("\n--- Recent PM feedback (act on these) ---")
            for f in feedbacks[:5]:
                parts.append("[" + (f.get("sentiment") or "?") + "] "
                             + pm_verbatim(f.get("feedback_text"),
                                           fonte="pm_feedback.feedback_text"))

        # Decisioni recenti correlate al tuo dominio
        decisions = self.get_recent_decisions(n=10)
        if decisions:
            parts.append("\n--- Recent recommendations and their status ---")
            for d in decisions[:7]:
                line = "[" + (d.get("status") or "?") + "] " + (d.get("action") or "?") + " " + (d.get("ticker") or "?")
                if d.get("eur_amount"):
                    line += " (EUR " + "{:+,.0f}".format(d["eur_amount"]) + ")"
                if d.get("outcome_pct") is not None:
                    line += " outcome: " + "{:+.1f}%".format(d["outcome_pct"])
                if d.get("pm_feedback"):
                    # 21/08: qui c'era una SECONDA resa dello stesso campo con un cap
                    # DIVERSO (`[:100]`) — il numero-in-due-posti in miniatura, e la
                    # copia piu' stretta delle due. Se il testo e' gia' reso INTERO
                    # nel blocco BINDING qui sopra basta il puntatore; se quella
                    # decisione li' non c'e' (fuori finestra, oppure riga con veto
                    # che rende `veto_reason`), il puntatore sarebbe una frase FALSA
                    # e il testo va reso qui — dalla policy, non da un letterale.
                    if d.get("id") in _pm_resi:
                        line += " | il PM ha commentato (testo intero nel blocco BINDING sopra)"
                    else:
                        line += " | PM: " + pm_verbatim(d["pm_feedback"], d.get("id"))
                parts.append(line)

        # TRACK RECORD personale (#211): hit-rate delle call firmate da questo
        # specialista, calcolato in codice dallo scorekeeper (guarded).
        try:
            from bellomberg.agents.scorekeeper import get_track_record_for_specialist
            _tr = get_track_record_for_specialist(specialist_name)
            if _tr:
                parts.append("\n" + _tr)
        except Exception:
            pass

        ctx = "\n".join(parts)
        # review 21/07: taglio al cap DICHIARATO, mai zitto (regola 14/07) — la coda
        # persa puo' contenere track record o feedback vincolanti.
        # 21/08: «coda persa» diceva CHE aveva tagliato, non COSA — e cio' che cade
        # per primo e' il TRACK RECORD. Ora le sezioni perse si NOMINANO.
        if len(ctx) > max_chars:
            ctx = ctx[:max_chars - 230] + _marcatore_troncamento(parts, ctx, max_chars)
        return ctx

    # ========================================================
    # IMPORT FROM PORTFOLIO.JSON (preferred over Excel)
    # ========================================================

    def reset_positions(self):
        """Cancella TUTTE le posizioni e trade history. Usa con cautela."""
        with self._conn() as conn:
            conn.execute("DELETE FROM positions")
            conn.execute("DELETE FROM trade_history")
            conn.execute("DELETE FROM position_prices")
        return {"reset": True}

    def import_from_json(self, json_path):
        """Importa portafoglio da portfolio.json (ticker veri + tesi + temi).
        Usa il peso_percentuale + prezzo_medio_carico + cash totale per ricostruire le quantita.
        """
        if not os.path.exists(json_path):
            return {"error": "File non trovato: " + json_path}

        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        all_positions = []
        for p in data.get("etf", []):
            all_positions.append({**p, "type": "etf"})
        for p in data.get("azioni", []):
            all_positions.append({**p, "type": "azione"})

        cash = data.get("cash_disponibile_eur", 0)

        # Calcolo NAV totale presumibile dal peso e dal market_value implicito.
        # Strategia: il peso% in portfolio.json e' rispetto al market value (no cash).
        # Quindi serve un anchor: la prima posizione con prezzo_medio disponibile.
        # PERO': portfolio.json NON ha quantita esplicite per ETF. Importiamo solo metadata + tesi + prezzo medio.
        # Le quantita verranno aggiornate dopo, con price_updater + un calcolo di "implied quantity = (NAV * weight) / prezzo_medio".
        # Per ora salviamo metadata + prezzo_medio_carico (dove disponibile) + qty placeholder.

        imported = 0
        errors = []
        for p in all_positions:
            try:
                ticker = p.get("ticker", "").upper()
                if not ticker:
                    continue
                nome = p.get("nome", ticker)
                prezzo_medio = p.get("prezzo_medio_carico")
                tesi = p.get("tesi", "")
                temi = p.get("temi_monitoraggio", [])
                peso_pct = p.get("peso_percentuale", 0)

                # La valuta esplicita vince; un ticker ambiguo non autorizza USD implicito.
                etichetta_valuta = _classifica_valuta(ticker, p.get("valuta"))
                if etichetta_valuta.valore is None:
                    raise ValueError(etichetta_valuta.dichiarazione)
                valuta = etichetta_valuta.valore

                # Quantita placeholder = 1 (verra aggiornata da price_updater dopo)
                # NB: il peso reale sara calcolato da market_value / NAV dopo update prezzi
                # Tracciamo il peso_intended in note per riferimento
                note = "Imported from portfolio.json. Target weight: " + str(peso_pct) + "%"

                # Per ora salviamo qty=0 se non disponibile, sara da aggiungere manualmente o calcolato
                # quando si ha prezzo live + cash assumption
                qty = 1.0  # placeholder, da aggiornare

                self.add_or_update_position(
                    ticker=ticker, nome=nome, quantita=qty,
                    prezzo_medio=prezzo_medio,
                    valuta=valuta,
                    tesi=tesi,
                    temi_monitoraggio=temi,
                    note=note,
                )
                imported += 1
            except Exception as e:
                errors.append(p.get("ticker", "?") + ": " + str(e))

        return {"imported": imported, "errors": errors, "cash_eur": cash,
                "note": "Quantita = 1.0 placeholder. Usa 'add' per ogni posizione con qty reale, oppure 'rebuild-from-weights' per derivare qty dai pesi target."}




    def import_hybrid(self, excel_path, json_path="portfolio.json"):
        """Import HYBRID: quantita da Excel + ticker/tesi da portfolio.json.
        Match per nome fuzzy. Cancella positions esistenti prima.
        """
        if not os.path.exists(json_path):
            return {"error": "portfolio.json non trovato"}
        if not os.path.exists(excel_path):
            return {"error": "Excel non trovato: " + excel_path}

        try:
            import openpyxl
        except ImportError:
            return {"error": "openpyxl required"}

        # 1. Leggi portfolio.json
        with open(json_path, "r", encoding="utf-8") as f:
            j = json.load(f)
        json_positions = []
        for p in j.get("etf", []) + j.get("azioni", []):
            json_positions.append({
                "ticker": p.get("ticker", "").upper(),
                "nome": p.get("nome", ""),
                "tesi": p.get("tesi", ""),
                "temi": p.get("temi_monitoraggio", []),
                "prezzo_medio_json": p.get("prezzo_medio_carico"),
                "valuta": p.get("valuta"),
            })

        # 2. Leggi Excel
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        target_sheet = None
        for s in wb.sheetnames:
            if s.lower() in ("holdings", "portafoglio", "positions"):
                target_sheet = s; break
        if not target_sheet:
            target_sheet = wb.sheetnames[0]
        ws = wb[target_sheet]

        HEADER_KW = ["ticker", "nome", "q.t", "quantit", "prezzo", "val. mercato"]
        header_row = None
        for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
            if not row: continue
            txt = " ".join(str(c).lower() for c in row if c)
            if sum(1 for k in HEADER_KW if k in txt) >= 2:
                header_row = ri; break
        if not header_row:
            return {"error": "Header non trovato in Excel"}
        headers = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        col_map = {}
        for i, h in enumerate(headers):
            if not h: continue
            hl = str(h).lower().strip()
            if "nome" in hl or "name" in hl: col_map["nome"] = i
            elif "q.t" in hl or "quantit" in hl: col_map["quantita"] = i
            elif "prezzo medio" in hl or "avg" in hl or "costo medio" in hl: col_map["prezzo_medio"] = i

        excel_rows = []
        for row in ws.iter_rows(min_row=header_row+1, values_only=True):
            if not row or "nome" not in col_map: continue
            nome_raw = row[col_map["nome"]]
            if not nome_raw: continue
            nome = str(nome_raw).strip()
            if nome.upper() in ("TOTALE", "TOTAL", "TOT", ""): continue
            try:
                qty = float(row[col_map["quantita"]]) if "quantita" in col_map and row[col_map["quantita"]] else 0
                pm = float(row[col_map["prezzo_medio"]]) if "prezzo_medio" in col_map and row[col_map["prezzo_medio"]] else None
            except Exception:
                qty, pm = 0, None
            if qty > 0:
                excel_rows.append({"nome": nome, "quantita": qty, "prezzo_medio_excel": pm})

        # 3. Fuzzy match: per ogni excel_row trova il ticker piu simile in json
        def normalize(s):
            return "".join(c.lower() for c in s if c.isalnum())[:20]

        # Reset positions
        self.reset_positions()

        imported = []
        unmatched = []
        errors = []
        for er in excel_rows:
            excel_norm = normalize(er["nome"])
            best_match = None
            best_score = 0

            # 1. Check manual overrides first
            override_ticker = None
            for key, tk in NAME_TO_TICKER_OVERRIDES.items():
                if key in excel_norm:
                    override_ticker = tk
                    break
            if override_ticker:
                for jp in json_positions:
                    if jp["ticker"].upper() == override_ticker:
                        best_match = jp
                        best_score = 100  # forced match
                        break

            # 2. Fuzzy match if no override
            if not best_match:
                for jp in json_positions:
                    jp_norm = normalize(jp["nome"])
                    score = 0
                    for a, b in zip(excel_norm, jp_norm):
                        if a == b: score += 1
                        else: break
                    if score > best_score and score >= 4:
                        best_score = score
                        best_match = jp
            if best_match:
                ticker = best_match["ticker"]
                etichetta_valuta = _classifica_valuta(ticker, best_match.get("valuta"))
                if etichetta_valuta.valore is None:
                    errors.append(ticker + ": " + etichetta_valuta.dichiarazione)
                    continue
                valuta = etichetta_valuta.valore
                # Prezzo medio: preferisci Excel se disponibile, altrimenti JSON
                pm = er["prezzo_medio_excel"] if er["prezzo_medio_excel"] else best_match["prezzo_medio_json"]
                self.add_or_update_position(
                    ticker=ticker, nome=best_match["nome"],
                    quantita=er["quantita"], prezzo_medio=pm,
                    valuta=valuta, tesi=best_match["tesi"],
                    temi_monitoraggio=best_match["temi"],
                )
                imported.append({"ticker": ticker, "nome": best_match["nome"],
                                  "quantita": er["quantita"], "prezzo_medio": pm,
                                  "valuta": valuta, "match_score": best_score})
            else:
                unmatched.append(er["nome"])

        return {"imported": imported, "unmatched": unmatched, "errors": errors,
                "n_imported": len(imported), "n_unmatched": len(unmatched),
                "cash_eur": j.get("cash_disponibile_eur", 0)}



    def build_capo_memory_context(self, max_chars=5000):
        """Memoria per il Capo: tutte le decisioni recenti + feedback + memo precedente sintetizzato."""
        parts = ["=== CAPO MEMORY (from previous weekly runs) ==="]

        # Last memo summary (per continuity)
        # audit 11/09: il memo precedente e' l'ultimo NON marcato DUPLICATO (memos[0] e' il
        # segnaposto della run in corso); senza duplicati il comportamento e' quello di sempre
        memos = self.get_recent_memos(n=6)
        if memos:
            fingerprint_corrente = _fingerprint_mandato_corrente()
            _precedenti = [m for m in memos[1:] if not e_duplicato(m)]
            last = _precedenti[0] if _precedenti else memos[0]
            title = last.get("title", "")
            ts = last.get("timestamp", "")
            parts.append(f"[{ts}] {title}")
            parts.append(_etichetta_mandato_storico(
                last.get("full_markdown"), fingerprint_corrente))
            full_md = (last.get("full_markdown") or "")[:1500]
            parts.append(full_md)

        # Decisioni recenti CON STATUS (questa e' la "pagina Decisioni" che il PM aggiorna:
        # EXECUTED = ha eseguito la tua raccomandazione; SKIPPED = NON eseguita QUESTA volta
        # (timing/liquidita'/priorita'), NON un rifiuto definitivo; PENDING = ancora aperta).
        try:
            decisions = self.get_recent_decisions(n=20)
            # audit 11/09: le decisioni di una run marcata DUPLICATO non sono pendenti ne'
            # decadute — il PM non le ha mai valutate — e non entrano nella memoria
            decisions = [d for d in decisions if not e_duplicato(d)]
            # audit 11/09: quanto e' stato eseguito DAVVERO (legame esplicito o inferito)
            try:
                _esec = self.esecuzioni_delle_decisioni(decisions)
            except Exception as _ee:
                print("[MEMORY_DB] esecuzioni delle decisioni non lette (dichiarato): " + str(_ee))
                _esec = {}
            if decisions:
                ex = [d for d in decisions if (d.get("status") or "").upper() == "EXECUTED"]
                sk = [d for d in decisions if (d.get("status") or "").upper() == "SKIPPED"]
                pe = [d for d in decisions if (d.get("status") or "").upper() == "PENDING"]
                # review 16/07: con l'auto-archivio (>7g -> EXPIRED) le decadute sparivano
                # da OGNI bucket: la regola "riproponila ricordando che era gia' suggerita"
                # e l'anti-insistenza non potevano vederle. Ora hanno la loro riga.
                xp = [d for d in decisions if (d.get("status") or "").upper() == "EXPIRED"]
                parts.append("\n--- DECISIONI: come il PM ha risposto alle tue raccomandazioni (ultime 20) ---")

                def _fmt(d):
                    amt = d.get("eur_amount")
                    amt_s = ("{:+,.0f} EUR".format(amt)) if isinstance(amt, (int, float)) else "-"
                    out = d.get("outcome_pct")
                    note = (d.get("rationale") or "").strip()  # #31: qui scrive la guardia book-aware
                    # 21/07 (memo #46): il feedback del PM vive nel blocco dedicato
                    # "PAROLE DIRETTE DEL PM" qui sotto (non inline: il cap 6800/8500
                    # del contesto non tollera doppioni — review 21/07).
                    # audit 11/09: l'importo qui era SOLO quello proposto; l'eseguito
                    # (quantita' x prezzo dei trade, e la sua quota del proposto) non
                    # arrivava mai al Capo
                    _es = (" | " + self._frase_esecuzione(_esec[d["id"]], d.get("eur_amount"))
                           if d.get("id") in _esec else "")
                    return ("#" + str(d.get("id")) + " " + str(d.get("action")) + " " +
                            str(d.get("ticker")) + " " + amt_s + _es +
                            (" | outcome=" + str(out) if out is not None else "") +
                            ((" | " + note[:100]) if note else ""))
                if ex:
                    parts.append("ESEGUITE dal PM: " + "; ".join(_fmt(d) for d in ex[:10]))
                # audit 11/09: lo status PARTIAL non aveva un bucket (8 righe invisibili al Capo)
                pa = [d for d in decisions if (d.get("status") or "").upper() == "PARTIAL"]
                if pa:
                    parts.append("ESEGUITE IN PARTE dal PM (PARTIAL): " + "; ".join(_fmt(d) for d in pa[:10]))
                if sk:
                    parts.append("NON eseguite questa volta (SKIPPED, non rifiutate per sempre): "
                                 + "; ".join(_fmt(d) for d in sk[:10]))
                if pe:
                    parts.append("ANCORA PENDENTI: " + "; ".join(_fmt(d) for d in pe[:10]))
                if xp:
                    parts.append("DECADUTE (EXPIRED: mai eseguite, auto-archiviate dopo 7 giorni "
                                 "o archiviate dal PM — contano come non-esecuzione, al pari delle "
                                 "SKIPPED, per la regola anti-insistenza qui sotto): "
                                 + "; ".join(_fmt(d) for d in xp[:10]))
                # query dedicata, NON filtrata dalla finestra n=20 (v. get_decisions_with_pm_feedback)
                vet = self.get_decisions_with_pm_feedback(n=MAX_RIGHE_FEEDBACK_PM)
                if vet:
                    def _fmt_pm(d):
                        if d.get("veto"):
                            return ("#" + str(d.get("id")) + " " + str(d.get("action")) + " "
                                    + str(d.get("ticker")) + " -> "
                                    + pm_verbatim(
                                        d.get("veto_reason"), d.get("id"), virgolette=True,
                                        prefisso="⛔ VETO dal "
                                                 + str(d.get("veto_at") or "")[:10]
                                                 + " (eterno finche' non revocato): "))
                        # audit 11/09: lo stato accanto alle parole — un commento su una
                        # decisione SKIPPED non e' un ordine, e su una ESEGUITA e' cosa fatta
                        return ("#" + str(d.get("id")) + " " + str(d.get("action")) + " "
                                + str(d.get("ticker")) + " [" + self.stato_decisione(d, _esec)
                                + "] -> "
                                + pm_verbatim(d.get("pm_feedback"), d.get("id"),
                                              virgolette=True))
                    # niente slice [:10]: la parte feedback e' gia' limitata dalla query,
                    # i VETI non si tagliano MAI (review 23/07, finding MEDIA-3)
                    parts.append("PAROLE DIRETTE DEL PM sulle decisioni (VINCOLANTI, prevalgono su ogni "
                                 "regola di riproposta qui sotto; le righe ⛔ VETO sono divieti "
                                 "PERMANENTI finche' il PM non li revoca)"
                                 + self._fuori_finestra_pm(vet) + ": "
                                 + "; ".join(_fmt_pm(d) for d in vet))
                parts.append("REGOLA: una decisione SKIPPED NON va abbandonata. Se la tesi regge ancora, "
                             "RIPROPONILA (a size invariata o ridotta, mai maggiorata), ricordando che era "
                             "gia' stata suggerita. ECCEZIONE VINCOLANTE (21/07, lezione memo #46): le "
                             "PAROLE DIRETTE DEL PM prevalgono su questa regola — un rifiuto (es. 'basta "
                             "proporla') vieta OGNI variante sullo stesso ticker (HOLD->HEDGE compreso), "
                             "salvo fatti NUOVI verificabili citando il veto. MA (16/07): se risulta gia' proposta e NON ESEGUITA piu' di "
                             "una volta (stesso ticker e stessa azione, contando SKIPPED e DECADUTE di "
                             "questo elenco), la non-esecuzione ripetuta e' un segnale del PM — dichiara apertamente il disaccordo nella sezione della "
                             "decisione, chiedi una decisione esplicita, e presentala DECLASSATA (size "
                             "ridotta o RESEARCH a 0), mai maggiorata. Se la tesi e' decaduta, lasciala "
                             "cadere e spiega perche'. Dai continuita' anche alle ESEGUITE. "
                             "REGOLA (audit 11/09): una decisione ESEGUITA di recente (riga "
                             "'ESEGUITA il GG/MM', entro 7 giorni) NON va riproposta nello stesso "
                             "verso come se fosse nuova — il PM l'ha gia' fatta. La si ripropone SOLO "
                             "con un fatto NUOVO verificabile, scritto nella colonna Timing della "
                             "ACTION TABLE col tag 'NOVITA':' (come 'SOPRA POLICY'); altrimenti la "
                             "posizione si tiene (HOLD) e si aggiornano fair value o ricerca. Non e' "
                             "un divieto sul titolo: e' il divieto di ricomprare cio' che il PM ha "
                             "appena comprato senza dirgli perche'.")
        except Exception:
            pass

        # TRACK RECORD (#190/#211): esito di MERCATO delle call passate, calcolato
        # in codice — il pezzo che trasforma "ricordo cosa ho detto" in "so se
        # avevo ragione". Guarded: un guasto dello scorekeeper non tocca la memoria.
        try:
            from bellomberg.agents.scorekeeper import get_track_record_for_capo
            _tr = get_track_record_for_capo()
            if _tr:
                parts.append("\n" + _tr)
        except Exception:
            pass

        # LEZIONE #210 (reflection outcome-grounded): la lezione generata a fine
        # run precedente, ancorata allo scorekeeper. Guarded come sopra.
        try:
            from bellomberg.agents.reflection import get_latest_lesson_block
            _ls = get_latest_lesson_block()
            if _ls:
                parts.append("\n" + _ls)
        except Exception:
            pass

        # Trade ESEGUITI dal PM con commenti (cosi' il consigliere sa cosa hai fatto davvero e perche')
        try:
            trades = self.get_recent_trades(n=12)
            if trades:
                # audit 11/09: il legame trade->decisione INFERITO (v. esecuzioni_delle_decisioni)
                _link_inf = {}
                try:
                    for _did, _info in (_esec or {}).items():
                        if _info.get("inferito"):
                            for _tid in _info.get("trade_ids") or []:
                                _link_inf[_tid] = _did
                except Exception:
                    _link_inf = {}
                parts.append("\n--- Trade ESEGUITI dal PM questa/e settimana/e (con i tuoi commenti) ---")
                for tr in trades:
                    cm = (tr.get("pm_rationale") or tr.get("note") or "").strip()
                    dt = (tr.get("data") or "")[:10]
                    line = ("[" + dt + "] " + str(tr.get("action")) + " " +
                            str(tr.get("quantita")) + " " + str(tr.get("ticker")) + " @ " +
                            str(tr.get("prezzo")) + " " + str(tr.get("valuta")))
                    if tr.get("linked_decision_id"):
                        line += " (esegue decisione #" + str(tr.get("linked_decision_id")) + ")"
                    elif tr.get("link_origin") == "none":
                        line += " (senza decisione, dichiarato dal PM)"
                    elif tr.get("id") in _link_inf:
                        line += (" (esegue decisione #" + str(_link_inf[tr["id"]])
                                 + ", legame inferito: stesso titolo e verso entro 7 giorni)")
                    if cm:
                        line += " | COMMENTO PM: " + cm[:200]
                    parts.append(line)
                parts.append("Tieni conto di QUESTI movimenti e dei commenti del PM: confermano, "
                             "modificano o contraddicono le tue raccomandazioni precedenti.")
        except Exception:
            pass

        # PM feedback
        try:
            feedbacks = self.get_recent_pm_feedback(n=10)
            if feedbacks:
                parts.append("\n--- Recent PM feedback ---")
                for fb in feedbacks[:10]:
                    ts = (fb.get("timestamp", "") or "")[:10]
                    parts.append("[" + ts + "] " + str(fb.get("sentiment", "NEUTRAL")) + ": "
                                 + pm_verbatim(fb.get("feedback_text"),
                                               fonte="pm_feedback.feedback_text"))
        except Exception:
            pass

        result = "\n".join(parts)
        # review 21/07: taglio al cap DICHIARATO, mai zitto (regola 14/07) — misurato
        # che la coda persa conteneva i trade eseguiti dal PM e la lezione #210.
        # 21/08: il marcatore ora NOMINA le sezioni cadute (v. build_specialist_*).
        if len(result) > max_chars:
            result = result[:max_chars - 230] + _marcatore_troncamento(parts, result, max_chars)
        return result

def _classifica_valuta(ticker, valuta_posizione=None):
    """Porta unica alla classificazione, riletta a ogni chiamata."""
    from bellomberg.storage.classificazione import valuta
    return valuta(ticker, valuta_posizione=valuta_posizione)


def _infer_currency(ticker, valuta_posizione=None):
    """Compatibilita' per i due importatori legacy: puo' rendere None, mai USD implicito."""
    return _classifica_valuta(ticker, valuta_posizione).valore
