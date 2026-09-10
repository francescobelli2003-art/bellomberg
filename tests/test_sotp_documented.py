"""Actual heterogeneous child adapters; parent-only bridge and complete coverage."""
from copy import deepcopy
from pathlib import Path
import json
import pytest
from bellomberg.valuation import dcf_engine,sector_analysis
from test_sector_operating_drivers import operating_records,DAY,SPAN
from test_sector_bank_capital import bank_records,bank_bundle
from test_real_estate_valuation import property_records


def remap_entities(value,names):
    if isinstance(value,dict):return {names.get(k,k):remap_entities(v,names) for k,v in value.items()}
    if isinstance(value,list):return [remap_entities(v,names) for v in value]
    return names.get(value,value) if isinstance(value,str) else value


def sotp_records():
    from bellomberg.valuation.sotp_adapter import POLICY,claim_targets
    bank=bank_bundle(remap_entities(bank_records(),{'SYNTH-GROUP':'BANK-GROUP','SYNTH-PARENT':'BANK-PARENT','LEGAL-A':'BANK-LEGAL'}),symbol='SYNTH-BANK')
    prop=bank_bundle(remap_entities(property_records(),{'SYNTH-GROUP':'PROPERTY-GROUP'}),profile='property_owner',symbol='SYNTH-REIT')
    children={'BANK':bank,'PROPERTY':prop}
    rows=[r for r in operating_records() if r['driver'] in ('perimeter','calendar','quotation','shares')]
    for r in rows:
        r['entity']='HOLDING-GROUP'
        if r['driver']=='perimeter':r['value']['entity']='HOLDING-GROUP'
        if r['driver']=='quotation':r['value']['price']=50.
    def add(driver,field,value,scenario='model',unit='contract',basis='sotp',period=SPAN):
        rows.append(dict(driver=driver,field=field,value=deepcopy(value),unit=unit,accounting_basis=basis,
            scenario=scenario,period=period,entity='HOLDING-GROUP',source_id='https://example.org/synthetic/holding-segment-and-parent-statements',
            as_of=DAY,valid_until='2027-01-01',kind='analyst_estimate',rationale='Synthetic complete holding perimeter and sourced ownership/claims'))
    add('policy','segment_perimeters',POLICY,basis='scope')
    add('segments','segment_perimeters',{'BANK':{'entity':'BANK-GROUP','legal_entities':['BANK-GROUP','BANK-PARENT','BANK-LEGAL'],
        'profile_id':'bank','held_common_shares_m':6.,'share_class':'ordinary'},
        'PROPERTY':{'entity':'PROPERTY-GROUP','legal_entities':['PROPERTY-GROUP'],
        'profile_id':'property_owner','held_common_shares_m':8.,'share_class':'ordinary'}},period='2025-12-31')
    add('children','segment_valuations',children,period='2025-12-31')
    add('parent_balance','group_equity_bridge',{'legal_entity':'HOLDING-PARENT','cash':30.,'debt':20.,'preferred':0.,
        'other_liabilities':5.,'investments_book':{'BANK':60.,'PROPERTY':80.},'other_assets':0.,'common_equity':145.,
        'basis':'standalone_parent_excluding_child_assets_debt_income_and_funding','cash_basis':'unrestricted',
        'source':'Synthetic standalone parent balance sheet'},period='2025-12-31')
    for sc in ('bear','base','bull'):
        add('segment_fx','segment_valuations',{'BANK':{'currency':'EUR','to_group_rate':1.,'as_of':'2025-12-31'},
            'PROPERTY':{'currency':'EUR','to_group_rate':1.,'as_of':'2025-12-31'}},sc,period='2025-12-31')
        claims=[]
        for name,child in children.items():
            for index,target in enumerate(claim_targets(child,sc)):
                claims.append({'claim_id':name+'-CLAIM-'+str(index),'allocation':name,'legal_entity':target['entity'],
                    'currency':'EUR','amount':target['amount'],'role':target['role'],
                    'measurement_basis':'consumed_model_bridge','record_ref':target['record_ref'],
                    'source':'Synthetic claim inventory tied to child inputs'})
        for role,amount in [('cash',30.),('debt',20.),('preferred',0.),('other_liabilities',5.)]:
            claims.append({'claim_id':'PARENT-'+role,'allocation':'PARENT','legal_entity':'HOLDING-PARENT','currency':'EUR',
                'amount':amount,'role':role,'measurement_basis':'standalone_parent_settlement','record_ref':None,
                'source':'Synthetic standalone parent instrument register'})
        add('claims','group_equity_bridge',{'items':claims},sc,period='2025-12-31')
        add('central_cash_cost','central_costs',[2.,2.],sc,'EUR million','cash')
        add('central_discount_rate','discount_rate_inputs',.1,sc,'ratio','valuation')
        add('central_growth','central_costs',0.,sc,'ratio','valuation')
    return rows


def sotp_bundle(rows=None):
    records=sotp_records() if rows is None else rows
    return sector_analysis.prepare_sector_analysis('SYNTH-HOLD',as_of=DAY,providers={
        'profile':lambda *a,**k:{'status':'ok','source_id':'synthetic','as_of':DAY,'data':{'evidence':[
            {'field':f,'value':v,'source_id':'synthetic','as_of':DAY} for f,v in (
                ('instrument','equity'),('business_model','mixed_business'),
                ('segments',[{'id':name,'business_model':segment['profile_id']} for name,segment in
                    next(r for r in records if r['driver']=='segments')['value'].items()]))]}},
        'method_inputs':lambda *a,**k:{'status':'ok','source_id':'synthetic','as_of':DAY,'records':deepcopy(records)}},
        user_context={'analysis_context':{'scenario_rationale':{s:'Synthetic matched holding scenario' for s in ('bear','base','bull')}}})


def generate(tmp_path,rows=None):
    return dcf_engine.generate_valuation('SYNTH-HOLD',prepared_bundle=sotp_bundle(rows),output_dir=str(tmp_path))


def test_sotp_calls_actual_bank_and_property_adapters_with_raw_equity_values(tmp_path):
    rows=sotp_records();before=json.dumps(rows,sort_keys=True)
    result=generate(tmp_path,rows)
    assert result['valuation_usability']['usable'],result.get('error')
    # Bank14,4+118 cash; property545 equity. Ownership6/10 and8/10, parent
    # cash30-debt20-other5, central aftertax cost2/.1. No child debt deduction.
    bank_equity=14./1.1+122./1.21
    expected=(.6*bank_equity+.8*545.+30.-20.-5.-20.)/10.
    assert result['fair_value_base']==round(expected,2)
    assert json.dumps(rows,sort_keys=True)==before
    segments=result['calculation_details']['scenarios']['base']['segments']
    assert {s['method_id'] for s in segments}=={'bank_residual_income','property_nav'}
    assert all(Path(s['path']).is_file() and s['generation_id'] for s in segments)


@pytest.mark.parametrize('fault',['missing_segment','incomplete_child','duplicate_claim','wrong_allocation','unmapped_debt',
    'overlap','overowned','wrong_fx','wrong_cutoff','missing_cost','consolidated_parent','intercompany','opening_claims_in_cost','unconsumed'])
def test_sotp_never_exposes_total_for_missing_or_double_counted_components(tmp_path,fault):
    rows=sotp_records();get=lambda d:next(r for r in rows if r['driver']==d)
    segments=get('segments')['value'];children=get('children')['value'];claims=get('claims')['value']['items']
    if fault=='missing_segment':children.pop('PROPERTY')
    if fault=='incomplete_child':
        child=children['BANK'];new=[r for r in child['case']['records'] if r['driver']!='capital.subsidiaries.0.required_statutory_capital']
        children['BANK']=bank_bundle(new,symbol='SYNTH-BANK')
    if fault=='duplicate_claim':claims[-1]['claim_id']=claims[0]['claim_id']
    if fault=='wrong_allocation':claims[0]['allocation']='PARENT'
    if fault=='unmapped_debt':claims[0]['amount']+=1.
    if fault=='overlap':segments['PROPERTY']['legal_entities'].append('BANK-PARENT')
    if fault=='overowned':segments['PROPERTY']['held_common_shares_m']=11.
    if fault=='wrong_fx':get('segment_fx')['value']['BANK']['to_group_rate']=1.2
    if fault=='wrong_cutoff':get('segment_fx')['value']['BANK']['as_of']='2026-01-01'
    if fault=='missing_cost':rows.remove(get('central_cash_cost'))
    if fault=='consolidated_parent':get('parent_balance')['value']['basis']='consolidated_including_children'
    if fault=='intercompany':get('policy')['value']['intercompany']='parent_loan_to_child'
    if fault=='opening_claims_in_cost':
        get('policy')['value']['central_costs']='includes_settlement_of_opening_accrued_expense'
        get('central_cash_cost')['value']=[5.,0.]
    if fault=='unconsumed':get('parent_balance')['value']['default_holding_premium']=.1
    result=generate(tmp_path,rows)
    assert not result['valuation_usability']['usable'] and result.get('fair_value_base') is None
    assert result['acquisition_tasks']


@pytest.mark.parametrize('fault',['child_generation','child_snapshot','segment_generation','child_removed','child_blocked','evidence'])
def test_sotp_consumer_rejects_mismatched_child_provenance(tmp_path,fault):
    from bellomberg.valuation.dcf_quality import normalize_valuation_payload
    result=generate(tmp_path);assert result['valuation_usability']['usable']
    child=result['child_valuations']['BANK']
    if fault=='child_generation':child['generation_id']='different-generation'
    if fault=='child_snapshot':child['acquisition_snapshot']['case']['ticker']='WRONG'
    if fault=='segment_generation':result['calculation_details']['scenarios']['base']['segments'][0]['generation_id']='wrong'
    if fault=='child_removed':result['child_valuations'].pop('BANK')
    if fault=='child_blocked':child['sanity']['severity']='BLOCK'
    if fault=='evidence':next(r for r in result['analytical_quality']['rows'] if r['driver']=='children')['values']={}
    blocked=normalize_valuation_payload(result,as_of=DAY)
    assert not blocked['valuation_usability']['usable']
    assert blocked['fair_value_base'] is None


def test_sotp_is_independent_of_child_quote_unit_and_adr_ratio(tmp_path):
    rows=sotp_records();children=next(r for r in rows if r['driver']=='children')['value']
    original=generate(tmp_path,rows);assert original['valuation_usability']['usable']
    bank=children['BANK'];records=deepcopy(bank['case']['records'])
    quote=next(r for r in records if r['driver']=='quotation')['value']
    # The bank remains EUR equity. Its quoted security is now GBP pennies,
    # two ordinary shares per quoted ADR; the parent owns ordinary shares.
    quote.update(quote_currency='GBP',quote_unit='GBX',quote_units_per_currency=100,
                 financial_to_quote_rate=.8,shares_per_quote=2,price=quote['price']*160)
    children['BANK']=bank_bundle(records,symbol='SYNTH-BANK')
    changed=generate(tmp_path,rows)
    assert changed['valuation_usability']['usable'],changed.get('error')
    assert changed['fair_value_base']==original['fair_value_base']
    assert changed['child_valuations']['BANK']['fair_value_base']!=original['child_valuations']['BANK']['fair_value_base']


def test_documented_excel_preserves_long_contracts_and_complete_horizon(tmp_path):
    from openpyxl import Workbook,load_workbook
    from bellomberg.valuation.dcf_quality_sheet import append_quality_sheet
    text='=SOURCE '+('\U0001f600'*18000)+' END'
    years=[str(y) for y in range(2026,2034)];path=[float(i) for i in range(8)]
    report={'record_adapter':True,'method_id':'mixed_business_sotp','status':'DOCUMENTATA',
        'issues':[],'snapshot':{'forecast_years':years},'rows':[
            {'driver':'cash','scenario':'base','values':path,'evidence':{'unit':'EUR million','source':'synthetic'}},
            {'driver':'children','scenario':'model','values':{'A':{'proof':text}},'evidence':{'source':'synthetic'}}]}
    wb=Workbook();append_quality_sheet(wb,report);file=tmp_path/'quality.xlsx';wb.save(file);wb.close()
    wb=load_workbook(file);ws=wb['Qualita e revisioni'];rows=list(ws.values)
    assert all(year in next(r for r in rows if r[0]=='Scenario') for year in years)
    assert list(next(r for r in rows if r[1]=='cash')[2:10])==path
    first=next(r for r in rows if r[1]=='children.A.proof')[2]
    rebuilt=first+''.join(r[2] for r in rows if r[0]=='TEXT CONTINUATION')
    assert rebuilt==text and all(c.data_type!='f' for row in ws for c in row)
    wb.close()


@pytest.mark.parametrize('held,usable',[(9.,True),(10.,False)])
def test_sotp_digital_ownership_uses_issued_common_not_unexercised_warrants(tmp_path,held,usable):
    from test_sector_nav_drivers import nav_records,nav_bundle
    from bellomberg.valuation.sotp_adapter import claim_targets
    rows=sotp_records();get=lambda d:next(r for r in rows if r['driver']==d)
    child=nav_bundle(remap_entities(nav_records('dat'),{'SYNTH-GROUP':'PROPERTY-GROUP'}),'dat')
    get('quotation')['value']['price']=15.
    get('children')['value']['PROPERTY']=child
    get('segments')['value']['PROPERTY'].update(profile_id='dat',held_common_shares_m=held)
    for row in rows:
        if row['driver']!='claims':continue
        items=[i for i in row['value']['items'] if i['allocation']!='PROPERTY']
        for index,target in enumerate(claim_targets(child,row['scenario'])):
            items.append({'claim_id':'DAT-'+str(index),'allocation':'PROPERTY','legal_entity':target['entity'],
                'currency':'EUR','amount':target['amount'],'role':target['role'],'measurement_basis':'consumed_model_bridge',
                'record_ref':target['record_ref'],'source':'Synthetic issued capital and treasury claims'})
        row['value']['items']=items
    result=generate(tmp_path,rows)
    assert result['valuation_usability']['usable'] is usable,result.get('error')
    if usable:
        part=result['calculation_details']['scenarios']['base']['segments'][1]
        assert part['owned_equity_value']==94.5 and part['valuation_shares_m']==10.
    else:assert result.get('fair_value_base') is None and result['acquisition_tasks']


@pytest.mark.parametrize('commitment',[True,False])
def test_sotp_undrawn_parent_commitment_is_support_even_without_modeled_calls(tmp_path,commitment):
    from test_resources_valuation import resource_records,resource_bundle
    from bellomberg.valuation.sotp_adapter import _has_calls
    rows=resource_records()
    for r in rows:
        if r['driver']=='opening_cash':r['value']=50.
        if r['driver']=='opening_cash_buffer':r['value']=20.
        if r['driver']=='minimum_cash':r['value']=[20.,10.]
        if r['driver']=='funding' and not commitment:r['value']={'capacities':[0.,0.],'commitments':{}}
    child=resource_bundle(rows)
    result=dcf_engine.generate_valuation(child['case']['ticker'],prepared_bundle=child,output_dir=str(tmp_path))
    assert result['valuation_usability']['usable'],result.get('error')
    ledger=result['calculation_details']['scenarios']['base']['ledger']
    assert len(ledger)==2 and sum(r['funding_at_start']+r['funding_at_end'] for r in ledger)==0.
    assert _has_calls(result,child,'base') is commitment
