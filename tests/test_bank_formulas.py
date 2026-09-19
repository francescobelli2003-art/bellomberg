"""Synthetic bank workbook: separate engine baseline and linked legal ledgers."""
from copy import deepcopy
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from bellomberg.valuation.bank_formulas import apply_bank_formulas
from bellomberg.valuation.dcf_engine import generate_valuation
from test_sector_bank_capital import bank_bundle, bank_records


def _case():
    with TemporaryDirectory() as directory:
        payload = generate_valuation('SYNTH-BANK', prepared_bundle=bank_bundle(), output_dir=directory)
    assert payload['valuation_usability']['usable']
    return payload


def test_bank_schedules_preserve_engine_payload_and_link_three_scenarios():
    payload = _case()
    original = deepcopy(payload)
    book = Workbook()
    assert apply_bank_formulas(book, payload)
    assert payload == original
    assert all(name in book for name in ('Model Inputs', 'Bank bear', 'Bank base', 'Bank bull',
                                          'Input Checks', 'Model Checks', 'Summary'))
    for name in ('Bank bear', 'Bank base', 'Bank bull'):
        ws = book[name]
        assert ws['D9'].data_type == 'f'       # consolidated NI
        assert ws['D10'].data_type == 'f'      # common book roll-forward
        assert ws['D21'].data_type == 'f'      # statutory capital roll-forward
        assert ws['D23'].data_type == 'f'      # legal-entity liquidity
        assert ws['D14'].data_type == 'f'      # shareholder distribution
        assert ws['F10'].data_type == 'f'      # terminal P/TBV
    assert book['Summary']['E8'].data_type == 'f'
    assert book['Model Checks']['D9'].data_type == 'f'


def test_bank_key_driver_dependencies_and_failure_guards():
    payload = _case()
    book = Workbook()
    assert apply_bank_formulas(book, payload)
    inputs = book['Model Inputs']
    labels = {str(inputs.cell(r, 2).value): r for r in range(1, inputs.max_row + 1)}
    assert inputs.cell(labels['capital.shares m'], 4).value > 0
    assert inputs.cell(labels['capital.ke'], 4).value > 0
    assert inputs.cell(labels['net interest income'], 4).value > 0
    for label in ('capital.terminal equity', 'capital.subsidiaries.0.required statutory capital',
                  'capital.subsidiaries.0.liquidity before transfers'):
        assert str(inputs.cell(labels[label], 4).value).startswith('=IFERROR(')
    checks = book['Input Checks']
    expressions = '\n'.join(str(checks.cell(r, 3).value) for r in range(1, checks.max_row + 1))
    for fragment in ('Shares, Ke, ROE and growth', 'funding'):
        assert any(fragment in str(checks.cell(r, 2).value) for r in range(1, checks.max_row + 1))
    assert 'COUNTIF' in expressions
    assert '"KO"' in expressions
    assert '"n.d."' in book['Summary']['E8'].value


def test_bank_unusable_payload_does_not_offer_a_live_model():
    payload = _case()
    payload['valuation_usability']['usable'] = False
    book = Workbook()
    assert not apply_bank_formulas(book, payload)
    assert book.sheetnames == ['Sheet']


def _positive_growth_case():
    rows = bank_records()
    for scenario in ('bear', 'base', 'bull'):
        get = lambda name: next(row['value'] for row in rows
                                if row['scenario'] == scenario and row['driver'] == name)
        next(row for row in rows if row['scenario'] == scenario and row['driver'] == 'terminal_growth')['value'] = .02
        forecast = get('capital_constraints')['LEGAL-A']['constraints']
        forecast.append({'id': 'Secondary forecast requirement', 'exposure': [500., 500.],
                         'ratio': [.1, .1], 'buffer': [0., 0.], 'absolute_floor': [0., 0.],
                         'terminal_requirement': 50.})
        terminal = get('terminal_ledger')
        capital = terminal['capital']
        capital['parent_cash_minimum'] = [10.2]
        capital['terminal_debt'] = 5.1
        capital['parent_cash_flows']['debt_issued'] = [.1]
        capital['parent_cash_flows']['other_cash_receipts'] = [1.9]
        sub = capital['subsidiaries'][0]
        sub['proposed_distribution'] = [7.64]
        sub['other_statutory_movements'] = [.04]
        sub['liquidity_before_transfers'] = [23.96]
        terminal['liquidity_bridge']['LEGAL-A']['operating_cash'] = [12.96]
        constraints = terminal['capital_constraints']['LEGAL-A']['constraints']
        constraints[0]['terminal_requirement'] = 112.2
        constraints.insert(0, {'id': 'Extra continuing requirement', 'exposure': [300.],
                               'ratio': [.1], 'buffer': [0.], 'absolute_floor': [0.],
                               'terminal_requirement': 30.6})
        constraints.append({'id': 'Continuing floor', 'exposure': [0.], 'ratio': [0.],
                            'buffer': [0.], 'absolute_floor': [20.],
                            'terminal_requirement': 20.4})
    with TemporaryDirectory() as directory:
        payload = generate_valuation('SYNTH-BANK', prepared_bundle=bank_bundle(rows), output_dir=directory)
    assert payload['valuation_usability']['usable'], payload.get('error')
    return payload


def test_positive_growth_terminal_requirements_allow_different_constraint_lists():
    payload = _positive_growth_case()
    book = Workbook()
    assert apply_bank_formulas(book, payload)
    checks = book['Input Checks']
    labels = [str(checks.cell(row, 2).value) for row in range(1, checks.max_row + 1)]
    assert any('continuing next requirement 2' in label for label in labels)
    assert payload['fair_value_base'] == 11.36
