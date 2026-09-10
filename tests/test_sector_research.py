"""Research candidates use the acquired case independently of holdings."""
from datetime import date
import json
import sqlite3
from types import SimpleNamespace

from bellomberg.valuation import sector_analysis as sector
from test_sector_analysis import providers_for, DAY


def test_research_reuses_acquisition_across_rounds_and_declares_gaps(tmp_path, monkeypatch):
    from bellomberg.core import current_facts
    from bellomberg.storage import memory_db
    database = tmp_path / "research.db"
    with sqlite3.connect(database) as conn:
        conn.execute("CREATE TABLE decisions(id INTEGER, ticker TEXT, timestamp TEXT, memo_id INTEGER, timing TEXT, rationale TEXT, action TEXT, status TEXT)")
        conn.execute("CREATE TABLE decision_notes(decision_id INTEGER, autore TEXT, testo TEXT, timestamp TEXT)")
        conn.execute("INSERT INTO decisions VALUES(1,'OUTSIDE.EU',?,2,'filing','synthetic idea','RESEARCH','PENDING')", (date.today().isoformat(),))
    monkeypatch.setattr(memory_db, "SQLITE_PATH", str(database))
    calls, bundles, links = [], {}, {}
    providers = providers_for("insurance_pc", calls=calls)
    first = current_facts.research_block(sector_bundles=bundles, providers=providers, as_of=DAY, decision_links=links)
    second = current_facts.research_block(sector_bundles=bundles, providers=providers, as_of=DAY, decision_links=links)
    assert first == second
    assert len(calls) == 2
    assert "insurance_pc_distributable_equity" in first and "integrated" in first
    assert "legal_entity_capital" in first and "data_missing" in first
    assert links == {"OUTSIDE.EU": 1}
    assert bundles["OUTSIDE.EU"]["case"]["ticker"] == "OUTSIDE.EU"
    with sqlite3.connect(database) as conn:
        assert conn.execute("SELECT count(*) FROM decisions").fetchone()[0] == 1


def test_specialist_dispatch_receives_exact_research_bundle(monkeypatch):
    from bellomberg.agents.specialists.base import Specialist
    from bellomberg.agents import chat_tools
    bundle = sector.prepare_sector_analysis("OUTSIDE", as_of=DAY, providers=providers_for("insurance_pc"))
    specialist = Specialist.__new__(Specialist)
    specialist.name = "fundamentals"
    specialist._sector_bundles = {"OUTSIDE": bundle}
    specialist.blackboard = SimpleNamespace(memory_db=None)
    seen = []
    monkeypatch.setattr(chat_tools, "dispatch", lambda name, data, **kw: seen.append(kw) or {"ok": False})
    specialist._execute_meta_tool("get_valuation", {"ticker": "OUTSIDE"})
    assert seen[0]["prepared_bundle"] is bundle


def test_valuation_failure_does_not_invoke_alternate_dispatcher(monkeypatch):
    from bellomberg.agents.specialists.base import Specialist
    from bellomberg.agents import chat_tools
    specialist = Specialist.__new__(Specialist)
    specialist.name = "fundamentals"
    specialist.blackboard = SimpleNamespace(memory_db=None)
    def failure(*a, **kw):
        raise ValueError("snapshot mismatch")
    monkeypatch.setattr(chat_tools, "dispatch", failure)
    result = specialist._execute_meta_tool("get_valuation", {"ticker": "OUTSIDE"})
    assert result["exclude_from_action_table"] and "snapshot mismatch" in result["error"]


def test_stamped_tool_result_reaches_committee_and_research_link(monkeypatch):
    import threading
    from bellomberg.agents.specialists.base import Specialist
    from bellomberg.agents import chat_tools
    bundle = sector.prepare_sector_analysis("OUTSIDE", as_of=DAY, providers=providers_for("insurance_pc"))
    revised = sector.revise_sector_analysis(bundle, assumptions={"variant_view": "New evidence"})
    payload = {"ticker": "OUTSIDE", "snapshot_id": revised["snapshot_id"], "generation_id": "synthetic-id",
               "valuation_decision": revised["decision"], "acquisition_snapshot": revised}
    linked = []
    db = SimpleNamespace(link_valuation_snapshot=lambda *a, **kw: linked.append((a, kw)))
    specialist = Specialist.__new__(Specialist)
    specialist.name = "fundamentals"
    specialist._sector_bundles = {"OUTSIDE": bundle}
    specialist._research_decision_links = {"OUTSIDE": 11}
    specialist.blackboard = SimpleNamespace(memory_db=db, valuation_results={}, _lock=threading.RLock())
    stamped = chat_tools._stamp(payload, "synthetic valuation")
    monkeypatch.setattr(chat_tools, "dispatch", lambda *a, **kw: stamped)
    assert specialist._execute_meta_tool("get_valuation", {"ticker": "OUTSIDE"}) is stamped
    assert linked == [((revised["snapshot_id"],), {"generation_id": "synthetic-id", "decision_id": 11})]
    assert specialist._sector_bundles["OUTSIDE"] == revised
    row = json.loads(sector.valuation_results_block(specialist.blackboard.valuation_results).splitlines()[-1])
    assert row["method_id"] == revised["decision"]["method_id"]
    assert row["snapshot_id"] == revised["snapshot_id"]


def test_committee_view_does_not_promote_legacy_or_secondary_fv():
    block = sector.valuation_results_block({"OUTSIDE": {"fair_value_base": 999, "fair_value_sotp": 111,
                                                       "price": 10, "sanity": {"severity": "OK"}}})
    row = json.loads(block.splitlines()[-1])
    assert row["fair_value"] is None
    assert row["valuation_usability"]["usable"] is False
    assert "999" not in block and "111" not in block


def test_excel_starts_with_same_incomplete_decision():
    from openpyxl import Workbook
    from bellomberg.valuation.dcf_quality_sheet import append_sector_quality_sheet
    bundle = sector.prepare_sector_analysis("OUTSIDE", as_of=DAY, providers=providers_for("insurance_pc"))
    wb = Workbook()
    ws = append_sector_quality_sheet(wb, {"ticker": "OUTSIDE", "valuation_decision": bundle["decision"],
        "snapshot_id": bundle["snapshot_id"], "acquisition_snapshot": bundle,
        "acquisition_tasks": bundle["acquisition_tasks"]})
    assert wb.worksheets[0] == ws
    assert ws["B2"].value == "NON UTILIZZABILE / BOZZА".replace("А", "A")
    assert ws["B3"].value == bundle["decision"]["method_id"]
    assert "legal_entity_capital" in ws["C5"].value


def test_fundamentals_no_universal_operating_or_us_filter():
    from bellomberg.agents.specialists.fundamentals import FundamentalsSpecialist
    prompt = FundamentalsSpecialist.system_prompt
    assert "NON chiamare MAI get_valuation senza growth_path" not in prompt
    assert "solo su US/ADR" not in prompt
    assert "valuation_usability" in prompt and "acquisition_tasks" in prompt


def test_score_rejects_workbook_changed_after_valid_generation(tmp_path, monkeypatch):
    from bellomberg.agents import specialist_scores
    from bellomberg.storage import classificazione
    from test_sector_valuation_integration import synthetic_bundle, documented_payload, SYMBOL
    workbook = tmp_path / ("VAL_" + SYMBOL + ".xlsx")
    workbook.write_bytes(b"synthetic original workbook")
    bundle = synthetic_bundle(day=date.today().isoformat())
    payload = documented_payload(bundle, workbook)
    payload["_timestamp"] = date.today().isoformat()
    sidecar = workbook.with_suffix(".payload.json")
    sidecar.write_text(json.dumps(payload), encoding="utf-8")
    monkeypatch.setattr(specialist_scores, "REPORT_DIR", tmp_path)
    monkeypatch.setattr(classificazione, "carica_veicoli", lambda: {"origine": "synthetic", "veicoli": {}, "motivo": None})
    portfolio = {"positions": [{"ticker": SYMBOL, "peso_pct": 100}]}
    assert specialist_scores.fundamentals_score(portfolio) is not None
    workbook.write_bytes(b"changed by analyst without revalidation")
    assert specialist_scores.fundamentals_score(portfolio) is None


def test_external_planned_case_crosses_real_tool_storage_and_f17(tmp_path, monkeypatch):
    from bellomberg.storage import memory_db
    from bellomberg.agents import chat_tools
    from test_sector_valuation_api import endpoint
    monkeypatch.setattr(memory_db.MemoryDB, "_init_chroma", lambda self: None)
    db = memory_db.MemoryDB(str(tmp_path / "chain.db"), str(tmp_path / "chroma"))
    monkeypatch.setattr(memory_db, "MemoryDB", lambda: db)
    bundle = sector.prepare_sector_analysis("OUTSIDE.EU", as_of=DAY, providers=providers_for("insurance_pc"))
    result = chat_tools.dispatch("get_valuation", {"ticker": "OUTSIDE.EU"}, prepared_bundle=bundle)["data"]
    assert result["_thesis_saved"]["thesis_id"]
    saved = db.get_valuation_history("OUTSIDE.EU")[0]
    assert saved["fair_value"] is None
    assert saved["snapshot_id"] == bundle["snapshot_id"]
    assert saved["valuation_payload"]["valuation_decision"] == result["valuation_decision"]
    model = endpoint(tmp_path, db.get_latest_valuation_snapshots())["models"][0]
    assert model["ticker"] == "OUTSIDE.EU" and model["file"] == ""
    assert model["snapshot_id"] == result["snapshot_id"]
    assert model["generation_id"] == result["generation_id"]
    assert model["valuation_usability"]["usable"] is False and model["fair_value"] is None
    assert model["acquisition_tasks"] == result["acquisition_tasks"]
