from copy import deepcopy

import pytest
from openpyxl import load_workbook

from bellomberg.valuation.documented_formulas import operating_formula_ready
from bellomberg.valuation.documented_inputs import build_documented_workbook
from bellomberg.valuation.dcf_engine import generate_valuation
from tests.test_documented_cashflow_bridge import _bundle


@pytest.mark.parametrize('years', [4, 10, 15])
def test_live_workbook_keeps_full_calendar_and_original_engine_result(tmp_path, years):
    result = generate_valuation('SYNTH-EXT', prepared_bundle=_bundle(years), output_dir=str(tmp_path))
    assert operating_formula_ready(result)
    wb = load_workbook(result['path'])
    for scenario in ('bear', 'base', 'bull'):
        sheet = wb[scenario]
        assert sheet.cell(7, years + 3).value == f'{2025+years}-12-31'
        assert sheet.cell(26, years + 3).data_type == 'f'
        assert sheet['D54'].data_type == 'f'
    assert wb['Valuation']['B2'].value == result['fair_value_base']
    assert wb['Sensitivity']['E7'].font.color.rgb.endswith('FFFFFF')
    assert wb.calculation.fullCalcOnLoad is True
    assert max(len(cell.value) for s in wb for row in s for cell in row if cell.data_type == 'f') < 8192
    wb.close()


def test_missing_engine_result_cannot_be_rebuilt_silently_from_inputs(tmp_path):
    result = generate_valuation('SYNTH-EXT', prepared_bundle=_bundle(4), output_dir=str(tmp_path))
    payload = deepcopy(result)
    payload['calculation_details']['scenarios']['base']['rows']['ufcf'][2] = None
    assert not operating_formula_ready(payload)
    wb = load_workbook(build_documented_workbook(payload, tmp_path))
    assert wb['base']['F11'].value == 'n.d.'
    assert 'BOZZA' in wb['Summary']['B5'].value
    assert not any(c.data_type == 'f' for s in wb for row in s for c in row)
    wb.close()


@pytest.mark.parametrize('status', ['fx_not_rolled', 'stale', 'currency_mismatch'])
def test_excel_does_not_restore_a_blocked_price_comparison(tmp_path, status):
    from test_market_quote import generate
    payload = generate(tmp_path)
    payload['market_quote'].update(status=status, message='Synthetic comparison unavailable')
    for scenario in ('bear','base','bull'):
        payload['market_quote']['upside_'+scenario+'_pct'] = None
    wb = load_workbook(build_documented_workbook(payload, tmp_path))
    assert all(wb['Summary'].cell(9,c).value == 'n.d.' for c in (4,5,6))
    assert wb['Sensitivity']['D20'].value == 'n.d.'
    assert 'Synthetic comparison unavailable' in wb['Summary']['B26'].value
    assert wb['Summary']['E8'].data_type == 'f'
    assert wb['Sensitivity']['E10'].data_type == 'f'
    wb.close()


def test_valid_snapshot_quote_keeps_interactive_comparison(tmp_path):
    from test_market_quote import generate
    payload = generate(tmp_path)
    assert payload['market_quote']['status'] == 'ok'
    wb = load_workbook(payload['path'])
    assert wb['Summary']['E9'].data_type == 'f'
    assert wb['Sensitivity']['D20'].data_type == 'f'
    wb.close()
