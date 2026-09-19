"""Linked NAV workbook checks; no Excel or provider access is required."""
import sys
from copy import deepcopy

from openpyxl import load_workbook

from bellomberg.valuation import dcf_engine
from bellomberg.valuation.nav_formulas import apply_nav_formulas

sys.path.insert(0, "tests")
from test_sector_nav_drivers import nav_bundle


def _result(tmp_path, profile):
    return dcf_engine.generate_valuation(
        "SYNTH-NAV", prepared_bundle=nav_bundle(profile=profile), output_dir=str(tmp_path)
    )


def test_fund_nav_links_inputs_and_preserves_raw_valuation(tmp_path):
    result = _result(tmp_path, "cef")
    workbook = load_workbook(result["path"])
    valuation_b2 = workbook["Valuation"]["B2"].value
    raw = {name: workbook[name]["B2"].value for name in ("bear", "base", "bull")}
    assert apply_nav_formulas(workbook, result) is True
    assert workbook["Valuation"]["B2"].value == valuation_b2
    assert {name: workbook[name]["B2"].value for name in raw} == raw
    assert {"Summary", "NAV Model", "NAV Inputs", "Model Checks"} <= set(workbook.sheetnames)
    assert "'NAV Inputs'!D13" in workbook["NAV Model"]["D8"].value
    assert "'NAV Inputs'!D16" in workbook["NAV Model"]["D12"].value
    assert "'NAV Inputs'!D20" in workbook["NAV Model"]["D12"].value
    assert 'Model Checks' in workbook['NAV Model']['D18'].value
    assert "*'NAV Inputs'!D12" in workbook["NAV Model"]["D19"].value
    assert workbook.calculation.fullCalcOnLoad is True
    workbook.close()


def test_digital_nav_fd_and_itm_formulas_are_explicit(tmp_path):
    result = _result(tmp_path, "dat")
    workbook = load_workbook(result["path"])
    assert apply_nav_formulas(workbook, result) is True
    model = workbook["NAV Model"]
    assert "ownership" in " ".join(str(c.value) for row in workbook["Model Checks"].iter_rows() for c in row)
    assert 'D24' in model['D14'].value
    assert "'NAV Inputs'!D14-(0)" in model["D10"].value
    assert "*'NAV Inputs'!D28" in model["D8"].value
    assert workbook['NAV Inputs']['D8'].data_type == 'f'
    assert 'Warrant ITM' in [row[1].value for row in workbook['Model Checks']]
    workbook.close()


def test_invalid_nav_is_declared_nd_and_returns_false(tmp_path):
    result = _result(tmp_path, "cef")
    invalid = deepcopy(result)
    invalid["valuation_usability"] = {"usable": False}
    from bellomberg.valuation.documented_inputs import build_documented_workbook
    workbook = load_workbook(build_documented_workbook(invalid, tmp_path / 'invalid'))
    assert apply_nav_formulas(workbook, invalid) is False
    assert 'NAV Model' not in workbook
    workbook.close()
