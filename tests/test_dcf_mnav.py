"""Motore mNAV V5 Lotto 1 (design audit/18, decisioni PM D1-D4 23/07): funzioni
pure di dcf_mnav su FIXTURE FINTE dei tool ufficiali (dat_metrics/cef_nav) —
ZERO rete, date esplicite nei casi STALE (mai date.today() implicita).
Aspettative calcolate A MANO nei commenti (mai copiate dall'output del codice).

05/09 (Fable 5.1, classificazione lotto 2b): il perimetro mNAV viene dal NEGOZIO dei veicoli
(`MNAV_KINDS` derivato dal `sottostante` delle voci dat; i fondi chiusi da `cef_nav.CEF_SOURCES`,
a sua volta derivato): qui il negozio e' un file finto in tmp_path con simboli inventati, e il
messaggio di rifiuto porta un conteggio, mai l'elenco dei coperti.
"""
import json
import re
from datetime import date

import pytest

from bellomberg.valuation import cef_nav
import bellomberg.storage.classificazione as cl
from bellomberg.valuation import dcf_mnav
from bellomberg.valuation.dcf_mnav import (RECORD_STALE_DAYS_BTC, build_mnav_model, build_mnav_spec,
                      compute_mnav_values)

OGGI = date(2026, 7, 23)
INFO_USD = {"currency": "USD", "longName": "Test Vehicle Inc"}
CHIAVE_CEF = next(iter(cef_nav.FETCH_NAV))     # dominio del fetcher vero, letto dal registro

VOCI = {
    "TESORO": {"tipo": "dat", "provenienza": "dichiarato", "sottostante": "BTC-USD", "kind": "dat_bitcoin"},
    "HYPEX": {"tipo": "dat", "provenienza": "dichiarato", "sottostante": "HYPE", "kind": "dat_hype"},
    "TESORO2": {"tipo": "dat", "provenienza": "dichiarato", "sottostante": "ETH-USD"},   # non supportato
    "TESORO3": {"tipo": "dat", "provenienza": "dichiarato"},                            # sottostante assente
    "FONDO.L": {"tipo": "cef", "provenienza": "dichiarato", "nav_fonte": CHIAVE_CEF, "nav_valuta": "USD",
                "nome": "Fondo Chiuso Finto"},
    "FONDO2.L": {"tipo": "cef", "provenienza": "dichiarato"},                            # senza fonte
    "FONDO3.L": {"tipo": "cef", "provenienza": "dichiarato", "nav_fonte": "sito-finto.example",
                 "nav_valuta": "USD"},
    "METALLO": {"tipo": "commodity", "provenienza": "dichiarato"},
}


@pytest.fixture(scope="module")
def NEG(tmp_path_factory):
    p = tmp_path_factory.mktemp("negozio") / "v.json"
    p.write_text(json.dumps(VOCI), encoding="utf-8")
    return cl.carica_veicoli(str(p))


# --- fixture DAT bitcoin: btc_nav = 600k x 100k = 60 mld; mktcap = 300M x 200 = 60 mld
# mnav_equity = 1,000; EV = 60+8+12-0,5 = 79,5 mld -> mnav_ev = 1,325
# NAV equity/azione = (60-8-12+0,5) mld / 300M = 40,5 mld / 300M = 135,00
BTC_PAYLOAD = {
    "ticker": "TESORO",
    "latest": {"btc_holdings": 600_000, "basic_shares_outstanding": 300_000_000,
               "debt": 8_000_000_000, "pref": 12_000_000_000, "cash": 500_000_000,
               "as_of_date": "2026-07-20"},
    "derived_mnav": {
        "mnav_equity_basic": 1.000, "mnav_ev": 1.325,
        "inputs": {"price_mstr": 200.0, "price_btc": 100_000,
                   "record_as_of": "2026-07-20",
                   "computed_at_utc": "2026-07-23T10:00:00Z", "cache_ttl_min": 30},
        "leggimi": "giudicare su mnav_ev"},
}

# --- fixture DAT hype ($M): hype_val = 30 x 40 = 1200; tax 21% -> dtl = -0,21x(1200-900)
# = -63; dtl_change = 0 - 63 = -63; adj_nav = 1000-20+30-50-900+1200-63 = 1197
# warrant: strike 5 ITM -> (12-5)x10/12 = 5,8333; strike 20 OTM -> 0
# fd = 105,8333; anav_ps = 1197/105,8333 = 11,3102; mnav = 12/11,3102 = 1,061
# anav_dtl = (1197+63)/105,8333 = 11,9055; mnav_dtl = 12/11,9055 = 1,008
HYPE_PAYLOAD = {
    "ticker": "HYPEX",
    "dat_fundamentals_musd": {"bookNAV": 1000.0, "cashFromOps": 20.0,
                              "cashFromFin": 30.0, "treasuryDeploy": 50.0,
                              "reportedDigital": 900.0, "hypeHeld": 30.0,
                              "taxRate": 21.0, "taxBasis": 900.0,
                              "reportedDTL": 0.0, "basicShares": 100.0},
    "warrants": [{"strike": 5.0, "amount": 10.0}, {"strike": 20.0, "amount": 5.0}],
    "vintage": {"effective_date": "2026-06-30", "next_update": "2026-09-30",
                "lag_note": "dati di bilancio Q2"},
    "purr_quote": {"last": 12.0, "fonte": "yfinance (Nasdaq)"},
    "hype_quote": {"last": 40.0, "fonte": "hyperliquid allMids"},
    "derived": {"mnav": 1.061},
}

# --- fixture CEF: px 4200 GBp -> 56,28 USD (tasso implicito 56,28/42 = 1,34)
# sconto = 56,28/80 - 1 = -29,65%; FV a target 0,80 = 64,00
CEF_PAYLOAD = {
    "ticker": "FONDO.L", "name": "Fondo Chiuso Finto",
    "nav_per_share_usd": 80.0, "nav_as_of": "2026-07-17", "nav_age_days": 6,
    "mtd_return_pct": 1.2, "qtd_return_pct": 3.4, "ytd_return_pct": 10.0,
    "nav_source": "sito-finto.example (NAV settimanale ufficiale)",
    "market_price": 4200.0, "market_price_currency": "GBp",
    "market_price_in_nav_ccy": 56.28,
    "fx_conversion": "GBp/100 -> GBP -> USD a GBPUSD=1.3400 [src: yfinance]",
    "discount_to_nav_pct": -29.7,
    "discount_note": "negativo = sconto",
}


# --------------------------------------------------------------------------
# il perimetro viene dal negozio (lotto 2b)
# --------------------------------------------------------------------------

def test_mnav_kinds_dichiarato_nel_negozio(NEG):
    assert dcf_mnav.mnav_kinds_del_negozio(NEG) == {"TESORO": "dat_bitcoin", "HYPEX": "dat_hype"}
    sk = dcf_mnav.dat_senza_kind_del_negozio(NEG)
    assert set(sk) == {"TESORO2", "TESORO3"}
    assert "ETH-USD" in sk["TESORO2"] and "sottostante" in sk["TESORO3"]
    assert dcf_mnav.KIND_DA_SOTTOSTANTE == {"BTC-USD": "dat_bitcoin", "HYPE": "dat_hype"}


def test_sottostante_supportato_non_inventa_kind_mancante(NEG):
    import copy
    n = copy.deepcopy(NEG)
    n["veicoli"]["TESORO"]["kind"] = None
    assert "TESORO" not in dcf_mnav.mnav_kinds_del_negozio(n)
    assert "kind" in dcf_mnav.dat_senza_kind_del_negozio(n)["TESORO"]
    spec = dcf_mnav.build_mnav_spec("TESORO", {}, negozio=n, tool_payload={})
    assert "kind" in spec["error"]


def test_vista_mnav_rifiuta_kind_incoerente_anche_in_un_esito_iniettato(NEG):
    import copy
    n = copy.deepcopy(NEG)
    n["veicoli"]["TESORO"]["kind"] = "dat_hype"
    assert "TESORO" not in dcf_mnav.mnav_kinds_del_negozio(n)
    assert "incoerente" in dcf_mnav.dat_senza_kind_del_negozio(n)["TESORO"]


def test_le_viste_all_import_sono_derivate_dal_negozio_corrente():
    corrente = cl.carica_veicoli()
    if corrente["origine"] in ("assente", "illeggibile"):
        pytest.skip("negozio dei veicoli %s: la misura sarebbe vacua ({} == {})" % corrente["origine"])
    assert dcf_mnav.MNAV_KINDS == dcf_mnav.mnav_kinds_del_negozio(corrente)


def test_il_cablaggio_ai_tool_passa_il_negozio(monkeypatch, NEG):
    """Review 05/09 (ALTA): il ramo `tool_payload is None` non era misurato — togliere
    `negozio=n` ai due tool, o scambiarli, lasciava tutto verde. Qui i tool sono registratori."""
    from bellomberg.valuation import dat_metrics
    import bellomberg.valuation.cef_nav as cn
    chiamate = []
    monkeypatch.setattr(dat_metrics, "get_dat_metrics",
                        lambda t, negozio=None: (chiamate.append(("dat", t, negozio)) or dict(BTC_PAYLOAD)))
    monkeypatch.setattr(cn, "get_cef_nav",
                        lambda t, negozio=None: (chiamate.append(("cef", t, negozio)) or dict(CEF_PAYLOAD)))
    s = build_mnav_spec("TESORO", INFO_USD, today=OGGI, negozio=NEG)
    assert "error" not in s and s["kind"] == "dat_bitcoin"
    s2 = build_mnav_spec("FONDO.L", INFO_USD, today=OGGI, negozio=NEG)
    assert "error" not in s2 and s2["kind"] == "cef_nav"
    assert chiamate == [("dat", "TESORO", NEG), ("cef", "FONDO.L", NEG)]
    assert chiamate[0][2] is NEG and chiamate[1][2] is NEG     # lo STESSO negozio, non riletto dal disco


def test_dat_con_sottostante_non_supportato_rifiutata_dichiarata(NEG):
    s = build_mnav_spec("TESORO2", INFO_USD, tool_payload=BTC_PAYLOAD, today=OGGI, negozio=NEG)
    assert "error" in s and "sottostante" in s["error"] and "ETH-USD" in s["error"]
    assert "non supportato" in s["error"] and "perimetro" in s["error"]
    s3 = build_mnav_spec("TESORO3", INFO_USD, tool_payload=BTC_PAYLOAD, today=OGGI, negozio=NEG)
    assert "error" in s3 and "sottostante" in s3["error"]


def test_cef_senza_fonte_riconosciuta_rifiutato_dichiarato(NEG):
    s = build_mnav_spec("FONDO2.L", INFO_USD, tool_payload=CEF_PAYLOAD, today=OGGI, negozio=NEG)
    assert "error" in s and "fonte NAV" in s["error"] and "perimetro" in s["error"]
    s3 = build_mnav_spec("FONDO3.L", INFO_USD, tool_payload=CEF_PAYLOAD, today=OGGI, negozio=NEG)
    assert "error" in s3 and "sito-finto.example" in s3["error"] and "FETCH_NAV" in s3["error"]


def test_fuori_perimetro_conta_e_non_elenca(NEG):
    for t in ("METALLO", "ZZZQ.XX"):
        s = build_mnav_spec(t, INFO_USD, tool_payload={"x": 1}, today=OGGI, negozio=NEG)
        assert "error" in s and "perimetro" in s["error"], s
        assert re.search(r"\b2 DAT\b", s["error"]) and re.search(r"\b1 fond", s["error"]), s["error"]
        for altro in set(VOCI) - {t}:
            assert altro not in s["error"], (t, altro)


def test_negozio_assente_o_illeggibile_dichiarato_nel_rifiuto(tmp_path):
    assente = cl.carica_veicoli(str(tmp_path / "manca.json"))
    s = build_mnav_spec("TESORO", INFO_USD, tool_payload=BTC_PAYLOAD, today=OGGI, negozio=assente)
    assert "error" in s and "ASSENTE" in s["error"] and "perimetro" in s["error"]
    (tmp_path / "rotto.json").write_text("{", encoding="utf-8")
    rotto = cl.carica_veicoli(str(tmp_path / "rotto.json"))
    s2 = build_mnav_spec("TESORO", INFO_USD, tool_payload=BTC_PAYLOAD, today=OGGI, negozio=rotto)
    assert "error" in s2 and "ILLEGGIBILE" in s2["error"]


# --------------------------------------------------------------------------
# la meccanica di ieri, sui simboli inventati
# --------------------------------------------------------------------------

def test_rifiuto_fuori_perimetro_e_fonte_rotta(NEG):
    # un veicolo dichiarato di un altro tipo NON entra (estensione = voce nel negozio + fonte)
    s = build_mnav_spec("METALLO", INFO_USD, tool_payload={"x": 1}, today=OGGI, negozio=NEG)
    assert "error" in s and "perimetro" in s["error"]
    # fonte con errore = rifiuto DICHIARATO, mai un canonico senza la fonte
    s2 = build_mnav_spec("TESORO", INFO_USD, tool_payload={"error": "sito giu'"}, today=OGGI, negozio=NEG)
    assert "error" in s2 and "RIFIUTATO" in s2["error"]
    # derived_mnav assente (prezzi live mancanti) = rifiuto
    s3 = build_mnav_spec("TESORO", INFO_USD, today=OGGI, negozio=NEG,
                         tool_payload={"latest": BTC_PAYLOAD["latest"],
                                       "mnav_error": "prezzo n.d."})
    assert "error" in s3 and "derived_mnav" in s3["error"]


def test_btc_mirror_e_fv_solo_con_target(NEG):
    # senza target: scheda completa, FV n.d. DICHIARATO (D2)
    s = build_mnav_spec("TESORO", INFO_USD, tool_payload=BTC_PAYLOAD, today=OGGI, negozio=NEG)
    assert "error" not in s
    assert s["kind"] == "dat_bitcoin" and s["profile_key"] == "dat_bitcoin"
    assert s["currency"] == "USD" and s["price"] == pytest.approx(200.0)
    mv = compute_mnav_values(s)
    assert mv["mnav_equity"] == pytest.approx(1.000)
    assert mv["mnav_ev"] == pytest.approx(1.325)
    assert mv["nav_per_share"] == pytest.approx(135.00)
    assert mv.get("fair_value_nav") is None
    assert "nessun nav_target" in mv["fv_note"]
    # con target 1,2 su mNAV EV: FV = (1,2x60 - 8 - 12 + 0,5) mld / 300M = 175,00
    s2 = build_mnav_spec("TESORO", INFO_USD, nav_target=1.2,
                         tool_payload=BTC_PAYLOAD, today=OGGI, negozio=NEG)
    mv2 = compute_mnav_values(s2)
    assert mv2["fair_value_nav"] == pytest.approx(175.00)


def test_btc_record_stale_rifiutato(NEG):
    p = dict(BTC_PAYLOAD)
    p["latest"] = dict(p["latest"], as_of_date="2026-05-01")
    p["derived_mnav"] = dict(p["derived_mnav"],
                             inputs=dict(p["derived_mnav"]["inputs"],
                                         record_as_of="2026-05-01"))
    # 01/05 -> 23/07 = 83 giorni > soglia: STALE = rifiuto dichiarato
    assert 83 > RECORD_STALE_DAYS_BTC
    s = build_mnav_spec("TESORO", INFO_USD, tool_payload=p, today=OGGI, negozio=NEG)
    assert "error" in s and "STALE" in s["error"]


def test_btc_campi_ev_assenti_dichiarati(NEG):
    # campo assente != zero (review 21/07 del tool): mNAV EV / NAV / FV n.d. dichiarati
    p = dict(BTC_PAYLOAD)
    p["latest"] = dict(p["latest"], debt=None)
    s = build_mnav_spec("TESORO", INFO_USD, nav_target=1.0, tool_payload=p, today=OGGI, negozio=NEG)
    assert "error" not in s and s["ev_missing"] == ["debt"]
    mv = compute_mnav_values(s)
    assert mv["mnav_equity"] == pytest.approx(1.000)   # l'equity resta calcolabile
    assert mv["mnav_ev"] is None and mv["nav_per_share"] is None
    assert mv.get("fair_value_nav") is None
    assert "assente" in mv["fv_note"] or "campi" in mv["fv_note"]
    assert any("assente != zero" in w for w in mv["warnings"])


def test_parita_mirror_tool_rotta_urla(NEG):
    # il tool dichiara un mNAV diverso dal mirror sugli stessi input = bug: WARN
    p = dict(BTC_PAYLOAD)
    p["derived_mnav"] = dict(p["derived_mnav"], mnav_equity_basic=1.10)
    s = build_mnav_spec("TESORO", INFO_USD, tool_payload=p, today=OGGI, negozio=NEG)
    mv = compute_mnav_values(s)
    assert any("PARITA' ROTTA" in w for w in mv["warnings"])
    from bellomberg.valuation.dcf_mnav import _mnav_sanity
    san = _mnav_sanity(s, mv)
    assert san["severity"] == "WARN" and "PARITA'" in san["headline"]


def test_hype_treasury_method_e_mirror(NEG):
    s = build_mnav_spec("HYPEX", INFO_USD, nav_target=1.0,
                        tool_payload=HYPE_PAYLOAD, today=OGGI, negozio=NEG)
    assert "error" not in s and s["kind"] == "dat_hype"
    mv = compute_mnav_values(s)
    assert mv["warrants_itm"] == 1                       # solo strike 5 < prezzo 12
    assert mv["fd_shares_m"] == pytest.approx(105.833, abs=0.001)
    assert mv["adjusted_nav_musd"] == pytest.approx(1197.0)
    assert mv["nav_per_share"] == pytest.approx(11.3102, abs=0.0001)
    assert mv["mnav"] == pytest.approx(1.061, abs=0.001)
    assert mv["mnav_dtl_addback"] == pytest.approx(1.008, abs=0.001)
    assert mv["fair_value_nav"] == pytest.approx(11.31, abs=0.01)
    # parita' col derived del tool: nessun warning di parita'
    assert not any("PARITA'" in w for w in mv["warnings"])


def test_hype_prezzo_mancante_rifiutato(NEG):
    p = dict(HYPE_PAYLOAD, purr_quote={"last": None, "fonte": "yfinance ko"})
    s = build_mnav_spec("HYPEX", INFO_USD, tool_payload=p, today=OGGI, negozio=NEG)
    assert "error" in s and "prezzo live mancante" in s["error"]
    # input di bilancio incompleti = rifiuto
    p2 = dict(HYPE_PAYLOAD, dat_fundamentals_musd={"hypeHeld": 30.0})
    s2 = build_mnav_spec("HYPEX", INFO_USD, tool_payload=p2, today=OGGI, negozio=NEG)
    assert "error" in s2 and "incompleti" in s2["error"]


def test_cef_fx_del_tool_e_sconto(NEG):
    s = build_mnav_spec("FONDO.L", INFO_USD, nav_target=0.80,
                        tool_payload=CEF_PAYLOAD, today=OGGI, negozio=NEG)
    assert "error" not in s and s["kind"] == "cef_nav"
    # prezzo GIA' convertito dal tool (stessa misura, mai un secondo fetch)
    assert s["price"] == pytest.approx(56.28)
    assert s["fx_rate"] == pytest.approx(1.34)           # implicito: 56,28/(4200/100)
    mv = compute_mnav_values(s)
    assert mv["discount_to_nav_pct"] == pytest.approx(-29.65, abs=0.01)
    assert mv["fair_value_nav"] == pytest.approx(64.00)
    assert not any("PARITA'" in w for w in mv["warnings"])


def test_cef_nav_stale_rifiutato(NEG):
    p = dict(CEF_PAYLOAD, nav_staleness="STALE: NAV di 20 giorni fa (soglia 14g)")
    s = build_mnav_spec("FONDO.L", INFO_USD, tool_payload=p, today=OGGI, negozio=NEG)
    assert "error" in s and "STALE" in s["error"]
    # prezzo convertito assente (FX giu') = sconto non calcolabile = rifiuto
    p2 = {k: v for k, v in CEF_PAYLOAD.items() if k != "market_price_in_nav_ccy"}
    p2["discount"] = "n.d. (FX GBPUSD non disponibile)"
    s2 = build_mnav_spec("FONDO.L", INFO_USD, tool_payload=p2, today=OGGI, negozio=NEG)
    assert "error" in s2 and "RIFIUTATO" in s2["error"]


def test_target_bande_warn_mai_block(NEG):
    from bellomberg.valuation.dcf_mnav import _mnav_sanity
    # target legale ma fuori banda di buon senso: WARN dichiarato, MAI BLOCK (D2)
    s = build_mnav_spec("FONDO.L", INFO_USD, nav_target=0.20,
                        tool_payload=CEF_PAYLOAD, today=OGGI, negozio=NEG)
    mv = compute_mnav_values(s)
    assert mv["fair_value_nav"] == pytest.approx(16.00)  # 80 x 0,2: view estrema ma view
    assert any("banda di buon senso" in w for w in mv["warnings"])
    san = _mnav_sanity(s, mv)
    assert san["severity"] == "WARN"
    assert san["exclude_from_action_table"] is False     # mai esclusione automatica
    # fuori hard bound: input rotto -> FV n.d. dichiarato
    s2 = build_mnav_spec("FONDO.L", INFO_USD, nav_target=5.0,
                         tool_payload=CEF_PAYLOAD, today=OGGI, negozio=NEG)
    mv2 = compute_mnav_values(s2)
    assert mv2.get("fair_value_nav") is None
    assert any("hard bound" in w for w in mv2["warnings"])
    # non numerico: idem, dichiarato
    s3 = build_mnav_spec("FONDO.L", INFO_USD, nav_target="boh",
                         tool_payload=CEF_PAYLOAD, today=OGGI, negozio=NEG)
    assert compute_mnav_values(s3).get("fair_value_nav") is None


def _labels_col_a(ws):
    return {str(c.value): c.row for row in ws.iter_rows(min_col=1, max_col=1)
            for c in row if c.value is not None}


def test_workbook_btc_offline(tmp_path, NEG):
    from openpyxl import load_workbook
    s = build_mnav_spec("TESORO", INFO_USD, nav_target=1.2, variant_view="test",
                        tool_payload=BTC_PAYLOAD, today=OGGI, negozio=NEG)
    out = str(tmp_path / "VAL_TESORO.xlsx")
    r = build_mnav_model(s, out)
    assert r["ok"] is True and r["engine"] == "mnav"
    assert r["payload_currency"] == "USD"
    assert r["mnav_equity"] == pytest.approx(1.000)
    assert r["mnav_ev"] == pytest.approx(1.325)
    assert r["fair_value_nav"] == pytest.approx(175.00)
    # alias per la catena tesi/sanity storica: stessa cifra, dichiarato
    assert r["fair_value_base"] == r["fair_value_nav"]
    assert "alias" in r["fv_alias_note"]
    assert r["sanity"]["severity"] == "OK"
    wb = load_workbook(out)
    assert {"Thesis & Assumptions", "Fonti & Vintage", "Sensitivity"} <= set(wb.sheetnames)
    ws = wb["Thesis & Assumptions"]
    # formule VIVE nei punti chiave (mai valori morti)
    assert ws["B12"].value == "=B10*B11"                 # BTC-NAV
    assert ws["B21"].value == "=B15/B12"                 # mNAV equity
    assert ws["B22"].value == "=(B15+B16+B17-B18)/B12"   # mNAV EV
    lbl = _labels_col_a(ws)
    t_row = lbl["Target premio/sconto (ANALISTA)"]
    fv_row = lbl["FAIR VALUE per azione"]
    assert str(ws[f"B{fv_row}"].value) == f"=(B{t_row}*B12-B16-B17+B18)/B13"
    # Sensitivity: shock vivi dal Thesis, FV riga col target VIVO (mai cablato)
    sen = wb["Sensitivity"]
    assert str(sen["C4"].value).startswith("='Thesis & Assumptions'!B11*")
    assert f"B{t_row}" in str(sen["C8"].value)


def test_workbook_btc_senza_target_e_campi_assenti(tmp_path, NEG):
    from openpyxl import load_workbook
    p = dict(BTC_PAYLOAD)
    p["latest"] = dict(p["latest"], pref=None)
    s = build_mnav_spec("TESORO", INFO_USD, tool_payload=p, today=OGGI, negozio=NEG)
    out = str(tmp_path / "VAL_TESORO.xlsx")
    r = build_mnav_model(s, out)
    assert r["ok"] is True
    assert r.get("fair_value_nav") is None and r.get("fair_value_base") is None
    assert r["sanity"]["status"] == "n/d"
    # review V5 M2: mai severity "OK" senza fair value (badge verde su niente)
    assert r["sanity"]["severity"] is None
    wb = load_workbook(out)
    ws = wb["Thesis & Assumptions"]
    lbl = _labels_col_a(ws)
    assert any(k.startswith("mNAV EV: n.d. DICHIARATO") for k in lbl)
    assert any(k.startswith("FV n.d. DICHIARATO") for k in lbl)
    # Sensitivity senza campi EV: riga n.d. dichiarata, niente formule rotte
    sen = wb["Sensitivity"]
    lbl_s = _labels_col_a(sen)
    assert any("n.d. DICHIARATO" in k for k in lbl_s)


@pytest.mark.parametrize("language,anav_label,mnav_prefix,target_label", [
    ("it", "NAV rettificato ($M)", "mNAV = prezzo", "Target premio/sconto (ANALISTA)"),
    ("en", "Adjusted NAV ($M)", "mNAV = share price", "Target premium/discount (ANALYST)"),
])
def test_workbook_hype_offline(tmp_path, NEG, language, anav_label, mnav_prefix, target_label):
    from openpyxl import load_workbook
    s = build_mnav_spec("HYPEX", INFO_USD, nav_target=1.0, variant_view="test",
                        tool_payload=HYPE_PAYLOAD, today=OGGI, negozio=NEG)
    out = str(tmp_path / "VAL_HYPEX.xlsx")
    r = build_mnav_model(s, out, language=language)
    assert r["ok"] is True
    assert r["mnav"] == pytest.approx(1.061, abs=0.001)
    assert r["fair_value_nav"] == pytest.approx(11.31, abs=0.01)
    wb = load_workbook(out)
    ws = wb["Thesis & Assumptions"]
    lbl = _labels_col_a(ws)
    # formula del sito in celle VIVE: mNAV e Adjusted NAV sono formule
    mnav_row = next(r for k, r in lbl.items() if k.startswith(mnav_prefix))
    assert str(ws[f"B{mnav_row}"].value).startswith("=")
    anav_row = lbl[anav_label]
    assert str(ws[f"B{anav_row}"].value).startswith("=")
    # warrant a treasury method: formula IF ITM (mai stringhe vuote nei rami)
    w_rows = [row for lbl_, row in lbl.items() if lbl_.startswith("warrant ")]
    assert w_rows and all(str(ws[f"D{w}"].value).startswith("=IF(") for w in w_rows)
    assert '"' not in str(ws[f"D{w_rows[0]}"].value)     # rami numerici, mai ""
    sen = wb["Sensitivity"]
    assert str(sen["C4"].value).startswith("='Thesis & Assumptions'!")
    # FV in sensitivity: riferimento VIVO alla cella target del Thesis
    t_row = lbl[target_label]
    assert f"B{t_row}" in str(sen["C10"].value)


def test_hype_zero_warrant_e_anav_degenerato(tmp_path, NEG):
    from openpyxl import load_workbook
    # 0 tranche warrant: FD = basic, dichiarato (nessuna formula rotta)
    p = dict(HYPE_PAYLOAD, warrants=[])
    s = build_mnav_spec("HYPEX", INFO_USD, tool_payload=p, today=OGGI, negozio=NEG)
    mv = compute_mnav_values(s)
    assert mv["fd_shares_m"] == pytest.approx(100.0)
    # ANAV negativo (review V5 M3): adj_nav = -3000-20+30-50-900+1200-63 = -2803
    # -> mirror n.d. E il foglio NON mostra un mNAV negativo vivo come headline
    fund = dict(HYPE_PAYLOAD["dat_fundamentals_musd"], bookNAV=-3000.0)
    p2 = dict(HYPE_PAYLOAD, dat_fundamentals_musd=fund, derived={})
    s2 = build_mnav_spec("HYPEX", INFO_USD, nav_target=1.0, tool_payload=p2, today=OGGI, negozio=NEG)
    mv2 = compute_mnav_values(s2)
    assert mv2["adjusted_nav_musd"] == pytest.approx(-2803.0)
    assert mv2["mnav"] is None and mv2.get("fair_value_nav") is None
    out = str(tmp_path / "VAL_HYPEX.xlsx")
    r = build_mnav_model(s2, out)
    assert r["ok"] is True and r.get("mnav") is None
    ws = load_workbook(out)["Thesis & Assumptions"]
    lbl = _labels_col_a(ws)
    assert any(k.startswith("mNAV: n.d. DICHIARATO") for k in lbl)


def test_hype_parita_non_verificabile_dichiarata(NEG):
    # review V5 B3: derived assente nel tool = cross-check impossibile, dichiarato
    p = dict(HYPE_PAYLOAD, derived={})
    s = build_mnav_spec("HYPEX", INFO_USD, tool_payload=p, today=OGGI, negozio=NEG)
    mv = compute_mnav_values(s)
    assert mv["mnav"] is not None                        # il mirror calcola comunque
    assert any("NON verificabile" in w for w in mv["warnings"])


def test_workbook_cef_offline(tmp_path, NEG):
    from openpyxl import load_workbook
    s = build_mnav_spec("FONDO.L", INFO_USD, nav_target=0.80, variant_view="test",
                        tool_payload=CEF_PAYLOAD, today=OGGI, negozio=NEG)
    out = str(tmp_path / "VAL_FONDO_L.xlsx")
    r = build_mnav_model(s, out)
    assert r["ok"] is True
    assert r["discount_to_nav_pct"] == pytest.approx(-29.65, abs=0.01)
    assert r["fair_value_nav"] == pytest.approx(64.00)
    assert r["price_quote"] == pytest.approx(4200.0)
    assert r["price_quote_currency"] == "GBp"
    wb = load_workbook(out)
    ws = wb["Thesis & Assumptions"]
    # conversione GBp VIVA in cella col tasso del tool (stessa misura del payload)
    assert ws["B13"].value == "=B11/100*B12"
    assert ws["B14"].value == "=B13/B10-1"               # sconto vivo
    lbl = _labels_col_a(ws)
    fv_row = lbl["FAIR VALUE per azione"]
    t_row = lbl["Target premio/sconto (ANALISTA)"]
    assert str(ws[f"B{fv_row}"].value) == f"=B10*B{t_row}"
    sen = wb["Sensitivity"]
    assert str(sen["B4"].value).startswith("='Thesis & Assumptions'!B10*")
    assert str(sen["B5"].value).endswith("-1")
    # Fonti & Vintage: la fonte del NAV DEL PAYLOAD e il vintage in chiaro
    fon = wb["Fonti & Vintage"]
    vals = [str(c.value) for row in fon.iter_rows() for c in row if c.value]
    assert any("sito-finto.example" in v for v in vals)


def test_prepare_canonical_path(tmp_path):
    # helper condiviso operating/banca/rab/mnav: nome canonico senza timestamp,
    # punto del ticker -> underscore (nessun file da spostare in questo test)
    from bellomberg.valuation.dcf_engine import _prepare_canonical_path
    p = _prepare_canonical_path(str(tmp_path), "FONDO.L")
    assert p.endswith("VAL_FONDO_L.xlsx")
    assert str(tmp_path) in p
