"""Synthetic managed-care linked model across fiscal calendar and legal ledger."""
from copy import deepcopy
from tempfile import TemporaryDirectory

from openpyxl import Workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
import pytest

from bellomberg.valuation.managed_care_formulas import apply_managed_care_formulas
from test_managed_care_integration import generate, make_bundle, records_for


def _payload(**kwargs):
    with TemporaryDirectory() as directory:
        payload = generate(make_bundle(**kwargs), directory)
    assert payload['valuation_usability']['usable'], payload.get('error')
    return payload


@pytest.mark.parametrize('kind,non_solar,count', [
    ('FY', False, 1), ('interim', False, 3),
    ('FY', True, 6), ('interim', True, 2),
])
def test_calendar_earnings_capital_and_terminal_link_for_every_horizon(kind, non_solar, count):
    payload = _payload(kind=kind, non_solar=non_solar, count=count)
    original = deepcopy(payload)
    wb = Workbook()
    assert apply_managed_care_formulas(wb, payload)
    assert payload == original
    for scenario in ('bear', 'base', 'bull'):
        ws = wb['Care ' + scenario]
        assert ws['D9'].data_type == 'f'  # first segment premium
        assert ws['D10'].data_type == 'f'  # first segment medical cost
        statutory_row = next(row for row in range(1, ws.max_row + 1)
                             if 'statutory income' in str(ws.cell(row, 2).value))
        assert ws.cell(statutory_row, 4).data_type == 'f'
        assert ws.cell(7, count + 3).value is not None
        last = ws.print_area.split(':')[-1].replace('$', '')
        assert column_index_from_string(coordinate_from_string(last)[0]) >= count + 3
        assert len(wb._model_link['calls'][scenario]) == count
        assert payload['managed_care']['scenarios'][scenario]['fair_value_per_share'] > 0
    assert wb['Summary']['E8'].data_type == 'f'
    assert wb['Model Checks']['D9'].data_type == 'f'


def test_coordinated_earnings_and_legal_cash_change_matches_engine():
    records, context = records_for()
    for record in records:
        if record['scenario'] != 'base':
            continue
        if record['driver'] == 'other_revenue':
            record['value'][0] = 2.
        elif record['driver'] in ('capital.subsidiaries.0.gaap_net_income',
                                  'capital.subsidiaries.0.proposed_distribution'):
            record['value'][0] = 3.9
    payload = _payload(records=records, context=context)
    assert payload['fair_value_base'] == pytest.approx(10.0265, abs=.01)
    wb = Workbook()
    assert apply_managed_care_formulas(wb, payload)
    assert wb['Care base']['D9'].data_type == 'f'


def test_unusable_snapshot_has_no_linked_sheets():
    payload = _payload()
    payload['valuation_usability']['usable'] = False
    wb = Workbook()
    assert not apply_managed_care_formulas(wb, payload)
    assert wb.sheetnames == ['Sheet']
