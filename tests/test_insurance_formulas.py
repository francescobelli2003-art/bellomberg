"""Insurance workbooks preserve the engine baseline and link product to capital."""
from copy import deepcopy
from tempfile import TemporaryDirectory

from openpyxl import Workbook
import pytest

from bellomberg.valuation.dcf_engine import generate_valuation
from bellomberg.valuation.insurance_formulas import apply_insurance_formulas
from test_insurance_valuation import insurance_bundle, pc_records
from test_insurance_life import life_bundle, life_records


def _value(bundle):
    with TemporaryDirectory() as directory:
        result = generate_valuation('SYNTH-INS', prepared_bundle=bundle, output_dir=directory)
    assert result['valuation_usability']['usable'], result.get('error')
    return result


@pytest.mark.parametrize('bundle,name,expected', [
    (insurance_bundle, 'PC', 25.0),
    (life_bundle, 'Life', 10.33),
])
def test_insurance_product_book_cash_capital_and_terminal_are_linked(bundle, name, expected):
    payload = _value(bundle())
    original = deepcopy(payload)
    wb = Workbook()
    assert apply_insurance_formulas(wb, payload)
    assert payload == original
    assert payload['fair_value_base'] == pytest.approx(expected)
    ws = wb[name + ' base']
    assert ws['D9'].data_type == 'f'  # investment income or exposed policies
    assert ws['D15'].data_type == 'f'  # product economics
    assert ws['D57'].data_type == 'f'  # legal capital schedule
    assert ws['F58'].data_type == 'f'  # continuing legal capital schedule
    assert wb['Summary']['E8'].data_type == 'f'
    assert wb['Model Checks']['D9'].data_type == 'f'
    assert ('$G$' if name == 'Life' else '$F$') in ws.print_area
    inputs = wb['Model Inputs']
    labels = [str(inputs.cell(row, 2).value) for row in range(1, inputs.max_row + 1)]
    assert 'capital.terminal equity' in labels
    assert 'capital.subsidiaries.0.gaap net income' in labels


def test_pc_operating_driver_change_reconciles_to_independent_engine_value():
    rows = pc_records()
    for record in rows:
        if record['scenario'] != 'base':
            continue
        name = record['driver']
        if name == 'insurance.0.investment_yield':
            record['value'] = [.06, .06]
        elif name == 'continuing_economics':
            record['value']['LEGAL-A']['investment_yield'] = .06
        elif name in ('capital.subsidiaries.0.proposed_distribution',
                      'capital.subsidiaries.0.permitted_distribution'):
            record['value'] = [26., 26.]
        elif name == 'capital.terminal_equity':
            record['value'] = 260.
        elif name == 'terminal_ledger':
            sub = record['value']['capital']['subsidiaries'][0]
            sub['proposed_distribution'] = [26.]
            sub['permitted_distribution'] = [26.]
            sub['gaap_net_income'] = [26.]
            record['value']['liquidity_bridge']['LEGAL-A']['operating_cash'] = [26.]
            sub['liquidity_before_transfers'] = [36.]
        elif name == 'liquidity_bridge':
            record['value']['LEGAL-A']['operating_cash'] = [26., 26.]
        elif name == 'consolidated_income':
            record['value'] = [26., 26.]
        elif name == 'terminal_income':
            record['value'] = 26.
        elif name == 'capital.subsidiaries.0.gaap_net_income':
            record['value'] = [26., 26.]
        elif name == 'capital.subsidiaries.0.liquidity_before_transfers':
            record['value'] = [36., 36.]
    payload = _value(insurance_bundle(rows))
    assert payload['fair_value_base'] == 26.0
    wb = Workbook()
    assert apply_insurance_formulas(wb, payload)


def test_unusable_insurance_payload_has_no_live_sheets():
    payload = _value(life_bundle())
    payload['valuation_usability']['usable'] = False
    wb = Workbook()
    assert not apply_insurance_formulas(wb, payload)
    assert wb.sheetnames == ['Sheet']
