"""A holding workbook embeds live child schedules without external links."""
from openpyxl import load_workbook
from test_sotp_documented import generate
from bellomberg.valuation.sotp_formulas import _translate


def test_sotp_formula_references_stay_internal_and_sources_literal(tmp_path):
    payload = generate(tmp_path)
    assert payload['valuation_usability']['usable'], payload.get('error')
    wb = load_workbook(payload['path'])
    assert wb['SOTP base']['D12'].data_type == 'f'
    assert wb['C1-Bank base']['D9'].data_type == 'f'
    assert wb['C2-Property base']['D10'].data_type == 'f'
    assert wb['Model Checks']['D9'].data_type == 'f'
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    assert '[' not in cell.value, (ws.title, cell.coordinate)
    assert any("'C2-" in cell.value for row in wb['SOTP base'] for cell in row if cell.data_type=='f')
    wb.close()


def test_formula_translation_preserves_string_constants_and_escaped_names():
    expression = 'IF(\'O\'\'Brien\'!D9="literal!keep",SUM(D8:D9),\'Model Inputs\'!D8)'
    result = _translate(expression, {"O'Brien":'C1-Cash','Model Inputs':'C1-Inputs'})
    assert result == 'IF(\'C1-Cash\'!D9="literal!keep",SUM(D8:D9),\'C1-Inputs\'!D8)'


def test_parent_debt_settlement_premium_remains_an_independent_input(tmp_path):
    from test_sotp_documented import sotp_records
    rows=sotp_records()
    for r in rows:
        if r['driver']=='claims':
            next(c for c in r['value']['items'] if c['allocation']=='PARENT' and c['role']=='debt')['amount']=25.
    result=generate(tmp_path, rows)
    assert result['valuation_usability']['usable'],result.get('error')
    assert result['calculation_details']['scenarios']['base']['parent_claims']['debt']==25.
    wb=load_workbook(result['path'])
    inputs=wb['Model Inputs']
    cells=[row[3] for row in inputs if row[1].value and str(row[1].value).endswith(' / amount') and row[3].value==25.]
    assert len(cells)==3 and all(c.data_type=='n' for c in cells)
    wb.close()
