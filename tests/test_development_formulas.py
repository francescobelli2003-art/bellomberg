"""Finite branch workbook matches documented rNPV and original-holder financing."""
from copy import deepcopy
from tempfile import TemporaryDirectory

from openpyxl import Workbook
import pytest

from bellomberg.valuation.dcf_engine import generate_valuation
from bellomberg.valuation.development_formulas import apply_development_formulas
from test_development_valuation import development_bundle, development_records


def _case(rows):
    with TemporaryDirectory() as directory:
        payload = generate_valuation('SYNTH-INS', prepared_bundle=development_bundle(rows),
                                     output_dir=directory)
    assert payload['valuation_usability']['usable'], payload.get('error')
    return payload


@pytest.mark.parametrize('probabilities,market,expected', [
    ((.5, .4), 2., 2.23),
    ((0., .4), 1., .91),
    ((1., 1.), 10., 11.03),
])
def test_probability_extremes_and_intermediate_have_live_finite_branches(probabilities, market, expected):
    rows = development_records(probabilities=probabilities)
    next(row for row in rows if row['driver'] == 'quotation')['value']['price'] = market
    payload = _case(rows)
    baseline = deepcopy(payload)
    wb = Workbook()
    assert apply_development_formulas(wb, payload)
    assert payload == baseline
    assert payload['fair_value_base'] == pytest.approx(expected)
    for scenario in ('bear', 'base', 'bull'):
        assert all(name in wb for name in ('Dev ' + scenario, 'Dev ' + scenario + ' 1',
                                           'Dev ' + scenario + ' 2', 'Dev ' + scenario + ' S'))
        assert len(wb._model_link['calls'][scenario]) == 14
        assert wb['Dev ' + scenario]['D19'].data_type == 'f'
        assert wb['Dev ' + scenario + ' S']['E22'].data_type == 'f'  # issued/closing shares
    assert wb['Summary']['E8'].data_type == 'f'
    assert wb['Model Checks']['D9'].data_type == 'f'


@pytest.mark.parametrize('price,expected,closing_shares', [(2., 2.09, 22.5), (4., 2.51, 16.25)])
def test_fixed_issue_price_dilutes_original_holders(price, expected, closing_shares):
    payload = _case(development_records(outside_price=price))
    outcome = payload['calculation_details']['scenarios']['base']['outcomes'][-1]
    assert payload['fair_value_base'] == pytest.approx(expected)
    assert outcome['ledger'][1]['closing_shares'] == pytest.approx(closing_shares)
    wb = Workbook()
    assert apply_development_formulas(wb, payload)
    assert wb['Dev base S']['E22'].data_type == 'f'
    assert wb['Dev base S']['E20'].data_type == 'f'  # original-holder start cash


def test_unreachable_branch_does_not_require_funding():
    rows = development_records(probabilities=(0., .4))
    for row in rows:
        if row['driver'] == 'quotation':
            row['value']['price'] = 1.
        if row['driver'] == 'minimum_cash':
            row['value'] = [5.] * 4
        if row['driver'] == 'funding':
            row['value'] = {'capacities': [0.] * 4, 'commitments': {}}
    payload = _case(rows)
    assert payload['fair_value_base'] == .86
    wb = Workbook()
    assert apply_development_formulas(wb, payload)
    checks = wb['Input Checks']
    assert any('committed funding' in str(checks.cell(row, 2).value)
               for row in range(1, checks.max_row + 1))


def test_unusable_payload_does_not_offer_a_live_tree():
    payload = _case(development_records())
    payload['valuation_usability']['usable'] = False
    wb = Workbook()
    assert not apply_development_formulas(wb, payload)
    assert wb.sheetnames == ['Sheet']


def test_six_stages_keep_probabilities_separate_from_values():
    rows = development_records(probabilities=(1.,1.))
    periods = [dict(start=f'{y}-01-01',end=f'{y}-12-31') for y in range(2026,2034)]
    span = '|'.join(p['start']+'/'+p['end'] for p in periods)
    names = [f'STAGE-{i}' for i in range(1,7)]
    for r in rows:
        if '|' in r['period']: r['period']=span
        key=r['driver']
        if key=='calendar': r['value']['periods']=periods
        elif key=='quotation': r['value']['price']=10.
        elif key=='asset': r['value']['stage_order']=names
        elif key=='life':
            for k in ('patent_expiry','contract_expiry','economic_expiry'): r['value'][k]='2033-12-31'
        elif key=='stage_schedule':
            r['value']={name:dict(period=i+1,cost=1.,success_received=1.,success_paid=0.,failure_close_cost=0.) for i,name in enumerate(names)}
        elif key=='probabilities': r['value']={name:1. for name in names}
        elif key=='commercial':
            r['value']=dict(units=[0.]*6+[10.,10.],net_price_per_unit=[100.]*8,incoming_royalty_rate=[.1]*8,
                            outgoing_royalty_rate=[0.]*8,cash_operating_cost=[0.]*8)
        elif key in ('central_cash_cost','expected_tax_shield'): r['value']=[0.]*8
        elif key=='minimum_cash': r['value']=[5.]*8
        elif key=='tax_rate': r['value']=[.2]*8
        elif key=='funding': r['value']={'capacities':[0.]+[2.]*5+[0.,0.],
            'commitments':{str(i):{'id':f'Synthetic staged funding {i}','available_date':f'{2026+i}-01-01',
                          'terms':'draw_when_needed_existing_shareholders_no_new_shares'} for i in range(1,6)}}
    payload = _case(rows)
    wb=Workbook()
    assert apply_development_formulas(wb,payload)
    ws=wb['Dev base']
    assert all(ws.cell(row,4).data_type=='f' for row in range(9,16))
    assert str(ws['D15'].value).startswith('=IFERROR(')
    assert ws['D27'].data_type=='f'  # result follows all reach rows
    assert len(wb._model_link['calls']['base'])==2*(sum(range(1,7))+8)
