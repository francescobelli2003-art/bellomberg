"""Synthetic current-quote evidence must survive delivery without replacing model price."""
import json
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

import pytest

from bellomberg.valuation import dcf_engine, sector_analysis


def quote():
    return {"contract": "market_quote/1", "status": "ok", "source_id": "synthetic",
            "source_status": "ok", "acquired_as_of": "2026-09-18", "information_cutoff": "2026-09-18",
            "symbol": "SYNTH", "info_symbol": "SYNTH", "exchange": "Synthetic market",
            "exchange_timezone": "Europe/Berlin", "freshness_policy": "previous_weekday; holidays_not_modelled",
            "quote_source_name": "Synthetic", "delayed_minutes": 0, "currency": "EUR",
            "price": 12., "observed_at": "2026-09-18T15:30:00Z", "observed_local_date": "2026-09-18",
            "price_model": 10., "price_model_as_of": "2026-06-30",
            "price_move_since_valuation_pct": 20., "fv_basis": "valuation_date_no_rollforward",
            "sanity_basis": "price_model", "upside_bear_pct": 0., "upside_base_pct": 25.,
            "upside_bull_pct": 50., "message": "Synthetic evidence"}


def test_sidecar_preserves_both_upside_bases(tmp_path):
    workbook = tmp_path / "VAL_SYNTH.xlsx"
    workbook.write_bytes(b"synthetic workbook placeholder; this test only checks sidecar fields")
    payload = {"path": str(workbook), "ticker": "SYNTH", "price": 10., "upside_pct": 50.,
               "market_quote": quote(), "sanity": {"severity": "OK"}}
    dcf_engine._write_payload_sidecar(payload)
    saved = json.loads(workbook.with_suffix(".payload.json").read_text(encoding="utf-8"))
    assert saved.get("market_quote") == quote()
    assert saved.get("upside_pct") == 50.
    assert saved["price"] == 10.


def test_committee_row_declares_historical_price_even_without_current_quote(monkeypatch):
    from bellomberg.valuation import dcf_quality
    payload = {"price": 10., "upside_pct": 50., "valuation_date": "2026-06-30",
               "fair_value_base": 15., "valuation_usability": {"usable": True}}
    monkeypatch.setattr(dcf_quality, "normalize_valuation_payload", lambda p: p)
    text = sector_analysis.valuation_results_block({"SYNTH": payload})
    row = json.loads(text.splitlines()[-1])
    assert row.get("price_model") == 10.
    assert row.get("price_model_as_of") == "2026-06-30"
    assert row.get("upside_model_pct") == 50.
    assert row.get("market_quote", {}).get("status_at_read") == "data_missing"
    assert row.get("upside_today_pct") is None


def test_workbook_quote_rows_do_not_shift_existing_value():
    from openpyxl import Workbook
    from bellomberg.reporting.valuation_quote import append_market_quote_rows
    wb = Workbook()
    ws = wb.active
    ws.append(["Model", "SYNTH"])
    ws.append(["Fair value", 15.])
    append_market_quote_rows(ws, quote(), usable=True)
    assert ws["B2"].value == 15.
    cells = list(ws.values)
    assert any(row[1] == 12. for row in cells)
    assert any(row[1] == 25. for row in cells)
    assert any(row[1] == "2026-06-30" for row in cells)
    assert any(row[1] == "2026-09-18T15:30:00Z" for row in cells)
    wb.close()


def test_report_declares_unrepresentable_legacy_price_without_breaking_delivery():
    from bellomberg.reporting.valuation_quote import quote_comparison_text
    from bellomberg.core.language import language_context
    with language_context("en"):
        rendered = quote_comparison_text({"price": 10 ** 1000, "valuation_usability": {"usable": False}})
    assert "model price n/a" in rendered
    assert "observed price n/a" in rendered


def calculated_payload(tmp_path):
    from test_sector_operating_drivers import DAY, operating_records
    info = {"symbol": "SYNTH", "currency": "EUR", "regularMarketPrice": 12.,
            "regularMarketTime": datetime(2026, 9, 10, 15, 30, tzinfo=timezone.utc).timestamp(),
            "exchangeTimezoneName": "Europe/Berlin", "fullExchangeName": "Synthetic exchange",
            "quoteSourceName": "Synthetic quote", "exchangeDataDelayedBy": 15}
    bundle = sector_analysis.prepare_sector_analysis("SYNTH", as_of=DAY, providers={
        "profile": lambda *a, **k: {"status": "ok", "source_id": "synthetic", "as_of": DAY,
            "data": {"info": info, "evidence": [{"field": key, "value": value,
                "source_id": "synthetic", "as_of": DAY} for key, value in
                (("instrument", "equity"), ("business_model", "manufacturing"))]}},
        "method_inputs": lambda *a, **k: {"status": "ok", "source_id": "synthetic", "as_of": DAY,
            "records": operating_records()}},
        user_context={"analysis_context": {"scenario_rationale": {
            key: "Synthetic comparison of independently dated prices" for key in ("bear", "base", "bull")}}})
    return dcf_engine.generate_valuation("SYNTH", prepared_bundle=bundle, output_dir=str(tmp_path))


@pytest.mark.parametrize("read_day,status,current_upside", [
    ("2026-09-10", "ok", 17.8), ("2026-09-19", "stale", None)])
def test_real_workbook_sidecar_api_committee_and_email_keep_two_dates(tmp_path, monkeypatch, read_day, status, current_upside):
    from bellomberg.valuation import market_quote
    from bellomberg.reporting.email_sender import corpo_valutazioni
    from bellomberg.core.language import language_context
    from test_sector_valuation_api import endpoint
    from openpyxl import load_workbook
    original_view = market_quote.market_quote_view
    monkeypatch.setattr(market_quote, "market_quote_view", lambda block, *, as_of=None, usable=True:
        original_view(block, as_of=as_of or read_day, usable=usable))
    payload = calculated_payload(tmp_path)
    assert payload["valuation_usability"]["usable"], payload["valuation_usability"]
    before = deepcopy(payload)
    assert payload["fair_value_base"] == 14.13
    assert payload["price"] == 10.
    assert payload["upside_pct"] == 41.3
    assert payload["market_quote"]["upside_base_pct"] == 17.8
    sidecar = Path(payload["path"]).with_suffix(".payload.json")
    saved = json.loads(sidecar.read_text(encoding="utf-8"))
    assert saved["market_quote"] == payload["market_quote"]
    assert saved["upside_pct"] == 41.3
    model = endpoint(tmp_path)["models"][0]
    assert model["fair_value"] == 14.13
    assert model["price_at_thesis"] == 10.
    assert model["price_model_as_of"] == "2025-12-31"
    assert model["upside_pct"] == 41.3
    assert model["market_quote"]["status_at_read"] == status
    assert model["upside_today_pct"] == current_upside
    assert model["detail"]["market_quote"]["upside_base_pct"] == 17.8
    row = json.loads(sector_analysis.valuation_results_block({"SYNTH": payload}).splitlines()[-1])
    assert row["upside_today_pct"] == current_upside
    assert row["upside_model_pct"] == 41.3
    with language_context("en"):
        mail = corpo_valutazioni({"SYNTH": payload}, [payload["path"]])
    assert "model upside 41.3%" in mail
    assert "Historical FV without rollforward" in mail
    assert "2026-09-10T15:30:00" in mail
    assert ("observed price 17.8%" in mail) is (status == "ok")
    workbook = load_workbook(payload["path"], data_only=True)
    try:
        assert workbook["Valuation"]["B2"].value == 14.13
        assert any(row[1] == 17.8 for row in workbook["Valuation"].values)
    finally:
        workbook.close()
    assert payload == before
    assert json.loads(sidecar.read_text(encoding="utf-8")) == saved
