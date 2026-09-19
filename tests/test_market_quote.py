"""Observed quotes are snapshot evidence, never a replacement for model-date prices."""
from copy import deepcopy
from datetime import datetime, timezone
from importlib import import_module
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bellomberg.valuation import dcf_engine, sector_analysis
from bellomberg.valuation.dcf_quality import normalize_valuation_payload
from test_sector_operating_drivers import bundle_for, operating_records


DAY = '2026-09-10'
SCENARIOS = ('bear', 'base', 'bull')


def api():
    return import_module('bellomberg.valuation.market_quote')


def timestamp(day=DAY, hour=15):
    return datetime.fromisoformat(day + f'T{hour:02}:30:00+00:00').timestamp()


def quote_info(**changes):
    return dict(symbol='SYNTH-EXT', regularMarketPrice=30., regularMarketTime=timestamp(),
                exchangeTimezoneName='Europe/Berlin', currency='EUR', fullExchangeName='Synthetic venue',
                quoteSourceName='Delayed Quote', exchangeDataDelayedBy=15, **changes)


def fixture():
    quotation = next(r['value'] for r in operating_records() if r['driver'] == 'quotation')
    bundle = {'case': {'ticker': 'SYNTH-EXT', 'as_of': DAY, 'info': quote_info(),
                      'sources': {'profile': {'status': 'ok', 'source_id': 'synthetic-profile', 'as_of': DAY}}}}
    return bundle, quotation, {s: 14.13 for s in SCENARIOS}


def acquired_bundle(bundle=None, info=None):
    bundle = bundle_for() if bundle is None else bundle
    sources = deepcopy(bundle['case']['sources'])
    sources['profile']['data']['info'] = quote_info() if info is None else info
    providers = {key: (lambda *a, value=value, **k: deepcopy(value)) for key, value in sources.items()}
    return sector_analysis.prepare_sector_analysis(bundle['case']['ticker'], as_of=DAY, providers=providers,
        user_context={'analysis_context': bundle['analysis_context']})


def generate(tmp_path, *, records=None, info=None):
    bundle = acquired_bundle(bundle_for(records=records), info)
    return dcf_engine.generate_valuation('SYNTH-EXT', prepared_bundle=bundle, output_dir=str(tmp_path))


def test_observed_price_is_additive_and_sanity_keeps_model_price(tmp_path):
    result = generate(tmp_path)
    assert result['valuation_usability']['usable'], result['valuation_usability']
    assert result['price'] == 10.
    assert result['fair_value_base'] == 14.13
    assert result['sanity']['severity'] == 'WARN'
    assert result['upside_pct'] == 41.3
    assert isinstance(result.get('market_quote'), dict)
    block = result['market_quote']
    assert block['price'] == 30. and block['upside_base_pct'] == -52.9
    assert block['price_move_since_valuation_pct'] == 200.
    assert block['status'] == 'ok' and block['sanity_basis'] == 'price_model'
    assert block['fv_basis'] == 'valuation_date_no_rollforward'
    sidecar = json.loads(Path(result['path']).with_suffix('.payload.json').read_text(encoding='utf-8'))
    assert sidecar['market_quote'] == block
    workbook = load_workbook(result['path'], data_only=True)
    assert workbook['Valuation']['B2'].value == 14.13
    assert -52.9 in [row[1] for row in workbook['Valuation'].values if len(row) > 1]
    workbook.close()


@pytest.mark.parametrize('field,value,status', [
    ('regularMarketPrice', None, 'data_missing'), ('regularMarketPrice', True, 'data_missing'),
    ('regularMarketPrice', float('inf'), 'data_missing'), ('regularMarketPrice', 0, 'data_missing'),
    ('regularMarketTime', None, 'data_missing'), ('regularMarketTime', True, 'data_missing'),
    ('regularMarketTime', 1e100, 'data_missing'), ('exchangeTimezoneName', 'Invalid/Zone', 'data_missing'),
    ('exchangeTimezoneName', None, 'data_missing'), ('symbol', 'OTHER', 'identity_mismatch'),
    ('symbol', None, 'identity_mismatch'), ('currency', 'USD', 'currency_mismatch'),
    ('currency', None, 'data_missing'), ('regularMarketTime', timestamp('2026-09-07'), 'stale'),
    ('regularMarketTime', timestamp('2026-09-12'), 'source_unavailable'),
])
def test_declared_gaps_have_no_price_fallback(field, value, status):
    bundle, quotation, fvs = fixture()
    bundle['case']['info'].update({field: value, 'currentPrice': 20., 'previousClose': 20., 'bid': 20.})
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] == status
    assert all(block['upside_' + s + '_pct'] is None for s in SCENARIOS)
    if field == 'regularMarketPrice' or status == 'identity_mismatch':
        assert block['price'] is None
    else:
        assert block['price'] == 30.
    json.dumps(block, allow_nan=False)


def test_only_profile_envelope_supplies_source_and_acquisition_date():
    bundle, quotation, fvs = fixture()
    bundle['case']['sources']['profile']['status'] = 'stale'
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] == 'source_unavailable' and block['source_status'] == 'stale'
    assert block['source_id'] == 'synthetic-profile' and block['upside_base_pct'] is None
    bundle['case']['sources']['profile'] = {'status': 'ok'}
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] == 'data_missing'


@pytest.mark.parametrize('source_status,stamp,status', [
    ('ok', None, 'data_missing'), ('source_error', timestamp(), 'source_unavailable'),
])
@pytest.mark.parametrize('symbol', ['OTHER', None])
def test_foreign_price_is_hidden_even_when_source_or_time_fails_first(source_status, stamp, status, symbol):
    bundle, quotation, fvs = fixture()
    bundle['case']['sources']['profile']['status'] = source_status
    bundle['case']['info'].update(symbol=symbol, regularMarketTime=stamp)
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] == status
    assert block['price'] is None
    assert api().market_quote_view(block, as_of=DAY)['price'] is None


def test_fx_is_not_rolled_and_pence_are_not_pounds():
    bundle, quotation, fvs = fixture()
    quotation['financial_currency'] = 'USD'
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] == 'fx_not_rolled' and block['upside_base_pct'] is None
    quotation.update(financial_currency='GBP', quote_currency='GBP', quote_unit='GBX', quote_units_per_currency=100.)
    bundle['case']['info']['currency'] = 'GBp'
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] == 'ok' and block['upside_base_pct'] == -52.9
    bundle['case']['info']['currency'] = 'GBP'
    assert api().build_market_quote(bundle, quotation, fvs)['status'] == 'currency_mismatch'


@pytest.mark.parametrize('reference,observation,status', [
    ('2026-09-14', '2026-09-11', 'ok'), ('2026-09-13', '2026-09-11', 'ok'),
    ('2026-09-13', '2026-09-10', 'stale'), ('2026-09-11', '2026-09-10', 'ok'),
])
def test_freshness_uses_previous_weekday(reference, observation, status):
    bundle, quotation, fvs = fixture()
    bundle['case']['as_of'] = bundle['case']['sources']['profile']['as_of'] = reference
    bundle['case']['info']['regularMarketTime'] = timestamp(observation)
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] == status
    assert 'holiday' in block['freshness_policy']


def test_exchange_date_uses_declared_timezone():
    bundle, quotation, fvs = fixture()
    bundle['case']['info']['regularMarketTime'] = timestamp('2026-09-08', 23)
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['observed_at'] == '2026-09-08T23:30:00Z'
    assert block['observed_local_date'] == '2026-09-09' and block['status'] == 'ok'


def test_builder_is_pure_and_uses_final_quoted_fair_values():
    args = fixture()
    original = deepcopy(args)
    args[2].update(bear=12., base=14.13, bull=None)
    original[2].update(args[2])
    block = api().build_market_quote(*args)
    assert args == original
    assert block['upside_bear_pct'] == -60.
    assert block['upside_base_pct'] == -52.9 and block['upside_bull_pct'] is None
    assert block['exchange_timezone'] == 'Europe/Berlin'


def test_nonfinite_metadata_and_ratios_never_enter_json():
    bundle, quotation, fvs = fixture()
    bundle['case']['info'].update(exchangeDataDelayedBy=float('nan'), regularMarketPrice=5e-324)
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['delayed_minutes'] is None and block['price'] == 5e-324
    assert block['upside_base_pct'] is None and 'non rappresentabile' in block['message']
    json.dumps(block, allow_nan=False)


@pytest.mark.parametrize('changes', [{'as_of': 'bad-date'}, {'as_of': '2026-09-11'}, {'source_id': ''}])
def test_acquisition_provenance_gaps_are_explicit(changes):
    bundle, quotation, fvs = fixture()
    bundle['case']['sources']['profile'].update(changes)
    block = api().build_market_quote(bundle, quotation, fvs)
    assert block['status'] in ('data_missing', 'source_unavailable')
    assert block['upside_base_pct'] is None


def test_view_rechecks_age_without_mutating_snapshot_or_recovering_upside():
    block = api().build_market_quote(*fixture())
    original = deepcopy(block)
    current = api().market_quote_view(block, as_of=DAY)
    stale = api().market_quote_view(block, as_of='2026-09-14')
    unusable = api().market_quote_view(block, as_of=DAY, usable=False)
    assert current['status_at_read'] == 'ok' and current['upside_base_pct'] == -52.9
    assert stale['status'] == 'ok' and stale['status_at_read'] == 'stale'
    assert stale['price'] == 30. and stale['observed_at'] == original['observed_at']
    assert stale['upside_base_pct'] is None and unusable['upside_base_pct'] is None
    assert block == original
    for missing in (None, {}, [], {'status': 'ok'}):
        view = api().market_quote_view(missing, as_of=DAY)
        assert view['status_at_read'] == 'data_missing' and view['upside_base_pct'] is None


@pytest.mark.parametrize('field,value', [
    ('price', 31.), ('observed_at', '2026-09-10T16:30:00Z'), ('currency', 'USD'),
    ('source_id', 'other'), ('info_symbol', 'OTHER'), ('price_model', 11.),
    ('price_model_as_of', '2026-01-01'), ('observed_local_date', '2026-09-11'),
    ('acquired_as_of', '2026-09-11'), ('status', 'stale'), ('upside_base_pct', 99.),
])
def test_tampered_quote_is_rejected_by_actual_gate(field, value, tmp_path):
    result = generate(tmp_path)
    result['market_quote'][field] = value
    normalized = normalize_valuation_payload(result)
    assert not normalized['valuation_usability']['usable']
    assert any('market_quote' in reason for reason in normalized['valuation_usability']['reasons'])
    assert normalized['fair_value_base'] is None
    assert normalized['market_quote']['upside_base_pct'] is None


def test_absent_legacy_block_and_free_message_do_not_reject_value(tmp_path):
    result = generate(tmp_path)
    result['market_quote']['message'] = 'A presentation wording change'
    assert normalize_valuation_payload(result)['valuation_usability']['usable']
    result.pop('market_quote')
    assert normalize_valuation_payload(result)['valuation_usability']['usable']


def test_quote_contract_cannot_be_silently_reinterpreted(tmp_path):
    result = generate(tmp_path)
    result['market_quote']['contract'] = 'market_quote/2'
    assert not normalize_valuation_payload(result)['valuation_usability']['usable']
    assert api().market_quote_view(result['market_quote'], as_of=DAY)['upside_base_pct'] is None


@pytest.mark.parametrize('status', [[], {}, None, 1, 'unknown'])
def test_malformed_view_status_is_declared_missing(status):
    block = api().build_market_quote(*fixture())
    block['status'] = status
    view = api().market_quote_view(block, as_of=DAY)
    assert view['status_at_read'] == 'data_missing' and view['upside_base_pct'] is None


def test_model_sanity_blocks_even_when_observed_price_would_pass(tmp_path):
    rows = operating_records()
    next(r for r in rows if r['driver'] == 'quotation')['value']['price'] = 5.
    info = quote_info()
    info['regularMarketPrice'] = 14.
    result = generate(tmp_path, records=rows, info=info)
    assert not result['valuation_usability']['usable']
    assert result['sanity']['severity'] == 'BLOCK'
    assert result['market_quote']['price'] == 14.
    assert result['market_quote']['upside_base_pct'] is None


def test_missing_current_quote_does_not_block_historical_valuation(tmp_path):
    result = generate(tmp_path, info={})
    assert result['valuation_usability']['usable']
    assert result['market_quote']['status'] == 'data_missing'
    assert result['market_quote']['upside_base_pct'] is None


def test_managed_care_finish_uses_same_contract(tmp_path):
    from test_managed_care_integration import make_bundle
    bundle = make_bundle()
    info = quote_info()
    info.update(symbol=bundle['case']['ticker'], currency='USD')
    bundle = acquired_bundle(bundle, info)
    result = dcf_engine.generate_valuation(bundle['case']['ticker'], prepared_bundle=bundle, output_dir=str(tmp_path))
    assert result['valuation_usability']['usable'], result['valuation_usability']
    assert isinstance(result.get('market_quote'), dict)
    assert result['market_quote']['status'] == 'ok'
    assert result['market_quote']['upside_base_pct'] == round((result['fair_value_base'] / 30. - 1) * 100, 1)
    workbook = load_workbook(result['path'], data_only=True)
    assert workbook['Managed care']['B2'].value == pytest.approx(result['fair_value_base'])
    assert result['market_quote']['upside_base_pct'] in [row[1] for row in workbook['Managed care'].values if len(row) > 1]
    workbook.close()
