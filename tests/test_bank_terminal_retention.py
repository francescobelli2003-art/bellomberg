"""Perpetual statutory coverage with explicitly growing flows, synthetic only."""
from copy import deepcopy

import pytest


def test_fixed_adjustments_need_not_force_statutory_stock_to_grow_at_g():
    from bellomberg.valuation.capital_inputs import continuing_capital_coverage
    # Common 200, cash 10 and fixed noncash adjustments 30 imply capital 160.
    # Supported flows retain 3.8; capital 163.8 differs from 160*1.02=163.2.
    result = continuing_capital_coverage(160., 163.8, 127.5, .02)
    assert result['sustainable']
    assert result['retained_capital'] == pytest.approx(3.8)
    assert result['perpetual_margin'] == pytest.approx(3.8*1.02-.02*127.5)


@pytest.mark.parametrize('growth,opening,closing,required,expected', [
    (.02,160.,160.5,127.5,False), # initial excess masks eventual shortage
    (.02,160.,162.5,127.5,True),  # boundary: retained flows fund required growth
    (.02,100.,102.5,127.5,False), # fails first year despite sufficient retention
    (0.,160.,159.,125.,False),   # eventually exhausts capital
    (0.,160.,160.,125.,True),
    (0.,160.,161.,125.,True),
    (-.02,160.,156.7,122.5,False), # negative long-run capital
    (-.02,160.,156.8,122.5,True),
    (-.02,160.,157.,122.5,True),
])
def test_perpetual_coverage_handles_growth_zero_and_contraction(growth,opening,closing,required,expected):
    from bellomberg.valuation.capital_inputs import continuing_capital_coverage
    assert continuing_capital_coverage(opening,closing,required,growth)['sustainable'] is expected


@pytest.mark.parametrize('args', [
    (True,160.,125.,.02), (160.,float('nan'),125.,.02),
    (160.,160.,-1.,.02), (160.,160.,125.,-1.), (160.,float('inf'),125.,.02),
    (1e308,1e308,1e308,1e308),
])
def test_nonfinite_or_invalid_coverage_is_rejected(args):
    from bellomberg.valuation.capital_inputs import continuing_capital_coverage
    with pytest.raises(ValueError):
        continuing_capital_coverage(*args)


def retention_records():
    from test_sector_bank_capital import bank_records
    from bellomberg.valuation.distributable_equity import CASH_SIGNS
    rows=bank_records()
    for scenario in ('bear','base','bull'):
        get=lambda driver: next(row for row in rows if row['driver']==driver and row['scenario']==scenario)
        get('terminal_growth')['value']=.02
        get('capital.terminal_equity')['value']=118.*(.1-.02)/(.1-.02)
        terminal=get('terminal_ledger')['value']
        terminal['statutory_projection']='retained_flows_at_g'
        cap=terminal['capital']
        cap['parent_cash_flows']={key:[0.] for key in CASH_SIGNS}
        cap['parent_cash_flows']['debt_issued']=[.1]
        cap['parent_cash_minimum']=[10.2];cap['terminal_debt']=5.1
        cap['parent_gaap_net_income']=[0.]
        sub=cap['subsidiaries'][0]
        sub.update(gaap_net_income=[11.8],proposed_distribution=[9.54],
                   permitted_distribution=[11.8],liquidity_before_transfers=[25.86])
        bridge=terminal['liquidity_bridge']['LEGAL-A']
        bridge.update(operating_cash=[9.86],parent_fees_paid=[0.],parent_tax_paid=[0.])
        terminal['capital_constraints']['LEGAL-A']['constraints'][0]['terminal_requirement']=112.2
        # Capital closes112.26, not112.2; retains2.26 vs required2.1568627.
    return rows


def test_explicit_retained_flows_generate_bank_model_and_keep_payload_immutable(tmp_path):
    from bellomberg.valuation.dcf_engine import generate_valuation
    from test_sector_bank_capital import bank_bundle
    rows=retention_records();before=deepcopy(rows)
    result=generate_valuation('SYNTH-BANK',prepared_bundle=bank_bundle(rows),output_dir=str(tmp_path))
    assert result['valuation_usability']['usable'],result.get('error')
    assert rows==before
    assert result['fair_value_base']==pytest.approx(11.36)


def test_legacy_terminal_does_not_silently_select_new_projection(tmp_path):
    from bellomberg.valuation.dcf_engine import generate_valuation
    from test_sector_bank_capital import bank_bundle
    rows=retention_records()
    for row in rows:
        if row['driver']=='terminal_ledger':row['value'].pop('statutory_projection')
    result=generate_valuation('SYNTH-BANK',prepared_bundle=bank_bundle(rows),output_dir=str(tmp_path))
    assert not result['valuation_usability']['usable']


def test_new_projection_is_linked_and_guarded_in_excel(tmp_path):
    from openpyxl import load_workbook
    from bellomberg.valuation.dcf_engine import generate_valuation
    from test_sector_bank_capital import bank_bundle
    payload=generate_valuation('SYNTH-BANK',prepared_bundle=bank_bundle(retention_records()),output_dir=str(tmp_path))
    assert payload['valuation_usability']['usable'],payload.get('error')
    wb=load_workbook(payload['path'])
    checks=wb['Input Checks']
    pairs=[(str(checks.cell(i,2).value),str(checks.cell(i,3).value)) for i in range(1,checks.max_row+1)]
    assert any('perpetual capital coverage' in label and 'IF(' in formula for label,formula in pairs)
    assert not any('continuing capital growth' in label for label,_ in pairs)
    assert any(word in wb['Bank base']['B4'].value for word in ('trattenuti','Retained'))


def test_first_year_capital_surplus_does_not_hide_perpetual_shortage(tmp_path):
    from bellomberg.valuation.dcf_engine import generate_valuation
    from test_sector_bank_capital import bank_bundle
    rows=retention_records()
    for row in rows:
        if row['driver']=='terminal_ledger' and row['scenario']=='base':
            row['value']['capital']['subsidiaries'][0]['gaap_to_statutory_income']=[-1.5]
    result=generate_valuation('SYNTH-BANK',prepared_bundle=bank_bundle(rows),output_dir=str(tmp_path))
    assert not result['valuation_usability']['usable']
    assert result.get('fair_value_base') is None
    assert any('prosecuzione perpetua' in task['reason'] for task in result['acquisition_tasks'])


@pytest.mark.parametrize('policy',['automatic',None,True])
def test_unknown_projection_is_not_a_fallback(tmp_path,policy):
    from bellomberg.valuation.dcf_engine import generate_valuation
    from test_sector_bank_capital import bank_bundle
    rows=retention_records()
    for row in rows:
        if row['driver']=='terminal_ledger':row['value']['statutory_projection']=policy
    result=generate_valuation('SYNTH-BANK',prepared_bundle=bank_bundle(rows),output_dir=str(tmp_path))
    assert not result['valuation_usability']['usable']
