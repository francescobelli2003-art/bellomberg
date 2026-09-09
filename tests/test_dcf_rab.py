"""Motore RAB V7 Lotto 1 (design audit/15 ok PM 21/07): funzioni pure di
dcf_rab + profili utilities in tassonomia. Zero rete: niente ticker_obj,
ancore regolatorie testate con `today` esplicito (mai date.today() implicita
nei casi STALE).
"""
from datetime import date

import pytest

from bellomberg.valuation.dcf_rab import (anchor_for, build_rab_spec, compute_fair_values_rab,
                     N_YEARS)
from bellomberg.market_data.sector_taxonomy import classify

WACC = {"rf": 0.035, "erp": 0.05, "crp": 0.01}
INFO_IT = {"currency": "EUR", "country": "Italy", "longName": "Rete Test",
           "sharesOutstanding": 2_000_000_000}
PROFILE_REG = classify("Utilities - Regulated Gas", "Utilities", "SRG.MI")
OGGI = date(2026, 7, 21)


def test_classify_profili_utilities():
    # keyword match (l'industry Yahoo e' "Utilities - Regulated Electric/Gas")
    p = classify("Utilities - Regulated Electric", "Utilities", "TRN.MI")
    assert p["engine"] == "rab" and p["_profile_key"] == "utilities - regulated"
    assert PROFILE_REG["engine"] == "rab"
    d = classify("Utilities - Diversified", "Utilities", "ENEL.MI")
    assert d["engine"] == "operating" and d["_profile_key"] == "utilities - diversified"
    r = classify("Utilities - Renewable", "Utilities", "ERG.MI")
    assert r["_profile_key"] == "utilities - renewable"
    # IPP merchant NON mappata di proposito: default dichiarato
    ipp = classify("Utilities - Independent Power Producers", "Utilities", "XXX")
    assert ipp["_profile_key"] == "default"


def test_classify_ipp_resta_operating_ma_dichiara_lo_sconosciuto():
    # il buco VOLUTO resta un buco (engine operating, prior generici) e diventa UDIBILE:
    # e' la sentinella contro chi un giorno «semplificasse» i due rami del ripiego in uno
    # solo, rompendo le operative vere con industry fuori dalle chiavi
    ipp = classify("Utilities - Independent Power Producers", "Utilities", "XXX")
    assert ipp["engine"] == "operating"
    assert ipp["_etichetta"].fonte == "nessuna"
    assert "SCONOSCIUT" in ipp["_etichetta"].dichiarazione.upper()


def test_rifiuto_senza_rab_base():
    # la RAB non e' su Yahoo: senza input analista il modello si RIFIUTA
    for rab in (None, {}, {"rab_base": 0}, {"rab_base": "boh"}):
        s = build_rab_spec("TRN.MI", INFO_IT, WACC, rab=rab,
                           profile=PROFILE_REG, today=OGGI)
        assert "error" in s and "rab_base" in s["error"]


def test_roll_forward_rab_deterministico():
    # paese SENZA ancora + allowed analista: ia=0 dichiarata -> aritmetica pura
    info = dict(INFO_IT, country="Spain")
    s = build_rab_spec("REE.MC", info, WACC,
                       rab={"rab_base": 10_000, "allowed_return": 0.07,
                            "capex_plan": [500] * N_YEARS, "da_pct_rab": 0.04},
                       profile=PROFILE_REG, today=OGGI)
    assert "error" not in s
    assert len(s["rab_vec"]) == N_YEARS
    assert s["rab_vec"][0] == pytest.approx(10_100, abs=0.1)   # 10000 + 500 - 400
    assert s["rab_vec"][1] == pytest.approx(10_196, abs=0.1)   # 10100 + 500 - 404
    assert "SENZA indicizzazione" in s["_sources"]["inflazione"]


def test_ancora_arera_e_stale():
    a, note = anchor_for("Italy", "distribuzione gas", today=OGGI)
    assert a is not None and a["regulator"] == "ARERA"
    assert a["wacc_real"] == 0.059 and not a["stale"]
    assert "513/2024" in note

    s = build_rab_spec("SRG.MI", INFO_IT, WACC,
                       rab={"rab_base": 20_000, "service": "distribuzione gas",
                            "net_debt": 12_000},
                       profile=PROFILE_REG, today=OGGI)
    # review V7 F1: nei CONTI va il tasso REALE (RAB indicizzata); il nominale
    # (reale 5,9% + ia 1,9%) resta solo come riferimento dichiarato
    assert s["allowed_return_calc"] == pytest.approx(0.059)
    assert s["allowed_return_nominal"] == pytest.approx(0.078)
    assert "REALE" in s["_sources"]["allowed_return"]
    assert "513/2024" in s["_sources"]["allowed_return"]
    # review V7 F4: aliquota teorica T del gross-up, non la tc dello scudo
    assert s["tax"] == pytest.approx(0.298)
    assert s["anchor_stale"] is False

    # periodo regolatorio scaduto -> STALE dichiarato + warning nel blend
    s_stale = build_rab_spec("SRG.MI", INFO_IT, WACC,
                             rab={"rab_base": 20_000, "service": "distribuzione gas",
                                  "net_debt": 12_000},
                             profile=PROFILE_REG, today=date(2028, 1, 10))
    assert s_stale["anchor_stale"] is True
    fv = compute_fair_values_rab(s_stale)
    assert any("STALE" in w for w in fv["warnings"])


def test_rifiuto_senza_ancora_ne_allowed_return():
    info = dict(INFO_IT, country="Spain")   # paese fuori dalle ancore in codice
    s = build_rab_spec("REE.MC", info, WACC, rab={"rab_base": 10_000},
                       profile=PROFILE_REG, today=OGGI)
    assert "error" in s and "RIFIUTATO" in s["error"]


def test_ev_rab_premium_e_banda():
    s = build_rab_spec("SRG.MI", INFO_IT, WACC,
                       rab={"rab_base": 20_000, "service": "distribuzione gas",
                            "net_debt": 12_000, "rab_premium": 1.2},
                       profile=PROFILE_REG, today=OGGI)
    fv = compute_fair_values_rab(s)
    # EV = 20000 x 1,2 = 24000; equity = 12000; 2000 mln azioni -> 6,00
    assert fv["fair_value_ev_rab"] == pytest.approx(6.00, abs=0.01)

    # premio accettato dai bound ma FUORI banda profilo -> warning dichiarato
    s2 = build_rab_spec("SRG.MI", INFO_IT, WACC,
                        rab={"rab_base": 20_000, "service": "distribuzione gas",
                             "net_debt": 12_000, "rab_premium": 2.0},
                        profile=PROFILE_REG, today=OGGI)
    fv2 = compute_fair_values_rab(s2)
    assert any("FUORI dalla banda" in w for w in fv2["warnings"])


def test_ddm_regolato_e_blend():
    s = build_rab_spec("SRG.MI", INFO_IT, WACC,
                       rab={"rab_base": 20_000, "service": "distribuzione gas",
                            "net_debt": 12_000, "rab_premium": 1.2},
                       profile=PROFILE_REG, today=OGGI)
    fv = compute_fair_values_rab(s)
    assert fv["fair_value_ddm_reg"] is not None and fv["fair_value_ddm_reg"] > 0
    assert set(fv["blend_methods"]) == {"EV/RAB", "DDM regolato"}
    assert fv["fair_value_blend"] is not None

    # senza net debt: EV/RAB e DDM saltati, buco DICHIARATO (mai numeri inventati)
    info_nodebt = dict(INFO_IT)
    s_nd = build_rab_spec("SRG.MI", info_nodebt, WACC,
                          rab={"rab_base": 20_000, "service": "distribuzione gas"},
                          profile=PROFILE_REG, today=OGGI)
    fv_nd = compute_fair_values_rab(s_nd)
    assert fv_nd["fair_value_ev_rab"] is None
    assert fv_nd["fair_value_blend"] is None
    assert any("Net debt n.d." in w for w in fv_nd["warnings"])


def test_peer_pe_su_utile_ammesso():
    s = build_rab_spec("SRG.MI", INFO_IT, WACC,
                       rab={"rab_base": 20_000, "service": "distribuzione gas",
                            "net_debt": 12_000},
                       profile=PROFILE_REG, today=OGGI)
    fv = compute_fair_values_rab(s, peer_pe_median=14.0)
    # review V7 F1+F4: NI anno 1 = (20000 x 5,9% REALE - 12000 x 3,6%) x (1-29,8%)
    # = 748 x 0,702 = 525,1 mln -> FV peer = 525,1 x 14 / 2000 azioni = 3,68
    assert fv["peer_method"] == "Peer P/E"
    assert fv["fair_value_peer"] == pytest.approx(3.68, abs=0.01)
    assert "Peer P/E" in fv["blend_methods"]


def test_workbook_rab_offline(tmp_path):
    from openpyxl import load_workbook

    from bellomberg.valuation.dcf_rab import build_rab_model
    s = build_rab_spec("SRG.MI", dict(INFO_IT, currentPrice=5.0), WACC,
                       rab={"rab_base": 20_000, "rab_base_year": 2025,
                            "service": "distribuzione gas", "net_debt": 12_000,
                            "capex_plan": [900] * 10},
                       variant_view="RAB dal piano industriale (test)",
                       profile=PROFILE_REG, today=OGGI)
    peers = [{"ticker": "TRN.MI", "name": "Terna", "pe": 16.0, "divy": 0.045,
              "in_band": True, "note": "ok"},
             {"ticker": "NG.L", "name": "National Grid", "pe": 45.0, "divy": 0.05,
              "in_band": False, "note": "P/E 45 FUORI banda: escluso dalla mediana"}]
    out = str(tmp_path / "VAL_SRG_MI.xlsx")
    r = build_rab_model(s, out, peers_data=peers, peers_note="peer di test")
    assert r["ok"] is True and r["engine"] == "rab"
    # il payload porta i metodi del mirror (contratto sidecar/F17)
    assert r["fair_value_ev_rab"] is not None
    assert r["fair_value_ddm_reg"] is not None
    assert r["fair_value_peer"] is not None          # mediana = solo Terna (in banda)
    assert r["fair_value_blend"] is not None
    wb = load_workbook(out)
    assert {"Thesis & Assumptions", "RAB Roll-forward", "DDM regolato",
            "EV-RAB & Peers", "Summary"} <= set(wb.sheetnames)
    # formule vive nei punti chiave (il bake COM e' del Lotto collaudo, non qui)
    assert str(wb["RAB Roll-forward"]["C13"].value).startswith("=")
    assert str(wb["DDM regolato"]["B22"].value).startswith("=")
    assert str(wb["Summary"]["B7"].value).startswith("=MEDIAN")
    # review V7 C4: la cella peer del Summary e' un riferimento VIVO, non un valore
    assert str(wb["Summary"]["B6"].value).startswith("='EV-RAB & Peers'!")


def test_workbook_capex_pct_vanilla_uk_e_gbp(tmp_path):
    # review V7: rami non coperti dal collaudo felice — capex % RAB, convenzione
    # vanilla Ofgem (DDM n.d. dichiarato), quotazione in pence (C1 ALTA)
    from openpyxl import load_workbook

    from bellomberg.valuation.dcf_rab import build_rab_model
    info_uk = {"currency": "GBp", "country": "United Kingdom",
               "longName": "Grid Test", "sharesOutstanding": 4_000_000_000,
               "currentPrice": 1050.0}
    s = build_rab_spec("NG.L", info_uk, WACC,
                       rab={"rab_base": 20_000, "service": "trasmissione elettrica",
                            "net_debt": 12_000, "capex_pct_rab": 0.06},
                       profile=PROFILE_REG, today=OGGI)
    # C1: input in GBP mln, prezzo portato da pence a GBP (dichiarato)
    assert s["currency"] == "GBP"
    assert s["price"] == pytest.approx(10.50)
    assert "pence" in s["_sources"]["prezzo"]
    # ancora Ofgem: tasso REALE nei conti, convenzione vanilla
    assert s["allowed_return_calc"] == pytest.approx(0.0446)
    assert s["convention"] == "cpih_real_vanilla"

    out = str(tmp_path / "VAL_NG_L.xlsx")
    r = build_rab_model(s, out, peers_data=[], peers_note="nessun peer (test)")
    # DDM n.d. dichiarato su vanilla; blend dal solo EV/RAB: (20000x1,15-12000)/4000
    assert r["fair_value_ddm_reg"] is None
    assert r["fair_value_blend"] == pytest.approx(2.75, abs=0.01)
    assert any("vanilla" in w for w in r["warnings"])
    wb = load_workbook(out)
    assert str(wb["DDM regolato"]["A3"].value).startswith("DDM regolato N.D.")
    assert wb["RAB Roll-forward"]["C11"].value == "=C10*$B$6"   # capex % RAB vivo
    assert wb["Summary"]["B5"].value == "n.d."                  # mai una MEDIAN sporca


def test_warnings_finanziabilita_leva_e_rab_stale():
    # review V7 F2b/F3/F5a: capex plan vero >> D&A con payout alto = dividendi
    # non finanziabili -> dichiarato e quantificato; leva calante e RAB vecchia dette
    s = build_rab_spec("SRG.MI", INFO_IT, WACC,
                       rab={"rab_base": 20_000, "rab_base_year": 2023,
                            "service": "distribuzione gas", "net_debt": 12_000,
                            "capex_plan": [2000] * 10},
                       profile=PROFILE_REG, today=OGGI)
    # F2a: g clampata a ROE implicito x retention (dichiarato in fonte)
    assert "COERENZA (review V7 F2)" in s["_sources"]["g"]
    fv = compute_fair_values_rab(s)
    assert any("NON FINANZIABILE" in w for w in fv["warnings"])
    assert any("leva implicita" in w for w in fv["warnings"])
    assert any("RAB base dell'anno 2023" in w for w in fv["warnings"])
