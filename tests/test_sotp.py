# -*- coding: utf-8 -*-
"""V7 Lotto 3 (Mattone C audit/15): layer segmenti -> SOTP. Test OFFLINE puri:
_compute_sotp (mirror Python) + _append_sotp_sheet (foglio a formule vive su un
workbook temporaneo). Nessuna rete, nessun DB, nessun bake COM.
Regola dei collaudi: i test si adattano all'API reale, mai il contrario."""
import openpyxl
import pytest

from bellomberg.valuation.dcf_engine import _SIDECAR_KEYS, _append_sotp_sheet, _compute_sotp


def _seg_reti(**kw):
    d = {"name": "Reti regolate", "engine": "rab", "rab_base": 10000.0,
         "rab_premium": 1.2, "src": "[src: piano industriale]", "ebitda": 1800.0}
    d.update(kw)
    return d


def _seg_gen(**kw):
    d = {"name": "Generazione", "engine": "multiple", "metric": "EBITDA 2026E",
         "value": 2000.0, "multiple": 8.0, "src": "[src: piano industriale]",
         "ebitda": 2000.0}
    d.update(kw)
    return d


def _sotp_base(segments=None, **kw):
    args = {"net_debt": 8000.0,
            "adjustments": [{"label": "Minority interest", "value_m": -500.0,
                             "commentary": "[src: bilancio] quota terzi"}],
            "shares_m": 1000.0, "headline_fv": 20.0,
            "consol_ebitda_m": 3800.0, "consol_revenue_m": None}
    args.update(kw)
    return _compute_sotp(segments if segments is not None
                         else [_seg_reti(), _seg_gen()], **args)


def test_sotp_none_senza_segments():
    for bad in (None, [], "Reti", 123, ["stringa", 4]):
        assert _compute_sotp(bad, net_debt=0, adjustments=None, shares_m=1,
                             headline_fv=None) is None


def test_sotp_due_segmenti_felice():
    s = _sotp_base()
    assert s["incomplete"] is False
    # EV: 10000*1.2 + 2000*8 = 28000; equity: 28000 - 8000 - 500 = 19500; fv 19.5
    assert s["ev_total"] == pytest.approx(28000.0)
    assert s["equity"] == pytest.approx(19500.0)
    assert s["fv_ps"] == pytest.approx(19.5)
    assert s["delta_pct"] == pytest.approx(-2.5)
    # copertura EBITDA 3800 vs 3800: nessuno scarto; ricavi mai dichiarati -> dichiarato
    assert not any("EBITDA" in w and "scarto" in w for w in s["warnings"])
    assert any("ricavi" in w and "nessun segmento" in w for w in s["warnings"])


def test_sotp_premio_default_dal_profilo_regolato():
    from bellomberg.market_data.sector_taxonomy import SUBSECTORS
    atteso = SUBSECTORS["utilities - regulated"]["rab_premium_default"]
    s = _sotp_base([_seg_reti(rab_premium=None), _seg_gen()])
    r0 = s["rows"][0]
    assert r0["mult"] == pytest.approx(atteso)
    assert "DEFAULT" in r0["method_note"]


def test_sotp_premio_fuori_banda_dichiarato_non_clampato():
    s = _sotp_base([_seg_reti(rab_premium=2.0), _seg_gen()])
    # niente clamp: l'EV usa il premio dell'analista, il foglio urla
    assert s["rows"][0]["ev"] == pytest.approx(20000.0)
    assert any("FUORI banda" in w for w in s["warnings"])


def test_sotp_incompleto_mai_somme_parziali():
    s = _sotp_base([_seg_reti(rab_base=None), _seg_gen()])
    assert s["incomplete"] is True
    assert s["rows"][0]["reason"] and "14/07" in s["rows"][0]["reason"]
    for k in ("ev_total", "equity", "fv_ps", "delta_pct"):
        assert s[k] is None
    assert "INCOMPLETO" in s["note"]


def test_sotp_stake_valida_e_fuori_range():
    s = _sotp_base([_seg_gen(stake=0.6), _seg_reti()])
    assert s["rows"][0]["ev"] == pytest.approx(2000.0 * 8.0 * 0.6)
    s2 = _sotp_base([_seg_gen(stake=1.5), _seg_reti()])
    assert s2["incomplete"] is True
    assert "stake" in s2["rows"][0]["reason"]


def test_sotp_riga_aggiustamento_negativa_dichiarata():
    adj = {"name": "Holding cost capitalizzato", "engine": "multiple",
           "metric": "costo annuo", "value": -300.0, "multiple": 8.0,
           "src": "[src: stima analista]"}
    s = _sotp_base([_seg_reti(), _seg_gen(), adj])
    assert s["rows"][2]["ev"] == pytest.approx(-2400.0)
    assert any("EV negativo" in w for w in s["warnings"])


def test_sotp_copertura_scarto_oltre_10pct():
    s = _sotp_base([_seg_reti(ebitda=1000.0), _seg_gen(ebitda=1000.0)],
                   consol_ebitda_m=3000.0)
    assert any("scarto" in w and "EBITDA" in w for w in s["warnings"])


def test_sotp_engine_non_valido_e_operating_a_multiplo():
    s = _sotp_base([_seg_gen(engine="dcf"), _seg_reti()])
    assert s["incomplete"] is True and "engine" in s["rows"][0]["reason"]
    s2 = _sotp_base([_seg_gen(engine="operating"), _seg_reti()])
    assert s2["incomplete"] is False
    assert "MULTIPLO" in s2["rows"][0]["method_note"]


def test_sotp_headline_mancante_delta_nd():
    s = _sotp_base(headline_fv=None)
    assert s["fv_ps"] == pytest.approx(19.5)
    assert s["delta_pct"] is None
    assert "delta n.d." in s["note"]


def test_sotp_delta_grande_warning_dichiarato():
    s = _sotp_base(headline_fv=10.0)  # fv 19.5 vs 10 -> +95%
    assert s["delta_pct"] == pytest.approx(95.0)
    assert any("storie diverse" in w for w in s["warnings"])


def test_sotp_net_debt_nd_dichiarato():
    s = _sotp_base(net_debt=None)
    assert s["ev_total"] == pytest.approx(28000.0)
    assert s["equity"] is None and s["fv_ps"] is None
    assert "net debt" in s["note"]


def test_sidecar_contract_sotp():
    for k in ("fair_value_sotp", "sotp_delta_pct", "sotp_ev_total",
              "sotp_n_segments", "sotp_incomplete", "sotp_note", "sotp_warnings"):
        assert k in _SIDECAR_KEYS


def _wb_finto(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Thesis & Assumptions"
    p = str(tmp_path / "VAL_TEST.xlsx")
    wb.save(p)
    return p


def test_foglio_sotp_formule_vive(tmp_path):
    s = _sotp_base()
    p = _wb_finto(tmp_path)
    _append_sotp_sheet(p, s, "TEST", "EUR")
    wb = openpyxl.load_workbook(p)
    assert "SOTP (segmenti)" in wb.sheetnames
    ws = wb["SOTP (segmenti)"]
    # layout deterministico: segmenti da riga 5, formule = mirror di _compute_sotp
    assert ws["G5"].value == "=D5*E5*F5"
    assert ws["G6"].value == "=D6*E6*F6"
    assert ws["G8"].value == "=G5+G6"               # somma EV enumerata (solo base EV)
    assert ws["G9"].value == pytest.approx(-8000.0)  # net debt firmato
    assert ws["G10"].value == pytest.approx(-500.0)  # bridge minority
    assert ws["G11"].value == "=G8+SUM(G9:G10)"      # equity
    assert ws["G13"].value == "=G11/G12"             # fv/azione
    assert ws["G15"].value == "=G13/G14-1"           # delta vs consolidato (G5 Kairos)
    assert ws["G14"].value == pytest.approx(20.0)
    assert wb.calculation.fullCalcOnLoad
    # il foglio riporta le ATTENZIONI del mirror
    testo = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "ATTENZIONI" in testo and "ricavi" in testo


def test_foglio_sotp_incompleto_niente_formule_totali(tmp_path):
    s = _sotp_base([_seg_reti(rab_base=None), _seg_gen()])
    p = _wb_finto(tmp_path)
    _append_sotp_sheet(p, s, "TEST", "EUR")
    ws = openpyxl.load_workbook(p)["SOTP (segmenti)"]
    assert ws["G5"].value == "n.d."          # riga rifiutata: nessuna formula
    assert ws["G8"].value == "n.d."          # totale n.d., mai SUM parziale
    testo = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "INCOMPLETO" in testo


def test_foglio_sotp_idempotente(tmp_path):
    s = _sotp_base()
    p = _wb_finto(tmp_path)
    _append_sotp_sheet(p, s, "TEST", "EUR")
    _append_sotp_sheet(p, s, "TEST", "EUR")  # riscrittura: rimpiazza, non duplica
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames.count("SOTP (segmenti)") == 1


# --- fix da review avversariale pre-commit (finanza F1-F6, codice C1-C8) ---

def test_sotp_item_non_oggetto_dichiarato():
    # review codice C1: niente scarti muti — l'item nudo finisce nei warnings
    s = _sotp_base([_seg_reti(), "Generazione"])
    assert s is not None and s["n_segments"] == 1
    assert any("NON-oggetto SCARTATI" in w for w in s["warnings"])


def test_sotp_stake_null_e_default_100pct():
    # review codice C5: null/assente = 100%, non input insensato
    s = _sotp_base([_seg_gen(stake=None), _seg_reti()])
    assert s["incomplete"] is False
    assert s["rows"][0]["ev"] == pytest.approx(16000.0)


def test_sotp_rab_premium_zero_bocciato():
    # review codice C4: premio <=0 = riga n.d., mai una divisione azzerata nel totale
    s = _sotp_base([_seg_reti(rab_premium=0.0), _seg_gen()])
    assert s["incomplete"] is True
    assert "rab_premium" in s["rows"][0]["reason"]


def test_sotp_basis_equity_fuori_da_somma_ev():
    # review finanza F2: la quota di un'associata (equity value) si somma DOPO il debito
    assoc = {"name": "Associata 33%", "engine": "multiple", "basis": "equity",
             "metric": "valore quota (post-debito)", "value": 1000.0, "multiple": 1.0,
             "src": "[src: test]"}
    s = _sotp_base([_seg_reti(), _seg_gen(), assoc])
    assert s["ev_total"] == pytest.approx(28000.0)      # l'associata NON e' qui dentro
    assert s["equity_parts"] == pytest.approx(1000.0)
    assert s["equity"] == pytest.approx(28000.0 - 8000.0 - 500.0 + 1000.0)


def test_sotp_basis_equity_su_rab_bocciata():
    s = _sotp_base([_seg_reti(basis="equity"), _seg_gen()])
    assert s["incomplete"] is True
    assert "basis" in s["rows"][0]["reason"]


def test_sotp_metrica_equity_su_base_ev_warning():
    # review finanza F2: P/E in colonna EV = debito contato due volte -> urlo
    s = _sotp_base([_seg_gen(metric="P/E 2026E"), _seg_reti()])
    assert any("basis:'equity'" in w for w in s["warnings"])


def test_sotp_stake_su_base_ev_warning_consolidate():
    # review finanza F1: stake<1 su EV con debito consolidato al 100% -> dichiarato
    s = _sotp_base([_seg_reti(stake=0.75), _seg_gen()])
    assert any("NON consolidate" in w for w in s["warnings"])


def test_sotp_dedup_minorities_segmento_vs_bridge():
    # review finanza F3: segmento 'Minorities' + riga minority nel bridge = urlo
    minr = {"name": "Minorities LatAm", "engine": "multiple", "metric": "book",
            "basis": "equity", "value": -700.0, "multiple": 1.0, "src": "[src: test]"}
    s = _sotp_base([_seg_reti(), minr])
    assert any("DOPPIO CONTEGGIO" in w for w in s["warnings"])


def test_sotp_net_debt_proxy_dichiarato():
    # review finanza F4: la definizione del net debt e' dichiarata nei warnings
    s = _sotp_base()
    assert any("proxy contabile" in w for w in s["warnings"])


def test_sotp_copertura_parziale_su_sottoinsieme():
    # review finanza F6: un buco non spegne la guardia — sottoinsieme + esclusi elencati
    s = _sotp_base([_seg_reti(ebitda=1000.0), _seg_gen(ebitda=None)],
                   consol_ebitda_m=3000.0)
    assert any("PARZIALE" in w and "Generazione" in w for w in s["warnings"])


def test_foglio_sotp_niente_formula_injection(tmp_path):
    # review codice C8: un name che inizia con '=' resta TESTO, mai formula
    s = _sotp_base([_seg_reti(name="=1+1"), _seg_gen()])
    p = _wb_finto(tmp_path)
    _append_sotp_sheet(p, s, "TEST", "EUR")
    ws = openpyxl.load_workbook(p)["SOTP (segmenti)"]
    assert ws["A5"].value == " =1+1" and ws["A5"].data_type != "f"


def test_foglio_sotp_equity_row_nel_bridge(tmp_path):
    assoc = {"name": "Associata", "engine": "multiple", "basis": "equity",
             "metric": "valore quota", "value": 1000.0, "multiple": 1.0,
             "src": "[src: test]"}
    s = _sotp_base([_seg_reti(), _seg_gen(), assoc])
    p = _wb_finto(tmp_path)
    _append_sotp_sheet(p, s, "TEST", "EUR")
    ws = openpyxl.load_workbook(p)["SOTP (segmenti)"]
    # righe 5-7 segmenti; riga 9 somma SOLO le base EV (5,6); equity in riga 12 con +G7
    assert ws["G9"].value == "=G5+G6"
    assert ws["G12"].value == "=G9+SUM(G10:G11)+G7"
