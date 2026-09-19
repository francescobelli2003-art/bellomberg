"""Independent engine baselines and fail-closed workbook generation."""
from openpyxl import load_workbook
import pytest
from bellomberg.valuation.dcf_engine import generate_valuation
from test_sector_rab_drivers import rab_bundle
from test_real_estate_valuation import property_bundle
from test_property_development import developer_bundle
from test_resources_valuation import resource_bundle


@pytest.mark.parametrize('factory', [developer_bundle, resource_bundle])
def test_finite_models_keep_zero_terminal_and_original_timing(tmp_path, factory):
    bundle = factory()
    payload = generate_valuation(bundle['case']['ticker'], prepared_bundle=bundle, output_dir=str(tmp_path))
    assert payload['valuation_usability']['usable'], payload.get('error')
    wb = load_workbook(payload['path'])
    assert wb['Valuation']['B2'].value == payload['fair_value_base']
    for i,s in enumerate(('bear','base','bull'),8):
        assert wb['Model Checks'].cell(i,3).value == pytest.approx(payload['calculation_details']['scenarios'][s]['fair_value_per_share'])
        assert payload['calculation_details']['scenarios'][s]['terminal_value'] == 0
        assert 'Finite '+s in wb
    wb.close()


def test_property_links_cash_noi_nav_and_forward_liquidity_without_double_count(tmp_path):
    payload = generate_valuation('SYNTH-PROP', prepared_bundle=property_bundle(), output_dir=str(tmp_path))
    assert payload['valuation_usability']['usable'], payload.get('error')
    wb = load_workbook(payload['path'])
    assert wb['Valuation']['B2'].value == 54.5
    assert wb['Model Checks']['C9'].value == pytest.approx(54.5)
    assert wb['Summary']['E8'].data_type == 'f'
    assert 'Property base' in wb
    wb.close()


def test_rab_workbook_links_recognized_assets_cash_and_continuing_year(tmp_path):
    payload = generate_valuation('SYNTH-RAB', prepared_bundle=rab_bundle(), output_dir=str(tmp_path))
    assert payload['valuation_usability']['usable'], payload.get('error')
    wb = load_workbook(payload['path'])
    assert wb['Valuation']['B2'].value == 9
    assert wb['Model Checks']['C9'].value == pytest.approx(9)
    assert wb['RAB base']['D51'].data_type == 'f'
    assert wb['Summary']['E8'].data_type == 'f'
    assert wb['Input Checks'].max_row > 80
    assert all(len(cell.value) < 8192 for ws in wb for row in ws for cell in row if cell.data_type == 'f')
    wb.close()


def test_incomplete_rab_never_gets_a_synthetic_live_result(tmp_path):
    from test_sector_rab_drivers import rab_records
    rows = rab_records()
    rows = [r for r in rows if r['driver'] != 'cash_capex']
    payload = generate_valuation('SYNTH-RAB', prepared_bundle=rab_bundle(rows), output_dir=str(tmp_path))
    assert not payload['valuation_usability']['usable']
    wb = load_workbook(payload['path'])
    assert 'Model Checks' not in wb
    wb.close()
