"""Real managed-care adapter through the common pipeline; synthetic inputs only."""
from copy import deepcopy
from datetime import date, timedelta
import json
import os
from pathlib import Path
import subprocess
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from bellomberg.valuation import dcf_engine, sector_analysis
from test_managed_care_calendar import calendar_case
from test_sector_valuation_integration import isolated_tools
from test_valuation_snapshot_persistence import db

DAY = '2026-09-10'
SYMBOL = 'SYNTH.EXT'


def records_for(kind='interim', non_solar=False, count=3):
    data, context = calendar_case(kind, non_solar, count)
    records = []
    periods = data['calendar']['fiscal_periods']
    period = '|'.join(p['start'] + '/' + p['end'] for p in periods)
    cutoff = data['calendar']['valuation_date']
    future = '|'.join(((date.fromisoformat(cutoff)+timedelta(days=1)).isoformat() if i == 0 else p['start']) + '/' + p['end'] for i,p in enumerate(periods))
    def record(field, driver, value, *, scenario='model', entity='SYNTH-GROUP', unit='USD million', basis='GAAP', span=period):
        if driver == 'quotation': span = cutoff
        if driver.startswith('capital.'):
            span = cutoff if 'opening_' in driver or driver in ('capital.ke','capital.shares_m') else periods[-1]['end'] if driver.startswith('capital.terminal_') else future
        records.append({'field': field, 'driver': driver, 'scenario': scenario, 'value': deepcopy(value),
            'entity': entity, 'period': span, 'unit': unit, 'accounting_basis': basis,
            'source_id': 'https://example.org/synthetic/filing', 'as_of': '2026-08-01',
            'valid_until': '2026-12-31', 'kind': 'historical' if driver.startswith('actuals.') else 'analyst_estimate',
            'rationale': 'Explicit synthetic input; no issuer estimate or public default'})
    record('valuation_perimeter', 'perimeter', {'currency': 'USD', 'consolidated_entity': 'SYNTH-GROUP',
        'parent_entity': 'SYNTH-PARENT', 'subsidiaries': [{'id': 'SYNTH-LEGAL-A', 'regime': 'Synthetic statutory regime'}],
        'share_class': 'ordinary', 'share_basis': 'Diluted ordinary shares; annual EPS denominator separately stated'}, unit='contract', basis='scope')
    record('forecast_periods', 'calendar', {**data['calendar'], 'years': data['years'], 'actuals_year': data['actuals']['year']}, unit='contract', basis='calendar')
    record('quotation_units', 'quotation', {'financial_currency': 'USD', 'quote_currency': 'USD',
        'quote_unit': 'USD', 'quote_units_per_currency': 1., 'financial_to_quote_rate': 1.,
        'shares_per_quote': 1., 'share_class': 'ordinary', 'price': 10.,
        'price_as_of': data['calendar']['valuation_date']}, unit='contract', basis='quotation')
    fields = {'premium': 'premium_revenue', 'medical_costs': 'medical_costs', 'gna': 'operating_expenses', 'net_income': 'accounting_bridge'}
    for key, field in fields.items():
        record(field, 'actuals.' + key, data['actuals'][key], span=data['calendar']['actuals_start'] + '/' + data['actuals']['period_end'])
    for scenario, values in data['scenarios'].items():
        def add(field, driver, value, **kwargs):
            record(field, driver, value, scenario=scenario, **kwargs)
        for segment, paths in values['segments'].items():
            add('premium_revenue', f'segments.{segment}.premium', paths['premium'])
            add('medical_costs', f'segments.{segment}.mcr', paths['mcr'], unit='ratio')
        for key, value in values.items():
            if key in ('capital', 'segments', 'adjustments_after_tax'): continue
            field = {'other_premium':'premium_revenue', 'other_medical_costs':'medical_costs', 'diluted_shares_m':'diluted_shares'}.get(key, 'operating_expenses')
            unit = 'million shares' if key == 'diluted_shares_m' else 'ratio' if key in ('gna_ratio','investment_yield') else 'USD million'
            add(field, key, value, unit=unit)
        for key, value in values['adjustments_after_tax'].items():
            add('accounting_bridge', 'adjustments_after_tax.' + key, value)
        for key, value in values['capital'].items():
            if key == 'subsidiaries':
                for index, sub in enumerate(value):
                    for name, numbers in sub.items():
                        if name == 'id': continue
                        field = 'legal_entity_liquidity' if name in ('liquidity_before_transfers','minimum_liquidity') else 'legal_entity_capital'
                        add(field, f'capital.subsidiaries.{index}.{name}', numbers, entity=sub['id'],
                            basis='GAAP' if name in ('opening_gaap_equity','gaap_net_income') else 'statutory')
            elif key == 'parent_cash_flows':
                for name, numbers in value.items():
                    add('parent_ledger', 'capital.parent_cash_flows.' + name, numbers, entity='SYNTH-PARENT', basis='cash')
            else:
                field = {'ke':'discount_rate_inputs','shares_m':'diluted_shares','discount_periods':'forecast_periods'}.get(key,'parent_ledger')
                unit = 'text' if isinstance(value,str) else {'ke':'ratio','shares_m':'million shares','discount_periods':'years'}.get(key,'USD million')
                basis = 'GAAP' if key in ('parent_gaap_net_income','consolidation_adjustments') else 'cash' if field == 'parent_ledger' else 'valuation'
                add(field, 'capital.' + key, value, entity='SYNTH-PARENT' if field == 'parent_ledger' else 'SYNTH-GROUP', unit=unit, basis=basis)
    return records, {'scenario_rationale': context['scenario_rationale'], 'revisions': []}


def make_bundle(symbol=SYMBOL, *, records=None, context=None, kind='interim', non_solar=False, count=3):
    if records is None:
        records, context = records_for(kind, non_solar, count)
    providers = {'profile': lambda *a, **k: {'status':'ok','source_id':'synthetic','as_of':DAY,
        'data': {'info': {'shortName':'Synthetic health plan','currency':'USD'}, 'vehicle_registry':None,
                 'evidence':[{'field':f,'value':v,'source_id':'synthetic','as_of':DAY}
                    for f,v in [('instrument','equity'),('business_model','managed_care')]]}},
        'method_inputs': lambda *a, **k: {'status':'ok','source_id':'synthetic-method-records','as_of':DAY,'records':deepcopy(records)}}
    return sector_analysis.prepare_sector_analysis(symbol, as_of=DAY, providers=providers,
        user_context={'analysis_context':context or {}})


def generate(bundle, directory):
    return dcf_engine.generate_valuation(bundle['case']['ticker'], output_dir=str(directory), prepared_bundle=bundle)


@pytest.mark.parametrize('kind,non_solar,count', [('FY',False,1),('interim',False,3),('FY',True,6),('interim',True,2)])
def test_real_adapter_reaches_workbook_and_sidecar(tmp_path, monkeypatch, kind, non_solar, count):
    monkeypatch.setattr(dcf_engine, '_generate_valuation_legacy', lambda *a,**k: pytest.fail('No operating fallback'))
    bundle = make_bundle(kind=kind, non_solar=non_solar, count=count)
    result = generate(bundle, tmp_path)
    assert result['valuation_usability']['usable'], result['valuation_usability']
    assert result['valuation_decision']['support_status'] == 'integrated'
    assert result['input_consumption']['unconsumed_fields'] == []
    assert len(result['input_consumption']['consumed_records']) == len(bundle['case']['records'])
    assert result['managed_care']['scenarios']['base']['rows'][0]['remaining_net_income'] == (7 if kind == 'FY' else 5)
    sidecar = json.loads(Path(result['path']).with_suffix('.payload.json').read_text(encoding='utf8'))
    assert sidecar['fair_value_base'] == result['fair_value_base']
    assert sidecar['snapshot_id'] == result['snapshot_id'] == bundle['snapshot_id']
    assert sidecar['managed_care'] == result['managed_care']
    wb = load_workbook(result['path'], data_only=True)
    assert wb.worksheets[0]['B2'].value == 'UTILIZZABILE'
    assert wb['Managed care']['B2'].value == pytest.approx(result['fair_value_base'])
    assert wb['Qualita e revisioni'].max_column >= count + 3
    wb.close()


@pytest.mark.parametrize('failure', ['capital_missing','liquidity','regime','source','proxy','extra','unit','entity','period','duplicate','unreconciled','price_date','fx','shares'])
def test_incomplete_or_incompatible_input_never_becomes_a_fair_value(tmp_path, failure):
    records, context = records_for()
    find = lambda driver: next(r for r in records if r['driver'] == driver and r['scenario'] == 'base')
    if failure == 'capital_missing': records.remove(find('capital.subsidiaries.0.permitted_distribution'))
    if failure == 'liquidity': find('capital.subsidiaries.0.liquidity_before_transfers')['value'][0] = 0
    if failure == 'regime': records[0]['value']['subsidiaries'][0]['regime'] = ''
    if failure == 'source': find('capital.ke')['valid_until'] = '2026-08-31'
    if failure == 'proxy': find('capital.ke')['kind'] = 'proxy'
    if failure == 'extra': records.append({**find('capital.ke'), 'driver':'net_debt'})
    if failure == 'unit': find('segments.Health.premium')['unit'] = 'USD'
    if failure == 'entity': find('capital.subsidiaries.0.gaap_net_income')['entity'] = 'another entity'
    if failure == 'period': find('capital.parent_gaap_net_income')['period'] = 'FY2025'
    if failure == 'duplicate': records.append({**find('capital.ke'), 'value':.5})
    if failure == 'unreconciled': find('capital.subsidiaries.0.gaap_net_income')['value'][0] = 999
    if failure == 'price_date': records[2]['value']['price_as_of'] = DAY
    if failure == 'fx': records[2]['value']['financial_to_quote_rate'] = 100
    if failure == 'shares': records[2]['value']['share_class'] = 'different'
    result = generate(make_bundle(records=records, context=context), tmp_path)
    assert not result['valuation_usability']['usable']
    assert result.get('fair_value_base') is None
    assert result['acquisition_tasks']
    for scenario in result.get('managed_care',{}).get('scenarios',{}).values():
        assert scenario['fair_value_per_share'] is None
        assert scenario['capital']['fair_value_per_share'] is None


def test_symbol_does_not_select_assumptions_and_sensitivity_uses_consumed_driver(tmp_path):
    first = generate(make_bundle(), tmp_path/'a')
    renamed = generate(make_bundle('SYNTH_RENAMED'), tmp_path/'b')
    assert first['valuation_usability']['usable'] and renamed['valuation_usability']['usable']
    assert first['valuation_decision'] == renamed['valuation_decision']
    assert first['fair_value_base'] == renamed['fair_value_base']
    records, context = records_for()
    for row in records:
        if row['driver'] == 'capital.ke': row['value'] = .11
    changed = generate(make_bundle(records=records,context=context), tmp_path/'c')
    assert changed['valuation_usability']['usable'], changed['valuation_usability']
    assert changed['fair_value_base'] < first['fair_value_base']
    assert changed['snapshot_id'] != first['snapshot_id']


def test_real_adapter_chat_build_cache_and_committee_share_result(isolated_tools, monkeypatch):
    tools = isolated_tools
    monkeypatch.setattr(dcf_engine, 'REPORT_DIR', tools.directory)
    bundle = make_bundle()
    chat = tools.chat.dispatch('get_valuation', {'ticker':SYMBOL}, prepared_bundle=bundle)['data']
    build = tools.build(SYMBOL, prepared_bundle=bundle, as_of=DAY)
    assert chat['valuation_usability']['usable'], chat['valuation_usability']
    assert build['fair_value_base'] == chat['fair_value_base']
    assert build['valuation_decision'] == chat['valuation_decision']
    assert chat['snapshot_id'] == bundle['snapshot_id']
    tools.memory.history = [{'valuation_payload':deepcopy(build),'generation_id':build['generation_id']}]
    cached = tools.chat.dispatch('get_valuation', {'ticker':SYMBOL}, prepared_bundle=bundle)['data']
    assert cached['reused'] is True
    assert cached['generation_id'] == build['generation_id']
    row = json.loads(sector_analysis.valuation_results_block({SYMBOL:cached}).splitlines()[-1])
    assert row['fair_value'] == build['fair_value_base']
    assert row['valuation_usability']['usable']
    records, context = records_for()
    records = [r for r in records if r['driver'] != 'capital.subsidiaries.0.permitted_distribution']
    incomplete = tools.build(SYMBOL, prepared_bundle=make_bundle(records=records, context=context), as_of=DAY)
    assert not incomplete['valuation_usability']['usable'] and incomplete.get('fair_value_base') is None
    assert incomplete['path'] and Path(incomplete['path']).is_file()
    assert 'incompleto' in incomplete['engine_note']


def test_public_tool_arguments_acquire_records_without_prepared_bundle(isolated_tools, monkeypatch):
    tools = isolated_tools
    monkeypatch.setattr(dcf_engine, 'REPORT_DIR', tools.directory)
    bundle = make_bundle()
    sources = bundle['case']['sources']
    calls = []
    def profile(ticker, *, as_of):
        calls.append(ticker)
        return deepcopy(sources['profile'])
    records, context = records_for()
    providers = {'profile': profile}  # no injected economic calculator or prepared bundle
    chat = tools.chat.dispatch('get_valuation', {'ticker':SYMBOL, 'method_records':records,
        'analysis_context':context}, sector_providers=providers, as_of=DAY)['data']
    assert chat['valuation_usability']['usable'], chat
    build = tools.build(SYMBOL, method_records=records, analysis_context=context, sector_providers=providers, as_of=DAY)
    assert build['valuation_usability']['usable'], build
    assert build['fair_value_base'] == chat['fair_value_base']
    assert calls == [SYMBOL, SYMBOL]
    # A RESEARCH bundle acquired before the analyst supplied records is revised
    # through the same validator, without any further provider call.
    incomplete = sector_analysis.prepare_sector_analysis(SYMBOL, as_of=DAY, providers=providers)
    count = len(calls)
    revised = tools.chat.dispatch('get_valuation', {'ticker':SYMBOL, 'method_records':records,
        'analysis_context':context}, prepared_bundle=incomplete)['data']
    assert revised['valuation_usability']['usable'], revised
    assert len(calls) == count
    assert revised['snapshot_id'] != incomplete['snapshot_id']
    changed_records = deepcopy(records)
    for row in changed_records:
        if row['driver'] == 'capital.ke': row['value'] = .11
    changed = tools.chat.dispatch('get_valuation', {'ticker':SYMBOL, 'method_records':changed_records},
        prepared_bundle=revised['acquisition_snapshot'])['data']
    assert changed['valuation_usability']['usable'], changed
    assert changed['fair_value_base'] < revised['fair_value_base']
    assert len(changed['acquisition_snapshot']['case']['records']) == len(records)
    assert changed['acquisition_snapshot']['case']['sources']['method_inputs']['data']['previous_acquisition']['records'] == revised['acquisition_snapshot']['case']['sources']['method_inputs']['records']


def test_readiness_without_record_bindings_does_not_certify_a_value():
    from test_sector_usability import payload_for
    from bellomberg.valuation.dcf_quality import assess_valuation_usability
    assert not assess_valuation_usability(payload_for('managed_care'), as_of=DAY)['usable']


def test_shared_result_reaches_score_and_only_latest_workbook_is_attached(tmp_path, monkeypatch):
    from bellomberg.agents import specialist_scores, consigliere_multi
    from bellomberg.agents.specialists.fundamentals import FundamentalsSpecialist
    from bellomberg.storage import classificazione
    from datetime import datetime
    first = generate(make_bundle(), tmp_path)
    second = generate(make_bundle(), tmp_path)
    monkeypatch.setattr(classificazione, 'carica_veicoli', lambda: {'origine':'synthetic','veicoli':{},'motivo':None})
    score = specialist_scores.fundamentals_score({'positions':[{'ticker':SYMBOL,'peso_pct':100}]}, valuations={SYMBOL:second})
    assert score is not None
    seen = []
    monkeypatch.setattr(specialist_scores, 'fundamentals_score', lambda **kw: seen.append(kw) or 'score')
    specialist = FundamentalsSpecialist.__new__(FundamentalsSpecialist)
    specialist.blackboard = SimpleNamespace(valuation_results={SYMBOL:second})
    assert specialist.compute_score() == 'score'
    assert seen[0]['valuations'][SYMBOL]['snapshot_id'] == second['snapshot_id']
    monkeypatch.setattr(consigliere_multi, 'REPORT_DIR', tmp_path)
    monkeypatch.setattr(consigliere_multi, 'MODELS_DIR', tmp_path/'unused')
    os.utime(first['path'], (1,1))
    os.utime(second['path'], (2,2))
    files = consigliere_multi._collect_dcf_files(datetime(1970,1,1))
    assert files == [second['path']]
    assert Path(first['path']).is_file()


@pytest.mark.parametrize('symbol', [SYMBOL, 'SYNTH-CARE'])
def test_score_keeps_book_coverage_with_partial_research_and_never_revives_a_block(tmp_path, monkeypatch, symbol):
    from bellomberg.agents import specialist_scores, agent_tools
    from bellomberg.agents.specialists.fundamentals import FundamentalsSpecialist
    good = generate(make_bundle(symbol), tmp_path)
    monkeypatch.setattr(specialist_scores, 'REPORT_DIR', tmp_path)
    monkeypatch.setattr(specialist_scores.cl, 'carica_veicoli', lambda: {'origine':'synthetic','veicoli':{},'motivo':None})
    monkeypatch.setattr(agent_tools, 'tool_get_portfolio_live', lambda: {'positions':[{'ticker':symbol,'peso_pct':100}]})
    specialist = FundamentalsSpecialist.__new__(FundamentalsSpecialist)
    specialist.blackboard = SimpleNamespace(valuation_results={'RESEARCH':{'ticker':'RESEARCH'}})
    score = specialist.compute_score()
    assert score is not None and score['metrics']['n_valued'] == 1
    assert any(good['valuation_date'] in str(line) for line in score['lines'])
    # A current explicit KO for the holding overrides its older valid file.
    specialist.blackboard.valuation_results[symbol] = {'ticker':symbol, 'error':'missing capital'}
    assert specialist.compute_score() is None
    records, context = records_for()
    records = [row for row in records if row['driver'] != 'capital.subsidiaries.0.permitted_distribution']
    bad = generate(make_bundle(symbol, records=records, context=context), tmp_path)
    os.utime(good['path'].replace('.xlsx','.payload.json'), (1,1))
    os.utime(bad['path'].replace('.xlsx','.payload.json'), (2,2))
    specialist.blackboard.valuation_results = {'RESEARCH':{'ticker':'RESEARCH'}}
    assert specialist.compute_score() is None  # latest disk KO cannot revive older valid generation


@pytest.mark.parametrize('failure', ['parent_whitespace','group_is_sub','early_actual_source','historical_forecast','zero_bear','zero_bull'])
def test_adversarial_sources_identities_and_all_scenario_sanity(tmp_path, failure):
    records, context = records_for()
    if failure == 'parent_whitespace':
        records[0]['value']['parent_entity'] = 'SYNTH-LEGAL-A '
        for row in records:
            if row['entity'] == 'SYNTH-PARENT': row['entity'] = 'SYNTH-LEGAL-A '
    if failure == 'group_is_sub':
        records[0]['value']['consolidated_entity'] = 'SYNTH-LEGAL-A'
        for row in records:
            if row['entity'] == 'SYNTH-GROUP': row['entity'] = 'SYNTH-LEGAL-A'
    if failure == 'early_actual_source':
        for row in records:
            if row['driver'].startswith('actuals.'): row['as_of'] = '2026-05-01'
    if failure == 'historical_forecast':
        for row in records:
            if row['driver'].endswith('permitted_distribution'): row['kind'] = 'historical'
    if failure.startswith('zero_'):
        for row in records:
            if row['scenario'] != failure[5:]: continue
            if row['driver'] == 'capital.parent_opening_cash': row['value'] = 10.
            if row['driver'] == 'capital.terminal_equity': row['value'] = 0.
            if row['driver'] == 'capital.subsidiaries.0.proposed_distribution': row['value'] = [0.] * 3
    result = generate(make_bundle(records=records, context=context), tmp_path)
    assert not result['valuation_usability']['usable'], result
    assert result.get('fair_value_base') is None
    assert result['acquisition_tasks']


def test_real_research_database_f17_and_react_preserve_complete_and_missing_cases(tmp_path, monkeypatch, db):
    from bellomberg.core import current_facts
    from bellomberg.storage import memory_db
    from test_sector_valuation_api import endpoint
    with db._conn() as conn:
        decision_id = conn.execute("INSERT INTO decisions(timestamp, action, ticker, status) VALUES (?, 'RESEARCH', ?, 'PENDING')", (DAY, SYMBOL)).lastrowid
    monkeypatch.setattr(memory_db, 'SQLITE_PATH', db.db_path)
    source_bundle = make_bundle()
    providers = {key: (lambda ticker, *, as_of, source=value: deepcopy(source)) for key,value in source_bundle['case']['sources'].items()}
    # The research block acquires the same real method inputs outside holdings.
    bundles, links = {}, {}
    research = current_facts.research_block(sector_bundles=bundles, providers=providers, as_of=DAY, decision_links=links)
    assert SYMBOL in bundles, research
    result = generate(sector_analysis.revise_sector_analysis(bundles[SYMBOL], analysis_context=source_bundle['analysis_context']), tmp_path)
    assert result['valuation_usability']['usable'], result
    db.save_valuation_thesis(SYMBOL, valuation_payload=result, fair_value=result['fair_value_base'], price=result['price'])
    db.link_valuation_snapshot(result['snapshot_id'], generation_id=result['generation_id'], decision_id=decision_id)
    restored = db.get_valuation_snapshot(result['snapshot_id'], generation_id=result['generation_id'])
    assert restored['fair_value_base'] == result['fair_value_base']
    good = endpoint(tmp_path, db.get_latest_valuation_snapshots())['models'][0]
    assert good['ticker'] == SYMBOL and good['valuation_usability']['usable'], good['valuation_usability']
    assert good['detail']['valuation_date'] == result['valuation_date']
    records, context = records_for()
    records = [r for r in records if not (r['scenario'] == 'base' and r['driver'] == 'capital.subsidiaries.0.permitted_distribution')]
    broken = generate(make_bundle(records=records, context=context), tmp_path)
    db.save_valuation_thesis(SYMBOL, valuation_payload=broken)
    bad = endpoint(tmp_path, db.get_latest_valuation_snapshots())['models'][0]
    assert bad['fair_value'] is None and not bad['valuation_usability']['usable']
    assert Path(result['path']).is_file()  # old workbook stays in place
    assert db.get_valuation_snapshot(result['snapshot_id'], generation_id=result['generation_id'])['fair_value_base'] == result['fair_value_base']
    ui_fixture = tmp_path/'managed-care-ui.json'
    ui_fixture.write_text(json.dumps({'complete':good, 'incomplete':bad}), encoding='utf8')
    rendered = subprocess.run(['node','--test','tests/product/sector-valuation.cjs'], cwd=Path('app').resolve(),
        env={**os.environ,'SECTOR_VALUATION_FIXTURE':str(ui_fixture)}, text=True, capture_output=True)
    assert rendered.returncode == 0, rendered.stdout + rendered.stderr
