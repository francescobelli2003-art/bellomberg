"""DCF-Q1: observable workbook/tool/persistence contracts, offline only."""
import json
import sys
from datetime import datetime
from types import SimpleNamespace as NS

import openpyxl
import pytest


def _spec():
    return {"ticker": "SYNTH", "currency": "USD", "revenue_m": 1000,
            "shares_m": 100, "net_debt": 20, "gross_margin": .3,
            "growth_path": [.03] * 5, "wacc_inputs": {"wacc": .1},
            "terminal_g": .02}


def test_builder_exposes_automatic_quality_and_dated_snapshot(tmp_path):
    from bellomberg.valuation.dcf_buyside_v3 import build_model_v3
    result = build_model_v3(_spec(), str(tmp_path / "synth.xlsx"))
    assert result["ok"], result
    assert "analytical_quality" in result
    quality = result["analytical_quality"]
    assert quality["status"] == "BOZZA_AUTOMATICA"
    assert quality["revision_bridge"]["status"] == "n.d."
    wb = openpyxl.load_workbook(result["path"])
    assert "Qualita e revisioni" in wb.sheetnames
    assert "BOZZA_AUTOMATICA" in wb["Thesis & Output"]["A3"].value
    assert "non certifica" in wb["Qualita e revisioni"]["A3"].value.lower()
    wb.close()


def test_sidecar_preserves_quality_without_changing_block(tmp_path):
    from bellomberg.valuation.dcf_engine import _write_payload_sidecar
    path = tmp_path / "VAL_SYNTH_FLAGGED.xlsx"
    path.write_bytes(b"test path, no excel read in sidecar")
    quality = {"status": "INCOMPLETA", "snapshot": {"version": 1}}
    _write_payload_sidecar({"path": str(path), "ticker": "SYNTH",
        "analytical_quality": quality, "sanity": {"severity": "BLOCK"}})
    saved = json.loads(path.with_suffix(".payload.json").read_text(encoding="utf-8"))
    assert saved.get("analytical_quality") == quality
    assert saved["sanity"]["severity"] == "BLOCK"


def test_analysis_context_is_forwarded_and_never_reuses_old_valuation(monkeypatch):
    import bellomberg.agents.chat_tools as ct
    from bellomberg.valuation import dcf_engine
    calls = []
    monkeypatch.setitem(sys.modules, "bellomberg.storage.memory_db", NS(MemoryDB=lambda: NS(
        get_valuation_history=lambda *a, **k: [dict(date=datetime.now().isoformat(),
            fair_value=100, sanity_severity="OK")])) )
    monkeypatch.setattr(dcf_engine, "generate_valuation", lambda *a, **kw: calls.append(kw) or {})
    context = {"as_of": "2026-09-10"}
    providers = {"profile": lambda ticker, *, as_of: {
        "status": "ok", "source_id": "synthetic-profile", "as_of": as_of,
        "data": {"info": {"quoteType": "EQUITY", "industry": "Software - Application"},
                 "vehicle_registry": None}}}
    ct.dispatch("get_valuation", {"ticker": "SYNTH", "analysis_context": context},
                sector_providers=providers, as_of=context["as_of"])
    assert len(calls) == 1
    # S2 carries the exact context inside the acquired bundle, not a parallel kwarg.
    assert calls[0]["prepared_bundle"]["analysis_context"] == context


def test_previous_snapshot_reads_flagged_without_touching_bytes(tmp_path):
    from bellomberg.valuation import dcf_engine
    assert hasattr(dcf_engine, "_read_previous_valuation_snapshot")
    path = tmp_path / "VAL_SYNTH_FLAGGED.payload.json"
    snapshot = {"version": 1, "ticker": "SYNTH"}
    path.write_text(json.dumps({"analytical_quality": {"snapshot": snapshot}}), encoding="utf-8")
    before = path.read_bytes()
    previous, note = dcf_engine._read_previous_valuation_snapshot(str(tmp_path), "SYNTH")
    assert previous == snapshot and note is None
    assert path.read_bytes() == before
    (tmp_path / "VAL_SYNTH.payload.json").write_text("{}", encoding="utf-8")
    previous, note = dcf_engine._read_previous_valuation_snapshot(str(tmp_path), "SYNTH")
    assert previous is None and "ambigu" in note.lower()


def test_sheet_shows_global_gaps_and_watches_numeric_anchors():
    from bellomberg.valuation.dcf_quality_sheet import append_quality_sheet
    wb = openpyxl.Workbook()
    wb.active.title = "DCF"
    wb["DCF"]["B4"] = .1
    report = {"status": "INCOMPLETA", "issues": ["forecast_years discordanti dal modello"],
              "rows": [{"scenario": "model", "driver": "wacc", "values": .1,
                        "evidence": {}, "issues": []}], "snapshot": {}}
    ws = append_quality_sheet(wb, report)
    assert any("forecast_years discordanti" in str(c.value) for row in ws for c in row)
    assert any("'DCF'!B4" in str(c.value) for row in ws for c in row)


def test_sheet_never_interprets_source_text_as_excel_formula():
    from bellomberg.valuation.dcf_quality_sheet import append_quality_sheet
    wb = openpyxl.Workbook()
    ws = append_quality_sheet(wb, {"snapshot": {}, "status": "=HYPERLINK(\"https://bad.example\")",
         "rows": [{"scenario": "=1+1", "driver": "=2+2", "values": None,
                   "evidence": {"rationale": "=cmd()"}, "issues": []}]})
    assert all(c.data_type == "s" for row in ws for c in row
               if isinstance(c.value, str) and c.value.startswith("=") and c.coordinate != "A4")


def test_bridge_uses_real_dcf_and_reconciles_measured_effects():
    from copy import deepcopy
    from bellomberg.valuation.dcf_quality import assess_quality
    from bellomberg.valuation.dcf_buyside_v3 import (
        _default_scenarios, _merge_scenarios, compute_fair_values_v3)
    spec = _spec()
    spec["wacc_inputs"]["rf"] = .04
    spec["ronic_anchor"] = {"value": .12, "source": "synthetic reference"}
    current = _merge_scenarios(None, _default_scenarios(spec, None))
    context = {"as_of": "2026-09-10", "forecast_years": [2026, 2027, 2028, 2029, 2030]}
    common = dict(last_rev=1000, nwc0=30, compute_fv=compute_fair_values_v3, today="2026-09-10")
    old = assess_quality(spec, current, current, context, **common)["snapshot"]
    old["as_of"] = "2026-08-05"
    for scenario in old["scenarios"].values():
        scenario["gross_margin"] = [.2] * 5
        scenario["revenue_growth"] = [.01] * 5
    context["revisions"] = [{"metric": "synthetic forecast revision", "period": "FY2026",
        "basis": "synthetic reported consolidated revenue", "unit": "fraction", "previous_value": .2,
        "current_value": .3, "value_type": "point", "previous_source": "https://example.org/old",
        "previous_date": "2026-08-01", "source": "https://example.org/new", "source_date": "2026-09-01",
        "drivers": ["gross_margin", "revenue_growth"], "rationale": "synthetic scenario evidence"}]
    saved = deepcopy((spec, current, context, old))
    report = assess_quality(spec, current, current, context, old, **common)
    bridge = report["revision_bridge"]
    assert bridge["status"] == "CALCOLATO", bridge
    before = compute_fair_values_v3(spec, old["scenarios"], 1000, nwc0=30)["fair_value_weighted"]
    after = compute_fair_values_v3(spec, current, 1000, nwc0=30)["fair_value_weighted"]
    assert bridge["delta"] == pytest.approx(after - before)
    assert sum(s["delta"] for s in bridge["steps"]) == pytest.approx(bridge["delta"])
    assert bridge["residual"] == pytest.approx(0, abs=1e-10)
    assert bridge["currency"] == "USD"
    assert (spec, current, context, old) == saved


def test_documentation_does_not_change_fair_values_or_scenario_inputs(tmp_path):
    from bellomberg.valuation.dcf_buyside_v3 import build_model_v3
    plain = build_model_v3(_spec(), str(tmp_path / "plain.xlsx"))
    documented_spec = _spec()
    documented_spec["_analysis_context"] = {"as_of": "2026-09-10"}
    documented = build_model_v3(documented_spec, str(tmp_path / "documented.xlsx"))
    for key in ("fair_value_base", "fair_value_bear", "fair_value_bull", "fair_value_weighted"):
        assert plain[key] == documented[key]


def test_sheet_monitors_fully_diluted_share_cell_used_by_fair_value(tmp_path):
    from bellomberg.valuation.dcf_buyside_v3 import build_model_v3
    spec = _spec()
    spec["diluted_shares_m"] = 120
    result = build_model_v3(spec, str(tmp_path / "diluted.xlsx"))
    wb = openpyxl.load_workbook(result["path"])
    ws = wb["Qualita e revisioni"]
    row = next(c.row for cells in ws for c in cells if c.column == 2 and c.value == "shares_m")
    assert wb["DCF"]["B8"].value == ws.cell(row, 3).value == 120
    assert "'DCF'!B8" in ws.cell(row, 8).value
    wb.close()
