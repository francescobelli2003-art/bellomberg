from copy import deepcopy

import pytest
from openpyxl import load_workbook

from bellomberg.core.language import language_context
from bellomberg.valuation import dcf_engine as de, documented_inputs as di
from tests.test_sector_operating_drivers import bundle_for as bf
def _valuation(tmp_path, language="it"):
    with language_context(language):
        return de.generate_valuation("SYNTH-EXT", prepared_bundle=bf(), output_dir=str(tmp_path))
def _cells(wb):
    return [cell for sheet in wb for row in sheet.iter_rows() for cell in row]
def _book(payload, tmp_path, **opts):
    return load_workbook(di.build_documented_workbook(payload, str(tmp_path)), **opts)
def test_operating_presentation_preserves_payload_and_year_series(tmp_path):
    result = _valuation(tmp_path)
    payload = deepcopy(result)
    before = deepcopy(payload)
    wb = _book(payload, tmp_path, data_only=True)
    assert payload == before
    assert wb.sheetnames[0] == "Summary"
    assert wb['Summary'].freeze_panes is None
    assert all(selection.pane is None for selection in wb['Summary'].sheet_view.selection)
    assert {"Revenue Build", "Model Inputs", "Sensitivity", "Input Checks", "Model Checks"} <= set(wb.sheetnames)
    assert wb["Valuation"]["B2"].value == pytest.approx(result["fair_value_base"])
    base = wb["base"]
    years = result["analytical_quality"]["snapshot"]["forecast_years"]
    assert [base.cell(7, column).value for column in (4, 5)] == years
    # The immutable baseline is retained next to the live model, not inferred
    # from formula caches (Excel recalculates them on opening).
    baseline = dict(wb['Valuation'].values)
    for driver in ("revenue", "ebitda", "ebit", "ufcf", "research_amortization"):
        assert [baseline[f'base.rows.{driver}.{i}'] for i in (0, 1)] == pytest.approx(
            result['calculation_details']['scenarios']['base']['rows'][driver])
    wb.close()

    missing = deepcopy(result)
    missing["calculation_details"]["scenarios"]["base"]["rows"]["revenue"] = [123.4, None]
    missing["calculation_details"]["scenarios"]["base"]["rows"]["ebitda"] = [0.0, None]
    wb = _book(missing, tmp_path, data_only=True)
    base = wb["base"]
    assert base["D8"].value == pytest.approx(123.4)
    assert base["E8"].value == "n.d."
    assert base["D9"].value == 0.0
    assert base["E9"].value == "n.d."
    wb.close()
def test_unusable_summary_hides_fair_value_and_upside(tmp_path):
    payload = deepcopy(_valuation(tmp_path))
    payload["valuation_usability"] = {"usable": False, "reasons": ["missing"]}
    wb = _book(payload, tmp_path, data_only=True)
    summary = wb["Summary"]
    assert [summary.cell(row, column).value for row in (8, 9) for column in (4, 5, 6)] == ["n.d."] * 6
    wb.close()
def test_languages_keep_sheetnames_and_numbers_invariant(tmp_path):
    it = _valuation(tmp_path / "it", "it")
    en = _valuation(tmp_path / "en", "en")
    wi, we = load_workbook(it["path"], data_only=True), load_workbook(en["path"], data_only=True)
    assert wi.sheetnames == we.sheetnames
    ni = sorted(cell.value for cell in _cells(wi) if isinstance(cell.value, (int, float)))
    ne = sorted(cell.value for cell in _cells(we) if isinstance(cell.value, (int, float)))
    assert ni == pytest.approx(ne)
    wi.close(); we.close()
def test_unknown_method_keeps_legacy_workbook_shape(tmp_path):
    result = _valuation(tmp_path)
    legacy = deepcopy(result)
    legacy["method"] = "legacy_method"
    wb = _book(legacy, tmp_path, read_only=True)
    assert "Summary" not in wb.sheetnames
    assert "Valuation" in wb.sheetnames
    wb.close()
def test_injection_strings_are_literal_cells(tmp_path):
    result = _valuation(tmp_path)
    payload = deepcopy(result)
    payload["ticker"] = '=HYPERLINK("https://bad")'
    payload["valuation_basis"] = "source: =1+1"
    wb = _book(payload, tmp_path, data_only=False)
    cells = _cells(wb)
    assert any(cell.value == payload["ticker"] and cell.data_type == "s" for cell in cells)
    assert any(cell.value == payload["valuation_basis"] and cell.data_type == "s" for cell in cells)
    assert any(cell.data_type == "f" for cell in cells)
    assert not any(cell.data_type == "f" and 'HYPERLINK' in cell.value for cell in cells)
    wb.close()
