from openpyxl import Workbook, load_workbook
from bellomberg.core.language import language_context

from bellomberg.valuation import dcf_engine
from bellomberg.valuation.documented_analysis import present_analysis
from tests.test_sector_operating_drivers import bundle_for


def _nav_payload():
    evidence = [
        {"scenario": "model", "driver": "perimeter", "values": {"entity": "FUND", "currency": "EUR"},
         "evidence": {"source": "=HYPERLINK(\"https://source\")", "period": "2026-06-30", "kind": "historical"}},
        {"scenario": "model", "driver": "calendar", "values": {"valuation_date": "2026-06-30", "periods": [], "discount_convention": "snapshot"},
         "evidence": {"source": "official NAV", "period": "2026-06-30", "kind": "historical"}},
    ]
    return {"ticker": "FUND", "method": "fund_nav", "engine": "mnav", "valuation_date": "2026-06-30",
            "valuation_basis": "snapshot NAV", "input_consumption": {"status": "complete"},
            "valuation_usability": {"usable": True}, "analytical_quality": {"status": "DOCUMENTATA", "snapshot": {"forecast_years": []}, "rows": evidence},
            "acquisition_snapshot": {"analysis_context": {"scenario_rationale": {"base": "NAV published; no terminal DCF."}}}}


def test_analysis_distinguishes_long_fcff_from_snapshot_nav(tmp_path):
    result = dcf_engine.generate_valuation("SYNTH-EXT", prepared_bundle=bundle_for(), output_dir=str(tmp_path))
    workbook = load_workbook(result["path"])
    present_analysis(workbook, result)
    assert workbook.sheetnames[1] == "Analysis"
    assert workbook.sheetnames[0] in {"Summary", "Valuation"}
    assert any("FCFF:" in str(cell.value) for row in workbook["Analysis"].iter_rows() for cell in row)
    assert any("Valore" in str(cell.value) or "Value" in str(cell.value)
               for row in workbook["Analysis"].iter_rows() for cell in row)
    workbook.close()

    nav = _nav_payload()
    workbook = Workbook()
    present_analysis(workbook, nav)
    analysis = workbook["Analysis"]
    cells = [cell for row in analysis.iter_rows() for cell in row]
    assert any("Attività" in str(cell.value) and "claims" in str(cell.value) for cell in cells)
    assert any(cell.value == "n.d." for cell in cells)
    assert any("HYPERLINK(\"https://source\")" in str(cell.value) and cell.data_type == "s" for cell in cells)
    assert any("NAV published; no terminal DCF." in str(cell.value) for cell in cells)
    assert not any(cell.data_type == "f" for cell in cells)
    workbook.close()


def test_analysis_language_keeps_dates_and_snapshot_state(tmp_path):
    seen = []
    for language in ("it", "en"):
        with language_context(language):
            workbook = Workbook()
            present_analysis(workbook, _nav_payload())
        values = [str(cell.value) for row in workbook["Analysis"].iter_rows() for cell in row]
        seen.append(values)
        assert any("2026-06-30" in value for value in values)
        assert any("n=0" in value for value in values)
        assert any("snapshot" in value.lower() for value in values)
        workbook.close()
    assert all(any(token in value for value in values)
               for values in seen for token in ("2026-06-30", "n=0"))
