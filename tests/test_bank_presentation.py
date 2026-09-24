"""Bank Summary keeps the reviewed visual structure and guarded calculations."""
from copy import deepcopy

import pytest
from openpyxl import Workbook

from bellomberg.core.language import language_context
from bellomberg.valuation.bank_formulas import apply_bank_formulas
from bellomberg.valuation.market_quote import build_market_quote
from test_bank_formulas import _case
from test_market_quote import DAY, quote_info


def payload_with_quote():
    payload = _case()
    bundle = deepcopy(payload['acquisition_snapshot'])
    info = quote_info(); info['symbol'] = payload['ticker']
    bundle['case']['info'] = info
    bundle['case']['sources']['profile'] = {'status': 'ok', 'source_id': 'synthetic-quote', 'as_of': DAY}
    quotation = next(row['values'] for row in payload['analytical_quality']['rows'] if row['driver'] == 'quotation')
    payload['market_quote'] = build_market_quote(bundle, quotation,
        {s: payload['fair_value_' + s] for s in ('bear', 'base', 'bull')})
    assert payload['market_quote']['status'] == 'ok'
    return payload


def test_bank_summary_distinguishes_quote_dates_inputs_and_version_without_changing_engine():
    payload = payload_with_quote(); original = deepcopy(payload)
    book = Workbook()
    with language_context('it'):
        assert apply_bank_formulas(book, payload)
    assert payload == original
    ws = book['Summary']
    assert book.sheetnames[0] == 'Summary' and ws.freeze_panes is None
    assert ws['E17'].value == payload['valuation_date']
    assert ws['E18'].data_type == 'f' and 'Model Inputs' in ws['E18'].value
    assert ws['E19'].value == 30 and ws['E20'].value == payload['market_quote']['observed_at']
    assert ws['E22'].value == 'ok'
    for col in ('D', 'E', 'F'):
        assert 'Input Checks' in ws[col + '8'].value and 'n.d.' in ws[col + '8'].value
        assert '$E$19>0' in ws[col + '9'].value and col + '8/$E$19' in ws[col + '9'].value
        for row in (11, 12):
            assert 'Model Inputs' in ws[f'{col}{row}'].value
        assert 'Input Checks' in ws[col + '13'].value
    assert payload['generation_id'] in ws['B26'].value
    assert DAY in ws['B26'].value and '2027-01-01' in ws['B26'].value
    assert ws['B7'].fill.fgColor.rgb[-6:] == '17324D'
    assert len(ws._charts) == 1 and ws._charts[0].varyColors is False
    assert ws.print_title_rows is None and ws.page_setup.orientation == 'portrait'


@pytest.mark.parametrize('status', ['stale', 'data_missing', 'currency_mismatch', 'identity_mismatch', 'fx_not_rolled'])
def test_unusable_observed_quote_never_unlocks_summary_upside(status):
    payload = payload_with_quote(); payload['market_quote']['status'] = status
    book = Workbook(); assert apply_bank_formulas(book, payload)
    ws = book['Summary']
    assert ws['E22'].value == status
    assert all(ws[f'{c}9'].value == 'n.d.' and ws[f'{c}9'].data_type == 's' for c in ('D', 'E', 'F'))
    assert all(ws[f'{c}8'].data_type == 'f' for c in ('D', 'E', 'F'))


def test_missing_quote_keeps_no_fallback_price_and_no_invented_generation_date():
    payload = _case(); payload.pop('market_quote', None)
    book = Workbook(); assert apply_bank_formulas(book, payload)
    assert book['Summary']['E19'].value == 'n.d.'
    assert book['Summary']['E22'].value == 'data_missing'
    assert book['Summary']['E9'].value == 'n.d.'
    assert 'n.d.' in book['Summary']['E21'].value


def test_presentation_preserves_every_other_sheet_and_original_summary_guards(monkeypatch):
    from bellomberg.valuation import bank_presentation
    payload = payload_with_quote()
    original, presented = Workbook(), Workbook()
    with monkeypatch.context() as scoped:
        scoped.setattr(bank_presentation, 'apply_bank_summary', lambda model: None)
        assert apply_bank_formulas(original, payload)
    assert apply_bank_formulas(presented, payload)
    assert set(original.sheetnames) == set(presented.sheetnames)
    for name in original.sheetnames:
        if name == 'Summary':
            for row in (8, 13):
                assert [original[name].cell(row, c).value for c in (4, 5, 6)] == [
                    presented[name].cell(row, c).value for c in (4, 5, 6)]
        else:
            assert [(c.coordinate, c.value, c.data_type) for cells in original[name] for c in cells] == [
                (c.coordinate, c.value, c.data_type) for cells in presented[name] for c in cells]


def test_labels_are_translated_but_formulas_and_evidence_stay_literal():
    payload = payload_with_quote()
    payload['market_quote']['source_id'] = '=HYPERLINK("https://example.org")'
    books = []
    for language in ('it', 'en'):
        book = Workbook()
        with language_context(language):
            assert apply_bank_formulas(book, payload)
        assert book['Summary']['E21'].data_type == 's'
        books.append(book)
    assert books[0]['Summary']['B11'].value != books[1]['Summary']['B11'].value
    assert [books[0]['Summary'].cell(row, 5).value for row in (8, 9, 11, 12, 13)] == [
        books[1]['Summary'].cell(row, 5).value for row in (8, 9, 11, 12, 13)]
