"""Real synthetic workbooks: labels change, numeric cells and formulas do not."""
import copy
from datetime import date
from pathlib import Path

import openpyxl
import pytest


def _workbook(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    numeric = {}; formulas = {}; texts = []
    names = wb.sheetnames
    for ws in wb:
        for row in ws:
            for cell in row:
                key = (ws.title, cell.coordinate)
                if cell.data_type == 'f': formulas[key] = cell.value
                elif type(cell.value) in (int, float): numeric[key] = cell.value
                elif isinstance(cell.value, str): texts.append(cell.value)
    wb.close()
    return names, numeric, formulas, texts


def test_sotp_labels_do_not_change_segment_values_or_formulas(tmp_path):
    from bellomberg.valuation.dcf_engine import _append_sotp_sheet
    from test_sotp import _sotp_base
    data = _sotp_base()
    views = []
    for language in ('it', 'en'):
        path = tmp_path/(language + '.xlsx')
        wb = openpyxl.Workbook(); wb.save(path); wb.close()
        _append_sotp_sheet(path, copy.deepcopy(data), 'SYNTH', 'EUR', language=language)
        views.append(_workbook(path))
    assert views[0][:3] == views[1][:3] and views[0][2]
    assert 'Segmento' in views[0][3] and 'Segment' in views[1][3]


def _generate(kind, language, tmp_path, monkeypatch):
    from bellomberg.valuation import dcf_buyside_v3, dcf_bank, dcf_rab, dcf_modeler
    target = str(tmp_path / (kind + language + '.xlsx'))
    if kind == 'v3':
        from test_dcf_quality_integration import _spec
        spec = _spec(); spec['variant_view'] = 'Citazione personale invariata'; spec['price'] = 15
        result = dcf_buyside_v3.build_model_v3(spec, target, language=language)
    elif kind == 'bank':
        from test_dcf_bank import INFO, WACC, PROFILE
        spec = dcf_bank.build_bank_spec('SYNTHBANK', copy.deepcopy(INFO), WACC, profile=PROFILE,
                                       variant_view='Citazione personale invariata')
        result = dcf_bank.build_bank_model(spec, target, language=language)
    elif kind == 'rab':
        from test_dcf_rab import INFO_IT, WACC, PROFILE_REG, OGGI
        spec = dcf_rab.build_rab_spec('SYNTHRAB', dict(INFO_IT, country='Spain', currentPrice=5), WACC,
              rab={'rab_base': 10000, 'allowed_return': .07, 'capex_plan': [500]*10, 'da_pct_rab': .04,
                   'net_debt': 5000, 'kd': .04, 'tax_rate': .25}, profile=PROFILE_REG, today=OGGI,
              variant_view='Citazione personale invariata')
        assert 'error' not in spec, spec
        result = dcf_rab.build_rab_model(spec, target, language=language)
    else:
        monkeypatch.setattr(dcf_modeler, 'MODELS_DIR', str(tmp_path))
        data = {'ticker':'SYNTH', 'name':'Citazione personale invariata', 'current_price': 15,
                'revenue_ttm': 1e9, 'shares_outstanding': 1e8, 'ebitda_margin': .2, 'op_margin': .15,
                'net_margin': .1, 'gross_margin': .3, 'revenue_growth': .05, 'total_debt': 2e8, 'total_cash': 1e8}
        monkeypatch.setattr(dcf_modeler, '_fetch_fundamentals', lambda ticker: copy.deepcopy(data))
        result = dcf_modeler.build_dcf_excel('SYNTH', output_path=target, language=language)
    assert 'error' not in result, result
    return result['path']


@pytest.mark.parametrize('kind,it_label,en_label', [
    ('v3','Prezzo corrente','Current price'), ('bank','Prezzo corrente','Current price'),
    ('rab','Prezzo di mercato','Market price'), ('generic','SINTESI DELLA VALUTAZIONE','VALUATION SUMMARY'),
])
def test_legacy_builders_keep_all_numeric_and_formula_cells(kind, it_label, en_label, tmp_path, monkeypatch):
    it = _workbook(_generate(kind, 'it', tmp_path, monkeypatch))
    en = _workbook(_generate(kind, 'en', tmp_path, monkeypatch))
    assert it[:3] == en[:3]
    assert len(it[1]) > 10 and len(it[2]) > 10
    assert any(it_label in s for s in it[3]) and any(en_label in s for s in en[3])
    assert any('Citazione personale invariata' in s for s in it[3])
    assert any('Citazione personale invariata' in s for s in en[3])


@pytest.mark.parametrize('method', ['documented', 'managed'])
def test_snapshot_builders_preserve_structured_records_and_localise_labels(method, tmp_path):
    from bellomberg.valuation.documented_inputs import build_documented_workbook
    from bellomberg.valuation.managed_care_adapter import build_managed_care_workbook
    quality = {'record_adapter': True, 'status': 'BOZZA_AUTOMATICA', 'snapshot': {'forecast_years':[2027]},
               'rows':[{'scenario':'base','driver':'original_key','values':17.25,
                        'evidence':{'source':'Fonte originale invariata','rationale':'Motivazione invariata'}}]}
    payload = {'ticker':'SYNTH', 'method':'synthetic_method', 'generation_id':'synthetic_generation',
        'valuation_date':'2026-09-12', 'valuation_basis':'Original basis', 'fair_value_base':42.125,
        'analytical_quality':quality, 'acquisition_snapshot':{'case':{'as_of':'2026-09-12'}},
        'calculation_details':{'scenarios':{'base':{'original_key':17.25}}},
        'managed_care':{'scenarios':{'base':{'rows':[{'year':2027,'original_key':17.25}]}}}}
    builder = build_documented_workbook if method == 'documented' else build_managed_care_workbook
    original = copy.deepcopy(payload)
    it = _workbook(builder(payload, str(tmp_path/'it'), language='it'))
    en = _workbook(builder(payload, str(tmp_path/'en'), language='en'))
    assert it[:3] == en[:3] and it[1]
    assert 'VALUTAZIONE SETTORIALE' in it[3] and 'SECTOR VALUATION' in en[3]
    assert 'Fonte originale invariata' in it[3] and 'Fonte originale invariata' in en[3]
    assert payload == original


@pytest.mark.parametrize('ticker,payload_name', [('TESORO','BTC_PAYLOAD'), ('HYPEX','HYPE_PAYLOAD'), ('FONDO.L','CEF_PAYLOAD')])
def test_mnav_target_binding_and_all_three_workbooks_are_language_invariant(ticker, payload_name, tmp_path):
    import json
    import test_dcf_mnav as fixture
    from bellomberg.storage.classificazione import carica_veicoli
    from bellomberg.valuation.dcf_mnav import build_mnav_spec, build_mnav_model
    path = tmp_path/'vehicles.json'; path.write_text(json.dumps(fixture.VOCI), encoding='utf8')
    registry = carica_veicoli(str(path))
    info = dict(fixture.INFO_USD)
    spec = build_mnav_spec(ticker, info, nav_target=1.2, variant_view='Citazione personale invariata',
        tool_payload=copy.deepcopy(getattr(fixture,payload_name)), today=fixture.OGGI, negozio=registry)
    assert 'error' not in spec, spec
    it_result = build_mnav_model(copy.deepcopy(spec), str(tmp_path/'it.xlsx'), language='it')
    en_result = build_mnav_model(copy.deepcopy(spec), str(tmp_path/'en.xlsx'), language='en')
    assert 'error' not in it_result and 'error' not in en_result
    it = _workbook(it_result['path']); en = _workbook(en_result['path'])
    assert it[:3] == en[:3] and len(it[2]) > 5
    assert 'Target premio/sconto (ANALISTA)' in it[3]
    assert 'Target premium/discount (ANALYST)' in en[3]
