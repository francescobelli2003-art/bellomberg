"""Sourcebook uses only the frozen provider envelopes, with literal text."""
from copy import deepcopy
from io import BytesIO

from openpyxl import Workbook, load_workbook

from bellomberg.valuation.sourcebook_presentation import present_sourcebook


def _payload():
    return {'acquisition_snapshot': {'case': {'as_of': '2026-09-19', 'sources': {
        'financials': {'status': 'ok', 'source_id': '=bad-history', 'as_of': '2026-09-18',
                       'data': {'income_stmt': {'table': True, 'columns': ['2025-12-31', '2024-12-31'],
                                                'index': ['Revenue', '=malicious'],
                                                'data': [[0, -23.5], [None, 12]], 'missing_cells': 1},
                                'balance_sheet': None, 'cashflow': None,
                                'errors': {'cashflow': 'source error'}}},
        'consensus': {'status': 'stale', 'source_id': 'Yahoo', 'as_of': '2026-09-01',
                      'valid_until': '2026-09-10', 'message': 'Expired',
                      'data': {'eps_estimates': [{'period': '0y', 'avg': 0.0},
                                                 {'period': '+1y', 'avg': -1.2}],
                               'price_targets': '=bad-consensus'}},
        'guidance': {'status': 'data_missing', 'source_id': 'guidance registry',
                     'message': 'No active guidance', 'data': None},
    }}}}


def test_sourcebook_preserves_raw_values_labels_metadata_and_literal_strings():
    payload = _payload()
    before = deepcopy(payload)
    wb = Workbook()
    assert present_sourcebook(wb, payload)
    assert payload == before
    history = wb['Source History']
    assert history['D16'].value == 0
    assert history['E16'].value == -23.5
    assert history['D17'].value == 'n.d.'
    assert history['C16'].value == 'n.d.'  # unit not asserted from a Yahoo table
    assert history['D5'].value == '=bad-history'
    consensus = wb['Source Consensus']
    assert consensus['D8'].value == 'stale'
    assert consensus['D10'].value == 'Expired'
    values = [(consensus.cell(row, 2).value, consensus.cell(row, 3).value)
              for row in range(14, consensus.max_row + 1)]
    assert ('eps_estimates/0/period', '0y') in values
    assert ('eps_estimates/1/period', '+1y') in values
    assert ('eps_estimates/0/avg', 0.0) in values
    assert ('eps_estimates/1/avg', -1.2) in values
    assert wb['Source Guidance']['D8'].value == 'data_missing'
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    assert cell.data_type == 's'
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    reopened = load_workbook(stream)
    assert reopened['Source Consensus']['D8'].value == 'stale'
    assert reopened['Source History']['B17'].data_type == 's'


def test_sourcebook_requires_existing_snapshot():
    wb = Workbook()
    assert not present_sourcebook(wb, {})
    assert wb.sheetnames == ['Sheet']
